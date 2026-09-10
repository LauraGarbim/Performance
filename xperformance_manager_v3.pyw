

import os, sys, re, json, time, shutil, zipfile, threading, subprocess, traceback, atexit, sqlite3, contextlib, base64
from datetime import datetime, timedelta

_HERE = os.path.dirname(os.path.abspath(__file__))
_LOGF = os.path.join(_HERE, "xperformance_v2_erro.log")

def _checar_dependencias():
    """Verifica dependências SEM instalar nada (a instalação automática via pip foi
    removida por segurança). Se faltar algo, avisa e encerra pedindo a instalação
    única via requirements.txt."""
    req = [("webview", "pywebview"), ("pdfplumber", "pdfplumber"),
           ("openpyxl", "openpyxl"), ("pypdf", "pypdf")]
    if os.name == "nt":
        req.append(("win32com", "pywin32"))
    faltando = []
    for mod, pip in req:
        try:
            __import__(mod)
        except ImportError:
            faltando.append(pip)
    if faltando:
        msg = ("Faltam dependências para rodar o XPerformance:\n  " + ", ".join(faltando) +
               "\n\nAbra o Prompt de Comando nesta pasta e rode uma única vez:\n"
               "  pip install -r requirements.txt")
        try:
            with open(_LOGF, "a", encoding="utf-8") as f:
                f.write(msg + "\n")
        except Exception:
            pass
        if os.name == "nt":
            try:
                import ctypes
                ctypes.windll.user32.MessageBoxW(0, msg, "XPerformance — dependências", 0x10)
            except Exception:
                pass
        sys.exit(1)

_checar_dependencias()

import webview, pdfplumber, openpyxl, pypdf
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Criptografia em repouso (LGPD) ──────────────────────────────────────────────
# DPAPI do Windows: cifra vinculada à conta/máquina do usuário, sem senha manual.
# Protege os dados sensíveis nos arquivos internos (SQLite e JSONs) — uma cópia
# solta desses arquivos fica ILEGÍVEL fora da conta Windows da Laura.
_ENC_TXT = "XPENC1:"      # prefixo de valor de texto cifrado (colunas do SQLite)
_ENC_BIN = b"XPENC1\n"    # prefixo de arquivo cifrado (JSONs)
try:
    import win32crypt
    _CRYPTO_OK = (os.name == "nt")
except Exception:
    _CRYPTO_OK = False

def _dpapi_enc(b):
    return win32crypt.CryptProtectData(b, "xperformance", None, None, None, 0)

def _dpapi_dec(b):
    return win32crypt.CryptUnprotectData(b, None, None, None, 0)[1]

def _enc_txt(s):
    """Cifra um texto (str -> token 'XPENC1:...'). No-op se cripto indisponível."""
    if not _CRYPTO_OK or not s:
        return s
    try:
        return _ENC_TXT + base64.b64encode(_dpapi_enc(s.encode("utf-8"))).decode("ascii")
    except Exception:
        return s

def _dec_txt(s):
    """Decifra um token 'XPENC1:...'; devolve o próprio valor se não estiver cifrado
    (compatível com dados legados em texto puro — migram no próximo save)."""
    if not isinstance(s, str) or not s.startswith(_ENC_TXT):
        return s
    try:
        return _dpapi_dec(base64.b64decode(s[len(_ENC_TXT):])).decode("utf-8")
    except Exception:
        return s

def _ler_json_seguro(path, default):
    """Lê JSON cifrado (ou legado em texto puro) com DPAPI."""
    try:
        with open(path, "rb") as f:
            raw = f.read()
    except Exception:
        return default
    if raw.startswith(_ENC_BIN):
        try:
            raw = _dpapi_dec(raw[len(_ENC_BIN):])
        except Exception:
            return default
    try:
        return json.loads(raw.decode("utf-8"))
    except Exception:
        return default

def _gravar_json_seguro(path, obj, **kw):
    """Grava JSON cifrado (DPAPI) de forma atômica. Sem cripto, grava em texto puro."""
    try:
        data = json.dumps(obj, ensure_ascii=False, **kw).encode("utf-8")
        if _CRYPTO_OK:
            try:
                data = _ENC_BIN + _dpapi_enc(data)
            except Exception:
                pass
        tmp = path + ".tmp"
        with open(tmp, "wb") as f:
            f.write(data)
        os.replace(tmp, path)
    except Exception:
        pass

EXCEL_PATH   = os.path.join(_HERE, "XPerformance - Email Mensal v2.xlsx")
LEGACY_EXCEL = os.path.join(_HERE, "XPerformance - Email Mensal.xlsx")
PASTA_PDFS   = os.path.join(_HERE, "Relatórios Clientes")
CONFIG_PATH  = os.path.join(_HERE, "_config.json")
NOTIF_PATH   = os.path.join(_HERE, "_notificacoes.json")
ONDE_PATH    = os.path.join(_HERE, "_onde_investir.json")
DB_PATH      = os.path.join(_HERE, "xperformance_dados.db")
EMAIL_ABA    = "E-mail"
CLI_START    = 14
CLI_MAX      = CLI_START + 5000
COLS = ["Nome", "Conta", "Rentabilidade (%)", "Ganho (R$)", "CDI (%)", "Arquivo PDF"]
MESES = {"01":"Janeiro","02":"Fevereiro","03":"Março","04":"Abril","05":"Maio","06":"Junho",
         "07":"Julho","08":"Agosto","09":"Setembro","10":"Outubro","11":"Novembro","12":"Dezembro"}

ASSUNTO_DEF = "Atualização de Carteira — {mes_ref}"
CORPO_DEF = """\
Olá {primeiro_nome},

Fechamos {mes_ref} e preparei um resumo objetivo do seu relatório XPerformance — os números principais, com o contexto do que cada um significa.

{resumo_pontos}

{composicao_carteira}

{contexto_mercado}

Qualquer dúvida ou se quiser conversar sobre os próximos passos, estou à disposição!

Att,
Guilherme Enrico"""

def mes_bonito(aba):
    m = re.match(r"(\d{2})-(\d{4})", str(aba) or "")
    return f"{MESES.get(m.group(1), m.group(1))}/{m.group(2)}" if m else str(aba)

MERCADO_ABA = "Mercado"

STRATEGY_VARS = [
    ("Pós Fixado", "pos_fixado"), ("Pré Fixado", "pre_fixado"), ("Inflação", "inflacao"),
    ("Renda Variável Brasil", "rv_brasil"), ("Renda Variável Global", "rv_global"),
    ("Fundos Listados", "fundos_listados"), ("Multimercado", "multimercado"),
    ("Alternativo", "alternativo"), ("Caixa", "caixa"), ("Proventos", "proventos"),
]
def _norm_cls(s):
    return re.sub(r"\s+", " ", str(s or "").strip()).lower()

def _clean_cls(raw):
    """Limpa o nome da estratégia (o PDF cola texto do cabeçalho antes do nome,
    ex.: 'M Desde início Pós Fixado' -> 'Pós Fixado')."""
    raw = re.sub(r"\s+", " ", str(raw or "")).strip()
    low = raw.lower()
    best = None
    for nome, _ in STRATEGY_VARS:
        if low.endswith(nome.lower()) and (best is None or len(nome) > len(best)):
            best = nome
    if best:
        return best

    raw = re.sub(r".*?(?:Total investido|Desde in[íi]cio|Saldo Bruto|Estrat[ée]gias|"
                 r"M[êe]s Atual|\b\d+\s*M\b|\bAno\b|\bM\b)\s*", "", raw, flags=re.I).strip()
    parts = raw.split()
    return " ".join(parts[-3:]) if parts else raw
B3_TICKERS = [("PETR4", "Petrobras PN"), ("VALE3", "Vale ON"), ("ITUB4", "Itaú Unibanco PN"),
              ("BBDC4", "Bradesco PN"), ("BBAS3", "Banco do Brasil ON"), ("ABEV3", "Ambev ON"),
              ("B3SA3", "B3 ON"), ("WEGE3", "WEG ON")]

EVOL_EXTRA = ["Movimentações (R$)", "Rent. Total (%)", "Ganho Total (R$)", "Período Total", "Composição"]
COLS_FULL = COLS + ["Patrimônio (R$)"] + EVOL_EXTRA

class _SafeDict(dict):
    """Variáveis desconhecidas ficam como estão no texto, sem quebrar o envio."""
    def __missing__(self, k): return "{" + k + "}"

def F(h):  return PatternFill("solid", fgColor=h.lstrip("#"))
def ft(bold=False, color="202124", size=10):
    return Font(bold=bold, color=color.lstrip("#"), size=size, name="Segoe UI")
def bd():
    s = Side(style="thin", color="DADCE0"); return Border(left=s, right=s, top=s, bottom=s)
def al(h="left", wrap=False): return Alignment(horizontal=h, vertical="center", wrap_text=wrap)

def _rotacionar_log(max_bytes=1_000_000):
    """Impede o log de erros crescer sem limite (ele acumula identificadores de
    clientes ao longo do tempo). Mantém apenas a parte final."""
    try:
        if os.path.exists(_LOGF) and os.path.getsize(_LOGF) > max_bytes:
            with open(_LOGF, "rb") as f:
                f.seek(-(max_bytes // 2), os.SEEK_END)
                tail = f.read()
            with open(_LOGF, "wb") as f:
                f.write(b"[log truncado para conter o tamanho]\n" + tail)
    except Exception:
        pass

def carregar_config():
    c = _ler_json_seguro(CONFIG_PATH, {})
    return c if isinstance(c, dict) else {}

def _notif_store_load():
    st = _ler_json_seguro(NOTIF_PATH, {})
    return st if isinstance(st, dict) else {}

def _notif_store_save(st):
    _gravar_json_seguro(NOTIF_PATH, st, indent=1)

def _notif_ts(chegou):
    m = re.match(r"(\d{2})/(\d{2})/(\d{4}) (\d{2}):(\d{2})", str(chegou or ""))
    return (m.group(3) + m.group(2) + m.group(1) + m.group(4) + m.group(5)) if m else ""

def gravar_config(cfg):
    _gravar_json_seguro(CONFIG_PATH, cfg, indent=2)

def _wb_valido(path):
    """Confere se o .xlsx é um zip íntegro (valida CRC) sem segurar o handle do arquivo."""
    try:
        with zipfile.ZipFile(path) as z:
            return z.testzip() is None and "xl/workbook.xml" in z.namelist()
    except Exception:
        return False

def _novo_wb():
    wb = openpyxl.Workbook()
    if "Sheet" in wb.sheetnames: del wb["Sheet"]
    return wb

_read_lock = threading.RLock()
_live = {"wb": None, "disk": None, "dirty": False, "gen": 0, "timer": None, "pin": 0}

def _stat_key():
    try:
        st = os.stat(EXCEL_PATH)
        return (st.st_mtime_ns, st.st_size)
    except OSError:
        return None

def _carregar_live():
    """(Re)carrega o workbook do disco para a memória, com recuperação de corrompido."""
    wb = None
    if os.path.exists(EXCEL_PATH):
        try:
            wb = openpyxl.load_workbook(EXCEL_PATH)
        except Exception:

            bak = EXCEL_PATH + ".bak.xlsx"
            if os.path.exists(bak):
                try:
                    wb = openpyxl.load_workbook(bak)
                    try: os.replace(EXCEL_PATH, EXCEL_PATH + ".corrompido")
                    except Exception: pass
                except Exception:
                    wb = None
            if wb is None:

                try: os.replace(EXCEL_PATH, EXCEL_PATH + ".corrompido")
                except Exception: pass
                wb = _novo_wb()
    elif os.path.exists(LEGACY_EXCEL):
        try:
            wb = openpyxl.load_workbook(LEGACY_EXCEL)
        except Exception:
            wb = None
    if wb is None:
        wb = _novo_wb()
    _live["wb"] = wb
    _live["disk"] = _stat_key()
    _live["dirty"] = False
    _live["gen"] += 1
    return wb

def _tem_dados():
    """Há dados para ler? Vale o workbook vivo (pode ainda não ter ido ao disco)."""
    return _live["wb"] is not None or os.path.exists(EXCEL_PATH)

def get_wb():
    """Workbook vivo compartilhado. Leituras e gravações veem o mesmo estado na hora."""
    with _read_lock:
        wb = _live["wb"]
        if wb is None:
            return _carregar_live()

        if _live.get("pin", 0) == 0 and not _live["dirty"] and not _live.get("flushing") and _stat_key() != _live["disk"]:
            return _carregar_live()
        return wb

def _wb_cached():
    """Compatibilidade: mesmas leituras, agora direto do workbook vivo."""
    return get_wb()

@contextlib.contextmanager
def _wb_pin():
    """Congela a geração do workbook durante um endpoint de leitura, para que
    todas as *_cache() de uma mesma resposta venham do MESMO estado (Bug B:
    painel/lista deixam de misturar dados de gerações diferentes)."""
    with _read_lock:
        get_wb()                                # 1 reload de entrada, ainda com pin==0
        _live["pin"] = _live.get("pin", 0) + 1  # só então congela a geração
    try:
        yield
    finally:
        with _read_lock:
            _live["pin"] = max(0, _live.get("pin", 0) - 1)

_derived_cache = {"key": None, "data": {}}

def _cached(nome, fn):
    with _read_lock:
        get_wb()
        key = _live["gen"]
        if _derived_cache["key"] != key:
            _derived_cache["key"] = key
            _derived_cache["data"] = {}
        d = _derived_cache["data"]
        if nome not in d:
            d[nome] = fn()
        return d[nome]

def coletar_cache():   return _cached("agg",  lambda: coletar_por_cliente(_wb_cached()))
def cadastro_cache():  return _cached("cad",  lambda: ler_clientes_cadastro(_wb_cached()))
def aportes_cache():   return _cached("aportes", lambda: ler_aportes(_wb_cached()))
def contatos_cache():  return _cached("contatos", lambda: ler_contatos(_wb_cached()))
def receita_cache():   return _cached("receita", lambda: ler_receita(_wb_cached()))
def aloc_cache():      return _cached("aloc", lambda: ler_alocacao(_wb_cached()))
def onde_cache():      return _cached("onde", lambda: ler_onde_investir(_wb_cached()))
def mercado_cache():   return _cached("mercado", lambda: ler_mercado(_wb_cached()))

_save_lock = threading.Lock()
_FLUSH_DELAY = 0.8

def salvar_wb(wb):
    """Confirma a alteração EM MEMÓRIA imediatamente (o app inteiro já enxerga) e
    agenda a gravação do .xlsx em segundo plano. O clique nunca espera o disco."""
    base = os.path.basename(EXCEL_PATH); dir_ = os.path.dirname(EXCEL_PATH) or "."

    locks = {"~$" + base, "~$" + base[2:]}
    if any(os.path.exists(os.path.join(dir_, lk)) for lk in locks):
        raise PermissionError(f"A planilha está ABERTA no Excel. Feche o arquivo "
                              f"\"{base}\" e tente de novo.")
    with _read_lock:
        _live["wb"] = wb
        _live["dirty"] = True
        _live["gen"] += 1
        _derived_cache["key"] = None
        t = _live["timer"]
        if t is not None:
            try: t.cancel()
            except Exception: pass
        t = threading.Timer(_FLUSH_DELAY, _flush_disco)
        t.daemon = True
        _live["timer"] = t
        t.start()

def _flush_disco():
    """Grava o workbook vivo no .xlsx (thread de fundo, serializada e atômica)."""
    tmp = None
    try:
        with _save_lock, _MUT_LOCK:
            with _read_lock:
                if not _live["dirty"] or _live["wb"] is None:
                    return
                _live["flushing"] = True
                tmp = EXCEL_PATH + f".{os.getpid()}.{threading.get_ident()}.tmp.xlsx"
                _live["wb"].save(tmp)
                _live["dirty"] = False
            if not _wb_valido(tmp):
                raise IOError("Arquivo temporário inválido ao gravar — tente de novo.")

            # Conflito externo: o .xlsx mudou no disco (OneDrive/Excel) desde o último
            # load/flush. Preserva a versão do disco em .conflito antes de sobrescrever,
            # para o trabalho em memória do usuário nunca se perder E a versão externa
            # ficar recuperável.
            try:
                if _live["disk"] is not None and os.path.exists(EXCEL_PATH) and _stat_key() != _live["disk"]:
                    stamp = time.strftime("%Y%m%d-%H%M%S")
                    shutil.copy2(EXCEL_PATH, EXCEL_PATH + f".conflito-{stamp}.xlsx")
                    _db_log(f"_flush_disco: conflito externo detectado — versão do disco "
                            f"preservada em .conflito-{stamp}.xlsx")
            except Exception:
                pass

            gravou = False
            for _ in range(6):
                try:
                    os.replace(tmp, EXCEL_PATH)
                    gravou = True
                    break
                except PermissionError:
                    time.sleep(0.4)
            if not gravou:
                raise PermissionError("Não consegui gravar a planilha (Excel/OneDrive/antivírus "
                                      "segurando o arquivo).")
            with _read_lock:
                _live["disk"] = _stat_key()

            try: shutil.copy2(EXCEL_PATH, EXCEL_PATH + ".bak.xlsx")
            except Exception: pass
    except Exception:

        with _read_lock:
            _live["dirty"] = True
            t = threading.Timer(5.0, _flush_disco)
            t.daemon = True
            _live["timer"] = t
            t.start()
        try:
            with open(_LOGF, "a", encoding="utf-8") as f:
                f.write("flush em segundo plano falhou:\n" + traceback.format_exc())
        except Exception:
            pass
    finally:
        with _read_lock:
            _live["flushing"] = False
        if tmp and os.path.exists(tmp):
            try: os.remove(tmp)
            except Exception: pass

def _flush_agora():
    """Descarrega qualquer gravação pendente AGORA (chamado ao fechar o app)."""
    with _read_lock:
        t = _live["timer"]
        if t is not None:
            try: t.cancel()
            except Exception: pass
    for _ in range(3):
        _flush_disco()
        with _read_lock:
            if not _live["dirty"]:
                break
        time.sleep(1.0)

atexit.register(_flush_agora)

X_BG = "0A0A0C"; X_BAND = "16161B"; X_ROW1 = "111114"; X_ROW2 = "17171C"
X_BORDER = "26262E"; X_TXT = "F2F2F4"; X_MUT = "8E8E98"; X_YELLOW = "F2C029"
C_OK = "3DD68C"; C_WARN = "F0B429"; C_ERR = "F4717F"

CAD_HEADERS = ["Nome Completo", "Conta", "E-mail", "Termômetro", "Perfil",
               "Estado", "Cidade", "Nascimento",
               "Previdência", "Prev. valor", "Prev. onde",
               "Internacional", "Intl. valor", "Intl. onde",
               "Seguro", "Cobertura",
               "Plugado na mesa", "Mesa valor", "Mesa onde",
               "Grupo Familiar", "Objetivo", "Aposentar aos"]
CAD_NCOL = len(CAD_HEADERS)
CAD_LEFT = {1, 3, 7, 10, 11, 13, 14, 16, 18, 19, 20, 21}

def _cab_cli(ws):
    r = CLI_START - 1
    ult = get_column_letter(CAD_NCOL)
    try: ws.merge_cells(f"A{r}:{ult}{r}")
    except Exception: pass
    t = ws.cell(r, 1, "CADASTRO DE CLIENTES")
    t.font = Font(name="Segoe UI", bold=True, size=11, color=X_YELLOW)
    t.fill = F(X_BG); t.alignment = al("center")
    ws.row_dimensions[r].height = 26
    for col, h in enumerate(CAD_HEADERS, 1):
        c = ws.cell(CLI_START, col, h)
        c.font = ft(bold=True, color="202124"); c.fill = F("f4f4f5")
        c.alignment = al("center"); c.border = bd()
    ws.row_dimensions[CLI_START].height = 20

def garantir_email(wb):
    """Garante a aba E-mail e (re)aplica o visual XP sem mexer nos valores."""

    if EMAIL_ABA in wb.sheetnames:
        ws0 = wb[EMAIL_ABA]
        if ws0.max_row > CLI_MAX or ws0.max_column > 60:
            assunto0 = ws0["B1"].value; corpo0 = ws0["B3"].value
            cad0 = []
            for row in ws0.iter_rows(min_row=CLI_START + 1, max_row=CLI_MAX, values_only=True):
                if row and row[1]:
                    cad0.append(tuple(str(row[i] or "") if i < len(row) else "" for i in range(CAD_NCOL)))
            idx = wb.sheetnames.index(EMAIL_ABA)
            del wb[EMAIL_ABA]
            ws0 = wb.create_sheet(EMAIL_ABA, idx)
            ws0["B1"].value = assunto0 or ASSUNTO_DEF
            ws0["B3"].value = corpo0 or CORPO_DEF
            for ri, vals in enumerate(cad0, start=CLI_START + 1):
                for col, v in enumerate(vals, 1):
                    ws0.cell(ri, col, v)
    novo = EMAIL_ABA not in wb.sheetnames or not wb[EMAIL_ABA]["B3"].value
    if novo:
        if EMAIL_ABA in wb.sheetnames: del wb[EMAIL_ABA]
        ws = wb.create_sheet(EMAIL_ABA, 0)
        ws["B1"].value = ASSUNTO_DEF; ws["B3"].value = CORPO_DEF
    ws = wb[EMAIL_ABA]
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X_YELLOW
    for rng in (f"B1:C1", f"B3:C3"):
        try: ws.merge_cells(rng)
        except Exception: pass
    for addr, txt, vert in (("A1", "ASSUNTO", "center"), ("A3", "CORPO", "top")):
        c = ws[addr]; c.value = txt
        c.font = Font(name="Segoe UI", bold=True, size=10, color=X_YELLOW)
        c.fill = F(X_BG)
        c.alignment = Alignment(horizontal="center", vertical=vert)
    ws["B1"].font = ft(); ws["B1"].fill = F("ffffff")
    ws["B1"].border = bd(); ws["B1"].alignment = al("left")
    ws["B3"].font = Font(name="Segoe UI", size=10, color="202124")
    ws["B3"].fill = F("ffffff"); ws["B3"].border = bd()
    ws["B3"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    ws.row_dimensions[1].height = 26; ws.row_dimensions[3].height = 400
    ws.column_dimensions["A"].width = 34; ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 38
    ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 16
    ws.column_dimensions["F"].width = 10; ws.column_dimensions["G"].width = 18
    ws.column_dimensions["H"].width = 14
    for letra, larg in (("I", 13), ("J", 14), ("K", 22), ("L", 13), ("M", 14), ("N", 22),
                        ("O", 12), ("P", 16), ("Q", 16), ("R", 14), ("S", 22), ("T", 34)):
        ws.column_dimensions[letra].width = larg
    _cab_cli(ws)

    for row in ws.iter_rows(min_row=CLI_START + 1, max_row=min(ws.max_row, CLI_MAX)):
        if not row[1].value: continue
        ri = row[0].row
        for col in range(1, CAD_NCOL + 1):
            c = ws.cell(ri, col); c.font = ft(size=10)
            c.fill = F("ffffff") if ri % 2 == 0 else F("f7f7f8")
            c.alignment = al("left" if col in CAD_LEFT else "center"); c.border = bd()
        ws.row_dimensions[ri].height = 20

def buscar_email(wb, conta):
    if EMAIL_ABA not in wb.sheetnames: return ""
    for row in wb[EMAIL_ABA].iter_rows(min_row=CLI_START + 1, max_row=CLI_MAX, values_only=True):
        if row and _conta_str(row[1]) == _conta_str(conta):
            return str(row[2] or "").strip()
    return ""

def registrar_clientes(wb, lista):
    ws = wb[EMAIL_ABA]; existentes = {}
    for row in ws.iter_rows(min_row=CLI_START + 1, max_row=min(ws.max_row, CLI_MAX)):
        if row[1].value: existentes[_conta_str(row[1].value)] = row[0].row
    proxima = max(existentes.values(), default=CLI_START) + 1
    for d in lista:
        conta = str(d.get("conta", "")).strip()
        if not conta: continue
        if conta in existentes:
            if not ws.cell(existentes[conta], 1).value:
                ws.cell(existentes[conta], 1).value = d.get("nome", "")
        else:
            ri = proxima
            for col, val in enumerate([d.get("nome", ""), conta, ""], 1):
                c = ws.cell(ri, col, val); c.font = ft(size=10)
                c.fill = F("ffffff") if ri % 2 == 0 else F("f8f9fa")
                c.alignment = al("left" if col == 1 else "center"); c.border = bd()
            ws.row_dimensions[ri].height = 20
            existentes[conta] = proxima; proxima += 1

def _estado_flag(v):
    """Estado escolhido diretamente: 🟢 ok (Tem) · 🟡 rev (A revisar) · 🔴 nao (Não tem) · '' (não informado)."""
    s = str(v or "").strip().lower()
    if s in ("tem", "sim", "s", "ok", "1", "true", "x", "yes"): return "ok"
    if s in ("a revisar", "revisar", "parcial", "rev", "análise", "analise", "em análise", "em analise"): return "rev"
    if s in ("não tem", "nao tem", "não", "nao", "n", "0", "false", "no"): return "nao"
    return ""

def _is_sim(v):
    return _estado_flag(v) == "ok"

def _plugado(v):
    """'Plugado na mesa' = Tem ou A revisar (não conta 'Não tem' nem vazio)."""
    return _estado_flag(v) in ("ok", "rev")

def flags_cliente(c):
    """Status (escolhido diretamente no seletor) + valores das 4 informações."""
    def one(k, usa_onde=True):
        return {"tem": c.get(k + "_tem", ""), "valor": c.get(k + "_valor", ""),
                "onde": c.get(k + "_onde", "") if usa_onde else "",
                "status": _estado_flag(c.get(k + "_tem"))}
    return {"prev": one("prev"), "intl": one("intl"), "seg": one("seg", usa_onde=False), "mesa": one("mesa")}

CAD_EXTRA = [("prev_tem", 8), ("prev_valor", 9), ("prev_onde", 10),
             ("intl_tem", 11), ("intl_valor", 12), ("intl_onde", 13),
             ("seg_tem", 14), ("seg_valor", 15),
             ("mesa_tem", 16), ("mesa_valor", 17), ("mesa_onde", 18),
             ("grupo_familiar", 19),
             ("objetivo", 20), ("apos_idade", 21)]

def _gf_parse(raw):
    return [e.strip() for e in str(raw or "").split(";") if e.strip()]

def _gf_join(entradas):
    return "; ".join(entradas)

def _gf_split(entrada):
    """Separa 'ref|parentesco' -> (ref, parentesco)."""
    p = str(entrada or "").split("|", 1)
    return p[0].strip(), (p[1].strip() if len(p) > 1 else "")

_GF_INV = {}
def _gf_reg(chaves, masc, fem):
    for k in chaves: _GF_INV[k] = (masc, fem)
_gf_reg(("marido", "esposa"), "Marido", "Esposa")
_gf_reg(("cônjuge", "conjuge"), "Cônjuge", "Cônjuge")
_gf_reg(("companheiro", "companheira"), "Companheiro", "Companheira")
_gf_reg(("pai", "mãe", "mae"), "Filho", "Filha")
_gf_reg(("filho", "filha"), "Pai", "Mãe")
_gf_reg(("irmão", "irmao", "irmã", "irma"), "Irmão", "Irmã")
_gf_reg(("avô", "avó", "avo"), "Neto", "Neta")
_gf_reg(("neto", "neta"), "Avô", "Avó")
_gf_reg(("tio", "tia"), "Sobrinho", "Sobrinha")
_gf_reg(("sobrinho", "sobrinha"), "Tio", "Tia")
_gf_reg(("primo", "prima"), "Primo", "Prima")
_gf_reg(("genro", "nora"), "Sogro", "Sogra")
_gf_reg(("sogro", "sogra"), "Genro", "Nora")
_gf_reg(("cunhado", "cunhada"), "Cunhado", "Cunhada")
_gf_reg(("enteado", "enteada"), "Padrasto", "Madrasta")
_gf_reg(("padrasto", "madrasta"), "Enteado", "Enteada")
_gf_reg(("outro", "outra"), "Outro", "Outro")

def _gf_fem(nome):
    """Heurística de gênero pt-BR pelo primeiro nome (feminino ≈ termina em 'a')."""
    p = str(nome or "").strip().split()
    n = (p[0] if p else "").lower()
    if not n: return False
    fim = n[-1].translate(str.maketrans("áàâãä", "aaaaa"))
    return fim == "a"

def _gf_inverso(parentesco, nome_cliente=""):
    """Parentesco do CLIENTE ATUAL visto pelo outro lado do vínculo."""
    par = _GF_INV.get(str(parentesco or "").strip().lower())
    if not par: return ""
    return par[1] if _gf_fem(nome_cliente) else par[0]

def _gf_resolve(entradas, nome_map):
    """Converte as entradas do campo em objetos para a interface."""
    out = []
    for e in entradas:
        ref, par = _gf_split(e)
        if ref.startswith("#"):
            conta = ref[1:].strip()
            out.append({"tipo": "cliente", "conta": conta,
                        "nome": (nome_map.get(conta) or "Conta " + conta),
                        "parentesco": par})
        elif ref:
            out.append({"tipo": "texto", "nome": ref, "parentesco": par})
    return out

def _feriados_br(ano):
    """Feriados nacionais do ano (fixos + móveis via Páscoa, algoritmo de Meeus)."""
    a = ano % 19; b = ano // 100; c = ano % 100
    d = b // 4; e = b % 4; f = (b + 8) // 25; g = (b - f + 1) // 3
    h = (19 * a + b - d - g + 15) % 30
    i = c // 4; k = c % 4
    l = (32 + 2 * e + 2 * i - h - k) % 7
    m = (a + 11 * h + 22 * l) // 451
    mes = (h + l - 7 * m + 114) // 31
    dia = ((h + l - 7 * m + 114) % 31) + 1
    pascoa = datetime(ano, mes, dia).date()
    td = timedelta
    return {datetime(ano, 1, 1).date(),
            pascoa - td(days=48), pascoa - td(days=47),
            pascoa - td(days=2),
            datetime(ano, 4, 21).date(), datetime(ano, 5, 1).date(),
            pascoa + td(days=60),
            datetime(ano, 9, 7).date(), datetime(ano, 10, 12).date(),
            datetime(ano, 11, 2).date(), datetime(ano, 11, 15).date(),
            datetime(ano, 11, 20).date(), datetime(ano, 12, 25).date()}

def _dias_uteis(ini, fim, feriados):
    """Dias úteis entre ini e fim (inclusive)."""
    n = 0
    o = ini.toordinal()
    while o <= fim.toordinal():
        d = datetime.fromordinal(o).date()
        if d.weekday() < 5 and d not in feriados: n += 1
        o += 1
    return n

DIAS_SEMANA = ["segunda-feira", "terça-feira", "quarta-feira", "quinta-feira",
               "sexta-feira", "sábado", "domingo"]

def _parse_nasc(nasc):
    """Nascimento -> (dia, mês, ano|None). Aceita 'dd/mm/aaaa', 'dd/mm/aa',
    'dd/mm' e datas convertidas pelo Excel ('aaaa-mm-dd' / datetime).
    Nunca devolve ano no futuro nem idade impossível (nesses casos, ano=None)."""
    s = str(nasc or "").strip()
    if not s: return None
    ano_atual = datetime.now().year

    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", s)
    if m:
        ano, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if not (1 <= d <= 31 and 1 <= mo <= 12): return None
        if not (ano_atual - 120 <= ano <= ano_atual): ano = None
        return d, mo, ano

    m = re.match(r"^(\d{1,2})[/\-.](\d{1,2})(?:[/\-.](\d{4}|\d{2}))?$", s)
    if not m: return None
    d, mo = int(m.group(1)), int(m.group(2))
    if not (1 <= d <= 31 and 1 <= mo <= 12): return None
    ano = m.group(3)
    if ano is not None:
        ano = int(ano)
        if ano < 100:

            ano += 2000 if ano <= ano_atual % 100 else 1900
        if not (ano_atual - 120 <= ano <= ano_atual): ano = None
    return d, mo, ano

def _aniv_pack(nasc, hoje=None):
    """Idade e status de aniversário do cliente.
    aniv_status: 'hoje' = aniversário é hoje · 'fds' = hoje é sexta e o
    aniversário cai no fim de semana (aviso antecipado, para parabenizar no
    dia útil) · '' = nada a avisar."""
    hoje = hoje or datetime.now().date()
    out = {"idade": None, "idade_nova": None, "aniv_status": "",
           "aniv_data": "", "aniv_dsem": ""}
    p = _parse_nasc(nasc)
    if not p: return out
    d, mo, ano = p
    def _data(a):
        try: return hoje.replace(year=a, month=mo, day=d)
        except ValueError: return hoje.replace(year=a, month=mo, day=28)
    prox = _data(hoje.year)
    if prox < hoje: prox = _data(hoje.year + 1)
    if ano:
        out["idade"] = hoje.year - ano - (1 if (hoje.month, hoje.day) < (mo, d) else 0)
        out["idade_nova"] = prox.year - ano

        if not (0 <= out["idade"] <= 120):
            out["idade"] = out["idade_nova"] = None
    out["aniv_data"] = "%02d/%02d" % (d, mo)
    out["aniv_dsem"] = DIAS_SEMANA[prox.weekday()]
    if prox == hoje:
        out["aniv_status"] = "hoje"
    elif hoje.weekday() == 4 and prox.weekday() in (5, 6) and (prox - hoje).days <= 2:
        out["aniv_status"] = "fds"
    return out

_db_conn = None
_db_lock = threading.RLock()   # serializa criação/uso da conexão única compartilhada

def _db():
    global _db_conn
    with _db_lock:
        if _db_conn is None:
            _db_conn = sqlite3.connect(DB_PATH, check_same_thread=False)
            _db_conn.execute("PRAGMA journal_mode=WAL")
            _db_conn.execute("CREATE TABLE IF NOT EXISTS cadastro(conta TEXT PRIMARY KEY, dados TEXT)")
            _db_conn.execute("CREATE TABLE IF NOT EXISTS tabelas(nome TEXT PRIMARY KEY, dados TEXT)")
            _db_conn.execute("CREATE TABLE IF NOT EXISTS tombstones(conta TEXT PRIMARY KEY)")
            _db_conn.execute("CREATE TABLE IF NOT EXISTS kv(k TEXT PRIMARY KEY, v TEXT)")
            _db_conn.commit()
        return _db_conn

def _db_log(txt):
    try:
        with open(_LOGF, "a", encoding="utf-8") as f:
            f.write(txt + "\n")
    except Exception:
        pass

def _db_save_cadastro(lista, tombstone_ausentes=False):
    """Espelha o cadastro no banco (rede de segurança). Merge NÃO-destrutivo:
    nunca sobrescreve um campo bom do banco com valor vazio. A lápide (tombstone)
    só é criada em exclusão intencional (tombstone_ausentes=True) — assim uma
    perda acidental pode ser restaurada por _db_reconciliar na próxima abertura."""
    try:
        db = _db()
        atuais = {_conta_str(c.get("conta")) for c in (lista or []) if _conta_str(c.get("conta"))}
        antigos = {r[0] for r in db.execute("SELECT conta FROM cadastro")}
        existentes = {}
        for conta, dados in db.execute("SELECT conta,dados FROM cadastro").fetchall():
            try: existentes[conta] = json.loads(_dec_txt(dados))
            except Exception: existentes[conta] = {}
        for c in (lista or []):
            k = _conta_str(c.get("conta"))
            if not k: continue
            d = {kk: vv for kk, vv in c.items() if kk not in ("ativo", "patrimonio") and isinstance(kk, str)}
            merged = dict(existentes.get(k) or {})   # preserva campos bons já no banco
            for kk, vv in d.items():
                if str(vv or "").strip():            # só sobrescreve com valor não-vazio
                    merged[kk] = vv
            merged["conta"] = k
            db.execute("INSERT OR REPLACE INTO cadastro(conta,dados) VALUES(?,?)",
                       (k, _enc_txt(json.dumps(merged, ensure_ascii=False, default=str))))
            db.execute("DELETE FROM tombstones WHERE conta=?", (k,))
        if tombstone_ausentes:
            for k in antigos - atuais:
                db.execute("DELETE FROM cadastro WHERE conta=?", (k,))
                db.execute("INSERT OR REPLACE INTO tombstones(conta) VALUES(?)", (k,))
        db.commit()
    except Exception:
        _db_log("_db_save_cadastro:\n" + traceback.format_exc())

def _db_save_tabela(nome, rows):
    try:
        db = _db()
        db.execute("INSERT OR REPLACE INTO tabelas(nome,dados) VALUES(?,?)",
                   (nome, _enc_txt(json.dumps(rows or [], ensure_ascii=False, default=str))))
        db.commit()
    except Exception:
        _db_log("_db_save_tabela(%s):\n%s" % (nome, traceback.format_exc()))

def _db_load_tabela(nome):
    try:
        r = _db().execute("SELECT dados FROM tabelas WHERE nome=?", (nome,)).fetchone()
        return json.loads(_dec_txt(r[0])) if r and r[0] else []
    except Exception:
        return []

def _db_kv_set(k, v):
    try:
        db = _db(); db.execute("INSERT OR REPLACE INTO kv(k,v) VALUES(?,?)", (k, _enc_txt(str(v or "")))); db.commit()
    except Exception:
        _db_log("_db_kv_set:\n" + traceback.format_exc())

def _db_kv_get(k):
    try:
        r = _db().execute("SELECT v FROM kv WHERE k=?", (k,)).fetchone()
        return _dec_txt(r[0]) if r else ""
    except Exception:
        return ""

def _conta_str(v):
    """Conta como string CANÔNICA: sem espaços e sem '.0' de célula numérica.
    Toda junção cadastro ↔ relatórios usa esta forma — é o que faz as abas
    'conversarem' mesmo se o Excel converter a conta em número."""
    if v is None: return ""
    if isinstance(v, float) and v.is_integer(): v = int(v)
    return str(v).strip()

def ler_clientes_cadastro(wb):
    if EMAIL_ABA not in wb.sheetnames: return []
    rows = []
    for row in wb[EMAIL_ABA].iter_rows(min_row=CLI_START + 1, max_row=CLI_MAX, values_only=True):
        if row and row[1]:
            g = lambda i: str(row[i] or "") if i < len(row) else ""
            d = {"nome": g(0), "conta": _conta_str(row[1]), "email": g(2),
                 "termometro": g(3).strip().upper(), "perfil": g(4).strip(),
                 "estado": g(5).strip(), "cidade": g(6).strip(), "nascimento": g(7).strip()}
            for chave, idx in CAD_EXTRA:
                d[chave] = g(idx).strip()
            rows.append(d)
    return rows

def salvar_clientes_cadastro(wb, lista, permitir_reducao=False):
    """Grava o cadastro na aba E-mail. Linha de cliente inativo fica vermelha.
    Proteções contra perda de dados (Bug A):
    - um valor VAZIO nunca sobrescreve nome/e-mail bons já gravados;
    - fora de exclusão intencional (permitir_reducao=False) NENHUMA conta some:
      contas presentes na planilha e ausentes da lista são re-injetadas.
    permitir_reducao=True (Deletar Inativos / excluir cliente) permite reduzir."""
    if EMAIL_ABA not in wb.sheetnames: garantir_email(wb)
    ws = wb[EMAIL_ABA]

    # snapshot canônico do que já existe na planilha, por conta
    prev = {}
    for row in ws.iter_rows(min_row=CLI_START + 1, max_row=min(ws.max_row, CLI_MAX),
                            values_only=True):
        if row and row[1]:
            prev[_conta_str(row[1])] = row

    # (1) valor vazio nunca sobrescreve nome/e-mail bom de uma conta já existente
    for d in lista:
        old = prev.get(_conta_str(d.get("conta", "")))
        if not old: continue
        og = lambda i: str(old[i] or "").strip() if i < len(old) else ""
        if not str(d.get("nome", "") or "").strip() and og(0):
            d["nome"] = og(0)
        if not str(d.get("email", "") or "").strip() and og(2):
            d["email"] = og(2)

    # (2) fora de exclusão intencional, preservar contas ausentes da lista
    if not permitir_reducao:
        novas = {_conta_str(d.get("conta", "")) for d in lista if _conta_str(d.get("conta", ""))}
        perdidas = [k for k in prev if k not in novas]
        if perdidas:
            for k in perdidas:
                row = prev[k]
                g = lambda i: str(row[i] or "") if i < len(row) else ""
                d = {"nome": g(0), "conta": _conta_str(row[1]), "email": g(2),
                     "termometro": g(3).strip().upper(), "perfil": g(4).strip(),
                     "estado": g(5).strip(), "cidade": g(6).strip(),
                     "nascimento": g(7).strip(), "ativo": True}
                for chave, idx in CAD_EXTRA:
                    d[chave] = g(idx).strip()
                lista.append(d)
            try:
                with open(_LOGF, "a", encoding="utf-8") as f:
                    f.write(f"salvar_clientes_cadastro: {len(perdidas)} conta(s) preservada(s) "
                            "(ausentes da lista sem exclusão intencional)\n")
            except Exception:
                pass
    fim = min(max(ws.max_row, CLI_START + 1), CLI_MAX)
    for row in ws.iter_rows(min_row=CLI_START + 1, max_row=fim, max_col=CAD_NCOL):
        for c in row:
            c.value = None; c.fill = PatternFill()
    for ri, d in enumerate(lista, start=CLI_START + 1):
        inativo = not d.get("ativo", True)
        vals = [d.get("nome", ""), _conta_str(d.get("conta", "")), d.get("email", ""),
                str(d.get("termometro", "") or "").strip().upper(), d.get("perfil", ""),
                d.get("estado", ""), d.get("cidade", ""), d.get("nascimento", "")]
        vals += [d.get(chave, "") for chave, _ in CAD_EXTRA]
        for col, val in enumerate(vals, 1):
            c = ws.cell(ri, col, val)
            c.font = ft(size=10, color="C0392B" if inativo else "202124")
            c.fill = F("F8D7DA") if inativo else (F("ffffff") if ri % 2 == 0 else F("f8f9fa"))
            c.alignment = al("left" if col in CAD_LEFT else "center"); c.border = bd()
        ws.row_dimensions[ri].height = 20
    _db_save_cadastro(lista, tombstone_ausentes=permitir_reducao)

def _cdi_color(val):
    try:
        v = float(str(val).replace(",", ".").replace("%", "").strip())
        if v >= 100: return C_OK
        if v >= 80:  return C_WARN
        return C_ERR
    except Exception:
        return X_TXT

def _dark_border():
    s = Side(style="thin", color=X_BORDER)
    return Border(left=s, right=s, top=s, bottom=s)

def _br(v, nd=2):
    s = f"{v:,.{nd}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")

def escrever_mes(wb, aba, lista):
    if aba in wb.sheetnames: del wb[aba]
    ws = wb.create_sheet(aba); ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X_BAND
    ncols = len(COLS_FULL); last_col = get_column_letter(ncols)

    def _fv(s):
        try: return float(str(s))
        except Exception: return None
    rents  = [v for v in (_fv(d.get("rentabilidade")) for d in lista) if v is not None]
    ganhos = [v for v in (_fv(d.get("ganho")) for d in lista) if v is not None]
    cdis   = [v for v in (_fv(d.get("cdi")) for d in lista) if v is not None]
    kpi = f"{len(lista)} cliente(s)"
    if rents:  kpi += f"   ·   Rentabilidade média {_br(sum(rents)/len(rents))}%"
    if ganhos: kpi += f"   ·   Ganho total R$ {_br(sum(ganhos))}"
    if cdis:   kpi += f"   ·   CDI médio {_br(sum(cdis)/len(cdis), 1)}%"

    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]; t.value = f"  XPerformance — {mes_bonito(aba)}"
    t.font = Font(name="Segoe UI", bold=True, size=14, color=X_YELLOW)
    t.fill = F(X_BG); t.alignment = al("left"); ws.row_dimensions[1].height = 38

    ws.merge_cells(f"A2:{last_col}2")
    k = ws["A2"]; k.value = "  " + kpi
    k.font = Font(name="Segoe UI", size=10, color=X_MUT)
    k.fill = F(X_BG); k.alignment = al("left"); ws.row_dimensions[2].height = 22

    for col, h in enumerate(COLS_FULL, 1):
        c = ws.cell(3, col, h)
        c.font = Font(name="Segoe UI", bold=True, size=10, color=X_YELLOW)
        c.fill = F(X_BAND); c.alignment = al("center"); c.border = _dark_border()
    ws.row_dimensions[3].height = 26

    col_mov, col_rt, col_gt, col_pt, col_comp = 8, 9, 10, 11, 12
    alns = ["left", "center", "center", "center", "center", "left", "center",
            "center", "center", "center", "center", "left"]
    for ri, d in enumerate(lista, 4):
        bg = X_ROW1 if ri % 2 == 0 else X_ROW2
        r2, g2, c2 = d.get("rentabilidade", ""), d.get("ganho", ""), d.get("cdi", "")
        aloc = d.get("aloc") or {}
        vals = [d["nome"], d["conta"], _fv(r2) if _fv(r2) is not None else "",
                _fv(g2) if _fv(g2) is not None else "",
                _fv(c2) if _fv(c2) is not None else "", d["arquivo"],
                _fv(d.get("patrimonio")) if _fv(d.get("patrimonio")) is not None else "",
                _fv(d.get("movimentacoes")) if _fv(d.get("movimentacoes")) is not None else "",
                _fv(d.get("rent_total")) if _fv(d.get("rent_total")) is not None else "",
                _fv(d.get("ganho_total")) if _fv(d.get("ganho_total")) is not None else "",
                d.get("periodo_total", ""),
                json.dumps(aloc, ensure_ascii=False) if aloc else ""]
        for col, (val, aln_) in enumerate(zip(vals, alns), 1):
            c = ws.cell(ri, col, val)
            if   col == 5: cor = _cdi_color(c2)
            elif col in (3, col_rt): cor = C_ERR if (isinstance(val, float) and val < 0) else X_TXT
            elif col == col_mov: cor = C_ERR if (isinstance(val, float) and val < 0) else X_MUT
            elif col in (2, 6, col_gt, col_pt, col_comp): cor = X_MUT
            else: cor = X_TXT
            c.font = Font(name="Segoe UI", size=10 if col not in (6, col_comp) else 8,
                          bold=(col in (3, 5, col_rt)), color=cor)
            c.fill = F(bg); c.alignment = al(aln_); c.border = _dark_border()
            if col in (3, col_rt): c.number_format = '0.00"%"'
            if col in (4, 7, col_mov, col_gt): c.number_format = 'R$ #,##0.00'
            if col == 5: c.number_format = '0.0"%"'
        ws.row_dimensions[ri].height = 20

    last = 3 + len(lista)
    ws.auto_filter.ref = f"A3:{last_col}{last}"
    for col, w in enumerate([34, 12, 18, 16, 10, 50, 18, 16, 14, 18, 14, 44], 1):
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "A4"

def ler_template(wb):
    if EMAIL_ABA not in wb.sheetnames: return ASSUNTO_DEF, CORPO_DEF
    ws = wb[EMAIL_ABA]
    return (ws["B1"].value or ASSUNTO_DEF, ws["B3"].value or CORPO_DEF)

def ler_clientes(wb, aba):
    if aba not in wb.sheetnames: return []

    i_mov, i_rt, i_gt, i_pt, i_comp = 7, 8, 9, 10, 11
    def _f(v):
        try: return float(v)
        except Exception: return None
    rows = []
    for row in wb[aba].iter_rows(min_row=3, max_row=50000, values_only=True):
        if not row or not row[1]: continue
        nome = str(row[0] or "")
        if nome.startswith("⚠"): continue
        if not _conta_str(row[1]).isdigit(): continue
        r, g, c = row[2], row[3], row[4]
        pat = row[6] if len(row) > 6 else None
        mov = row[i_mov] if len(row) > i_mov else None
        rt = row[i_rt] if len(row) > i_rt else None
        gt = row[i_gt] if len(row) > i_gt else None
        pt = row[i_pt] if len(row) > i_pt else None
        aloc = {}
        comp = row[i_comp] if len(row) > i_comp else None
        if comp:
            try: aloc = {k: float(v) for k, v in json.loads(comp).items()}
            except Exception: aloc = {}
        def _rs(v):
            return f"{float(v):,.2f}".replace(",", "X").replace(".", ",").replace("X", ".") if v not in (None, "") else ""
        def _pc(v, nd=2):
            return f"{float(v):.{nd}f}".replace(".", ",") if v not in (None, "") else ""
        rows.append({
            "nome": nome, "conta": _conta_str(row[1]),

            "rentabilidade": _pc(r), "ganho": _rs(g), "cdi": _pc(c, 2),
            "patrimonio": _rs(pat), "movimentacoes": _rs(mov),
            "rentabilidade_total": _pc(rt), "ganho_total": _rs(gt),
            "periodo_total": str(pt or ""), "aloc": aloc,

            "rent_num": _f(r), "ganho_num": _f(g), "cdi_num": _f(c), "pat_num": _f(pat),
            "mov_num": _f(mov), "rent_total_num": _f(rt), "ganho_total_num": _f(gt),
            "arquivo": str(row[5] or "") if len(row) > 5 else "", "email": "", "mes_ref": aba})
    return rows

def listar_meses(wb):
    return sorted([s for s in wb.sheetnames if re.match(r"^\d{2}-\d{4}$", s)],
                  key=lambda s: (s.split("-")[1], s.split("-")[0]), reverse=True)

def _ordem_mes(s):
    p = str(s).split("-")
    return (p[1], p[0]) if len(p) == 2 else (s, "")

def coletar_por_cliente(wb):
    """Agrega todos os relatórios mensais por conta.
    Retorna {conta: {conta, nome, email, meses:{aba: dados}}}."""
    cad = {_conta_str(c["conta"]): c for c in ler_clientes_cadastro(wb)}
    out = {}
    for aba in listar_meses(wb):
        for c in ler_clientes(wb, aba):
            conta = _conta_str(c["conta"])
            o = out.setdefault(conta, {"conta": conta, "nome": "", "email": "", "meses": {}})
            o["meses"][aba] = c
            if c.get("nome") and not o["nome"]:
                o["nome"] = c["nome"]
    for conta, o in out.items():
        info = cad.get(conta, {})
        if info.get("nome"): o["nome"] = info["nome"]
        o["email"] = info.get("email", "")
        o["perfil"] = info.get("perfil", "")
        o["termometro"] = info.get("termometro", "")
        o["cad"] = info
    return out

def _sanear_cadastro():
    """Cura o cadastro: normaliza contas ('123 ' e 123.0 viram '123') e FUNDE
    linhas duplicadas da mesma conta, campo a campo, preservando o que está
    preenchido. Roda uma vez a cada abertura do app; só grava se mudou algo."""
    try:
        with _MUT_LOCK:
            if not os.path.exists(EXCEL_PATH):
                return
            wb = get_wb()
            cad = ler_clientes_cadastro(wb)
            por, ordem, mudou, fundidas = {}, [], False, 0
            for c in cad:
                k = _conta_str(c.get("conta"))
                if not k:
                    mudou = True; continue
                if c.get("conta") != k:
                    c["conta"] = k; mudou = True
                if k in por:
                    alvo = por[k]; mudou = True; fundidas += 1
                    for campo, val in c.items():
                        if campo == "conta": continue
                        if str(val or "").strip() and not str(alvo.get(campo) or "").strip():
                            alvo[campo] = val
                else:
                    por[k] = c; ordem.append(k)
            if mudou:
                for k in ordem: por[k].setdefault("ativo", True)
                salvar_clientes_cadastro(wb, [por[k] for k in ordem])
                salvar_wb(wb)
                try:
                    with open(_LOGF, "a", encoding="utf-8") as f:
                        f.write(f"_sanear_cadastro: cadastro normalizado ({fundidas} duplicata(s) fundida(s))\n")
                except Exception: pass
    except Exception:
        pass

def _db_reconciliar():
    """Roda a cada abertura. Regra de ouro:
    1) O que você digitou vive no BANCO — se sumir da planilha (releitura de PDFs,
       edição errada, corrupção), volta sozinho: linhas e campos são restaurados.
    2) A planilha continua mandando no dia a dia: valores editados NELA entram no
       banco. Cliente excluído PELO APP tem lápide e não ressuscita.
    3) Tabelas (aportes/contatos/receita/controle): restauradas apenas em perda
       total (aba sumiu/vazia); fora isso, a planilha é a verdade."""
    try:
        with _MUT_LOCK:
            db = _db(); wb = get_wb(); mudou = False

            cad = ler_clientes_cadastro(wb)
            por = {_conta_str(c.get("conta")): c for c in cad}
            tomb = {r[0] for r in db.execute("SELECT conta FROM tombstones")}
            restaurados = 0
            for conta, dados in db.execute("SELECT conta,dados FROM cadastro").fetchall():
                if conta in tomb: continue
                try: d = json.loads(_dec_txt(dados))
                except Exception: continue
                c = por.get(conta)
                if c is None:
                    d.setdefault("nome", "")
                    _ident = ("nome", "email", "perfil", "termometro", "estado", "cidade")
                    if not any(str(d.get(k) or "").strip() for k in _ident):
                        continue  # registro fantasma (só conta): não reinsere linha sem nome
                    d["conta"] = conta; d["ativo"] = True
                    cad.append(d); por[conta] = d
                    mudou = True; restaurados += 1
                else:
                    for k, v in d.items():
                        if k in ("conta", "ativo"): continue
                        if str(v or "").strip() and not str(c.get(k) or "").strip():
                            c[k] = v; mudou = True
            if mudou:
                agg = coletar_por_cliente(wb)
                for c in cad:
                    c["ativo"] = _conta_str(c.get("conta")) in agg
                salvar_clientes_cadastro(wb, cad)
                if restaurados:
                    _db_log(f"_db_reconciliar: {restaurados} cliente(s) restaurado(s) do banco")
            else:
                _db_save_cadastro(cad)

            def excel_de(nome):
                if nome == "aportes":
                    return [{k: a.get(k) for k in ("conta", "data", "valor", "obs")} for a in ler_aportes(wb)]
                if nome == "contatos":
                    return [{k: c_.get(k) for k in ("conta", "data", "tipo", "obs")} for c_ in ler_contatos(wb)]
                if nome == "receita":
                    return [{"mes": "%02d/%d" % (mm, aa), "valor": v} for (aa, mm), v in sorted(ler_receita(wb).items())]
                return ler_tabela_ctrl(wb, nome[5:])
            def grava_excel(nome, rows):
                if nome == "aportes":
                    ws = garantir_aportes(wb)
                    for r_ in rows: ws.append([r_.get("conta", ""), r_.get("data", ""), r_.get("valor", ""), r_.get("obs", "")])
                elif nome == "contatos":
                    ws = garantir_contatos(wb)
                    for r_ in rows: ws.append([r_.get("conta", ""), r_.get("data", ""), r_.get("tipo", ""), r_.get("obs", "")])
                elif nome == "receita":
                    ws = garantir_receita(wb)
                    for r_ in rows: ws.append([r_.get("mes", ""), r_.get("valor", "")])
                else:
                    salvar_tabela_ctrl(wb, nome[5:], rows)
            for nome in ("aportes", "contatos", "receita", "ctrl_mensal", "ctrl_montagem", "ctrl_ativos"):
                exc = excel_de(nome)
                salvo = _db_load_tabela(nome)
                if not exc and salvo:
                    grava_excel(nome, salvo); mudou = True
                    _db_log(f"_db_reconciliar: tabela '{nome}' restaurada do banco ({len(salvo)} linha(s))")
                else:
                    _db_save_tabela(nome, exc)

            if EMAIL_ABA in wb.sheetnames:
                ws = wb[EMAIL_ABA]
                for cel, chave in (("B1", "assunto"), ("B3", "corpo")):
                    atual = str(ws[cel].value or "").strip()
                    salvo = _db_kv_get(chave)
                    if not atual and salvo:
                        ws[cel].value = salvo; mudou = True
                    elif atual:
                        _db_kv_set(chave, ws[cel].value)
            if mudou:
                salvar_wb(wb)
    except Exception:
        _db_log("_db_reconciliar:\n" + traceback.format_exc())

def _cadastro_enriquecido(wb):
    """Cadastro + patrimônio (último relatório) + flag 'ativo' (tem relatório)."""
    cad = cadastro_cache()
    agg = coletar_cache()
    for c in cad:
        o = agg.get(_conta_str(c["conta"]))
        c["ativo"] = bool(o)
        c["patrimonio"] = None
        if o and o["meses"]:
            ult = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)[0]
            c["patrimonio"] = o["meses"][ult].get("pat_num")
    return cad

APORTES_ABA = "Aportes"

def _data_br(s):
    """'dd/mm/aaaa' (ou datetime/ISO do Excel) -> date | None."""
    m = re.match(r"^\s*(\d{1,2})/(\d{1,2})/(\d{2,4})\s*$", str(s or ""))
    if m:
        d, mo, a = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if a < 100: a += 2000
    else:
        m2 = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", str(s or ""))
        if not m2: return None
        a, mo, d = int(m2.group(1)), int(m2.group(2)), int(m2.group(3))
    try: return datetime(a, mo, d).date()
    except ValueError: return None

def garantir_aportes(wb):
    if APORTES_ABA in wb.sheetnames: return wb[APORTES_ABA]
    ws = wb.create_sheet(APORTES_ABA)
    ws.sheet_properties.tabColor = X_YELLOW
    ws.sheet_view.showGridLines = False
    for col, (h, w) in enumerate((("Conta", 14), ("Data", 14), ("Valor", 16), ("Observação", 42)), 1):
        c = ws.cell(1, col, h)
        c.font = ft(bold=True, color="202124"); c.fill = F("f4f4f5")
        c.alignment = al("center"); c.border = bd()
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws

def ler_aportes(wb):
    """Todos os aportes registrados: [{id, conta, data, valor, obs}]."""
    if APORTES_ABA not in wb.sheetnames: return []
    ws_a = wb[APORTES_ABA]
    if ws_a.max_row < 2: return []
    out = []
    for row in ws_a.iter_rows(min_row=2, max_row=ws_a.max_row):
        conta = _conta_str(row[0].value)
        if not conta: continue
        v = row[2].value
        try:
            valor = float(v)
        except Exception:
            s = str(v or "").replace("R$", "").strip()
            if "," in s: s = s.replace(".", "").replace(",", ".")
            try: valor = float(s)
            except Exception: continue
        raw = row[1].value
        dstr = raw.strftime("%d/%m/%Y") if hasattr(raw, "strftime") else str(raw or "").strip()
        obs = str(row[3].value or "") if len(row) > 3 else ""
        out.append({"id": row[0].row, "conta": conta, "data": dstr,
                    "valor": round(valor, 2), "obs": obs})
    return out

def _aportes_do_cliente(conta):
    conta = str(conta).strip()
    lst = [dict(a) for a in aportes_cache() if a["conta"] == conta]
    for a in lst: a["_d"] = _data_br(a["data"])
    lst.sort(key=lambda a: a["_d"] or datetime(1900, 1, 1).date(), reverse=True)
    return lst

def _aporte_lim():
    try: return int(carregar_config().get("aporte_alerta_dias") or 90)
    except Exception: return 90

def _aportes_stats(lst, alerta_dias):
    """Total, contagem, último aporte e dias sem aportar (alerta = acima do limite)."""
    hoje = datetime.now().date()
    total = round(sum(a["valor"] for a in lst), 2)
    n = len(lst)
    ultimo = ""; dias = None
    datas = [a.get("_d") for a in lst if a.get("_d")]
    if datas:
        du = max(datas); ultimo = du.strftime("%d/%m/%Y"); dias = (hoje - du).days
    alerta = (n == 0) or (dias is None) or (dias > alerta_dias)
    return {"total": total, "n": n, "ultimo": ultimo, "dias": dias, "alerta": alerta}

CONTATOS_ABA = "Contatos"
CONTATO_TIPOS = ["Ligação", "Mensagem", "E-mail", "Reunião presencial", "Evento", "Outro"]

def garantir_contatos(wb):
    if CONTATOS_ABA in wb.sheetnames: return wb[CONTATOS_ABA]
    ws = wb.create_sheet(CONTATOS_ABA)
    ws.sheet_properties.tabColor = X_YELLOW
    ws.sheet_view.showGridLines = False
    for col, (h, w) in enumerate((("Conta", 14), ("Data", 14), ("Tipo", 20), ("Observação", 42)), 1):
        c = ws.cell(1, col, h)
        c.font = ft(bold=True, color="202124"); c.fill = F("f4f4f5")
        c.alignment = al("center"); c.border = bd()
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws

def ler_contatos(wb):
    """Todos os contatos registrados: [{id, conta, data, tipo, obs}]."""
    if CONTATOS_ABA not in wb.sheetnames: return []
    ws_c = wb[CONTATOS_ABA]
    if ws_c.max_row < 2: return []
    out = []
    for row in ws_c.iter_rows(min_row=2, max_row=ws_c.max_row):
        conta = _conta_str(row[0].value)
        if not conta: continue
        raw = row[1].value
        dstr = raw.strftime("%d/%m/%Y") if hasattr(raw, "strftime") else str(raw or "").strip()
        out.append({"id": row[0].row, "conta": conta, "data": dstr,
                    "tipo": str(row[2].value or "").strip(),
                    "obs": str(row[3].value or "") if len(row) > 3 else ""})
    return out

def _contatos_do_cliente(conta):
    conta = str(conta).strip()
    lst = [dict(c) for c in contatos_cache() if c["conta"] == conta]
    for c in lst: c["_d"] = _data_br(c["data"])
    lst.sort(key=lambda c: c["_d"] or datetime(1900, 1, 1).date(), reverse=True)
    return lst

def _contato_stats(lst):
    hoje = datetime.now().date()
    ultimo = ""; tipo = ""; dias = None
    com_data = [c for c in lst if c.get("_d")]
    if com_data:
        c0 = max(com_data, key=lambda c: c["_d"])
        ultimo = c0["_d"].strftime("%d/%m/%Y"); tipo = c0.get("tipo", ""); dias = (hoje - c0["_d"]).days
    return {"n": len(lst), "ultimo": ultimo, "tipo": tipo, "dias": dias}

RECEITA_ABA = "Receita"

def garantir_receita(wb):
    if RECEITA_ABA in wb.sheetnames: return wb[RECEITA_ABA]
    ws = wb.create_sheet(RECEITA_ABA)
    ws.sheet_properties.tabColor = X_YELLOW
    ws.sheet_view.showGridLines = False
    for col, (h, w) in enumerate((("Mês", 12), ("Valor", 18)), 1):
        c = ws.cell(1, col, h)
        c.font = ft(bold=True, color="202124"); c.fill = F("f4f4f5")
        c.alignment = al("center"); c.border = bd()
        ws.column_dimensions[get_column_letter(col)].width = w
    return ws

def _mes_aa(s):
    """'mm/aaaa' -> (ano, mes) | None."""
    m = re.match(r"^\s*(\d{1,2})/(\d{2,4})\s*$", str(s or ""))
    if not m: return None
    mm = int(m.group(1)); aa = int(m.group(2))
    if aa < 100: aa += 2000
    return (aa, mm) if 1 <= mm <= 12 else None

def ler_receita(wb):
    """{(ano, mes): valor} da aba Receita."""
    if RECEITA_ABA not in wb.sheetnames: return {}
    ws = wb[RECEITA_ABA]
    if ws.max_row < 2: return {}
    out = {}
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        k = _mes_aa(row[0].value)
        if not k: continue
        v = row[1].value
        try:
            val = float(v)
        except Exception:
            s2 = str(v or "").replace("R$", "").strip()
            if "," in s2: s2 = s2.replace(".", "").replace(",", ".")
            try: val = float(s2)
            except Exception: continue
        out[k] = round(val, 2)
    return out

def _db_snapshot_aportes(wb):
    _db_save_tabela("aportes", [{k: a.get(k) for k in ("conta", "data", "valor", "obs")}
                                for a in ler_aportes(wb)])
def _db_snapshot_contatos(wb):
    _db_save_tabela("contatos", [{k: c.get(k) for k in ("conta", "data", "tipo", "obs")}
                                 for c in ler_contatos(wb)])
def _db_snapshot_receita(wb):
    _db_save_tabela("receita", [{"mes": "%02d/%d" % (mm, aa), "valor": v}
                                for (aa, mm), v in sorted(ler_receita(wb).items())])

CTRL_SHEETS = {
    "ofertas":  ("Ofertas",     ["conta", "ativo", "pct_carteira", "pct_alocar", "valor_alocar",
                                  "status", "data_liq", "origem"]),
    "mensal":   ("Aloc Mensal",  ["mes", "conta", "ativo", "valor", "pct", "roa", "retorno", "data_liq"]),
    "montagem": ("Montagem",     ["conta", "direcao", "ativo", "valor"]),
    "ativos":   ("Carteira Ativos", ["ativo", "indexador", "segmento", "financeiro", "yield",
                                     "conta", "carteira"]),
}

def ler_tabela_ctrl(wb, chave):
    nome, cols = CTRL_SHEETS[chave]
    if nome not in wb.sheetnames:
        return []
    ws = wb[nome]; out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row:
            continue
        vals = [("" if (i >= len(row) or row[i] is None) else str(row[i])) for i in range(len(cols))]
        if all(v.strip() == "" for v in vals):
            continue
        out.append({cols[i]: vals[i] for i in range(len(cols))})
    return out

def salvar_tabela_ctrl(wb, chave, rows):
    nome, cols = CTRL_SHEETS[chave]
    if nome in wb.sheetnames:
        del wb[nome]
    ws = wb.create_sheet(nome); ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X_YELLOW
    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c)
        cell.font = ft(bold=True, color="202124"); cell.fill = F("f4f4f5")
        cell.alignment = al("center"); cell.border = bd()
        ws.column_dimensions[get_column_letter(j)].width = 18
    for i, d in enumerate(rows or [], start=2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(i, j, (d.get(c, "") if isinstance(d, dict) else ""))
            cell.font = ft(size=10); cell.border = bd()
            cell.fill = F("ffffff") if i % 2 == 0 else F("f8f9fa"); cell.alignment = al("left")
    ws.row_dimensions[1].height = 20
    _db_save_tabela("ctrl_" + chave,
                    [{c: (d.get(c, "") if isinstance(d, dict) else "") for c in cols}
                     for d in (rows or [])])

PERFIL_KEY = {"conservadora": "cons", "conservador": "cons",
              "moderada": "mod", "moderado": "mod",
              "sofisticada": "sof", "sofisticado": "sof", "arrojado": "sof", "agressivo": "sof"}

def _perfil_nome(p):
    p = _norm_cls(p)
    if p.startswith("conserv"): return "Conservadora"
    if p.startswith("moder"): return "Moderada"
    if p.startswith("sofist") or p.startswith("arroj") or p.startswith("agress"): return "Sofisticada"
    return ""

def _idade(nasc):
    """Idade aproximada a partir de uma data de nascimento (procura o ano de 4 dígitos)."""
    m = re.search(r"(19\d{2}|20\d{2})", str(nasc or ""))
    if not m:
        return None
    ano = int(m.group(1))
    idade = datetime.now().year - ano
    return idade if 0 <= idade <= 120 else None

def _money(s):
    """'3.926.785,98' -> '3926785.98' (string pronta p/ float)."""
    return s.replace(".", "").replace(",", ".")

def _pct(s):
    """'2,19' -> '2.19'."""
    return s.replace(".", "").replace(",", ".") if s.count(",") and s.count(".") else s.replace(",", ".")

_LIGADURAS = str.maketrans({"ﬂ": "fl", "ﬁ": "fi", "ﬀ": "ff", "ﬃ": "ffi", "ﬄ": "ffl"})

def extrair(pdf_path):
    """Lê o relatório XPerformance (Evolução Patrimonial) com pypdf — rápido.
    Só as primeiras páginas (cartões/RESUMO e composição) e usa a tabela
    RESUMO DE INFORMAÇÕES DA CARTEIRA, que é estável e precisa."""
    fn = os.path.basename(pdf_path)
    mf = re.search(r"XPerformance\s*-\s*(\d+)\s*-\s*Ref\.\d+\.(\d+)\.pdf", fn, re.I)
    conta = mf.group(1) if mf else ""; mes = mf.group(2) if mf else ""
    aba = f"{mes.zfill(2)}-{datetime.now().year}" if mes else "Sem-data"
    d = dict(nome="", conta=conta, rentabilidade="", ganho="", cdi="", movimentacoes="",
             patrimonio="", rent_total="", ganho_total="", periodo_total="",
             aloc={}, mes_ref=aba, arquivo=fn, email="")
    try:
        reader = pypdf.PdfReader(pdf_path)
        n = len(reader.pages)
        _cache = {}
        def _pg(i):
            if i not in _cache:
                _cache[i] = reader.pages[i].extract_text() or ""
            return _cache[i]
        def _txt(idxs):
            return "\n".join(_pg(i) for i in idxs if 0 <= i < n).translate(_LIGADURAS)

        texto = _txt([1, 4])
        tem_resumo = ("PATRIM" in texto.upper() and "TOTAL BRUTO" in texto.upper()) or \
                     bool(re.search(r"(M[ÊE]S|IN[ÍI]CIO)\s+-?\s*R\$\s*[\d.,]+\s+-?[\d.,]+%", texto))
        tem_comp = bool(re.search(r"\(-?\d{1,3},\d{1,2}%\)\s*R\$", texto))
        if not (tem_resumo and tem_comp):
            texto = _txt(range(min(8, n)))
    except Exception as e:
        return d, str(e)

    mref = re.search(r"Data de refer[êe]ncia:?\s*(\d{1,2})/(\d{1,2})/(\d{4})", texto, re.I)
    if not mref:
        mref = re.search(r"Data de Refer[êe]ncia\D{0,40}?(\d{1,2})/(\d{1,2})/(\d{4})", texto, re.I | re.S)
    if mref:
        d["mes_ref"] = f"{mref.group(2).zfill(2)}-{mref.group(3)}"
    else:
        _meses_inv = {v.lower(): k for k, v in MESES.items()}
        mnome = re.search(r"\b(" + "|".join(MESES.values()) + r")\s*/\s*(\d{4})", texto, re.I)
        if mnome:
            d["mes_ref"] = f"{_meses_inv[mnome.group(1).lower()]}-{mnome.group(2)}"

    mp = re.search(r"PATRIM[ÔO]NIO TOTAL BRUTO:\s*R\$\s*([\d.,]+)", texto, re.I)
    if mp:
        d["patrimonio"] = _money(mp.group(1))

    resumo = {}
    for mm in re.finditer(r"(M[ÊE]S|ANO|12M|24M|IN[ÍI]CIO)\s+(-?)\s*R\$\s*([\d.,]+)\s+"
                          r"(-?[\d.,]+)%\s+(-?[\d.,]+)%\s+(-?)\s*R\$\s*([\d.,]+)", texto):
        k = mm.group(1).upper().replace("INICIO", "INÍCIO")
        resumo[k] = {"ganho": ("-" if mm.group(2) else "") + _money(mm.group(3)),
                     "rent": _pct(mm.group(4)), "cdi": _pct(mm.group(5)),
                     "mov": ("-" if mm.group(6) else "") + _money(mm.group(7))}
    mrow = resumo.get("MÊS") or resumo.get("MES")
    if mrow:
        d["ganho"] = mrow["ganho"]; d["rentabilidade"] = mrow["rent"]
        d["cdi"] = mrow["cdi"]; d["movimentacoes"] = mrow["mov"]
    if not d["rentabilidade"]:
        m = re.search(r"RENTABILIDADE M[ÊE]S:\s*(-?[\d.,]+)%", texto, re.I)
        if m: d["rentabilidade"] = _pct(m.group(1))
    if not d["ganho"]:
        m = re.search(r"GANHO M[ÊE]S:\s*R\$\s*(-?[\d.,]+)", texto, re.I)
        if m: d["ganho"] = _money(m.group(1))

    ml = re.search(r"GANHO M[ÊE]S:\s*R\$\s*-?[\d.,]+\s*RENTABILIDADE\s*([^\n:]*)", texto, re.S | re.I)
    pk = None
    if ml:
        up = ml.group(1).upper()
        if "INÍCIO" in up or "INICIO" in up:
            pk = "INÍCIO"; d["periodo_total"] = "Início"
        else:
            mm = re.search(r"(\d+)\s*M", up)
            if mm: pk = mm.group(1) + "M"; d["periodo_total"] = pk
    if not pk:

        m2 = re.search(r"RENTABILIDADE\s*\(?\s*(?:DESDE\s+O\s+|[ÚU]LTIMOS?\s+)?(IN[ÍI]CIO|\d+\s*M)\b",
                       texto, re.I)
        if m2:
            g = m2.group(1).upper().replace(" ", "")
            if g.startswith("IN"):
                pk = "INÍCIO"; d["periodo_total"] = "Início"
            else:
                pk = g; d["periodo_total"] = g
    if (not pk or pk not in resumo) and resumo:

        cands = ["INÍCIO"] + sorted((k for k in resumo if re.fullmatch(r"\d+M", k)),
                                    key=lambda x: -int(x[:-1])) + ["ANO"]
        for k2 in cands:
            if k2 in resumo:
                pk = k2
                d["periodo_total"] = "Início" if k2 == "INÍCIO" else ("Ano" if k2 == "ANO" else k2)
                break
    if pk and pk in resumo:
        d["rent_total"] = resumo[pk]["rent"]; d["ganho_total"] = resumo[pk]["ganho"]

    if not d["rent_total"]:
        m = re.search(r"RENTABILIDADE\s*\(?\s*(?:DESDE\s+O\s+|[ÚU]LTIMOS?\s+)?"
                      r"(IN[ÍI]CIO|\d+\s*M|ANO)\s*\)?\s*:?\s*(-?[\d.,]+)%", texto, re.I)
        if m:
            g = m.group(1).upper().replace(" ", "")
            if not d["periodo_total"]:
                d["periodo_total"] = "Início" if g.startswith("IN") else ("Ano" if g == "ANO" else g)
            d["rent_total"] = _pct(m.group(2))
    if not d["ganho_total"]:
        m = re.search(r"GANHO\s*\(?\s*(?:DESDE\s+O\s+|[ÚU]LTIMOS?\s+)?"
                      r"(?:IN[ÍI]CIO|\d+\s*M|ANO)\s*\)?\s*:?\s*R\$\s*(-?[\d.,]+)", texto, re.I)
        if m: d["ganho_total"] = _money(m.group(1))

    aloc = {}
    for mm in re.finditer(r"(?m)^(.+?)\s*\((-?\d{1,3},\d{1,2})%\)\s*R\$", texto):
        nome_cls = _clean_cls(mm.group(1))
        if nome_cls:
            try: aloc[nome_cls] = float(_money(mm.group(2)))
            except Exception: pass
    if aloc:
        d["aloc"] = aloc
    return d, None

ALOC_SHEET = "Alocação"
MESES_NUM = {"JANEIRO":"01","FEVEREIRO":"02","MARÇO":"03","MARCO":"03","ABRIL":"04","MAIO":"05",
             "JUNHO":"06","JULHO":"07","AGOSTO":"08","SETEMBRO":"09","OUTUBRO":"10",
             "NOVEMBRO":"11","DEZEMBRO":"12"}

ALOC_CLASSES = [
    ("Pós-fixado","Renda Fixa Brasil"), ("Inflação","Renda Fixa Brasil"), ("Prefixado","Renda Fixa Brasil"),
    ("Multimercados","Multimercados"), ("Renda Variável Brasil","Renda Variável Brasil"),
    ("Fundos Listados","Fundos Listados"), ("Alternativos","Alternativos"),
    ("Global - Renda Fixa","Global"), ("Global - Renda Variável","Global"),
]

ALOC_PERSPS = [
    ("Pós-fixado","Pós-fixado"), ("Inflação","Inflação"), ("Prefixado","Prefixado"),
    ("MULTIMERCADOS","Multimercados"), ("RENDA VARIÁVEL BRASIL","Renda Variável Brasil"),
    ("FUNDOS LISTADOS","Fundos Listados"), ("ALTERNATIVOS","Alternativos"),
    ("RENDA VARIÁVEL GLOBAL","Renda Variável Global"),
]
ALOC_CHART_CUT = re.compile(r"(Fonte:|Retorno anualizado|Composição setorial|Seleção de sub|"
                            r"Mercados com maior|desde 2017|Data-base|Elaboração)", re.I)

CLIENT_TO_ALOC = {
    "pós fixado":"Pós-fixado", "pos fixado":"Pós-fixado",
    "inflação":"Inflação", "inflacao":"Inflação", "pré fixado":"Prefixado", "pre fixado":"Prefixado",
    "multimercado":"Multimercados", "multimercados":"Multimercados",
    "renda variável brasil":"Renda Variável Brasil", "renda variavel brasil":"Renda Variável Brasil",
    "fundos listados":"Fundos Listados", "alternativo":"Alternativos", "alternativos":"Alternativos",
    "renda variável global":"Global - Renda Variável", "renda variavel global":"Global - Renda Variável",
}

def _resumir(txt, max_chars=300):
    """Resumo extrativo: primeiras frases completas até ~max_chars (nunca corta no meio)."""
    txt = re.sub(r"(clicando aqui|Veja mais an[áa]lises de Aloca[çc][ãa]o|"
                 r"Perspectivas por Classe de Ativo)", "", txt, flags=re.I)
    txt = re.sub(r"\s{2,}", " ", txt)
    txt = re.sub(r"\s+\d+\s*$", "", txt).strip()
    sents = re.split(r"(?<=[.!?])\s+", txt)
    out = ""
    for s in sents:
        s = s.strip()
        if not s:
            continue
        letras = sum(c.isalpha() for c in s); digitos = sum(c.isdigit() for c in s)
        if letras < 10 or digitos > letras:
            continue
        if out and len(out) + len(s) + 1 > max_chars:
            break
        out += (" " if out else "") + s
    if out and out[-1] not in ".!?":
        out += "."
    return out.strip()

def _brn(v, nd=1):
    s = f"{float(v):.{nd}f}".replace(".", ",")
    return s.rstrip("0").rstrip(",") if "," in s else s

def gerar_insights(d):
    """Insights gerenciais derivados dos NÚMEROS do relatório (não altera recomendações)."""
    cl = {c["classe"]: c for c in d.get("classes", [])}
    def soma(nomes, k): return round(sum(cl.get(n, {}).get(k, 0.0) for n in nomes), 1)
    rf = ["Pós-fixado", "Inflação", "Prefixado", "Global - Renda Fixa"]
    rv = ["Renda Variável Brasil", "Global - Renda Variável"]
    dv = ["Multimercados", "Fundos Listados", "Alternativos"]
    ins = []
    pf = cl.get("Pós-fixado")
    if pf:
        ins.append(f"Postura defensiva: o Pós-fixado (CDI) é o maior peso recomendado — "
                   f"{_brn(pf['cons'])}% / {_brn(pf['mod'])}% / {_brn(pf['sof'])}% (Conservadora/Moderada/Sofisticada).")
    if cl:
        ins.append(f"Renda fixa concentra a carteira: {_brn(soma(rf,'cons'))}% / {_brn(soma(rf,'mod'))}% / "
                   f"{_brn(soma(rf,'sof'))}% do total.")
        ins.append(f"A bolsa (Brasil + global) acompanha o apetite a risco: {_brn(soma(rv,'cons'))}% → "
                   f"{_brn(soma(rv,'mod'))}% → {_brn(soma(rv,'sof'))}%.")
        ins.append(f"Diversificadores (multimercados, fundos listados e alternativos) somam "
                   f"{_brn(soma(dv,'cons'))}% / {_brn(soma(dv,'mod'))}% / {_brn(soma(dv,'sof'))}%.")
    pos = []
    for p in d.get("persp", []):
        t = p["texto"].lower()
        if "sobrealocad" in t: pos.append(f"{p['classe']} (sobrealocado)")
        elif "subalocad" in t: pos.append(f"{p['classe']} (subalocado)")
        elif "neutra" in t or "neutro" in t: pos.append(f"{p['classe']} (neutro)")
    if pos:
        ins.append("Posicionamento citado pela XP: " + ", ".join(pos) + ".")
    dur = d.get("duration", {})
    if dur:
        ins.append("Duration recomendada: " + ", ".join(f"{k} {_brn(v)} anos" for k, v in dur.items()) + ".")
    return ins

def _page_text_gap(page):
    """Reconstrói o texto da página inserindo espaços por gap entre caracteres
    (corrige o 'colar' de palavras em texto justificado)."""
    linhas = {}
    for c in page.chars:
        linhas.setdefault(round(c["top"] / 2) * 2, []).append(c)
    out = []
    for top in sorted(linhas):
        cs = sorted(linhas[top], key=lambda c: c["x0"]); s = ""; prev = None
        widths = [c["x1"] - c["x0"] for c in cs if c["text"].strip()]
        avg = sum(widths) / len(widths) if widths else 3
        for c in cs:
            if prev is not None and (c["x0"] - prev["x1"]) > avg * 0.28 and not s.endswith(" "):
                s += " "
            s += c["text"]; prev = c
        out.append(s)
    return "\n".join(out)

def extrair_alocacao(pdf_path):
    fn = os.path.basename(pdf_path)
    d = {"arquivo": fn, "mes_ref": "", "mes_label": "", "nota": "", "classes": [],
         "retorno": {}, "vol": {}, "duration": {}, "persp": []}
    try:
        with pdfplumber.open(pdf_path) as pdf:
            clean = "\n".join(_page_text_gap(p) for p in pdf.pages)
    except Exception as e:
        return d, str(e)
    m = re.search(r"(JANEIRO|FEVEREIRO|MAR[ÇC]O|ABRIL|MAIO|JUNHO|JULHO|AGOSTO|SETEMBRO|"
                  r"OUTUBRO|NOVEMBRO|DEZEMBRO)\s+DE\s+(\d{4})", clean, re.I)
    if m:
        d["mes_ref"] = f"{MESES_NUM.get(m.group(1).upper(), '00')}-{m.group(2)}"
        d["mes_label"] = f"{m.group(1).title()}/{m.group(2)}"
    sec = re.search(r"Carteiras Recomendadas(.*?)(?:Ap[êe]ndice|Bandas de Aloca|\Z)", clean, re.S | re.I)
    seg = sec.group(1) if sec else clean

    nota = re.search(r"Carteiras Recomendadas\s*(.*?)(?:C\s*O\s*N\s*S\s*E\s*R\s*V\s*A\s*D|"
                     r"P[óo]s-?\s*fixado)", clean, re.S | re.I)
    if nota:
        d["nota"] = _resumir(re.sub(r"\s+", " ", nota.group(1)).strip(), 700)
    trios = re.findall(r"(\d{1,3},\d{1,2})%\s+(\d{1,3},\d{1,2})%\s+(\d{1,3},\d{1,2})%", seg)
    for i, (cls, grp) in enumerate(ALOC_CLASSES):
        if i < len(trios):
            c0, m0, s0 = trios[i]
            d["classes"].append({"classe": cls, "grupo": grp,
                "cons": float(c0.replace(",", ".")), "mod": float(m0.replace(",", ".")),
                "sof": float(s0.replace(",", "."))})
    ret = re.search(r"Retorno Esperado\s*CDI\s*\+\s*(\d+)%\s*CDI\s*\+\s*(\d+)%\s*CDI\s*\+\s*(\d+)%", seg)
    if ret:
        d["retorno"] = {"cons": f"CDI + {ret.group(1)}%", "mod": f"CDI + {ret.group(2)}%",
                        "sof": f"CDI + {ret.group(3)}%"}
    if len(trios) >= 10:
        v = trios[-1]
        d["vol"] = {"cons": float(v[0].replace(",", ".")), "mod": float(v[1].replace(",", ".")),
                    "sof": float(v[2].replace(",", "."))}
    pos = re.search(r"Posicionamento por Classe de Ativo(.*?)Carteiras Recomendadas", clean, re.S | re.I)
    if pos:
        for cls, key in [("Inflação", "Inflação"), ("Prefixado", "Prefixado"),
                         ("Renda Fixa", "Global - Renda Fixa")]:
            mm = re.search(re.escape(cls) + r"\s+(\d,\d)", pos.group(1))
            if mm:
                d["duration"][key] = float(mm.group(1).replace(",", "."))
    pers_sec = re.search(r"Perspectivas por Classe de Ativo(.*?)Posicionamento por Classe", clean, re.S | re.I)
    pt = pers_sec.group(1) if pers_sec else clean
    pos_hdr = []
    for hdr, label in ALOC_PERSPS:
        mm = re.search(r"(?m)^\s*" + re.escape(hdr) + r"\s*$", pt) or re.search(re.escape(hdr), pt)
        if mm:
            pos_hdr.append((mm.start(), hdr, label))
    pos_hdr.sort()
    for idx, (st, hdr, label) in enumerate(pos_hdr):
        end = pos_hdr[idx + 1][0] if idx + 1 < len(pos_hdr) else len(pt)
        bloco = pt[st + len(hdr):end]
        cut = ALOC_CHART_CUT.search(bloco)
        if cut:
            bloco = bloco[:cut.start()]
        bloco = re.sub(r"\s+", " ", bloco).strip()
        bloco = re.sub(r"(JUNHO DE \d{4}|JANEIRO DE \d{4}|Veja mais.*?aqui|"
                       r"Coment[áa]rios por classe.*?ativo|RENDA FIXA BRASIL|^GLOBAL)", "",
                       bloco, flags=re.I).strip()
        if len(bloco) > 40:
            d["persp"].append({"classe": label, "texto": _resumir(bloco, 300)})
    d["insights"] = gerar_insights(d)
    return d, None

def escrever_alocacao(wb, d):
    """Grava os dados do relatório de alocação (substitui os anteriores).
    Mantém uma tabela legível + o JSON completo para o app."""
    idx = wb.sheetnames.index(ALOC_SHEET) if ALOC_SHEET in wb.sheetnames else None
    if idx is not None:
        del wb[ALOC_SHEET]
    ws = wb.create_sheet(ALOC_SHEET) if idx is None else wb.create_sheet(ALOC_SHEET, idx)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X_YELLOW
    ws["A1"] = f"Carteiras Recomendadas — {d.get('mes_label', '')}"
    ws["A1"].font = Font(name="Segoe UI", bold=True, size=13, color=X_YELLOW)
    hdr = ["Classe", "Grupo", "Conservadora", "Moderada", "Sofisticada"]
    for col, h in enumerate(hdr, 1):
        c = ws.cell(3, col, h)
        c.font = ft(bold=True, color="202124"); c.fill = F("f4f4f5"); c.border = bd(); c.alignment = al("center")
    r = 4
    for cl in d.get("classes", []):
        ws.cell(r, 1, cl["classe"]); ws.cell(r, 2, cl["grupo"])
        for col, k in ((3, "cons"), (4, "mod"), (5, "sof")):
            cc = ws.cell(r, col, cl[k] / 100.0); cc.number_format = "0.0%"; cc.alignment = al("center")
        r += 1
    rr = d.get("retorno", {}); vv = d.get("vol", {})
    ws.cell(r, 1, "Retorno esperado")
    for col, k in ((3, "cons"), (4, "mod"), (5, "sof")):
        ws.cell(r, col, rr.get(k, "")); ws.cell(r, col).alignment = al("center")
    r += 1
    ws.cell(r, 1, "Volatilidade alvo")
    for col, k in ((3, "cons"), (4, "mod"), (5, "sof")):
        if vv.get(k) is not None:
            cc = ws.cell(r, col, vv[k] / 100.0); cc.number_format = "0.00%"; cc.alignment = al("center")
    r += 2
    ws.cell(r, 1, "Nota tática:"); ws.cell(r, 2, d.get("nota", ""))
    for col, w in enumerate([26, 22, 14, 14, 14], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws["A100"] = "DADOS_APP_JSON"
    ws["B100"] = json.dumps(d, ensure_ascii=False)

def ler_alocacao(wb):
    if ALOC_SHEET not in wb.sheetnames:
        return {}
    raw = wb[ALOC_SHEET]["B100"].value
    if not raw:
        return {}
    try:
        return json.loads(raw)
    except Exception:
        return {}

def comparar_carteira(aloc_cliente, dados_aloc, perfil_forcado=None):
    """Mapeia a carteira atual do cliente para as classes do relatório e compara
    com as 3 carteiras recomendadas. Não altera nenhuma recomendação.
    perfil_forcado ('cons'/'mod'/'sof'): se definido, é o perfil-alvo; senão usa o mais próximo."""
    classes = [c["classe"] for c in dados_aloc.get("classes", [])]
    atual = {c: 0.0 for c in classes}
    extras = {}
    for nome, pct in (aloc_cliente or {}).items():
        alvo = CLIENT_TO_ALOC.get(_norm_cls(nome))
        if alvo and alvo in atual:
            atual[alvo] += float(pct)
        else:
            extras[nome] = extras.get(nome, 0.0) + float(pct)
    linhas, dist = [], {"cons": 0.0, "mod": 0.0, "sof": 0.0}
    for cl in dados_aloc.get("classes", []):
        nm = cl["classe"]; at = round(atual.get(nm, 0.0), 2)
        linhas.append({"classe": nm, "grupo": cl["grupo"], "atual": at,
                       "cons": cl["cons"], "mod": cl["mod"], "sof": cl["sof"], "extra": False})
        for k in dist:
            dist[k] += abs(at - cl[k])
    for nome, val in extras.items():
        v = round(val, 2)
        linhas.append({"classe": nome, "grupo": "", "atual": v,
                       "cons": None, "mod": None, "sof": None, "extra": True})
        for k in dist:
            dist[k] += v
    dist = {k: round(v, 2) for k, v in dist.items()}
    auto = min(dist, key=dist.get) if dados_aloc.get("classes") else None

    closest = perfil_forcado if perfil_forcado in ("cons", "mod", "sof") else auto
    _rb = {c["classe"]: c for c in dados_aloc.get("classes", [])}
    aderencia = (round(sum(min(l["atual"], _rb.get(l["classe"], {}).get(closest, 0.0))
                           for l in linhas if not l.get("extra")), 1) if closest else None)
    return {"linhas": linhas, "dist": dist, "closest": closest, "perfil_auto": auto,
            "aderencia": aderencia,
            "perfil_definido": bool(perfil_forcado in ("cons", "mod", "sof")),
            "nao_mapeadas": extras, "retorno": dados_aloc.get("retorno", {}),
            "vol": dados_aloc.get("vol", {}), "persp": dados_aloc.get("persp", []),
            "nota": dados_aloc.get("nota", ""), "mes_label": dados_aloc.get("mes_label", "")}

def gerar_analise(comp, dados_aloc, oi):
    """Análise profunda da carteira: aderência às recomendações + recomendações
    inteligentes, com a fundamentação vinda dos relatórios de mercado."""
    nomes = {"cons": "Conservadora", "mod": "Moderada", "sof": "Sofisticada"}
    linhas = comp.get("linhas", []); closest = comp.get("closest")
    rec_by = {c["classe"]: c for c in dados_aloc.get("classes", [])}

    aderencias = {}
    for k in ("cons", "mod", "sof"):
        ov = sum(min(l["atual"], rec_by.get(l["classe"], {}).get(k, 0.0))
                 for l in linhas if not l.get("extra"))
        aderencias[k] = round(ov, 1)
    aderencia = aderencias.get(closest, 0.0)
    perfil = nomes.get(closest, "")

    persp = {_norm_cls(p["classe"]): p["texto"] for p in dados_aloc.get("persp", [])}
    alias = {"global - renda variável": "renda variável global", "global - renda fixa": "pós-fixado"}
    def motivo(classe):
        k = _norm_cls(classe); return persp.get(alias.get(k, k), "")

    recs = []
    for l in linhas:
        if l.get("extra"):
            recs.append({"classe": l["classe"], "acao": "Realocar", "delta": l["atual"],
                         "atual": l["atual"], "alvo": 0.0,
                         "motivo": f"Classe fora da carteira recomendada — considerar realocar para classes do perfil {perfil}."})
            continue
        alvo = rec_by.get(l["classe"], {}).get(closest, 0.0)
        diff = round(l["atual"] - alvo, 2)
        if abs(diff) >= 2.0:
            recs.append({"classe": l["classe"], "acao": "Reduzir" if diff > 0 else "Aumentar",
                         "delta": round(abs(diff), 2), "atual": l["atual"], "alvo": alvo,
                         "motivo": motivo(l["classe"])})
    recs.sort(key=lambda x: x["delta"], reverse=True)
    grau = "alta" if aderencia >= 85 else ("moderada" if aderencia >= 70 else "baixa")
    if recs:
        resumo = (f"A carteira está {_brn(aderencia)}% aderente ao perfil {perfil} (aderência {grau}). "
                  f"Os maiores desajustes frente à recomendação estão em "
                  + ", ".join(r["classe"] for r in recs[:3]) + ".")
    else:
        resumo = (f"A carteira está {_brn(aderencia)}% aderente ao perfil {perfil} "
                  f"(aderência {grau}) — bem alinhada, sem ajustes relevantes.")

    mes_a = dados_aloc.get("mes_label", "")
    mes_o = (oi or {}).get("mes_label", "")
    ctx = ("As recomendações desta análise vêm do relatório de Carteiras Recomendadas "
           "(Alocação XP" + (" — " + mes_a if mes_a else "") + "): os percentuais-alvo por perfil "
           "estão na seção “Carteiras Recomendadas” e a justificativa de cada classe na seção "
           "“Perspectivas por Classe de Ativo”.")
    if oi and oi.get("secoes"):
        ctx += (" O cenário e o posicionamento tático vêm do relatório “Onde Investir”"
                + (" — " + mes_o if mes_o else "") + " (seção “Como se posicionar nesse cenário”).")
    if dados_aloc.get("nota"):
        ctx += " Destaque tático do mês: " + dados_aloc["nota"]
    return {"perfil": perfil, "closest": closest, "aderencia": aderencia,
            "aderencias": aderencias, "resumo": resumo, "contexto": ctx,
            "recomendacoes": recs[:7], "classes_rec": dados_aloc.get("classes", []),
            "retorno": dados_aloc.get("retorno", {}), "vol": dados_aloc.get("vol", {}),
            "mes_label": dados_aloc.get("mes_label", "")}

OI_SHEET = "OndeInvestir"
OI_SECOES = ["Cenário global", "Brasil: riscos se intensificam", "Como se posicionar nesse cenário",
             "Renda fixa doméstica", "Renda variável doméstica", "Fundos Listados",
             "Renda variável global"]
OI_RUIDO = re.compile(r"(ONDE INVESTIR.*?\d{4}|Em colabora[çc][ãa]o com[^\n.]*|Fonte:[^\n]*|"
                      r"Performance mensal[^\n]*|Composi[çc][ãa]o setorial[^\n]*|"
                      r"Carteiras recomendadas: principais destaques do m[êe]s)", re.I)

def extrair_onde_investir(pdf_path):
    """Lê o relatório 'Onde Investir' e resume os pontos-chave (resumo extrativo)."""
    fn = os.path.basename(pdf_path)
    d = {"arquivo": fn, "mes_label": "", "resumo": [], "secoes": []}
    try:
        with pdfplumber.open(pdf_path) as pdf:
            clean = "\n".join(_page_text_gap(p) for p in pdf.pages)
    except Exception as e:
        return d, str(e)
    m = re.search(r"Onde Investir:?\s*([A-Za-zçãéêíóô]+\s+de\s+\d{4})", clean, re.I)
    if m:
        d["mes_label"] = m.group(1)

    rs = re.search(r"\bResumo\b(.*?)Cen[áa]rio global", clean, re.S)
    if rs:
        for b in rs.group(1).split("•"):
            b = OI_RUIDO.sub("", re.sub(r"\s+", " ", b)).strip()
            if len(b) > 40:
                d["resumo"].append(b)

    pos = []
    for sec in OI_SECOES:
        mm = re.search(r"(?m)^\s*" + re.escape(sec), clean)
        if mm:
            pos.append((mm.start(), sec))
    pos.sort()
    for i, (st, sec) in enumerate(pos):
        end = pos[i + 1][0] if i + 1 < len(pos) else len(clean)
        bloco = clean[st:end]
        titulo = bloco.splitlines()[0].strip()
        corpo = OI_RUIDO.sub("", "\n".join(bloco.splitlines()[1:]))
        corpo = re.sub(r"\s+", " ", corpo).strip()
        d["secoes"].append({"titulo": titulo, "texto": _resumir(corpo, 520)})
    if not d["resumo"] and not d["secoes"]:
        return d, "Não reconheci o conteúdo de 'Onde Investir' neste PDF."
    return d, None

def escrever_onde_investir(wb, d):
    """Guarda o resumo 'Onde Investir' no BANCO DE DADOS e remove a aba
    antiga da planilha, se existir — o Excel fica só com os dados de trabalho."""
    _db_kv_set("onde_investir", json.dumps(d, ensure_ascii=False))
    if wb is not None and OI_SHEET in wb.sheetnames:
        del wb[OI_SHEET]

def ler_onde_investir(wb=None):
    try:
        raw = _db_kv_get("onde_investir")
        if raw:
            return json.loads(raw) or {}
    except Exception:
        pass
    try:
        if wb is not None and OI_SHEET in wb.sheetnames:
            raw = wb[OI_SHEET]["B100"].value
            if raw:
                return json.loads(raw)
    except Exception:
        pass
    return {}

def migrar_onde_investir():
    """Migração única: leva o resumo para o banco e apaga o JSON antigo e a
    aba OndeInvestir de planilhas antigas."""
    try:
        if os.path.exists(ONDE_PATH):
            if not _db_kv_get("onde_investir"):
                try:
                    with open(ONDE_PATH, encoding="utf-8") as f:
                        _db_kv_set("onde_investir", f.read())
                except Exception:
                    pass
            try: os.remove(ONDE_PATH)
            except Exception: pass
        if not os.path.exists(EXCEL_PATH):
            return
        wb = get_wb()
        if OI_SHEET not in wb.sheetnames:
            return
        if not _db_kv_get("onde_investir"):
            try:
                raw = wb[OI_SHEET]["B100"].value
                if raw:
                    _db_kv_set("onde_investir", raw)
            except Exception:
                pass
        del wb[OI_SHEET]
        salvar_wb(wb)
    except Exception:
        pass

def detectar_tipo_relatorio(pdf_path):
    """Identifica o tipo de relatório pelo conteúdo: 'onde', 'aloc', 'cliente' ou None."""
    try:
        with pdfplumber.open(pdf_path) as pdf:
            txt = " ".join((pdf.pages[i].extract_text() or "")
                           for i in range(min(3, len(pdf.pages))))
    except Exception:
        return None
    up = re.sub(r"\s+", " ", txt).upper()
    if "ONDE INVESTIR" in up:
        return "onde"
    if "CARTEIRAS RECOMENDADAS" in up or "PERSPECTIVAS POR CLASSE DE ATIVO" in up \
       or ("CONSERVADORA" in up and "SOFISTICADA" in up):
        return "aloc"
    if "XPERFORMANCE" in up or "EVOLUÇÃO PATRIMONIAL" in up \
       or ("RELAT" in up and "INVESTIMENTOS" in up and "REFER" in up):
        return "cliente"
    return None

def _http_json(url, timeout=15):
    import urllib.request
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))

def _sgs(serie, n=1):
    return _http_json(f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{serie}/dados/ultimos/{n}?formato=json")

def _sgs_per(serie, ini, fim):
    return _http_json(f"https://api.bcb.gov.br/dados/serie/bcdata.sgs.{serie}/dados"
                      f"?formato=json&dataInicial={ini.strftime('%d/%m/%Y')}&dataFinal={fim.strftime('%d/%m/%Y')}")

def _focus(recurso, filtro):
    """API de Expectativas (Boletim Focus) do Banco Central."""
    import urllib.parse
    q = urllib.parse.quote(filtro, safe="() ',").replace(" ", "%20").replace("'", "%27")
    url = ("https://olinda.bcb.gov.br/olinda/servico/Expectativas/versao/v1/odata/"
           f"{recurso}?$top=1&$filter={q}&$orderby=Data%20desc&$format=json")
    return _http_json(url).get("value") or []

def fetch_mercado():
    """Coleta indicadores do dia, do mês (atual e anterior) e previsões (Focus/BCB)."""
    hoje = datetime.now(); erros = []
    def pega(serie, label):
        try: return float(_sgs(serie)[-1]["valor"])
        except Exception as e: erros.append(f"{label}: {e}"); return None
    selic  = pega(432, "Selic")
    cdi_aa = pega(4389, "CDI a.a.")
    dolar  = pega(1, "Dólar")
    poup   = pega(196, "Poupança")

    prim_atual = hoje.replace(day=1)
    fim_ant = prim_atual - timedelta(days=1)
    prim_ant = fim_ant.replace(day=1)

    def coleta_mes(ini, fim, label):
        r = {"mes": label, "cdi_mes": None, "ipca": None, "poupanca": None,
             "selic_fim": None, "dolar_fim": None, "ibov_var": None}
        try:
            ds = _sgs_per(12, ini, fim)
            if ds:
                acc = 1.0
                for v in ds: acc *= 1 + float(v["valor"]) / 100
                r["cdi_mes"] = round((acc - 1) * 100, 4)
        except Exception as e: erros.append(f"CDI {label}: {e}")
        try:
            ds = _sgs_per(1, ini, fim)
            if ds: r["dolar_fim"] = float(ds[-1]["valor"])
        except Exception as e: erros.append(f"Dólar {label}: {e}")
        try:
            ds = _sgs_per(432, ini, fim)
            if ds: r["selic_fim"] = float(ds[-1]["valor"])
        except Exception as e: erros.append(f"Selic {label}: {e}")
        return r
    mes_at = coleta_mes(prim_atual, hoje, hoje.strftime("%m-%Y"))
    mes_an = coleta_mes(prim_ant, fim_ant, prim_ant.strftime("%m-%Y"))

    for serie, campo in ((433, "ipca"), (196, "poupanca")):
        try:
            for x in _sgs(serie, 4):
                dd, mm, yy = x["data"].split("/")
                lbl = f"{mm}-{yy}"
                for r_ in (mes_at, mes_an):
                    if r_["mes"] == lbl: r_[campo] = float(x["valor"])
        except Exception as e: erros.append(f"{campo} mensal: {e}")

    try:
        y = _http_json("https://query1.finance.yahoo.com/v8/finance/chart/%5EBVSP?range=3mo&interval=1mo")
        res = y["chart"]["result"][0]
        seq = [(datetime.fromtimestamp(t).strftime("%m-%Y"), c)
               for t, c in zip(res["timestamp"], res["indicators"]["quote"][0]["close"]) if c]
        for i in range(1, len(seq)):
            var = round((seq[i][1] / seq[i-1][1] - 1) * 100, 2)
            for r_ in (mes_at, mes_an):
                if r_["mes"] == seq[i][0]: r_["ibov_var"] = var
    except Exception as e: erros.append(f"Ibov mensal: {e}")

    prox = (prim_atual + timedelta(days=32)).replace(day=1)
    prev = {"mes": prox.strftime("%m-%Y"), "ipca": None, "cdi": None,
            "selic": None, "dolar": None, "poupanca": None,
            "fonte": "Boletim Focus/BCB (mediana das instituições)"}
    ref = prox.strftime("%m/%Y")
    try:
        v = _focus("ExpectativaMercadoMensais",
                   f"Indicador eq 'IPCA' and DataReferencia eq '{ref}' and baseCalculo eq 0")
        if v: prev["ipca"] = float(v[0]["Mediana"])
    except Exception as e: erros.append(f"Focus IPCA: {e}")
    try:
        v = _focus("ExpectativaMercadoMensais",
                   f"Indicador eq 'Câmbio' and DataReferencia eq '{ref}' and baseCalculo eq 0")
        if v: prev["dolar"] = float(v[0]["Mediana"])
    except Exception as e: erros.append(f"Focus Câmbio: {e}")
    try:
        v = _focus("ExpectativasMercadoSelic", "baseCalculo eq 0")
        if v: prev["selic"] = float(v[0]["Mediana"])
    except Exception as e: erros.append(f"Focus Selic: {e}")
    if prev["selic"] is None and selic is not None:
        prev["selic"] = selic
    if prev["selic"] is not None:

        cdi_aa_prev = max(prev["selic"] - 0.10, 0.0)
        prev["cdi"] = round(((1 + cdi_aa_prev / 100) ** (21 / 252) - 1) * 100, 2)

        if prev["selic"] > 8.5:
            tr_est = max((poup if poup is not None else 0.67) - 0.5, 0.0)
            prev["poupanca"] = round(0.5 + tr_est, 2)
        else:
            prev["poupanca"] = round(prev["selic"] * 0.7 / 12, 2)
    def _yahoo(simbolo):
        y = _http_json(f"https://query1.finance.yahoo.com/v8/finance/chart/{simbolo}?range=5d&interval=1d")
        meta = y["chart"]["result"][0]["meta"]
        preco = float(meta["regularMarketPrice"])
        prev_ = float(meta.get("chartPreviousClose") or meta.get("previousClose") or 0)
        var = round((preco / prev_ - 1) * 100, 2) if prev_ else None
        return preco, var
    ibov = None
    try:
        ibov = round(_yahoo("%5EBVSP")[0])
    except Exception as e:
        erros.append(f"Ibovespa: {e}")
    acoes = []
    for tk, nome in B3_TICKERS:
        try:
            preco, var = _yahoo(tk + ".SA")
            acoes.append({"ticker": tk, "nome": nome, "preco": round(preco, 2),
                          "var": var, "data": hoje.strftime("%d/%m/%Y")})
        except Exception as e:
            erros.append(f"{tk}: {e}")
    dia = {"data": hoje.strftime("%d/%m/%Y"), "selic": selic, "cdi_aa": cdi_aa,
           "dolar": dolar, "ibov": ibov}
    return {"dia": dia, "meses": [mes_at, mes_an], "prev": prev,
            "acoes": acoes, "erros": erros}

DIA_HDR = ["Data", "Selic (a.a.)", "CDI (a.a.)", "Dólar (R$)", "Ibovespa (pts)"]
MES_HDR = ["Mês", "CDI acumulado", "IPCA", "Poupança", "Selic (fim)", "Dólar (fim)", "Ibovespa (var)"]
PREV_HDR = ["Mês previsto", "IPCA", "CDI (a.m.)", "Selic (a.a.)", "Dólar", "Poupança (a.m.)"]
DIA_KEYS = ["data", "selic", "cdi_aa", "dolar", "ibov"]
MES_KEYS = ["mes", "cdi_mes", "ipca", "poupanca", "selic_fim", "dolar_fim", "ibov_var"]
PREV_KEYS = ["mes", "ipca", "cdi", "selic", "dolar", "poupanca"]

def ler_mercado(wb):
    out = {"dias": [], "meses": [], "acoes": [], "prev": {}, "atualizado": ""}
    if MERCADO_ABA not in wb.sheetnames: return out
    ws = wb[MERCADO_ABA]; secao = None
    for row in ws.iter_rows(min_row=1, max_row=5000, values_only=True):
        a = str(row[0] or "").strip()
        if not a: continue
        if "INDICADORES DO DIA" in a:  secao = "dia";  continue
        if "INDICADORES DO MÊS" in a:  secao = "mes";  continue
        if "PREVISÕES" in a:           secao = "prev"; continue
        if "AÇÕES B3" in a:            secao = "acao"; continue
        if "CARTEIRA" in a.upper() and "RECOMENDADA" in a.upper():
            secao = None; continue
        if "Atualizado" in a: out["atualizado"] = a.strip(); continue
        if a.startswith("Fonte:"):
            if secao == "prev": out["prev"]["fonte"] = a[6:].strip()
            continue
        if a in ("Data", "Mês", "Ticker", "Classe", "Mês previsto"): continue
        if secao == "dia" and re.match(r"\d{2}/\d{2}/\d{4}", a):
            out["dias"].append({k: row[i] if i else a for i, k in enumerate(DIA_KEYS)})
        elif secao == "mes" and re.match(r"\d{2}-\d{4}", a):
            out["meses"].append({k: (row[i] if i < len(row) else None) if i else a
                                 for i, k in enumerate(MES_KEYS)})
        elif secao == "prev" and re.match(r"\d{2}-\d{4}", a):
            fonte = out["prev"].get("fonte", "")
            out["prev"] = {k: (row[i] if i < len(row) else None) if i else a
                           for i, k in enumerate(PREV_KEYS)}
            out["prev"]["fonte"] = fonte
        elif secao == "acao" and re.match(r"^[A-Z0-9]{5,6}$", a):
            out["acoes"].append({"ticker": a, "nome": row[1], "preco": row[2],
                                 "var": row[3], "data": row[4]})
    return out

def escrever_mercado(wb, dados):
    """Reescreve a aba Mercado com visual XP."""
    if MERCADO_ABA in wb.sheetnames: del wb[MERCADO_ABA]
    ws = wb.create_sheet(MERCADO_ABA); ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = X_YELLOW
    ncols = len(MES_HDR); last_col = get_column_letter(ncols)

    def banda(r, txt, sub=False):
        ws.merge_cells(f"A{r}:{last_col}{r}")
        c = ws.cell(r, 1, "  " + txt)
        c.font = Font(name="Segoe UI", bold=not sub, size=10 if sub else 12,
                      color=X_MUT if sub else X_YELLOW)
        c.fill = F(X_BG); c.alignment = al("left")
        ws.row_dimensions[r].height = 20 if sub else 30

    def cabecalho(r, headers):
        for col, h in enumerate(headers, 1):
            c = ws.cell(r, col, h)
            c.font = Font(name="Segoe UI", bold=True, size=10, color=X_YELLOW)
            c.fill = F(X_BAND); c.alignment = al("center"); c.border = _dark_border()
        ws.row_dimensions[r].height = 22

    def linha(r, vals, fmts):
        bg = X_ROW1 if r % 2 == 0 else X_ROW2
        for col, (val, fmt) in enumerate(zip(vals, fmts), 1):
            c = ws.cell(r, col, val if val is not None else "")
            c.font = Font(name="Segoe UI", size=10, color=X_TXT if col == 1 else X_MUT)
            c.fill = F(bg); c.alignment = al("center"); c.border = _dark_border()
            if fmt: c.number_format = fmt
        ws.row_dimensions[r].height = 19

    banda(1, "Dados de Mercado")
    banda(2, dados.get("atualizado") or "", sub=True)
    r = 4
    banda(r, "INDICADORES DO DIA"); r += 1
    cabecalho(r, DIA_HDR); r += 1
    for d in dados["dias"]:
        linha(r, [d.get(k) for k in DIA_KEYS],
              [None, '0.00"%"', '0.00"%"', 'R$ 0.0000', '#,##0']); r += 1
    r += 1
    banda(r, "INDICADORES DO MÊS"); r += 1
    cabecalho(r, MES_HDR); r += 1
    for d in dados["meses"]:
        linha(r, [d.get(k) for k in MES_KEYS],
              [None, '0.00"%"', '0.00"%"', '0.00"%"', '0.00"%"', 'R$ 0.0000', '0.00"%"']); r += 1
    r += 1
    banda(r, "PREVISÕES — PRÓXIMO MÊS"); r += 1
    cabecalho(r, PREV_HDR); r += 1
    pv = dados.get("prev") or {}
    if pv.get("mes"):
        linha(r, [pv.get(k) for k in PREV_KEYS],
              [None, '0.00"%"', '0.00"%"', '0.00"%"', 'R$ 0.0000', '0.00"%"']); r += 1
        banda(r, "Fonte: " + (pv.get("fonte") or "Boletim Focus/BCB"), sub=True); r += 1
    r += 1
    banda(r, "AÇÕES B3"); r += 1
    cabecalho(r, ["Ticker", "Empresa", "Cotação", "Variação", "Data"]); r += 1
    for ac in dados.get("acoes") or []:
        bg = X_ROW1 if r % 2 == 0 else X_ROW2
        var = ac.get("var")
        var_cor = C_OK if (var is not None and var >= 0) else C_ERR if var is not None else X_MUT
        vals = [(ac.get("ticker"), X_TXT, None), (ac.get("nome"), X_MUT, None),
                (ac.get("preco"), X_TXT, 'R$ 0.00'), (var, var_cor, '0.00"%"'),
                (ac.get("data"), X_MUT, None)]
        for col, (val, cor, fmt) in enumerate(vals, 1):
            c = ws.cell(r, col, val if val is not None else "")
            c.font = Font(name="Segoe UI", size=10, bold=(col == 4), color=cor)
            c.fill = F(bg); c.alignment = al("center" if col != 2 else "left")
            c.border = _dark_border()
            if fmt: c.number_format = fmt
        ws.row_dimensions[r].height = 19; r += 1
    for col, w in enumerate([22, 24, 14, 14, 14, 14, 16], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

def atualizar_mercado_excel(novo=None):
    """Mescla dados de mercado com o histórico e salva. Retorna dados + erros.
    `novo` pode vir pré-buscado (fetch feito fora do lock, para não congelar a UI);
    se None, busca aqui (uso standalone)."""
    if novo is None:
        novo = fetch_mercado()
    wb = get_wb(); garantir_email(wb)
    dados = ler_mercado(wb)

    dias = [d for d in dados["dias"] if d.get("data") != novo["dia"]["data"]]
    dias.insert(0, novo["dia"])
    dias.sort(key=lambda d: datetime.strptime(str(d["data"]), "%d/%m/%Y"), reverse=True)

    meses = {m["mes"]: m for m in dados["meses"]}
    for nm in novo["meses"]:
        alvo = meses.get(nm["mes"], {"mes": nm["mes"]})
        for k in MES_KEYS[1:]:
            if nm.get(k) is not None: alvo[k] = nm[k]
        meses[nm["mes"]] = alvo
    lm = sorted(meses.values(),
                key=lambda m: (str(m["mes"]).split("-")[1], str(m["mes"]).split("-")[0]),
                reverse=True)
    prev = novo.get("prev") or {}
    dados.update({"dias": dias[:90], "meses": lm,
                  "acoes": novo.get("acoes") or dados.get("acoes") or [],
                  "prev": prev if prev.get("mes") else dados.get("prev") or {},
                  "atualizado": "Atualizado em " + datetime.now().strftime("%d/%m/%Y %H:%M")
                                + "  ·  Fontes: Banco Central (SGS/Focus) e Yahoo Finance"})
    escrever_mercado(wb, dados)
    salvar_wb(wb)
    dados["erros"] = novo["erros"]
    return dados

def vars_mercado(dados, mes_ref=None):
    """Converte os dados de mercado em variáveis pt-BR para o template de e-mail."""
    v = {}
    def br(x, nd=2):
        if x is None: return ""
        s = f"{float(x):,.{nd}f}"
        return s.replace(",", "X").replace(".", ",").replace("X", ".")
    if dados["dias"]:
        d0 = dados["dias"][0]
        v["selic"] = br(d0.get("selic")); v["cdi_hoje"] = br(d0.get("cdi_aa"))
        v["dolar"] = br(d0.get("dolar"), 2)
        v["ibovespa"] = br(d0.get("ibov"), 0)
    alvo = None
    if mes_ref:
        alvo = next((m for m in dados["meses"] if str(m["mes"]) == str(mes_ref)), None)
    if alvo is None and dados["meses"]: alvo = dados["meses"][0]
    if alvo:
        v["cdi_mes"] = br(alvo.get("cdi_mes"))
        v["ipca"] = br(alvo.get("ipca"))
        v["poupanca"] = br(alvo.get("poupanca"))
    for ac in dados.get("acoes") or []:
        t = str(ac.get("ticker", "")).lower()
        if t and ac.get("preco") is not None:
            v[t] = br(ac.get("preco"))
    return v

def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))

def _dados_format(c):
    nome_full = (c.get("nome") or "").strip()
    primeiro = nome_full.split()[0].title() if nome_full else ""
    d = {k: v for k, v in c.items() if not isinstance(v, dict)}
    out = {**d, "nome": nome_full, "primeiro_nome": primeiro,
           "mes_ref": mes_bonito(c.get("mes_ref", ""))}

    norm = {_norm_cls(k): v for k, v in (c.get("aloc") or {}).items()}
    for nome_cls, slug in STRATEGY_VARS:
        v = norm.get(_norm_cls(nome_cls))
        out[slug] = (f"{float(v):.2f}".replace(".", ",") if v not in (None, "") else "")
    return out

def _extra_vars(c, mdados=None, perfis=None):
    """Variáveis de mercado para o template de e-mail."""
    if mdados is None:
        try:
            mdados = mercado_cache()
        except Exception:
            mdados = {"dias": [], "meses": [], "acoes": [], "prev": {}}
    v = vars_mercado(mdados, c.get("mes_ref"))
    v["patrimonio"] = c.get("patrimonio", "")
    return v

def _key_cls(s):
    """Chave de classe sem acentos/pontuação, p/ casar nomes vindos de fontes diferentes."""
    import unicodedata
    s = unicodedata.normalize("NFD", str(s or "").lower())
    s = "".join(ch for ch in s if unicodedata.category(ch) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", s).strip()

_PAPEL_CLASSE = {
    "pos fixado": "a âncora de estabilidade da carteira",
    "inflacao": "a proteção do poder de compra no longo prazo",
    "prefixado": "a trava de juros, que se beneficia da queda da Selic",
    "multimercados": "o motor de diversificação e descorrelação",
    "multimercado": "o motor de diversificação e descorrelação",
    "renda variavel brasil": "o componente de crescimento local",
    "global renda variavel": "a diversificação internacional",
    "renda variavel global": "a diversificação internacional",
    "global renda fixa": "a renda fixa internacional",
    "fundos listados": "renda recorrente e diversificação via ativos listados",
    "alternativos": "retorno com baixa correlação com o mercado tradicional",
    "alternativo": "retorno com baixa correlação com o mercado tradicional",
    "caixa": "a liquidez imediata para oportunidades e imprevistos",
    "proventos": "a geração de renda recorrente",
}

def _persps_cliente(c, max_n=2):
    """Perspectivas da XP (relatório de alocação) para as classes onde o
    cliente tem posição, da maior para a menor."""
    try:
        aloc = c.get("aloc") or {}
        ents = sorted(((k, float(v)) for k, v in aloc.items() if v not in (None, "")),
                      key=lambda x: -x[1])
        persp = (aloc_cache() or {}).get("persp") or []
        pmap = {_key_cls(p.get("classe", "")): p for p in persp if p.get("texto")}
        out = []
        for k, v in ents:
            p = pmap.get(_key_cls(k))
            if not p: continue
            t = str(p["texto"]).strip()
            if len(t) > 260: t = t[:257].rsplit(" ", 1)[0] + "…"
            out.append({"classe": str(p.get("classe") or k),
                        "pct": f"{v:.1f}".replace(".", ","), "texto": t})
            if len(out) >= max_n: break
        return out
    except Exception:
        return []

def _frases_analise(c, extra):
    """Frases explicativas prontas — viram VARIÁVEIS de texto do template
    ({analise_cdi}, {analise_ganho}, {analise_acumulado}, {maior_posicao},
    {maior_posicao_pct}, {papel_maior_posicao}, {perspectiva_1}, …).
    Sem dado suficiente, a variável vira texto vazio."""
    ex = extra or {}
    fr = {"analise_cdi": "", "analise_ganho": "", "analise_acumulado": "",
          "maior_posicao": "", "maior_posicao_pct": "", "papel_maior_posicao": "",
          "perspectiva_1": "", "perspectiva_1_classe": "",
          "perspectiva_2": "", "perspectiva_2_classe": ""}
    try:
        rent, cdi = c.get("rent_num"), c.get("cdi_num")
        cdi_mes_txt = (" O CDI do mês foi " + str(ex.get("cdi_mes")) + "%.") if ex.get("cdi_mes") else ""
        if rent is not None:
            if cdi is not None and cdi >= 100:
                fr["analise_cdi"] = ("o equivalente a " + c.get("cdi", "") + "% do CDI. Sua carteira rendeu acima do "
                                     "referencial básico de renda fixa — ou seja, o risco assumido está sendo bem "
                                     "remunerado." + cdi_mes_txt)
            elif cdi is not None and cdi >= 80:
                fr["analise_cdi"] = ("próxima do CDI (" + c.get("cdi", "") + "%). Em carteiras diversificadas isso é "
                                     "esperado em alguns meses: parte dos ativos protege, parte busca retorno maior "
                                     "no acumulado." + cdi_mes_txt)
            elif cdi is not None:
                fr["analise_cdi"] = ("abaixo do CDI neste mês (" + c.get("cdi", "") + "%). Oscilações mensais fazem "
                                     "parte da estratégia — o objetivo é superar o referencial no horizonte acumulado, "
                                     "não em cada mês isolado." + cdi_mes_txt)
            else:
                fr["analise_cdi"] = "resultado do mês na carteira consolidada." + cdi_mes_txt
        if c.get("ganho_num") is not None:
            mov = c.get("mov_num")
            if mov is None or abs(mov) < 0.005:
                fr["analise_ganho"] = ("todo esse resultado veio do rendimento dos ativos — não houve aportes nem "
                                       "resgates no período.")
            elif mov > 0:
                fr["analise_ganho"] = ("além do rendimento, houve aporte líquido de R$ " + c.get("movimentacoes", "") +
                                       " no período, que também contribuiu para a evolução do patrimônio.")
            else:
                fr["analise_ganho"] = ("resultado já considerando o resgate líquido do período — o rendimento dos "
                                       "ativos segue fazendo o patrimônio trabalhar.")
        if c.get("rent_total_num") is not None:
            per = str(c.get("periodo_total") or "").strip() or "desde o início"
            gt = (", com ganho total de R$ " + c.get("ganho_total", "")) if c.get("ganho_total") else ""
            fr["analise_acumulado"] = ("é nesse horizonte mais longo que a estratégia deve ser avaliada" + gt +
                                       ". Meses individuais oscilam; o acumulado (" + per + ") mostra a direção.")
        aloc = c.get("aloc") or {}
        ents = sorted(((k, float(v)) for k, v in aloc.items() if v not in (None, "")),
                      key=lambda x: -x[1])
        if ents:
            k0, v0 = ents[0]
            fr["maior_posicao"] = str(k0)
            fr["maior_posicao_pct"] = f"{v0:.1f}".replace(".", ",")
            fr["papel_maior_posicao"] = _PAPEL_CLASSE.get(_key_cls(k0), "")
        for i, p in enumerate(_persps_cliente(c), 1):
            fr["perspectiva_%d_classe" % i] = p["classe"]
            fr["perspectiva_%d" % i] = p["texto"]
    except Exception:
        pass
    return fr

def _blocos_email(c, extra):
    """Blocos visuais do template: {resumo_pontos}, {composicao_carteira} e
    {contexto_mercado}. Só aparecem se a variável estiver na mensagem; a linha
    do bloco some quando não há dados. Nunca derrubam o e-mail."""
    try:
        ex = extra or {}
        F = "font-family:Segoe UI,Arial,sans-serif;"
        blocos = {"{resumo_pontos}": "", "{composicao_carteira}": "", "{contexto_mercado}": ""}
        fr = _frases_analise(c, extra)
        mes = mes_bonito(c.get("mes_ref", ""))

        def sec(t):
            return ('<div style="' + F + 'font-size:11px;letter-spacing:.09em;'
                    'text-transform:uppercase;color:#8a8a8a;font-weight:700;'
                    'margin:20px 0 10px;">' + _esc(t) + '</div>')

        def ponto(cor, titulo, texto):
            return ('<table role="presentation" width="100%" cellpadding="0" cellspacing="0" '
                    'style="margin:0 0 12px;"><tr>'
                    '<td width="3" style="background:' + cor + ';border-radius:2px;font-size:0;">&nbsp;</td>'
                    '<td style="padding:2px 0 2px 12px;' + F + 'font-size:13.5px;line-height:1.6;color:#2d2d33;">'
                    '<b style="color:#111111;">' + titulo + '</b> — ' + texto + '</td></tr></table>')

        pts = []
        cdi = c.get("cdi_num")
        if fr["analise_cdi"]:
            cor = ("#0f8a3d" if (cdi is not None and cdi >= 100) else
                   ("#b07300" if (cdi is not None and cdi >= 80) else
                    ("#c0392b" if cdi is not None else "#b07300")))
            pts.append(ponto(cor, "Rentabilidade de " + _esc(c.get("rentabilidade", "")) + "% em " + _esc(mes),
                             _esc(fr["analise_cdi"])))
        if fr["analise_ganho"]:
            pts.append(ponto("#F2C029", "Ganho de R$ " + _esc(c.get("ganho", "")) + " no mês",
                             _esc(fr["analise_ganho"])))
        if fr["analise_acumulado"]:
            per = str(c.get("periodo_total") or "").strip() or "desde o início"
            pts.append(ponto("#3B82F6", "Acumulado (" + _esc(per) + "): " + _esc(c.get("rentabilidade_total", "")) + "%",
                             _esc(fr["analise_acumulado"])))
        if pts:
            blocos["{resumo_pontos}"] = sec("Principais pontos do seu relatório") + "".join(pts)

        aloc = c.get("aloc") or {}
        ents = sorted(((k, float(v)) for k, v in aloc.items() if v not in (None, "")),
                      key=lambda x: -x[1])
        if ents:
            vmax = max(v for _, v in ents[:6]) or 1
            rows = []
            for k, v in ents[:6]:
                pct = f"{v:.1f}".replace(".", ",") + "%"
                w = max(4, int(v / vmax * 100))
                rows.append('<tr>'
                    '<td style="' + F + 'font-size:12.5px;color:#2d2d33;padding:5px 10px 5px 0;white-space:nowrap;">' + _esc(k) + '</td>'
                    '<td width="100%" style="padding:5px 0;"><table role="presentation" width="100%" cellpadding="0" cellspacing="0">'
                    '<tr><td width="' + str(w) + '%" style="background:#F2C029;border-radius:4px;font-size:0;height:8px;line-height:8px;">&nbsp;</td>'
                    '<td style="font-size:0;">&nbsp;</td></tr></table></td>'
                    '<td align="right" style="' + F + 'font-size:12.5px;font-weight:700;color:#111111;padding:5px 0 5px 12px;">' + pct + '</td></tr>')
            comp = (sec("Como sua carteira está distribuída") +
                    '<table role="presentation" width="100%" cellpadding="0" cellspacing="0">' + "".join(rows) + '</table>')
            if fr["papel_maior_posicao"]:
                comp += ('<div style="' + F + 'font-size:12.5px;color:#6b6b73;margin-top:8px;line-height:1.55;">'
                         'A maior posição é <b style="color:#2d2d33;">' + _esc(fr["maior_posicao"]) + '</b> ('
                         + _esc(fr["maior_posicao_pct"]) + '%), que é ' + _esc(fr["papel_maior_posicao"]) + '.</div>')
            blocos["{composicao_carteira}"] = comp

        ind = []
        for lab, key, suf, pre in (("Selic", "selic", "% a.a.", ""), ("CDI no mês", "cdi_mes", "%", ""),
                                   ("IPCA no mês", "ipca", "%", ""), ("Dólar", "dolar", "", "R$ "),
                                   ("Ibovespa", "ibovespa", " pts", "")):
            if ex.get(key):
                ind.append('<td style="padding:8px 12px;background:#fafafa;border:1px solid #ececec;border-radius:8px;'
                           + F + 'font-size:12px;color:#2d2d33;white-space:nowrap;"><span style="color:#8a8a8a;">'
                           + lab + '</span>&nbsp; <b>' + pre + _esc(ex[key]) + suf + '</b></td>'
                           '<td style="width:8px;font-size:0;">&nbsp;</td>')
        persp_html = []
        for p in _persps_cliente(c):
            persp_html.append('<div style="' + F + 'font-size:12.5px;line-height:1.6;color:#2d2d33;'
                              'background:#fbf7ea;border:1px solid #f0e3b8;border-radius:8px;'
                              'padding:12px 14px;margin:0 0 10px;">'
                              '<b style="color:#111111;">' + _esc(p["classe"]) + '</b>'
                              '<span style="color:#8a8a8a;"> · você tem ' + _esc(p["pct"]) + '% da carteira aqui</span><br>'
                              + _esc(p["texto"]) + '</div>')
        if ind or persp_html:
            ctx = sec("O que o mercado diz — e o que isso tem a ver com você")
            if ind:
                ctx += ('<table role="presentation" cellpadding="0" cellspacing="0" style="margin:0 0 14px;"><tr>'
                        + "".join(ind) + '</tr></table>')
            blocos["{contexto_mercado}"] = ctx + "".join(persp_html)
        return blocos
    except Exception:
        return {"{resumo_pontos}": "", "{composicao_carteira}": "", "{contexto_mercado}": ""}

def render_email_html(c, corpo, extra=None):
    """Layout claro, limpo e seguro para qualquer leitor de e-mail (tabelas + CSS inline).
    A mensagem é 100% do template: frases de análise viram variáveis de texto e os
    blocos visuais ({resumo_pontos}, {composicao_carteira}, {contexto_mercado}) só
    entram onde a Laura os colocar na mensagem."""
    dados = {**_dados_format(c), **(extra or {}), **_frases_analise(c, extra)}
    try:
        texto = corpo.format_map(_SafeDict(dados))
    except Exception:
        texto = corpo
    blocos = _blocos_email(c, extra)
    linhas = []
    for ln in texto.split("\n"):
        if ln.strip() in blocos:
            if blocos[ln.strip()]:
                linhas.append(blocos[ln.strip()])
            continue
        e = _esc(ln)
        if ln.strip() in ("Resumo", "Destaques"):
            linhas.append(f'<div style="font-weight:700;color:#111111;margin:18px 0 4px;">{e}</div>')
        elif ln.strip().startswith("·"):
            linhas.append(f'<div style="margin:2px 0 2px 10px;">{e}</div>')
        elif not ln.strip():
            linhas.append('<div style="height:10px;line-height:10px;">&nbsp;</div>')
        else:
            linhas.append(f'<div>{e}</div>')
    corpo_html = "\n".join(linhas)

    def chip(label, valor, cor="#111111"):
        return f"""<td align="center" style="padding:14px 6px;background:#fafafa;border:1px solid #ececec;border-radius:8px;">
          <div style="font-size:11px;letter-spacing:.06em;text-transform:uppercase;color:#8a8a8a;font-family:Segoe UI,Arial,sans-serif;">{label}</div>
          <div style="font-size:19px;font-weight:700;color:{cor};font-family:Segoe UI,Arial,sans-serif;margin-top:3px;">{valor}</div>
        </td>"""

    try:
        cdi_v = float(str(c.get("cdi", "")).replace(",", "."))
        cdi_cor = "#0f8a3d" if cdi_v >= 100 else ("#b07300" if cdi_v >= 80 else "#c0392b")
    except Exception:
        cdi_cor = "#111111"
    rent = c.get("rentabilidade", "—"); ganho = c.get("ganho", "—"); cdi = c.get("cdi", "—")
    mes = mes_bonito(c.get("mes_ref", ""))
    return f"""<!DOCTYPE html>
<html><body style="margin:0;padding:0;background:#f4f5f7;">
<table role="presentation" width="100%" cellpadding="0" cellspacing="0" style="background:#f4f5f7;padding:28px 12px;">
<tr><td align="center">
<table role="presentation" width="620" cellpadding="0" cellspacing="0"
       style="max-width:620px;width:100%;background:#ffffff;border:1px solid #e8e8e8;border-radius:12px;overflow:hidden;">
  <tr><td style="height:5px;background:#F2C029;font-size:0;line-height:0;">&nbsp;</td></tr>
  <tr><td style="padding:26px 34px 6px;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr>
      <td style="font-family:Segoe UI,Arial,sans-serif;font-size:19px;font-weight:800;color:#111111;">
        XPerformance</td>
      <td align="right" style="font-family:Segoe UI,Arial,sans-serif;font-size:12px;color:#9a9a9a;">
        {_esc(mes)}</td>
    </tr></table>
  </td></tr>
  <tr><td style="padding:14px 34px 4px;">
    <table role="presentation" width="100%" cellpadding="0" cellspacing="0"><tr>
      {chip("Rentabilidade", _esc(rent) + "%")}
      <td style="width:10px;font-size:0;">&nbsp;</td>
      {chip("Ganho", "R$ " + _esc(ganho))}
      <td style="width:10px;font-size:0;">&nbsp;</td>
      {chip("vs CDI", _esc(cdi) + "%", cdi_cor)}
    </tr></table>
  </td></tr>
  <tr><td style="padding:18px 34px 28px;font-family:Segoe UI,Arial,sans-serif;font-size:15px;line-height:1.55;color:#2d2d33;">
    {corpo_html}
  </td></tr>
  <tr><td style="padding:14px 34px;background:#fafafa;border-top:1px solid #efefef;
       font-family:Segoe UI,Arial,sans-serif;font-size:11px;color:#a0a0a0;">
    Relatório gerado por Guilherme Enrico - Rico Investimentos
  </td></tr>
</table>
</td></tr></table>
</body></html>"""

def _achar_pdf(nome, pasta=None):
    """Resolve o caminho do PDF do relatório de forma robusta.
    Tenta: caminho absoluto → pasta escolhida na interface → pasta padrão →
    busca pelo nome do arquivo (subpastas incluídas). Retorna None se não achar."""
    if not nome:
        return None
    nome = str(nome).strip()
    if os.path.isabs(nome) and os.path.isfile(nome):
        return nome
    for base_dir in (pasta, PASTA_PDFS):
        if base_dir and os.path.isdir(base_dir):
            cand = os.path.join(base_dir, nome)
            if os.path.isfile(cand):
                return cand
    alvo = os.path.basename(nome).lower()
    for base_dir in (pasta, PASTA_PDFS):
        if not (base_dir and os.path.isdir(base_dir)):
            continue
        try:
            for raiz, _dirs, arquivos in os.walk(base_dir):
                for a in arquivos:
                    if a.lower() == alvo:
                        return os.path.join(raiz, a)
        except OSError:
            pass
    return None

def enviar_via_outlook(dest, assunto, html, anexo=None):
    import win32com.client
    ol = win32com.client.Dispatch("Outlook.Application")
    mail = ol.CreateItem(0)
    mail.To = dest; mail.Subject = assunto; mail.HTMLBody = html
    if anexo:
        anexo = os.path.abspath(anexo)
        if not os.path.isfile(anexo):
            raise FileNotFoundError("Anexo não encontrado: " + anexo)

        mail.Attachments.Add(anexo, 1)
    mail.Send()

def _push(fn, *args):
    try:
        payload = ",".join(
            json.dumps(a, ensure_ascii=False).replace(" ", "\\u2028").replace(" ", "\\u2029")
            for a in args)
        webview.windows[0].evaluate_js(f"{fn}({payload})")
    except Exception:
        pass

class Api:

    def boot(self):
        _rotacionar_log()
        migrar_onde_investir()
        _sanear_cadastro()
        _db_reconciliar()

        out = {"ok": True, "pasta": PASTA_PDFS, "meses": [],
               "assunto": ASSUNTO_DEF, "corpo": CORPO_DEF, "cadastro": [],
               "excel": os.path.basename(EXCEL_PATH), "aviso": "",
               "aniversariantes": []}
        try:
            if _tem_dados():
                wb = _wb_cached()
                out["assunto"], out["corpo"] = ler_template(wb)
                out["meses"] = listar_meses(wb)
                out["cadastro"] = _cadastro_enriquecido(wb)
                for c in out["cadastro"]:
                    p = _aniv_pack(c.get("nascimento"))
                    if p["aniv_status"]:
                        out["aniversariantes"].append({
                            "nome": c.get("nome") or ("Conta " + str(c.get("conta", ""))),
                            "status": p["aniv_status"], "dsem": p["aniv_dsem"],
                            "data": p["aniv_data"], "idade_nova": p["idade_nova"]})
        except Exception as e:
            out["aviso"] = f"Não consegui ler a planilha: {e}"
        return out

    def painel_data(self):
        """Dados da aba Painel: visão geral executiva + notificações reais."""
        try:
            cfg = carregar_config()
            def _meta(ch):
                try:
                    v = cfg.get(ch)
                    return float(v) if v not in (None, "") else None
                except Exception:
                    return None
            hoje = datetime.now().date()
            tri = (hoje.month - 1) // 3
            out = {"ok": True, "evolucao": [], "aum": 0, "total": 0,
                   "rent_media": None, "cdi_medio": None, "mes_ref": "",
                   "aloc": [], "enq": None, "notificacoes": [],

                   "metas": {"mes": _meta("painel_meta_mes"),
                             "tri": _meta("painel_meta_tri"),
                             "sem": _meta("painel_meta_sem")},
                   "meta_per": (cfg.get("painel_meta_per") or "tri"),
                   "aportes": {"mes": 0.0, "tri": 0.0, "sem": 0.0},
                   "periodos": {"mes": mes_bonito("%02d-%d" % (hoje.month, hoje.year)),
                                "tri": "%dº trimestre de %d" % (tri + 1, hoje.year),
                                "sem": "%dº semestre de %d" % (1 if hoje.month <= 6 else 2, hoje.year)}}

            fer = _feriados_br(hoje.year)
            def _fim(m):
                if m == 12: return datetime(hoje.year, 12, 31).date()
                return datetime.fromordinal(datetime(hoje.year, m + 1, 1).toordinal() - 1).date()
            per_ini_fim = {"mes": (hoje.replace(day=1), _fim(hoje.month)),
                           "tri": (datetime(hoje.year, tri * 3 + 1, 1).date(), _fim(tri * 3 + 3)),
                           "sem": (datetime(hoje.year, 1 if hoje.month <= 6 else 7, 1).date(),
                                   _fim(6 if hoje.month <= 6 else 12))}
            out["uteis"] = {k: {"rest": _dias_uteis(max(hoje, ini), fim, fer),
                                "total": _dias_uteis(ini, fim, fer)}
                            for k, (ini, fim) in per_ini_fim.items()}
            if not _tem_dados():
                return out
            agg = coletar_cache(); cad = cadastro_cache()

            per_meses = {"mes": {hoje.month},
                         "tri": {tri * 3 + 1, tri * 3 + 2, tri * 3 + 3},
                         "sem": set(range(1, 7)) if hoje.month <= 6 else set(range(7, 13))}
            for a in aportes_cache():
                d_ = _data_br(a.get("data"))
                if not d_ or d_.year != hoje.year: continue
                for k, ms in per_meses.items():
                    if d_.month in ms:
                        out["aportes"][k] += a["valor"]
            out["aportes"] = {k: round(v, 2) for k, v in out["aportes"].items()}

            try:
                recs_p = receita_cache()
            except Exception:
                recs_p = {}
            out["receitas"] = {"mes": 0.0, "tri": 0.0, "sem": 0.0}
            out["rec_meses"] = {}
            for (aa, mm), v in recs_p.items():
                if aa != hoje.year: continue
                out["rec_meses"][str(mm)] = round(v, 2)
                for k, ms in per_meses.items():
                    if mm in ms:
                        out["receitas"][k] = round(out["receitas"][k] + v, 2)
            try: dados_aloc = aloc_cache()
            except Exception: dados_aloc = {}
            meses = sorted({m for o in agg.values() for m in o["meses"]}, key=_ordem_mes)

            ap_all = []
            for a in aportes_cache():
                d_ = _data_br(a.get("data"))
                if d_: ap_all.append((d_, str(a["conta"]).strip(), a["valor"]))
            try: recs = receita_cache()
            except Exception: recs = {}
            if ap_all or recs:
                idx_h = hoje.year * 12 + hoje.month - 1
                idxs = ([d_.year * 12 + d_.month - 1 for d_, _, _ in ap_all]
                        + [aa * 12 + mm - 1 for (aa, mm) in recs])
                idx_min = max(min(idxs), idx_h - 35)
                soma = {}; quem = {}
                for d_, conta_a, val in ap_all:
                    j = d_.year * 12 + d_.month - 1
                    if idx_min <= j <= idx_h:
                        soma[j] = soma.get(j, 0) + val
                        quem.setdefault(j, set()).add(conta_a)
                for j in range(idx_min, idx_h + 1):
                    aa, mm = divmod(j, 12); mm += 1
                    out["evolucao"].append({"mes": "%02d-%d" % (mm, aa),
                                            "label": mes_bonito("%02d-%d" % (mm, aa)),
                                            "total": round(soma.get(j, 0), 2),
                                            "cli": len(quem.get(j, ())),
                                            "rec": recs.get((aa, mm), 0)})
                out["evo_cli"] = len(set().union(*quem.values())) if quem else 0

            pos = []
            for conta, o in agg.items():
                if not o["meses"]: continue
                mu = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)[0]
                pos.append((conta, o, o["meses"][mu]))
            out["total"] = len(pos)
            out["aum"] = round(sum((m.get("pat_num") or 0) for _, _, m in pos), 2)
            ult = meses[-1] if meses else None
            if ult:
                rents = [o["meses"][ult].get("rent_num") for o in agg.values() if ult in o["meses"]]
                rents = [r for r in rents if r is not None]
                cdis = [o["meses"][ult].get("cdi_num") for o in agg.values() if ult in o["meses"]]
                cdis = [x for x in cdis if x is not None]
                if rents: out["rent_media"] = round(sum(rents) / len(rents), 2)
                if cdis: out["cdi_medio"] = round(sum(cdis) / len(cdis), 0)
                out["mes_ref"] = mes_bonito(ult)

            grupos = [("Renda Fixa", ("pos fixado", "prefixado", "inflacao", "global renda fixa")),
                      ("Ações", ("renda variavel brasil", "global renda variavel",
                                 "renda variavel global", "fundos listados")),
                      ("Multimercados", ("multimercados", "multimercado")),
                      ("Liquidez", ("caixa",))]
            mapa = {k: g for g, ks in grupos for k in ks}
            soma = {g: 0.0 for g, _ in grupos}; soma["Outros"] = 0.0
            peso = 0.0
            aloc_n = 0
            for _, _, m in pos:
                pat = m.get("pat_num") or 0
                aloc = m.get("aloc") or {}
                if not aloc or pat <= 0: continue
                peso += pat; aloc_n += 1
                for k, v in aloc.items():
                    try: v = float(v)
                    except Exception: continue
                    soma[mapa.get(_key_cls(k), "Outros")] += v * pat
            if peso > 0:
                al = [{"grupo": g, "pct": round(s / peso, 1)} for g, s in soma.items() if s / peso >= 0.5]
                al.sort(key=lambda x: -x["pct"])
                out["aloc"] = al
            out["aloc_n"] = aloc_n

            ruins = []
            if dados_aloc.get("classes"):
                ads = []
                for conta, o, m in pos:
                    aloc = m.get("aloc") or {}
                    if not aloc: continue
                    pf = PERFIL_KEY.get(_norm_cls(o.get("perfil", "")))
                    try:
                        ad = comparar_carteira(aloc, dados_aloc, pf).get("aderencia")
                    except Exception:
                        ad = None
                    if ad is None: continue
                    ads.append(ad)
                    if ad < 60:
                        ruins.append({"conta": conta, "nome": o["nome"] or ("Conta " + conta),
                                      "aderencia": round(ad, 1), "perfil": o.get("perfil", "")})
                if ads:
                    score = int(round(sum(ads) / len(ads)))
                    out["enq"] = {"score": score,
                                  "label": ("Excelente" if score >= 85 else "Bom" if score >= 70
                                            else "Atenção" if score >= 50 else "Crítico"),
                                  "pct_bom": int(round(sum(1 for a in ads if a >= 80) / len(ads) * 100)),
                                  "risco": ("Baixo Risco" if not ruins else
                                            "Risco Moderado" if len(ruins) <= 2 else "Risco Alto")}

            nots = []
            ruins.sort(key=lambda x: x["aderencia"])
            for r in ruins[:3]:
                nots.append({"tipo": "rebal", "titulo": "Rebalanceamento necessário",
                             "texto": ("A carteira de " + r["nome"] + " está com "
                                       + f"{r['aderencia']:.0f}" + "% de aderência ao perfil"
                                       + ((" " + r["perfil"]) if r["perfil"] else "")
                                       + " — abaixo do recomendado (60%)."),
                             "acao": "cart", "conta": r["conta"], "botao": "Revisar carteira"})
            for c in cad:
                p = _aniv_pack(c.get("nascimento"))
                nome = c.get("nome") or ("Conta " + str(c.get("conta", "")))
                if p["aniv_status"] == "hoje":
                    nots.append({"tipo": "aniv", "titulo": "Aniversário hoje",
                                 "chave": "aniv|%s|%s" % (c.get("conta", ""), hoje.year),
                                 "texto": nome + (" faz " + str(p["idade_nova"]) + " anos hoje"
                                                  if p["idade_nova"] is not None else " faz aniversário hoje")
                                          + " — um bom motivo para uma mensagem.",
                                 "acao": "cli", "conta": str(c.get("conta", "")), "botao": "Abrir cliente"})
                elif p["aniv_status"] == "fds":
                    nots.append({"tipo": "aniv", "titulo": "Aniversário no fim de semana",
                                 "chave": "aniv|%s|%s" % (c.get("conta", ""), hoje.year),
                                 "texto": nome + " faz aniversário no " + p["aniv_dsem"] + " ("
                                          + p["aniv_data"] + ") — parabenize hoje, no dia útil.",
                                 "acao": "cli", "conta": str(c.get("conta", "")), "botao": "Abrir cliente"})

            try:
                lim_ap = _aporte_lim()
                aps_d = {}
                for a in aportes_cache():
                    aps_d.setdefault(a["conta"], []).append(_data_br(a.get("data")))
                atras = []
                for conta, o, _m in pos:
                    datas = [d_ for d_ in aps_d.get(conta, []) if d_]
                    nome = o["nome"] or ("Conta " + conta)
                    if datas:
                        dias = (hoje - max(datas)).days
                        if dias > lim_ap: atras.append((dias, conta, nome, max(datas)))
                    else:
                        atras.append((10**6, conta, nome, None))
                atras.sort(key=lambda x: -x[0])
                for dias, conta, nome, du in atras[:3]:
                    nots.append({"tipo": "aporte",
                                 "titulo": "Sem aporte há muito tempo" if du else "Nenhum aporte registrado",
                                 "texto": (nome + " está há " + str(dias) + " dias sem aportar (último: "
                                           + du.strftime("%d/%m/%Y") + ")." if du
                                           else nome + " ainda não tem nenhum aporte registrado."),
                                 "acao": "cli", "conta": conta, "botao": "Abrir cliente"})
            except Exception:
                pass
            sem_email = [c.get("nome") or str(c.get("conta", "")) for c in cad
                         if not str(c.get("email") or "").strip()]
            if sem_email:
                nots.append({"tipo": "pend", "titulo": "Cadastro incompleto",
                             "texto": (str(len(sem_email)) + " cliente(s) sem e-mail: "
                                       + ", ".join(sem_email[:4]) + ("…" if len(sem_email) > 4 else "")
                                       + ". Sem e-mail, ficam de fora do envio mensal."),
                             "acao": "cli", "botao": "Abrir cadastro"})
            inativos = [c.get("nome") or str(c.get("conta", "")) for c in cad
                        if str(c.get("conta", "")).strip() and str(c.get("conta", "")).strip() not in agg]
            if inativos:
                nots.append({"tipo": "warn", "titulo": "Clientes sem relatório",
                             "texto": (str(len(inativos)) + " cliente(s) do cadastro sem relatório importado: "
                                       + ", ".join(inativos[:4]) + ("…" if len(inativos) > 4 else "") + "."),
                             "acao": "pdf", "botao": "Ler PDFs"})

            agora = datetime.now().strftime("%d/%m/%Y %H:%M")
            store = _notif_store_load()
            ativos_ids = set()
            for n in nots:
                nid = n.get("chave") or (n["tipo"] + "|" + str(n.get("conta") or ""))
                n["id"] = nid; ativos_ids.add(nid)
                ent = store.get(nid)
                if ent is None:
                    store[nid] = {"chegou": agora, "ignorada": False,
                                  "n": {k: v for k, v in n.items() if k != "id"}}
                else:
                    ent["n"] = {k: v for k, v in n.items() if k != "id"}
                n["chegou"] = store[nid].get("chegou", agora)
                n["ignorada"] = bool(store[nid].get("ignorada"))
                n["ativa"] = True

            hist = []
            for nid, ent in store.items():
                if nid in ativos_ids: continue
                h = dict(ent.get("n") or {})
                if not h.get("titulo"): continue
                h.update({"id": nid, "chegou": ent.get("chegou", ""),
                          "ignorada": bool(ent.get("ignorada")), "ativa": False})
                hist.append(h)

            if len(store) > 200:
                velhas = sorted((k for k in store if k not in ativos_ids),
                                key=lambda k: _notif_ts(store[k].get("chegou")))
                for k in velhas[:len(store) - 200]:
                    store.pop(k, None)
            _notif_store_save(store)

            todas = nots + hist
            todas.sort(key=lambda x: _notif_ts(x.get("chegou")), reverse=True)
            out["notificacoes"] = todas
            return out
        except Exception as e:
            return {"ok": False, "erro": str(e), "notificacoes": []}

    def notif_ignorar(self, nid, flag=True):
        """Marca/desmarca uma notificação como ignorada — persiste entre sessões."""
        try:
            st = _notif_store_load()
            ent = st.get(str(nid))
            if ent is None:
                st[str(nid)] = {"chegou": datetime.now().strftime("%d/%m/%Y %H:%M"),
                                "ignorada": bool(flag), "n": {}}
            else:
                ent["ignorada"] = bool(flag)
            _notif_store_save(st)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def choose_folder(self):
        try:
            r = webview.windows[0].create_file_dialog(webview.FOLDER_DIALOG, directory=PASTA_PDFS)
            return r[0] if r else None
        except Exception:
            return None

    def process_pdfs(self, pasta):
        def run():
            try:
                if not os.path.isdir(pasta):
                    _push("uiLog", "pdf", "❌  Pasta não encontrada.", "err"); _push("uiBusy", False); return
                pdfs = sorted(f for f in os.listdir(pasta) if f.lower().endswith(".pdf"))
                if not pdfs:
                    _push("uiLog", "pdf", "❌  Nenhum PDF encontrado.", "err"); _push("uiBusy", False); return
                _push("uiLog", "pdf", f"📂  {len(pdfs)} PDF(s) encontrados", "info")
                por_mes = {}
                for i, fn in enumerate(pdfs, 1):
                    d, e = extrair(os.path.join(pasta, fn))
                    if e: _push("uiLog", "pdf", f"   ⚠  {e}", "warn")
                    cls = "ok" if d["nome"] and d["rentabilidade"] else "warn"
                    _push("uiLog", "pdf",
                          f"{d['nome'] or '(sem nome)'}  ·  conta {d['conta']}  ·  "
                          f"{d['rentabilidade']}%  ·  R$ {d['ganho']}  ·  CDI {d['cdi']}%", cls)
                    _push("uiProgress", round(i * 100 / len(pdfs)))
                    por_mes.setdefault(d["mes_ref"], []).append(d)
                _push("uiLog", "pdf", "Salvando na planilha…", "info")
                with _MUT_LOCK:
                    wb = get_wb(); garantir_email(wb)
                    todos = [d for lst in por_mes.values() for d in lst]

                    lidas = []
                    nome_pdf = {}
                    for d in todos:
                        c = str(d.get("conta", "")).strip()
                        if not c: continue
                        if c not in lidas: lidas.append(c)
                        if d.get("nome") and not nome_pdf.get(c): nome_pdf[c] = d["nome"]
                    lidas_set = set(lidas)
                    cad = ler_clientes_cadastro(wb)
                    merged, vistos = [], set()
                    for c in cad:
                        conta = str(c["conta"]).strip()
                        if not conta or conta in vistos: continue
                        c["ativo"] = conta in lidas_set
                        if not str(c.get("nome") or "").strip() and nome_pdf.get(conta):
                            c["nome"] = nome_pdf[conta]
                        merged.append(c); vistos.add(conta)
                    for conta in lidas:
                        if conta not in vistos:
                            merged.append({"nome": nome_pdf.get(conta, ""), "conta": conta,
                                           "email": "", "termometro": "", "perfil": "", "ativo": True})
                            vistos.add(conta)

                    salvar_clientes_cadastro(wb, merged)

                    nomes = {str(n["conta"]): n.get("nome", "") for n in merged}
                    for lst in por_mes.values():
                        for d in lst:
                            if not d.get("nome"):
                                d["nome"] = nomes.get(str(d.get("conta", "")), "")
                    for aba, lst in por_mes.items():
                        escrever_mes(wb, aba, lst)
                        _push("uiLog", "pdf", f"✅  Aba '{aba}' — {len(lst)} cliente(s)", "ok")
                    # 'ativo' com base em TODOS os meses presentes (não só o lote atual)
                    agg = coletar_por_cliente(wb)
                    for c in merged:
                        c["ativo"] = _conta_str(c["conta"]) in agg
                    salvar_clientes_cadastro(wb, merged)
                    salvar_wb(wb)
                _push("uiLog", "pdf", f"✅  Salvo em {os.path.basename(EXCEL_PATH)}", "ok")
                inativos = [(c.get("nome") or c["conta"]) for c in merged if not c.get("ativo")]
                if inativos:
                    _push("uiInativos", inativos)
                    _push("uiLog", "pdf",
                          f"⚠  {len(inativos)} cliente(s) sem relatório nesta base (marcados em vermelho).", "warn")
                _push("uiRefresh")
            except Exception as ex:
                _push("uiLog", "pdf", f"❌  {ex}", "err")
                _push("toast", f"Não salvou: {ex}")
                try:
                    with open(_LOGF, "a", encoding="utf-8") as f: f.write(traceback.format_exc())
                except Exception: pass
            finally:
                _push("uiBusy", False)
        threading.Thread(target=run, daemon=True).start()
        return True

    def month_info(self, mes):
        try:
            wb = _wb_cached()
            cl = ler_clientes(wb, mes)
            cad = {str(x["conta"]): x for x in ler_clientes_cadastro(wb)}
            for c in cl:
                info = cad.get(str(c["conta"]), {})
                c["email"] = info.get("email", "") or buscar_email(wb, c["conta"])
                if not c.get("nome") and info.get("nome"): c["nome"] = info["nome"]
            return {"ok": True, "clientes": cl}
        except Exception as e:
            return {"ok": False, "erro": str(e), "clientes": []}

    def clients_overview(self):
        """Lista de clientes com quantos relatórios cada um possui."""
        try:
            if not _tem_dados():
                return {"ok": True, "clientes": []}
            agg = coletar_cache()
            nome_map = {str(c.get("conta", "")).strip(): (c.get("nome") or "")
                        for c in cadastro_cache()}
            lim_ap = _aporte_lim()
            aps = {}
            for a in aportes_cache():
                b = dict(a); b["_d"] = _data_br(b["data"])
                aps.setdefault(b["conta"], []).append(b)
            cts = {}
            for c_ in contatos_cache():
                b = dict(c_); b["_d"] = _data_br(b["data"])
                cts.setdefault(b["conta"], []).append(b)
            cl = []
            for conta, o in agg.items():
                meses = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)
                pat = o["meses"][meses[0]].get("pat_num") if meses else None
                fl = flags_cliente(o.get("cad", {}))
                grupo = _gf_resolve(_gf_parse(o.get("cad", {}).get("grupo_familiar")), nome_map)

                grupo = [g for g in grupo if g.get("tipo") == "cliente"]
                item = {"conta": conta, "nome": o["nome"], "email": o.get("email", ""),
                        "termometro": o.get("termometro", ""), "perfil": o.get("perfil", ""),
                        "patrimonio": pat, "n": len(meses),
                        "ultimo": mes_bonito(meses[0]) if meses else "",
                        "flags": fl, "plugado_mesa": _plugado(fl["mesa"]["tem"]),
                        "grupo_n": len(grupo),
                        "grupo_txt": " · ".join(
                            g["nome"] + (" (" + g["parentesco"] + ")" if g.get("parentesco") else "")
                            for g in grupo)}
                item.update(_aniv_pack(o.get("cad", {}).get("nascimento")))

                cad_c = o.get("cad", {})
                apos_obj = str(cad_c.get("objetivo", "")).strip()
                apos_alvo = None; apos_falta = None
                if apos_obj.lower() == "aposentadoria":
                    try:
                        apos_alvo = int(float(str(cad_c.get("apos_idade", "")).strip())) or None
                    except Exception:
                        apos_alvo = None
                    idade_c = _idade(cad_c.get("nascimento", ""))
                    if apos_alvo and idade_c is not None:
                        apos_falta = apos_alvo - idade_c
                item.update({"apos_obj": apos_obj, "apos_alvo": apos_alvo, "apos_falta": apos_falta})
                st_ap = _aportes_stats(aps.get(conta, []), lim_ap)
                item.update({"ap_n": st_ap["n"], "ap_total": st_ap["total"],
                             "ap_ultimo": st_ap["ultimo"], "ap_dias": st_ap["dias"],
                             "ap_alerta": st_ap["alerta"], "ap_lim": lim_ap})
                st_ct = _contato_stats(cts.get(conta, []))
                item.update({"ct_n": st_ct["n"], "ct_ultimo": st_ct["ultimo"],
                             "ct_tipo": st_ct["tipo"], "ct_dias": st_ct["dias"]})
                cl.append(item)
            cl.sort(key=lambda c: (not c["nome"], (c["nome"] or c["conta"]).lower()))
            return {"ok": True, "clientes": cl}
        except Exception as e:
            return {"ok": False, "erro": str(e), "clientes": []}

    def client_detail(self, conta):
        """Evolução Patrimonial de um cliente, um conjunto de cartões por mês."""
        try:
            if not _tem_dados():
                return {"ok": False, "erro": "Planilha ainda não existe.", "relatorios": []}
            agg = coletar_cache()
            nome_map = {str(c.get("conta", "")).strip(): (c.get("nome") or "")
                        for c in cadastro_cache()}
            o = agg.get(str(conta))
            if not o:
                return {"ok": False, "erro": "Cliente sem relatórios.", "relatorios": []}
            flags = flags_cliente(o.get("cad", {}))
            grupo = _gf_resolve(_gf_parse(o.get("cad", {}).get("grupo_familiar")), nome_map)
            rel = []
            for aba in sorted(o["meses"].keys(), key=_ordem_mes, reverse=True):
                c = o["meses"][aba]
                rel.append({
                    "mes": aba, "mes_bonito": mes_bonito(aba),
                    "ganho_mes": c.get("ganho_num"), "rent_mes": c.get("rent_num"),
                    "cdi_mes": c.get("cdi_num"), "movimentacoes": c.get("mov_num"),
                    "patrimonio": c.get("pat_num"),
                    "rent_total": c.get("rent_total_num"), "ganho_total": c.get("ganho_total_num"),
                    "periodo_total": c.get("periodo_total", ""),
                    "aloc": c.get("aloc") or {}, "arquivo": c.get("arquivo", "")})
            pat = rel[0]["patrimonio"] if rel else None
            ret = {"ok": True, "conta": str(conta), "nome": o["nome"],
                   "email": o.get("email", ""), "termometro": o.get("termometro", ""),
                   "perfil": o.get("perfil", ""), "patrimonio": pat,
                   "relatorios": rel, "flags": flags, "grupo": grupo,
                   "nascimento": o.get("cad", {}).get("nascimento", ""),
                   "objetivo": o.get("cad", {}).get("objetivo", ""),
                   "apos_idade": o.get("cad", {}).get("apos_idade", "")}
            ret.update(_aniv_pack(ret["nascimento"]))
            return ret
        except Exception as e:
            return {"ok": False, "erro": str(e), "relatorios": []}

    def save_cliente_extra(self, conta, dados):
        """Salva as informações extras (previdência/internacional/seguro/plugado) de UM cliente."""
        try:
            wb = get_wb(); garantir_email(wb)
            cad = ler_clientes_cadastro(wb)
            conta = str(conta).strip()
            alvo = next((c for c in cad if str(c.get("conta", "")).strip() == conta), None)
            if alvo is None:
                alvo = {"conta": conta}; cad.append(alvo)
            for chave, _ in CAD_EXTRA:
                if chave in (dados or {}):
                    alvo[chave] = str(dados.get(chave, "") or "").strip()
            agg = coletar_por_cliente(wb)
            for c in cad:
                c["ativo"] = str(c.get("conta", "")).strip() in agg
            salvar_clientes_cadastro(wb, cad)
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def save_grupo_familiar(self, conta, entradas):
        """Salva o Grupo Familiar de UM cliente.
        entradas = lista de strings: '#<conta>' (vínculo com cliente da base) ou
        texto livre (nome do familiar). Vínculos entre clientes são BIDIRECIONAIS:
        ao vincular A → B, '#A' também é gravado no campo de B (e removido ao desfazer)."""
        try:
            wb = get_wb(); garantir_email(wb)
            cad = ler_clientes_cadastro(wb)
            conta = _conta_str(conta)
            por_conta = {_conta_str(c.get("conta", "")): c for c in cad}
            agg = coletar_por_cliente(wb)
            alvo = por_conta.get(conta)
            if alvo is None:
                alvo = {"conta": conta}; cad.append(alvo); por_conta[conta] = alvo

            novas = []; vistos = set(); avisos = []
            for e in (entradas or []):
                ref, par = _gf_split(str(e or "").strip().rstrip(";"))
                par = par.replace(";", " ").replace("|", " ").strip()
                if ref.startswith("#"):
                    c2 = _conta_str(ref[1:])
                    if not c2 or c2 == conta:
                        continue
                    if c2 not in por_conta:
                        o2 = agg.get(c2)
                        if o2 is not None:
                            novo = {"conta": c2, "nome": o2.get("nome", "")}
                            cad.append(novo); por_conta[c2] = novo
                        else:
                            avisos.append(f"Conta {c2} não encontrada — vínculo não salvo.")
                            continue
                    ref = "#" + c2
                else:
                    ref = ref.replace("|", " ").strip()
                if not ref or ref.lower() in vistos:
                    continue
                vistos.add(ref.lower())
                novas.append(ref + ("|" + par if par else ""))
            antigas = _gf_parse(alvo.get("grupo_familiar"))
            alvo["grupo_familiar"] = _gf_join(novas)

            novos_l = {}
            for x in novas:
                ref, par = _gf_split(x)
                if ref.startswith("#"): novos_l[ref[1:]] = par
            antigos_l = {_gf_split(x)[0][1:] for x in antigas
                         if _gf_split(x)[0].startswith("#")}
            nome_alvo = alvo.get("nome") or ""
            for c2, par in novos_l.items():
                o2 = por_conta.get(c2)
                if o2 is None: continue
                l2 = _gf_parse(o2.get("grupo_familiar"))
                inv = _gf_inverso(par, nome_alvo)
                ent = next((x for x in l2 if _gf_split(x)[0] == "#" + conta), None)
                if ent is None:

                    l2.append("#" + conta + ("|" + inv if inv else ""))
                    o2["grupo_familiar"] = _gf_join(l2)
                elif inv and not _gf_split(ent)[1]:

                    l2[l2.index(ent)] = "#" + conta + "|" + inv
                    o2["grupo_familiar"] = _gf_join(l2)
            for c2 in antigos_l - set(novos_l):
                o2 = por_conta.get(c2)
                if o2 is None: continue
                l2 = [x for x in _gf_parse(o2.get("grupo_familiar"))
                      if _gf_split(x)[0] != "#" + conta]
                o2["grupo_familiar"] = _gf_join(l2)
            for c in cad:
                c["ativo"] = _conta_str(c.get("conta", "")) in agg
            salvar_clientes_cadastro(wb, cad)
            salvar_wb(wb)
            nome_map = {k: (v.get("nome") or "") for k, v in por_conta.items()}
            return {"ok": True, "grupo": _gf_resolve(novas, nome_map), "avisos": avisos}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def aportes_get(self, conta):
        """Histórico de aportes de um cliente + estatísticas + evolução mensal (12m)."""
        try:
            lim = _aporte_lim()
            lst = _aportes_do_cliente(conta) if _tem_dados() else []
            st = _aportes_stats(lst, lim)
            hoje = datetime.now().date()
            idx = hoje.year * 12 + (hoje.month - 1)
            mensal = []
            for i in range(11, -1, -1):
                j = idx - i; aa, mm = divmod(j, 12); mm += 1
                tot = round(sum(a["valor"] for a in lst
                                if a.get("_d") and a["_d"].year == aa and a["_d"].month == mm), 2)
                mensal.append({"label": "%02d/%s" % (mm, str(aa)[2:]), "total": tot})
            itens = [{k: v for k, v in a.items() if k != "_d"} for a in lst]
            out = {"ok": True, "itens": itens, "mensal": mensal, "alerta_dias": lim}
            out.update(st)
            return out
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def aporte_add(self, conta, data, valor, obs=""):
        """Registra um aporte manual na aba Aportes (alimenta a meta do Painel)."""
        try:
            conta = _conta_str(conta)
            if not conta: return {"ok": False, "erro": "Conta inválida."}
            try:
                v = float(valor)
            except Exception:
                s = str(valor or "").replace("R$", "").strip()
                if "," in s: s = s.replace(".", "").replace(",", ".")
                v = float(s) if s else 0
            if not v or v <= 0:
                return {"ok": False, "erro": "Informe um valor de aporte válido."}
            dstr = str(data or "").strip() or datetime.now().date().strftime("%d/%m/%Y")
            if _data_br(dstr) is None:
                return {"ok": False, "erro": "Data inválida — use dd/mm/aaaa."}
            wb = get_wb(); ws = garantir_aportes(wb)
            ri = ws.max_row + 1
            for col, val in enumerate([conta, dstr, round(v, 2), str(obs or "").strip()], 1):
                c = ws.cell(ri, col, val)
                c.font = ft(size=10)
                c.alignment = al("left" if col == 4 else "center"); c.border = bd()
            _db_snapshot_aportes(wb)
            salvar_wb(wb)
            return self.aportes_get(conta)
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def aporte_del(self, conta, rid):
        """Remove um aporte registrado (id = linha na aba Aportes)."""
        try:
            wb = get_wb()
            if APORTES_ABA not in wb.sheetnames:
                return {"ok": False, "erro": "Sem aportes registrados."}
            ws = wb[APORTES_ABA]; rid = int(rid)
            if _conta_str(ws.cell(rid, 1).value) != _conta_str(conta):
                return {"ok": False, "erro": "Aporte não encontrado — atualize e tente de novo."}
            ws.delete_rows(rid, 1)
            _db_snapshot_aportes(wb)
            salvar_wb(wb)
            return self.aportes_get(conta)
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def contatos_get(self, conta):
        """Histórico de contatos com o cliente + último contato."""
        try:
            lst = _contatos_do_cliente(conta) if _tem_dados() else []
            st = _contato_stats(lst)
            itens = [{k: v for k, v in c.items() if k != "_d"} for c in lst]
            out = {"ok": True, "itens": itens, "tipos": CONTATO_TIPOS}
            out.update(st)
            return out
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def contato_add(self, conta, data, tipo, obs=""):
        """Registra um contato (Conta, Data, Tipo, Observação) na aba Contatos."""
        try:
            conta = _conta_str(conta)
            if not conta: return {"ok": False, "erro": "Conta inválida."}
            tipo = str(tipo or "").strip() or "Outro"
            dstr = str(data or "").strip() or datetime.now().date().strftime("%d/%m/%Y")
            if _data_br(dstr) is None:
                return {"ok": False, "erro": "Data inválida — use dd/mm/aaaa."}
            wb = get_wb(); ws = garantir_contatos(wb)
            ri = ws.max_row + 1
            for col, val in enumerate([conta, dstr, tipo, str(obs or "").strip()], 1):
                c = ws.cell(ri, col, val)
                c.font = ft(size=10)
                c.alignment = al("left" if col == 4 else "center"); c.border = bd()
            _db_snapshot_contatos(wb)
            salvar_wb(wb)
            return self.contatos_get(conta)
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def contato_del(self, conta, rid):
        """Remove um contato registrado (id = linha na aba Contatos)."""
        try:
            wb = get_wb()
            if CONTATOS_ABA not in wb.sheetnames:
                return {"ok": False, "erro": "Sem contatos registrados."}
            ws = wb[CONTATOS_ABA]; rid = int(rid)
            if _conta_str(ws.cell(rid, 1).value) != _conta_str(conta):
                return {"ok": False, "erro": "Contato não encontrado — atualize e tente de novo."}
            ws.delete_rows(rid, 1)
            _db_snapshot_contatos(wb)
            salvar_wb(wb)
            return self.contatos_get(conta)
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def receita_set(self, mes, valor):
        """Define a receita de um mês (mm/aaaa; vazio = mês atual). Upsert na aba Receita."""
        try:
            hoje = datetime.now().date()
            if not str(mes or "").strip():
                aa, mm = hoje.year, hoje.month
            else:
                k = _mes_aa(mes)
                if not k: return {"ok": False, "erro": "Mês inválido — use mm/aaaa."}
                aa, mm = k
            try:
                v = float(valor)
            except Exception:
                s = str(valor or "").replace("R$", "").strip()
                if "," in s: s = s.replace(".", "").replace(",", ".")
                v = float(s) if s else 0.0
            wb = get_wb(); ws = garantir_receita(wb)
            fim = ws.max_row
            alvo = None
            if fim >= 2:
                for row in ws.iter_rows(min_row=2, max_row=fim):
                    if _mes_aa(row[0].value) == (aa, mm):
                        alvo = row[0].row; break
            ri = alvo or fim + 1
            for col, val in ((1, "%02d/%d" % (mm, aa)), (2, round(v, 2))):
                c = ws.cell(ri, col, val)
                c.font = ft(size=10); c.alignment = al("center"); c.border = bd()
            _db_snapshot_receita(wb)
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def get_config(self):
        return {"ok": True, "config": carregar_config()}

    def set_config(self, chave, valor):
        try:
            cfg = carregar_config(); cfg[str(chave)] = valor; gravar_config(cfg)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def controle_get(self):
        try:
            if not _tem_dados():
                return {"ok": True, "mensal": [], "montagem": [], "ativos": [], "clientes": []}
            wb = _wb_cached()
            mensal = ler_tabela_ctrl(wb, "mensal")
            montagem = ler_tabela_ctrl(wb, "montagem")
            ativos = ler_tabela_ctrl(wb, "ativos")
            cad = _cadastro_enriquecido(wb)
            clientes = [{"conta": c["conta"], "nome": c.get("nome", ""), "perfil": c.get("perfil", ""),
                         "termometro": (c.get("termometro", "") or "").upper(),
                         "patrimonio": c.get("patrimonio"), "plugado": _plugado(c.get("mesa_tem"))}
                        for c in cad]
            clientes.sort(key=lambda c: (not c["nome"], (c["nome"] or c["conta"]).lower()))
            return {"ok": True, "mensal": mensal,
                    "montagem": montagem, "ativos": ativos, "clientes": clientes}
        except Exception as e:
            return {"ok": False, "erro": str(e), "mensal": [], "montagem": [], "ativos": [], "clientes": []}

    def controle_save(self, chave, rows):
        try:
            if chave not in CTRL_SHEETS:
                return {"ok": False, "erro": "tabela inválida"}
            wb = get_wb(); garantir_email(wb)
            salvar_tabela_ctrl(wb, chave, rows or [])
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def deletar_inativos(self):
        """Exclui DEFINITIVAMENTE os clientes inativos (sem relatório) da planilha
        e do banco de dados — incluindo aportes, contatos e tabelas do Controle.
        Recebem lápide no banco: não são restaurados pela reconciliação."""
        try:
            wb = get_wb(); garantir_email(wb)
            cad = ler_clientes_cadastro(wb)
            agg = coletar_por_cliente(wb)
            manter, removidos = [], []
            for c in cad:
                k = _conta_str(c.get("conta"))
                if k in agg:
                    c["ativo"] = True; manter.append(c)
                else:
                    removidos.append({"conta": k, "nome": c.get("nome", "")})
            if not removidos:
                return {"ok": True, "n": 0, "removidos": [], "cadastro": manter}
            contas = {r["conta"] for r in removidos}
            salvar_clientes_cadastro(wb, manter, permitir_reducao=True)

            for aba, ler, campos, snap in (
                    (APORTES_ABA, ler_aportes, ("conta", "data", "valor", "obs"), _db_snapshot_aportes),
                    (CONTATOS_ABA, ler_contatos, ("conta", "data", "tipo", "obs"), _db_snapshot_contatos)):
                if aba in wb.sheetnames:
                    rows = [r for r in ler(wb) if r["conta"] not in contas]
                    ws = wb[aba]
                    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                        for cell in row: cell.value = None
                    ri = 2
                    for r in rows:
                        for j, k in enumerate(campos, 1):
                            ws.cell(ri, j, r.get(k, ""))
                        ri += 1
                    snap(wb)

            for chave in ("mensal", "montagem", "ativos"):
                rows = [r for r in ler_tabela_ctrl(wb, chave)
                        if _conta_str(r.get("conta")) not in contas]
                salvar_tabela_ctrl(wb, chave, rows)
            salvar_wb(wb)
            _db_log(f"deletar_inativos: {len(removidos)} cliente(s) excluído(s): "
                    + ", ".join(r["conta"] for r in removidos))
            return {"ok": True, "n": len(removidos), "removidos": removidos, "cadastro": manter}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def save_template(self, assunto, corpo):
        try:
            if not _tem_dados(): return {"ok": False}

            wb = get_wb(); garantir_email(wb)
            ws = wb[EMAIL_ABA]; ws["B1"].value = assunto; ws["B3"].value = corpo
            _db_kv_set("assunto", assunto); _db_kv_set("corpo", corpo)
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def preview_email(self, cliente, assunto, corpo):
        try:
            extra = _extra_vars(cliente)
            dados = _SafeDict({**_dados_format(cliente), **extra})
            try: ass = assunto.format_map(dados)
            except Exception: ass = assunto
            return {"ok": True, "assunto": ass, "email": cliente.get("email", ""),
                    "html": render_email_html(cliente, corpo, extra)}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def send_emails(self, mes, assunto, corpo, pasta=None):
        def run():
            try:
                import pythoncom; pythoncom.CoInitialize()
            except Exception:
                pass
            try:
                with _MUT_LOCK:
                    wb = _wb_cached()
                    cl = ler_clientes(wb, mes)
                    cad = {str(x["conta"]): x for x in ler_clientes_cadastro(wb)}
                    for c in cl:
                        info = cad.get(str(c["conta"]), {})
                        c["email"] = info.get("email", "") or buscar_email(wb, c["conta"])
                        if not c.get("nome") and info.get("nome"): c["nome"] = info["nome"]
                    mdados = ler_mercado(wb)
                com = [c for c in cl if c["email"]]
                sem = [c["nome"] for c in cl if not c["email"]]
                _push("uiLog", "mail", f"📧  Enviando {len(com)} e-mail(s)…", "info")
                if sem:
                    _push("uiLog", "mail",
                          f"⏭  Sem e-mail (ignorados): {', '.join(sem[:5])}{'…' if len(sem) > 5 else ''}", "warn")

                por_email = {}
                for c in com:
                    por_email.setdefault(c["email"].strip().lower(), []).append(c)
                for em, cls_ in por_email.items():
                    if len(cls_) > 1:
                        _push("uiLog", "mail",
                              f"⚠  MESMO E-MAIL em {len(cls_)} clientes ({em}): "
                              + ", ".join(x["nome"] or x["conta"] for x in cls_)
                              + " — confira o cadastro (cada um receberá o PRÓPRIO relatório).",
                              "warn")
                ok = err = 0
                for c in com:
                    extra = _extra_vars(c, mdados)
                    dados = _SafeDict({**_dados_format(c), **extra})

                    anexo = _achar_pdf(c.get("arquivo"), pasta)
                    if c.get("arquivo") and not anexo:
                        _push("uiLog", "mail",
                              f"⚠  PDF não encontrado ({c.get('arquivo')}) — {c['nome']} será enviado SEM anexo.",
                              "warn")

                    if anexo:
                        m_ax = re.search(r"XPerformance\s*-\s*(\d+)\s*-", os.path.basename(anexo), re.I)
                        if m_ax and m_ax.group(1) != str(c.get("conta", "")).strip():
                            _push("uiLog", "mail",
                                  f"🛑  BLOQUEADO: o anexo '{os.path.basename(anexo)}' é da conta "
                                  f"{m_ax.group(1)}, não de {c['nome']} (conta {c.get('conta')}). "
                                  f"Enviado SEM anexo por segurança.", "err")
                            anexo = None
                    try:
                        enviar_via_outlook(c["email"], assunto.format_map(dados),
                                           render_email_html(c, corpo, extra), anexo)
                        marca = "  📎" if anexo else ""
                        _push("uiLog", "mail", f"✅  {c['email']}  ({c['nome']}){marca}", "ok"); ok += 1
                    except Exception as ex:
                        _push("uiLog", "mail", f"❌  {c['email']}: {ex}", "err"); err += 1
                _push("uiLog", "mail", f"{'✅' if err == 0 else '⚠'}  {ok} enviado(s), {err} erro(s).",
                      "ok" if err == 0 else "warn")
            except Exception as ex:
                _push("uiLog", "mail", f"❌  {ex}", "err")
            finally:
                _push("uiBusy", False)
        threading.Thread(target=run, daemon=True).start()
        return True

    def market_get(self):
        try:
            if not _tem_dados():
                return {"ok": True, "dias": [], "meses": [], "acoes": [], "prev": {},
                        "atualizado": ""}
            d = mercado_cache()
            d["ok"] = True
            return d
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def market_update(self):
        """Atualiza indicadores de mercado on-line. A busca HTTP (lenta, ~25 chamadas)
        roda em thread FORA do _MUT_LOCK para não congelar a interface; só a gravação
        entra no lock. O resultado chega à tela via _push('uiMercado', ...).
        (Endpoint atualmente não acionado pela interface — pronto para uso futuro.)"""
        def run():
            try:
                novo = fetch_mercado()                 # rede: fora de qualquer lock
                with _MUT_LOCK:
                    d = atualizar_mercado_excel(novo)  # só a gravação sob o lock
                d["ok"] = True
                _push("uiMercado", d)
            except Exception as ex:
                _push("uiMercado", {"ok": False, "erro": str(ex)})
        threading.Thread(target=run, daemon=True).start()
        return {"ok": True, "async": True}

    def aloc_get(self):
        try:
            if not _tem_dados():
                return {"ok": True, "dados": {}}
            d = aloc_cache()
            return {"ok": True, "dados": d}
        except Exception as e:
            return {"ok": False, "erro": str(e), "dados": {}}

    def aloc_load(self):
        """Abre um PDF de Carteiras Recomendadas, extrai e substitui os dados salvos."""
        try:
            r = webview.windows[0].create_file_dialog(
                webview.OPEN_DIALOG, directory=_HERE, file_types=("PDF (*.pdf)",))
            if not r:
                return {"ok": False, "cancel": True}
            path = r[0] if isinstance(r, (list, tuple)) else r
            d, e = extrair_alocacao(path)
            if e:
                return {"ok": False, "erro": e}
            if not d.get("classes"):
                return {"ok": False, "erro": "Não reconheci as carteiras recomendadas neste PDF."}
            wb = get_wb(); garantir_email(wb)
            escrever_alocacao(wb, d)
            if MERCADO_ABA in wb.sheetnames:
                del wb[MERCADO_ABA]
            salvar_wb(wb)
            return {"ok": True, "dados": d}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def oi_get(self):
        try:
            if not _tem_dados():
                return {"ok": True, "dados": {}}
            d = onde_cache()
            return {"ok": True, "dados": d}
        except Exception as e:
            return {"ok": False, "erro": str(e), "dados": {}}

    def oi_load(self):
        """Lê um PDF 'Onde Investir' e guarda o resumo — NÃO mexe nos dados de Alocação."""
        try:
            r = webview.windows[0].create_file_dialog(
                webview.OPEN_DIALOG, directory=_HERE, file_types=("PDF (*.pdf)",))
            if not r:
                return {"ok": False, "cancel": True}
            path = r[0] if isinstance(r, (list, tuple)) else r
            d, e = extrair_onde_investir(path)
            if e:
                return {"ok": False, "erro": e}
            wb = get_wb(); garantir_email(wb)
            escrever_onde_investir(wb, d)
            salvar_wb(wb)
            return {"ok": True, "dados": d}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def mercado_load(self):
        """Lê VÁRIOS PDFs de mercado de uma vez e roteia cada um para sua aba
        (Carteiras Recomendadas ou Onde Investir), sem um substituir o outro."""
        try:
            r = webview.windows[0].create_file_dialog(
                webview.OPEN_DIALOG, directory=_HERE, allow_multiple=True,
                file_types=("PDF (*.pdf)",))
            if not r:
                return {"ok": False, "cancel": True}
            paths = list(r) if isinstance(r, (list, tuple)) else [r]
            wb = get_wb(); garantir_email(wb)
            carregados, ignorados = [], []
            for p in paths:
                t = detectar_tipo_relatorio(p)
                nome = os.path.basename(p)
                if t == "aloc":
                    d, e = extrair_alocacao(p)
                    if e or not d.get("classes"):
                        ignorados.append(nome); continue
                    escrever_alocacao(wb, d)
                    if MERCADO_ABA in wb.sheetnames:
                        del wb[MERCADO_ABA]
                    carregados.append("Carteiras Recomendadas (" + (d.get("mes_label") or "") + ")")
                elif t == "onde":
                    d, e = extrair_onde_investir(p)
                    if e:
                        ignorados.append(nome); continue
                    escrever_onde_investir(wb, d)
                    carregados.append("Onde Investir (" + (d.get("mes_label") or "") + ")")
                else:
                    ignorados.append(nome)
            salvar_wb(wb)
            wb2 = _wb_cached()
            aloc = ler_alocacao(wb2); onde = ler_onde_investir(wb2); wb2.close()
            return {"ok": True, "carregados": carregados, "ignorados": ignorados,
                    "aloc": aloc, "onde": onde}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def carteira_comparativo(self, conta):
        """Compara a carteira atual do cliente com as 3 carteiras recomendadas."""
        try:
            if not _tem_dados():
                return {"ok": False, "erro": "Sem planilha."}
            dados_aloc = aloc_cache()
            agg = coletar_cache()
            if not dados_aloc.get("classes"):
                return {"ok": False, "erro": "Carregue o relatório de alocação na aba Mercado."}
            o = agg.get(str(conta))
            if not o:
                return {"ok": False, "erro": "Cliente sem relatórios."}
            meses = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)
            ultimo = o["meses"][meses[0]]
            pf = PERFIL_KEY.get(_norm_cls(o.get("perfil", "")))
            comp = comparar_carteira(ultimo.get("aloc") or {}, dados_aloc, pf)
            comp.update({"ok": True, "conta": str(conta), "nome": o["nome"],
                         "perfil_cadastro": o.get("perfil", ""), "termometro": o.get("termometro", ""),
                         "patrimonio": ultimo.get("pat_num"), "mes_cliente": mes_bonito(meses[0])})
            return comp
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def analise_carteira(self, conta):
        """Análise profunda: aderência da carteira às recomendações + ajustes sugeridos."""
        try:
            if not _tem_dados():
                return {"ok": False, "erro": "Sem planilha."}
            dados_aloc = aloc_cache()
            oi = onde_cache()
            agg = coletar_cache()
            if not dados_aloc.get("classes"):
                return {"ok": False, "erro": "Carregue o relatório de alocação na aba Mercado."}
            o = agg.get(str(conta))
            if not o:
                return {"ok": False, "erro": "Cliente sem relatórios."}
            meses = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)
            pf = PERFIL_KEY.get(_norm_cls(o.get("perfil", "")))
            comp = comparar_carteira(o["meses"][meses[0]].get("aloc") or {}, dados_aloc, pf)
            an = gerar_analise(comp, dados_aloc, oi)
            an.update({"ok": True, "conta": str(conta), "nome": o["nome"],
                       "perfil_cadastro": o.get("perfil", ""),
                       "perfil_definido": comp.get("perfil_definido", False),
                       "mes_cliente": mes_bonito(meses[0])})
            return an
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def dashboard_data(self, incluir_plugados=True):
        """Agrega indicadores da carteira de clientes para o painel executivo.
        incluir_plugados=False remove os clientes 'plugados na mesa' APENAS do
        gráfico de enquadramento das carteiras — os demais agregados sempre
        consideram todos os clientes."""
        try:
            if not _tem_dados():
                return {"ok": True, "total": 0}
            cad = cadastro_cache()
            agg = coletar_cache()
            dados_aloc = aloc_cache()
            tem_aloc = bool(dados_aloc.get("classes"))
            TO = ["A", "B", "C", "D", "E"]
            termo_qtd = {k: 0 for k in TO}; termo_rent = {k: [] for k in TO}
            perfil = {"Conservadora": 0, "Moderada": 0, "Sofisticada": 0, "—": 0}
            estado_qtd = {}; faixa = {"<30": 0, "30–44": 0, "45–59": 0, "60+": 0, "—": 0}
            enq = {"Bom": 0, "Médio": 0, "Ruim": 0, "Crítico": 0, "s/ dados": 0}
            pat_list = []; rents = []; rents_tot = []
            sem_email = sem_perfil = sem_termo = inativos = 0
            sem_seg = sem_prev = sem_intl = sem_mesa = 0
            for c in cad:
                conta = str(c["conta"]); o = agg.get(conta)
                if not o: inativos += 1
                if not c.get("email"): sem_email += 1

                if _estado_flag(c.get("seg_tem")) in ("nao", ""): sem_seg += 1
                if _estado_flag(c.get("prev_tem")) in ("nao", ""): sem_prev += 1
                if _estado_flag(c.get("intl_tem")) in ("nao", ""): sem_intl += 1
                if not _plugado(c.get("mesa_tem")): sem_mesa += 1
                pf_txt = c.get("perfil", "")
                if not pf_txt: sem_perfil += 1
                t = (c.get("termometro", "") or "").upper()
                if not t: sem_termo += 1
                if t in termo_qtd: termo_qtd[t] += 1
                pn = _perfil_nome(pf_txt); perfil[pn if pn else "—"] += 1
                uf = (c.get("estado", "") or "").strip().upper()
                if uf: estado_qtd[uf] = estado_qtd.get(uf, 0) + 1
                idade = _idade(c.get("nascimento", ""))
                if idade is None: faixa["—"] += 1
                elif idade < 30: faixa["<30"] += 1
                elif idade < 45: faixa["30–44"] += 1
                elif idade < 60: faixa["45–59"] += 1
                else: faixa["60+"] += 1
                if o and o["meses"]:
                    ult = sorted(o["meses"].keys(), key=_ordem_mes, reverse=True)[0]
                    m = o["meses"][ult]
                    pat = m.get("pat_num"); rent = m.get("rent_num")
                    rtot = m.get("rent_total_num"); aloc = m.get("aloc") or {}
                    if pat is not None: pat_list.append((c.get("nome") or conta, pat))
                    if rent is not None: rents.append(rent)
                    if rtot is not None: rents_tot.append(rtot)

                    if t in termo_rent and rtot is not None: termo_rent[t].append(rtot)
                    if tem_aloc:
                        if not incluir_plugados and _plugado(c.get("mesa_tem")):
                            pass
                        elif aloc:
                            comp = comparar_carteira(aloc, dados_aloc, PERFIL_KEY.get(_norm_cls(pf_txt)))
                            ad = comp.get("aderencia")
                            if ad is None: enq["s/ dados"] += 1
                            elif ad >= 80: enq["Bom"] += 1
                            elif ad >= 60: enq["Médio"] += 1
                            elif ad >= 40: enq["Ruim"] += 1
                            else: enq["Crítico"] += 1
                        else:
                            enq["s/ dados"] += 1

            ap_mes = []; ap_termo = {k: 0 for k in TO}; ap_cli = 0
            try:
                hoje_d = datetime.now().date()
                idx = hoje_d.year * 12 + (hoje_d.month - 1)
                termo_por_conta = {str(c["conta"]).strip(): (c.get("termometro", "") or "").upper()
                                   for c in cad}
                por_mes = {}; ult12 = set()
                for a in aportes_cache():
                    conta_a = str(a["conta"]).strip()
                    if conta_a not in termo_por_conta: continue
                    d_ = _data_br(a.get("data"))
                    if not d_: continue
                    j = d_.year * 12 + (d_.month - 1)
                    if idx - 11 <= j <= idx:
                        por_mes.setdefault(j, set()).add(conta_a)
                        ult12.add(conta_a)
                for i in range(11, -1, -1):
                    j = idx - i; aa, mm = divmod(j, 12); mm += 1
                    ap_mes.append({"label": "%02d/%s" % (mm, str(aa)[2:]), "n": len(por_mes.get(j, ()))})
                for conta_a in ult12:
                    t_ = termo_por_conta.get(conta_a, "")
                    if t_ in ap_termo: ap_termo[t_] += 1
                ap_cli = len(ult12)
            except Exception:
                pass
            n_rent_termo = sum(len(v) for v in termo_rent.values())
            termo_rent_avg = {k: (round(sum(v) / len(v), 2) if v else None) for k, v in termo_rent.items()}
            aum = round(sum(p for _, p in pat_list), 2)
            ticket = round(aum / len(pat_list), 2) if pat_list else 0
            top5 = sorted(pat_list, key=lambda x: -x[1])[:5]
            top5_pct = round(sum(p for _, p in top5) / aum * 100, 1) if aum else 0
            estados = sorted(estado_qtd.items(), key=lambda x: -x[1])
            rent_media = round(sum(rents) / len(rents), 2) if rents else None
            rent_media_tot = round(sum(rents_tot) / len(rents_tot), 2) if rents_tot else None
            return {"ok": True, "total": len(cad), "aum": aum, "ticket": ticket,
                    "rent_media": rent_media, "rent_media_tot": rent_media_tot,
                    "termo_qtd": termo_qtd, "termo_rent": termo_rent_avg, "perfil": perfil,
                    "estados": estados, "faixa": faixa, "enq": enq, "tem_aloc": tem_aloc,
                    "top5": [{"nome": n, "pat": p} for n, p in top5], "top5_pct": top5_pct,
                    "pend": {"email": sem_email, "perfil": sem_perfil, "termo": sem_termo, "inativos": inativos,
                             "seg": sem_seg, "prev": sem_prev, "intl": sem_intl, "mesa": sem_mesa},
                    "ap_mes": ap_mes, "ap_termo": ap_termo, "ap_cli": ap_cli,
                    "n_rent_termo": n_rent_termo}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def save_cadastro(self, rows):
        try:
            wb = get_wb(); garantir_email(wb)
            agg = coletar_por_cliente(wb)
            for r in rows:
                r["ativo"] = str(r.get("conta", "")).strip() in agg
            salvar_clientes_cadastro(wb, rows)
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def delete_cliente(self, conta):
        """Exclusão INTENCIONAL de um cliente (remove da planilha e cria lápide
        no banco). É o único caminho de rotina que reduz o cadastro — as demais
        gravações preservam todas as contas (proteção do Bug A)."""
        try:
            wb = get_wb()
            if EMAIL_ABA not in wb.sheetnames:
                return {"ok": True}
            alvo = _conta_str(conta)
            cad = [c for c in ler_clientes_cadastro(wb)
                   if _conta_str(c.get("conta")) != alvo]
            agg = coletar_por_cliente(wb)
            for c in cad:
                c["ativo"] = _conta_str(c.get("conta")) in agg
            salvar_clientes_cadastro(wb, cad, permitir_reducao=True)
            salvar_wb(wb)
            return {"ok": True}
        except Exception as e:
            return {"ok": False, "erro": str(e)}

    def cadastro_get(self):
        """Recarrega o cadastro enriquecido (nome/conta/email/termômetro/perfil/patrimônio/ativo)."""
        try:
            if not _tem_dados():
                return {"ok": True, "cadastro": []}
            wb = _wb_cached()
            cad = _cadastro_enriquecido(wb)
            return {"ok": True, "cadastro": cad}
        except Exception as e:
            return {"ok": False, "erro": str(e), "cadastro": []}

_MUT_LOCK = threading.RLock()

def _serializar_api():
    def mk(f):
        def w(*a, **k):
            with _MUT_LOCK, _wb_pin():
                return f(*a, **k)
        w.__name__ = getattr(f, "__name__", "api")
        return w
    for nome, fn in list(vars(Api).items()):
        if not nome.startswith("_") and callable(fn):
            setattr(Api, nome, mk(fn))

_serializar_api()

HTML = r"""<!DOCTYPE html>
<html lang="pt-BR"><head><meta charset="utf-8">
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Manrope:wght@400;500;600;700;800&family=JetBrains+Mono:wght@400;500;600;700&display=swap" rel="stylesheet">
<style>
:root{
  --yellow:#F2C029; --yellow-hi:#FFD24D;
  --bg:#0A0A0C; --panel:#101013; --card:#16161B; --card2:#1C1C22;
  --border:#26262E; --border-hi:#34343E;
  --txt:#F2F2F4; --mut:#8E8E98; --mut2:#5C5C66;
  --ok:#10B981; --warn:#F59E0B; --err:#EF4444;
}
*{box-sizing:border-box;margin:0;padding:0}
html,body{height:100%}
body{background:var(--bg);color:var(--txt);font:14px/1.5 'Manrope',"Segoe UI",system-ui,sans-serif;
     display:flex;overflow:hidden;-webkit-font-smoothing:antialiased}
::-webkit-scrollbar{width:10px;height:10px}
::-webkit-scrollbar-thumb{background:rgba(255,255,255,.09);border-radius:6px;border:2px solid transparent;background-clip:padding-box}
::-webkit-scrollbar-thumb:hover{background:rgba(255,255,255,.16);background-clip:padding-box}
::-webkit-scrollbar-track{background:transparent}
::selection{background:rgba(242,192,41,.28);color:#fff}
@keyframes xpfade{from{opacity:.4;transform:translateY(7px)}to{opacity:1;transform:none}}

aside{width:236px;min-width:236px;background:linear-gradient(180deg,#0E0E12,#0B0B0E);
      border-right:1px solid rgba(255,255,255,.06);display:flex;flex-direction:column;padding:22px 14px 16px}
.logo{display:flex;align-items:center;gap:11px;padding:2px 8px 18px}
.logo-mark{width:40px;height:40px;border-radius:12px;background:#F2C029;
  display:flex;align-items:center;justify-content:center;font-weight:800;font-size:16px;color:#141414;
  letter-spacing:-.5px;box-shadow:0 6px 20px -6px rgba(242,192,41,.5)}
.logo-name{font-weight:700;font-size:14.5px;letter-spacing:.2px}
.logo-name small{display:block;font-weight:400;font-size:11px;color:var(--mut)}
nav{margin-top:8px;display:flex;flex-direction:column;gap:2px;flex:1;overflow-y:auto}
.nav-sec{font:700 10px/1 'Manrope',sans-serif;letter-spacing:.16em;text-transform:uppercase;color:#56565F;padding:14px 12px 8px}
.nav-sec:first-child{padding-top:4px}
.nav-item{display:flex;align-items:center;gap:12px;padding:10px 12px;border-radius:10px;
  color:#8C8C96;cursor:pointer;font-weight:600;font-size:13.5px;transition:all .15s;user-select:none}
.nav-item svg{width:17px;height:17px;stroke:currentColor;fill:none;stroke-width:1.8;flex:none}
.nav-item:hover{color:#E4E6EA;background:rgba(255,255,255,.035)}
.nav-item.sel{color:var(--yellow);background:rgba(242,192,41,.08);box-shadow:inset 2px 0 0 var(--yellow)}
.side-foot{margin-top:8px;padding:12px 10px 4px;border-top:1px solid rgba(255,255,255,.06);
  display:flex;align-items:center;gap:10px}
.side-av{width:34px;height:34px;border-radius:50%;flex:none;background:linear-gradient(150deg,#26262E,#16161B);
  display:flex;align-items:center;justify-content:center;font:700 12px 'Manrope';color:var(--yellow);
  border:1px solid rgba(255,255,255,.08)}
.side-nm{font:600 13px 'Manrope';color:#E4E6EA;line-height:1.3}
.side-nm small{display:block;font-weight:500;font-size:11px;color:#76767F}

main{flex:1;display:flex;flex-direction:column;overflow:hidden;min-width:0}
header{padding:22px 32px 18px;border-bottom:1px solid rgba(255,255,255,.06);display:flex;align-items:center;gap:24px;flex:none}
header h1{font-size:20px;font-weight:700;letter-spacing:-.01em}
header .sub{font-size:12.5px;color:#76767F;margin-top:2px}
.hsearch{display:flex;align-items:center;gap:7px;background:#1a1b1d;border:1px solid rgba(255,255,255,.06);
  border-radius:999px;padding:8px 16px;color:#76767F;width:240px;font-size:12.5px}
.hsearch input{border:none;background:transparent;color:var(--txt);padding:0;width:100%;font-size:12.5px}
.hsearch input::placeholder{color:#76767F}
.hsearch kbd{font:500 11px 'JetBrains Mono';color:#4C4C55;border:1px solid rgba(255,255,255,.1);
  border-radius:5px;padding:1px 5px;flex:none}
.page{flex:1;overflow-y:auto;padding:26px 32px 40px;display:none;
  background-image:radial-gradient(rgba(255,255,255,.018) 1px,transparent 1px);background-size:24px 24px}
.page.sel{display:flex;flex-direction:column;gap:24px;animation:xpfade .32s cubic-bezier(.22,.61,.36,1)}

.card{background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));border:1px solid rgba(255,255,255,.05);border-radius:28px;padding:24px;box-shadow:0 1px 2px rgba(0,0,0,.4),0 34px 70px -46px rgba(0,0,0,.95);transition:border-color .3s}
.card:hover{border-color:rgba(242,192,41,.12)}
.card h3{font-size:12px;text-transform:uppercase;letter-spacing:.08em;color:var(--mut);
  font-weight:600;margin-bottom:12px}
label.fld{display:block;font-size:12px;color:var(--mut);margin:0 0 5px 2px}
input,select,textarea{width:100%;background:#1a1b1d;border:1px solid rgba(255,255,255,.08);
  color:#E4E6EA;border-radius:14px;padding:9px 14px;font:inherit;outline:none;transition:border .15s,box-shadow .15s,background .15s}
input::placeholder,textarea::placeholder{color:#76767F}
input:hover,select:hover,textarea:hover{border-color:rgba(255,255,255,.16)}
input:focus,select:focus,textarea:focus{border-color:rgba(242,192,41,.55);box-shadow:0 0 0 3px rgba(242,192,41,.1)}
select option{background:#15151A;color:#E4E6EA}
  select option:checked{background:#F2C029;color:#141414}
  select{color-scheme:dark;appearance:none;background-image:url("data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='10' height='6'%3E%3Cpath d='M1 1l4 4 4-4' stroke='%238E8E98' fill='none' stroke-width='1.6'/%3E%3C/svg%3E");
  background-repeat:no-repeat;background-position:right 12px center;padding-right:30px;cursor:pointer}
textarea{resize:vertical;line-height:1.55}

.btn{display:inline-flex;align-items:center;justify-content:center;gap:8px;border:none;
  border-radius:999px;padding:9px 18px;font:600 13px 'Manrope',sans-serif;cursor:pointer;
  transition:all .2s;user-select:none;white-space:nowrap}
.btn:active{transform:scale(.97)}
.btn-primary{background:var(--yellow);color:#141414;font-weight:700}
.btn-primary:hover{background:var(--yellow-hi);box-shadow:0 0 15px rgba(242,192,41,.25)}
.btn-ghost{background:transparent;color:#C9CBD1;border:1px solid rgba(255,255,255,.1)}
.btn-ghost:hover{border-color:rgba(255,255,255,.2);color:#E4E6EA;background:rgba(255,255,255,.04)}
.btn-big{padding:13px 28px;font-size:14px;border-radius:999px}
.btn[disabled]{opacity:.45;pointer-events:none}

.console{background:rgba(13,14,16,.9);border:1px solid rgba(255,255,255,.05);border-radius:20px;padding:14px 18px;
  font:12.5px/1.65 Consolas,monospace;overflow-y:auto;flex:1;min-height:160px;white-space:pre-wrap}
.console .ok{color:var(--ok)} .console .err{color:var(--err)}
.console .warn{color:var(--warn)} .console .info{color:#3B82F6}
.progress{height:4px;background:var(--border);border-radius:3px;overflow:hidden;margin-top:12px;
  opacity:0;transition:opacity .2s}
.progress.on{opacity:1}
.progress>div{height:100%;width:0;background:var(--yellow);border-radius:3px;transition:width .25s}

.row{display:flex;gap:10px;align-items:center}
.grow{flex:1}
.chips{display:flex;flex-wrap:wrap;gap:7px;margin-top:10px}
.chip{background:rgba(242,192,41,.1);border:1px solid rgba(242,192,41,.3);color:var(--yellow);
  font:500 12px 'JetBrains Mono',Consolas,monospace;border-radius:8px;padding:4px 10px;cursor:pointer;transition:all .15s}
.chip:hover{background:rgba(242,192,41,.22)}
.hint{font-size:11.5px;color:var(--mut2);margin-top:8px}
.save-state{font-size:12px;color:var(--mut2);transition:color .2s}
.save-state.saved{color:var(--ok)}

.mail-grid{display:grid;grid-template-columns:minmax(380px,46%) 1fr;gap:16px;flex:1;min-height:0}
.mail-col{display:flex;flex-direction:column;gap:14px;min-height:0}
.preview-card{display:flex;flex-direction:column;flex:1;min-height:380px}
.preview-head{display:flex;align-items:center;gap:9px;margin-bottom:10px}
.preview-head select,.preview-head input{flex:1;min-width:0}
.cbx-list{position:fixed;z-index:9999;background:#1b1c1e;border:1px solid rgba(255,255,255,.1);
  border-radius:16px;max-height:280px;overflow-y:auto;box-shadow:0 18px 44px rgba(0,0,0,.6);padding:6px;display:none}
.cbx-it{display:flex;justify-content:space-between;align-items:center;gap:14px;
  padding:9px 13px;border-radius:10px;cursor:pointer;font-size:13.5px;color:#e3e2e4}
.cbx-it.on{background:#2a2b2e;color:var(--yellow)}
.cbx-it .cbx-nm{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.cbx-it .cbx-ct{font:600 11.5px 'Manrope';color:#9c948a;flex:none}
.cbx-it.on .cbx-ct{color:rgba(242,192,41,.75)}
.cbx-vazio{padding:12px 14px;color:#9c948a;font-size:13px}
.cf-it{display:flex;align-items:center;gap:9px;padding:7px 6px;border-radius:8px;cursor:pointer;font-size:13px;color:#e3e2e4}
.cf-it:hover{background:#2a2b2e}
.cf-it input{accent-color:var(--yellow);width:14px;height:14px;cursor:pointer;margin:0;flex:none}
.iconbtn{width:36px;height:36px;border-radius:999px;border:1px solid rgba(255,255,255,.1);background:transparent;
  color:var(--mut);cursor:pointer;font-size:15px;transition:all .15s;flex:none}
.iconbtn:hover{color:var(--yellow);border-color:var(--yellow)}
.preview-meta{font-size:12px;color:var(--mut);margin-bottom:8px;line-height:1.6;
  border-left:3px solid var(--yellow);padding-left:10px}
.preview-meta b{color:var(--txt);font-weight:600}
iframe{flex:1;width:100%;border:1px solid var(--border);border-radius:12px;background:#f4f5f7;min-height:300px}

.cli-grid{display:grid;grid-template-columns:minmax(190px,2fr) 84px 1fr 1.1fr 2.6fr 1fr 30px;align-items:center;gap:14px}
.clihdr{padding:0 16px 12px;border-bottom:1px solid rgba(255,255,255,.06)}
.clihdr span{font:600 11.5px 'Manrope',sans-serif;letter-spacing:.07em;text-transform:uppercase;color:#7A7A84}
.cli-row{padding:15px 16px;border-radius:12px;cursor:pointer;border:1px solid transparent;transition:background .15s}
.cli-row:hover{background:rgba(242,192,41,.05)}
.cli-av{width:42px;height:42px;border-radius:11px;flex:none;background:linear-gradient(150deg,#222229,#15151A);
  border:1px solid rgba(255,255,255,.08);display:flex;align-items:center;justify-content:center;font:700 14px 'Manrope';color:#C9CBD1}
#cliLista,#cliCad{background:linear-gradient(180deg,#15151A,#0F0F12)}
.chip-save{font-size:11.5px;color:#4C8C6A;background:rgba(5,150,105,.1);border:1px solid rgba(5,150,105,.2);
  border-radius:6px;padding:2px 8px;white-space:nowrap;flex:none}
table.cli{width:100%;border-collapse:separate;border-spacing:0}
table.cli th{font:600 11.5px 'Manrope',sans-serif;letter-spacing:.07em;text-transform:uppercase;color:#7A7A84;
  text-align:left;padding:0 12px 12px;border-bottom:1px solid rgba(255,255,255,.06)}
table.cli td{padding:11px 12px;border-bottom:1px solid rgba(255,255,255,.05);font-size:13px}
table.cli tbody tr{transition:background .15s}
table.cli tbody tr:hover td{background:rgba(242,192,41,.04)}
table.cli input{background:transparent;border:1px solid transparent;border-radius:7px;padding:7px 9px}
table.cli input:hover{border-color:rgba(255,255,255,.12);background:rgba(255,255,255,.025)}
table.cli input:focus{border-color:rgba(242,192,41,.55);background:#0C0C0F}
.del{color:var(--mut2);cursor:pointer;font-size:15px;padding:4px 8px;border-radius:7px;transition:all .15s}
.del:hover{color:var(--err);background:rgba(239,68,68,.1)}
.badge{display:inline-block;font-size:11.5px;border-radius:20px;padding:3px 11px;font-weight:600}
.badge.ok{background:rgba(16,185,129,.12);color:var(--ok)}
.badge.warn{background:rgba(245,158,11,.12);color:var(--warn)}

.pills{display:inline-flex;background:#1f2022;border:1px solid rgba(255,255,255,.06);border-radius:999px;padding:4px;gap:2px}
.pill{padding:8px 18px;border-radius:999px;font:600 12.5px 'Manrope',sans-serif;color:#9c948a;
  cursor:pointer;transition:all .2s;user-select:none;white-space:nowrap}
.pill:hover{color:#E4E6EA}
.pill.sel{background:#343537;color:#F2C029;box-shadow:0 1px 3px rgba(0,0,0,.4)}
.pills.seg{border-radius:10px;padding:3px;gap:2px}
.pills.seg .pill{padding:6px 13px;border-radius:7px;font-size:12.5px}
.kpis{display:grid;grid-template-columns:repeat(4,1fr);gap:14px}
@media(max-width:1100px){.kpis{grid-template-columns:repeat(2,1fr)}}
.kpi{background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));border:1px solid rgba(255,255,255,.05);border-radius:24px;
  padding:18px 20px;box-shadow:0 1px 2px rgba(0,0,0,.4),0 34px 70px -46px rgba(0,0,0,.95);
  transition:border-color .22s,transform .22s}
.kpi:hover{border-color:rgba(242,192,41,.16);transform:translateY(-2px)}
.kpi .lbl{font:700 10px 'Manrope',sans-serif;text-transform:uppercase;letter-spacing:.12em;color:#6A6A74}
.kpi .val{font:700 25px/1 'JetBrains Mono',monospace;margin-top:10px;letter-spacing:-.02em;color:#F2F3F5;white-space:nowrap}
.kpi .val small{font-size:12px;font-weight:400;color:var(--mut)}
.kpi.hl{background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));border-color:rgba(242,192,41,.2)}
.kpi.hl .lbl{color:#a8977b}
.kpi.hl .val{color:var(--yellow)}
.det-grid{display:grid;grid-template-columns:1.3fr 1fr;gap:24px;align-items:start}
@media(max-width:1100px){.det-grid{grid-template-columns:1fr}}
.bn-grid{display:grid;grid-template-columns:repeat(auto-fit,minmax(232px,1fr));gap:14px;margin-bottom:4px}
.bn-card{background:linear-gradient(180deg,#17171C,#111115);border:1px solid rgba(255,255,255,.07);
  border-radius:16px;padding:20px;position:relative;overflow:hidden;min-height:120px;
  box-shadow:0 1px 2px rgba(0,0,0,.4),0 24px 48px -34px rgba(0,0,0,.9)}
.bn-lbl{font:700 11px 'Manrope',sans-serif;text-transform:uppercase;letter-spacing:.13em;color:#7A7A84}
.bn-big{font:800 40px/1 'Manrope',sans-serif;letter-spacing:-.035em;color:var(--yellow);margin-top:12px}
.bn-sub{font-size:11px;color:var(--mut)}
.bn-track{height:7px;background:var(--border);border-radius:5px;overflow:hidden}
.bn-fill{height:100%;border-radius:5px;background:linear-gradient(90deg,var(--yellow),var(--yellow-hi))}
.bn-spark{margin-top:16px;height:34px;border-radius:9px;
  background:linear-gradient(180deg,rgba(242,192,41,.22),rgba(242,192,41,0))}
.bn-bars{display:flex;gap:9px;align-items:flex-end;height:76px;margin-top:12px}
.bn-bar{flex:1;display:flex;flex-direction:column;align-items:center;gap:5px}
.bn-bar .col{width:100%;display:flex;align-items:flex-end;height:48px}
.bn-bar .col>i{width:100%;border-radius:6px 6px 0 0;min-height:5px;display:block}
.kgrid{display:grid;grid-template-columns:repeat(auto-fit,minmax(210px,1fr));gap:14px;margin-bottom:4px}
.kpro{position:relative;overflow:hidden;border:1px solid var(--border);border-radius:16px;padding:16px 18px;
  background:linear-gradient(150deg,var(--card2),var(--card));min-height:118px;display:flex;flex-direction:column;gap:8px}
.kpro .ic{width:40px;height:40px;border-radius:12px;display:flex;align-items:center;justify-content:center}
.kpro .ic svg{width:21px;height:21px;stroke:#141414;fill:none;stroke-width:2}
.kpro .v{font-size:30px;font-weight:800;line-height:1;color:var(--txt)}
.kpro .l{font-size:11px;text-transform:uppercase;letter-spacing:.07em;color:var(--mut);font-weight:600}
.kpro .s{font-size:11px;color:var(--mut2)}
.dleg{display:flex;flex-direction:column;gap:7px}
.dleg .it{display:flex;align-items:center;gap:8px;font-size:12.5px}
.dleg .dot{width:10px;height:10px;border-radius:3px;flex:none}
.bn-story{font-size:12.5px;color:var(--txt);line-height:1.45;border-left:2px solid var(--yellow);
  padding-left:9px;margin:6px 0 14px}
.bn-story b{color:var(--yellow);font-weight:700}
.dgrid{display:grid;grid-template-columns:repeat(6,1fr);gap:18px;margin-bottom:4px}
.dgrid>.bn-card{padding:18px 20px 20px}
.sp2{grid-column:span 2}.sp3{grid-column:span 3}.sp4{grid-column:span 4}.sp6{grid-column:span 6}
@media(max-width:1100px){.dgrid{grid-template-columns:repeat(2,1fr)}
  .dgrid>.bn-card,.sp2,.sp3,.sp4{grid-column:span 1}.sp6{grid-column:span 2}}
.enq-hero{font-size:40px;font-weight:800;color:var(--yellow);line-height:1}
.enq-stack{display:flex;height:20px;border-radius:10px;overflow:hidden;background:var(--border);margin:10px 0 14px}
.hbar{display:flex;align-items:center;gap:10px;margin-bottom:11px}
.hbar .lab{width:104px;font-size:12.5px;flex:none}
.hbar .trk{flex:1;height:9px;background:var(--border);border-radius:5px;overflow:hidden}
.hbar .trk>i{display:block;height:100%;border-radius:5px}
.hbar .val{width:62px;text-align:right;font-size:12.5px;font-weight:700;flex:none}
#page-dash{font-family:'Manrope',"Segoe UI",system-ui,sans-serif;
  background-image:radial-gradient(rgba(255,255,255,.022) 1px,transparent 1px);background-size:22px 22px}
.dc-mono{font-family:'JetBrains Mono',Consolas,ui-monospace,monospace}
.dc-row4{display:grid;grid-template-columns:repeat(4,1fr);gap:18px;margin-bottom:18px}
.dc-row3{display:grid;grid-template-columns:repeat(3,1fr);gap:18px;margin-bottom:18px}
.dc-row21{display:grid;grid-template-columns:2fr 1fr;gap:18px;margin-bottom:18px}
@media(max-width:1080px){.dc-row4{grid-template-columns:repeat(2,1fr)}.dc-row3,.dc-row21{grid-template-columns:1fr}}
.dc-kpi{background:linear-gradient(180deg,#15171B,#111216);border:1px solid rgba(255,255,255,.07);
  border-radius:16px;padding:20px;box-shadow:0 1px 2px rgba(0,0,0,.35),0 18px 36px -28px rgba(0,0,0,.9);
  display:flex;flex-direction:column;min-height:150px;transition:border-color .2s,transform .2s;position:relative;overflow:hidden}
.dc-kpi:hover{border-color:rgba(255,255,255,.14);transform:translateY(-2px)}
.dc-kpi.hero{background:linear-gradient(180deg,#1A1712,#131210);border-color:rgba(217,119,6,.22)}
.dc-kpi .num{margin-top:auto;font-size:38px;font-weight:700;letter-spacing:-.02em;line-height:1;color:#F7F8FA}
.dc-kpi .lab{margin-top:11px;font-size:11px;font-weight:700;letter-spacing:.13em;text-transform:uppercase;color:#71757F}
.dc-kpi .sub{margin-top:4px;font-size:11.5px;color:#5A5E67;font-weight:500}
.dc-kpi .num2{margin-top:auto;font-family:'Manrope',"Segoe UI",system-ui,sans-serif;font-size:44px;
  font-weight:800;letter-spacing:-.035em;line-height:1;font-feature-settings:"tnum" 1,"ss01" 1}
.dc-ic2{display:flex;align-items:center;justify-content:center;margin:-2px 0 2px;filter:drop-shadow(0 2px 6px rgba(0,0,0,.45))}
.dc-ic2 svg{width:28px;height:28px;fill:none;stroke:currentColor;stroke-width:1.9;stroke-linecap:round;stroke-linejoin:round}
.dc-ico{width:40px;height:40px;border-radius:11px;border:1px solid rgba(255,255,255,.09);
  background:rgba(255,255,255,.03);display:flex;align-items:center;justify-content:center;color:#A6AAB3}
.dc-ico svg{width:20px;height:20px;fill:none;stroke:currentColor;stroke-width:1.75}
.dc-card{background:linear-gradient(180deg,#15171B,#111216);border:1px solid rgba(255,255,255,.07);
  border-radius:16px;padding:22px;box-shadow:0 1px 2px rgba(0,0,0,.35),0 18px 36px -28px rgba(0,0,0,.9);transition:border-color .2s}
.dc-card:hover{border-color:rgba(255,255,255,.12)}
.dc-h{display:flex;align-items:center;gap:8px;margin-bottom:14px}
.dc-h svg{width:14px;height:14px;fill:none;stroke:#71757F;stroke-width:2}
.dc-h span{font-size:11px;font-weight:700;letter-spacing:.14em;text-transform:uppercase;color:#71757F}
.dc-story{padding-left:12px;font-size:13px;line-height:1.55;color:#B4B8C0;font-weight:500;margin-bottom:18px;border-left:2px solid}
.dc-story b{font-weight:700}
.dc-lr{display:flex;justify-content:space-between;align-items:center;padding:9px 0;border-bottom:1px solid rgba(255,255,255,.05)}
.dc-lr:last-child{border-bottom:none}
.dc-trk{height:7px;border-radius:999px;background:rgba(255,255,255,.06);overflow:hidden}
.dc-trk>i{display:block;height:100%;border-radius:999px}
.dc-vcol{display:flex;justify-content:center}
.dc-vcol>i{width:60%;max-width:46px;border-radius:7px 7px 0 0;display:block}
.mesa-tgl-wrap{display:inline-flex;align-items:center;gap:8px;cursor:pointer;font-size:12.5px;color:var(--mut);
  background:var(--card);border:1px solid var(--border);border-radius:999px;padding:6px 13px;user-select:none;transition:border-color .15s,color .15s}
.mesa-tgl-wrap input{accent-color:var(--yellow);width:15px;height:15px;cursor:pointer;margin:0}
.mesa-tgl-wrap:hover{border-color:var(--border-hi);color:var(--txt)}
table.dat{width:100%;border-collapse:collapse}
table.dat th{font:600 11.5px 'Manrope',sans-serif;text-transform:uppercase;letter-spacing:.07em;color:#7A7A84;
  text-align:left;padding:0 12px 12px;border-bottom:1px solid rgba(255,255,255,.06);white-space:nowrap}
table.dat td{padding:12px;border-bottom:1px solid rgba(255,255,255,.05);font-size:13px;white-space:nowrap}
table.dat tbody tr{transition:background .15s}
table.dat tr:hover td{background:rgba(242,192,41,.04)}
.pos{color:var(--ok)} .neg{color:var(--err)} .neu{color:var(--mut)}
.mut{color:var(--mut)} .small{font-size:12px}

.fam-chip{display:inline-flex;align-items:center;gap:9px;background:rgba(255,255,255,.04);
  border:1px solid rgba(255,255,255,.1);color:#D2D5DB;font:600 12.5px 'Manrope',sans-serif;
  border-radius:999px;padding:6px 8px 6px 13px;transition:background .15s,border-color .15s}
.fam-chip.fam-cli{background:rgba(242,192,41,.08);border-color:rgba(242,192,41,.3);color:#F2C029}
.fam-chip.fam-cli:hover{background:rgba(242,192,41,.15)}
.fam-chip .fam-par{width:auto;flex:none;background:rgba(0,0,0,.22);border:1px solid rgba(255,255,255,.1);
  color:inherit;font:600 11.5px 'Manrope',sans-serif;padding:3px 20px 3px 8px;border-radius:999px;
  cursor:pointer;background-position:right 6px center}
.fam-chip .fam-par:hover{border-color:rgba(255,255,255,.22)}
.fam-chip .fam-par.fam-par-vazio{color:#8E8E98;font-style:italic;border-style:dashed}
.fam-x{cursor:pointer;color:inherit;opacity:.55;font-size:11px;padding:2px 7px;border-radius:999px;flex:none;transition:all .15s}
.fam-x:hover{opacity:1;background:rgba(239,68,68,.18);color:#EF4444}
.fam-badge{display:inline-flex;align-items:center;gap:5px;flex:none;background:rgba(242,192,41,.09);
  border:1px solid rgba(242,192,41,.26);color:#F2C029;font:700 11px 'Manrope',sans-serif;
  border-radius:999px;padding:2px 8px;cursor:default}
.fam-badge svg{width:11px;height:11px}

.bday-badge{display:inline-flex;align-items:center;gap:5px;flex:none;cursor:default;
  font:800 10.5px 'Manrope',sans-serif;letter-spacing:.05em;color:#141414;white-space:nowrap;
  background:linear-gradient(135deg,#FFD24D,#F2A129);border-radius:999px;padding:3px 10px;
  animation:bdaypulse 1.8s ease-out infinite}
.bday-badge.big{font-size:12.5px;padding:5px 14px;letter-spacing:.02em}
@keyframes bdaypulse{0%{box-shadow:0 0 0 0 rgba(242,192,41,.5)}
  70%{box-shadow:0 0 0 9px rgba(242,192,41,0)}100%{box-shadow:0 0 0 0 rgba(242,192,41,0)}}

header{position:relative}
.notif-badge{position:absolute;top:-6px;right:-6px;min-width:17px;height:17px;border-radius:999px;
  background:#EF4444;color:#fff;font:800 10px/17px 'Manrope',sans-serif;text-align:center;
  padding:0 4px;border:2px solid #0A0A0C;box-sizing:content-box}
.notif-drop{position:absolute;top:66px;right:34px;width:410px;max-height:480px;overflow-y:auto;z-index:70;
  background:var(--card2);border:1px solid var(--border-hi);border-radius:14px;padding:10px;
  box-shadow:0 24px 60px -18px rgba(0,0,0,.9);animation:pop .15s ease}
.notif-item{display:flex;gap:12px;padding:12px;border-radius:10px;cursor:pointer;transition:background .15s;align-items:flex-start}
.notif-item:hover{background:rgba(255,255,255,.045)}
.notif-ic{width:36px;height:36px;border-radius:10px;flex:none;display:flex;align-items:center;justify-content:center;
  background:rgba(242,192,41,.1);border:1px solid rgba(242,192,41,.2);color:var(--yellow)}
.notif-ic svg{width:17px;height:17px;fill:none;stroke:currentColor;stroke-width:1.9;stroke-linecap:round;stroke-linejoin:round}
.notif-tt{font:700 13px 'Manrope',sans-serif;color:#F2F3F5}
.notif-tx{font-size:12.5px;color:#9A9AA4;line-height:1.5;margin-top:2px}

.pn-card{background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));
  border:1px solid rgba(255,255,255,.06);border-radius:32px;position:relative;overflow:hidden;
  box-shadow:0 1px 2px rgba(0,0,0,.4),0 34px 70px -46px rgba(0,0,0,.95);
  transition:border-color .3s,transform .3s}
.pn-card:hover{border-color:rgba(242,192,41,.16)}
.pn-hero{background:
  radial-gradient(120% 90% at 85% 0%,rgba(242,192,41,.09),transparent 55%),
  linear-gradient(180deg,rgba(27,28,30,.92),rgba(13,14,16,.97))}
.pn-htitle{display:flex;align-items:center;gap:10px;color:#F2C029;font:500 17px 'Manrope',sans-serif}
.pn-htitle svg{width:20px;height:20px;fill:none;stroke:currentColor;stroke-width:1.9;stroke-linecap:round;stroke-linejoin:round}
.pn-label{font:600 11px 'Manrope',sans-serif;letter-spacing:.08em;text-transform:uppercase;color:#9c948a}
.pn-chip{display:inline-flex;align-items:center;gap:8px;background:#1f2022;border:1px solid rgba(255,255,255,.06);
  border-radius:999px;padding:8px 16px;font:600 11.5px 'Manrope',sans-serif;letter-spacing:.05em;color:#b8b2a6}
.pn-pillbar{display:inline-flex;background:#1f2022;border:1px solid rgba(255,255,255,.06);border-radius:999px;padding:4px;gap:2px}
.pn-pill{padding:6px 16px;border-radius:999px;font:600 11.5px 'Manrope',sans-serif;letter-spacing:.04em;
  color:#9c948a;cursor:pointer;user-select:none;transition:all .2s;white-space:nowrap}
.pn-pill:hover{color:#e3e2e4}
.pn-pill.sel{background:#343537;color:#F2C029;box-shadow:0 1px 3px rgba(0,0,0,.4)}
.pn-glow{text-shadow:0 0 34px rgba(242,192,41,.28)}
.pn-track{width:100%;height:12px;background:#343537;border-radius:999px;overflow:hidden}
.pn-track.fina{height:6px}
.pn-fill{height:100%;border-radius:999px;background:linear-gradient(90deg,rgba(242,192,41,.5),#F2C029);position:relative}
.pn-fill:after{content:"";position:absolute;inset:0;background:rgba(255,255,255,.18);animation:pnpulse 2s ease-in-out infinite}
@keyframes pnpulse{0%,100%{opacity:.1}50%{opacity:.4}}
.pn-btn{display:inline-flex;align-items:center;justify-content:center;border-radius:999px;padding:10px 20px;
  font:600 12px 'Manrope',sans-serif;letter-spacing:.05em;cursor:pointer;transition:all .2s;white-space:nowrap;border:none;flex:none}
.pn-btn:active{transform:scale(.96)}
.pn-btn-ghost{background:transparent;border:1px solid rgba(255,255,255,.1);color:#e3e2e4}
.pn-btn-ghost:hover{background:rgba(255,255,255,.05)}
.pn-btn-primary{background:#F2C029;color:#141414;font-weight:700}
.pn-btn-primary:hover{box-shadow:0 0 15px rgba(242,192,41,.25)}
.pn-nic{width:48px;height:48px;border-radius:50%;flex:none;display:flex;align-items:center;justify-content:center;
  background:rgba(242,192,41,.07);color:#F2C029}
.pn-nic svg{width:22px;height:22px;fill:none;stroke:currentColor;stroke-width:1.8;stroke-linecap:round;stroke-linejoin:round}
.pcht-pt .tipg{opacity:0;pointer-events:none;transition:opacity .15s}
.pcht-pt:hover .tipg{opacity:1}
.pcht-pt{cursor:pointer}

#page-dash .dc-kpi,#page-dash .dc-card{border-radius:28px;
  background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));
  border-color:rgba(255,255,255,.05);
  box-shadow:0 1px 2px rgba(0,0,0,.4),0 34px 70px -46px rgba(0,0,0,.95)}
#page-dash .dc-card:hover{border-color:rgba(242,192,41,.14)}
#page-dash .dc-kpi:hover{border-color:rgba(242,192,41,.18)}
#page-dash .dc-h svg{stroke:#9c8a5e}
#page-dash .dc-h span{color:#9c8a5e}
#page-dash .dc-story{border-left-width:2px;color:#b8b2a6}
#page-dash .dc-lr{border-bottom-color:rgba(255,255,255,.04)}
#page-dash .dc-row4,#page-dash .dc-row3,#page-dash .dc-row21{gap:24px;margin-bottom:24px}
#page-dash .dc-kpi .lab{color:#9c8a5e}
#page-dash .dc-kpi .sub{color:#7a6f55}

#page-cli .card,#page-cli .bn-card,#page-cli .kpi{border-radius:28px;
  background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));
  border-color:rgba(255,255,255,.05);
  box-shadow:0 1px 2px rgba(0,0,0,.4),0 34px 70px -46px rgba(0,0,0,.95)}
#page-cli .card:hover,#page-cli .bn-card:hover,#page-cli .kpi:hover{border-color:rgba(242,192,41,.14)}
#page-cli .bn-big{color:#F2C029}
#page-cli .bn-lbl,#page-cli .card h3,#page-cli .clihdr span{color:#9c948a}
#page-cli .kpi.hl{background:linear-gradient(180deg,rgba(27,28,30,.92),rgba(16,17,19,.96));border-color:rgba(242,192,41,.18)}
#page-cli .kpi.hl .lbl{color:#a8977b}
#page-cli .kpi.hl .val{color:#F2C029}
#page-cli .pills{border-radius:999px;background:#1f2022;border-color:rgba(255,255,255,.06);padding:4px}
#page-cli .pill{border-radius:999px}
#page-cli .pill.sel{background:#343537;color:#F2C029}
#page-cli .cli-row{border-radius:18px}
#page-cli .cli-row:hover{background:rgba(242,192,41,.05)}
#page-cli .btn-primary{background:#F2C029;color:#141414}
#page-cli .btn-primary:hover{background:#ffe9c2;box-shadow:0 0 15px rgba(242,192,41,.25)}
#page-cli #cliBusca,#page-cli #cadBusca,#page-cli #cliFamInput{border-radius:999px;background:#1f2022;padding:9px 16px}
#page-cli .badge.warn{background:rgba(242,192,41,.12);color:#F2C029}
#page-cli .badge.ok{background:rgba(16,185,129,.12);color:#10B981}
#page-cli .chip-save{color:#4C8C6A;background:rgba(5,150,105,.1);border-color:rgba(5,150,105,.2)}
#page-cli .mesa-tgl-wrap{border-radius:999px}
#page-cli table.cli tr:hover td{background:rgba(242,192,41,.04)}
#page-cli .cli-av{border-radius:10px;width:34px;height:34px;font-size:12px}
/* resumo profissional: título no topo, conteúdo CENTRALIZADO no meio, selo no rodapé.
   Todos os cards ficam com o mesmo ritmo visual, sem vazio nem sobreposição. */
#page-cli .bn-card{min-height:0;padding:12px 16px 8px;display:flex;flex-direction:column}
#page-cli .bn-lbl{font-size:10.5px;margin-bottom:2px}
#page-cli .bn-card>:nth-child(2){flex:1 1 auto;width:100%;margin:2px 0;align-content:center}  /* miolo ocupa o card, centrado */
#page-cli .bn-card>div:last-child{margin-top:auto}                       /* selo no rodapé */
#page-cli .bn-big{font-size:32px;line-height:1;padding:0;margin:0;
  display:flex;align-items:center;justify-content:center;text-align:center}
#page-cli .bn-bars{height:auto;margin-top:6px;align-items:flex-end;gap:8px}
#page-cli .bn-bar .col{height:26px}
#page-cli .bn-card>div:last-child[title]{margin-top:auto}
#page-cli .bn-card div[style*="border-top"]{margin-top:8px !important;padding-top:5px !important}
#page-cli .bn-grid{margin-bottom:0;align-items:stretch}
#page-cli .cli-row{padding:7px 16px}

.overlay{position:fixed;inset:0;background:rgba(0,0,0,.62);display:none;align-items:center;
  justify-content:center;z-index:50;backdrop-filter:blur(2px)}
.overlay.on{display:flex}
.modal{background:var(--card2);border:1px solid var(--border-hi);border-radius:16px;padding:26px 28px;
  width:400px;max-width:92vw;animation:pop .18s ease}
@keyframes pop{from{transform:scale(.95);opacity:0}to{transform:scale(1);opacity:1}}
.modal h4{font-size:15.5px;margin-bottom:10px}
.modal p{font-size:13.5px;color:var(--mut);white-space:pre-line;margin-bottom:20px}
.modal .row{justify-content:flex-end}
.toast{position:fixed;bottom:22px;right:24px;background:var(--card2);border:1px solid var(--border-hi);
  border-left:3px solid var(--yellow);border-radius:11px;padding:12px 18px;font-size:13px;z-index:60;
  opacity:0;transform:translateY(8px);transition:all .25s;pointer-events:none}
.toast.on{opacity:1;transform:none}
</style></head>
<body>

<aside>
  <div class="logo">
    <div class="logo-mark">XP</div>
    <div class="logo-name">Performance<small>Manager</small></div>
  </div>
  <nav>
    <div class="nav-sec">Geral</div>
    <div class="nav-item sel" data-tab="painel" onclick="tab('painel')">
      <svg viewBox="0 0 24 24"><path d="m3 11 9-8 9 8"/><path d="M5 9.5V21h5v-6h4v6h5V9.5"/></svg>
      Painel</div>
    <div class="nav-item" data-tab="dash" onclick="tab('dash')">
      <svg viewBox="0 0 24 24"><rect x="3" y="3" width="7" height="9" rx="1"/><rect x="14" y="3" width="7" height="5" rx="1"/><rect x="14" y="12" width="7" height="9" rx="1"/><rect x="3" y="16" width="7" height="5" rx="1"/></svg>
      Dashboard</div>
    <div class="nav-sec">Clientes &amp; Carteiras</div>
    <div class="nav-item" data-tab="cli" onclick="tab('cli')">
      <svg viewBox="0 0 24 24"><path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/><path d="M23 21v-2a4 4 0 0 0-3-3.87"/><path d="M16 3.13a4 4 0 0 1 0 7.75"/></svg>
      Clientes</div>
    <div class="nav-item" data-tab="cart" onclick="tab('cart')">
      <svg viewBox="0 0 24 24"><path d="M21 16V8a2 2 0 0 0-1-1.73l-7-4a2 2 0 0 0-2 0l-7 4A2 2 0 0 0 3 8v8a2 2 0 0 0 1 1.73l7 4a2 2 0 0 0 2 0l7-4A2 2 0 0 0 21 16z"/><path d="m3.27 6.96 8.73 5.04 8.73-5.04"/><path d="M12 22.08V12"/></svg>
      Carteiras</div>
    <div class="nav-item" data-tab="controle" onclick="tab('controle')">
      <svg viewBox="0 0 24 24"><line x1="4" y1="21" x2="4" y2="14"/><line x1="4" y1="10" x2="4" y2="3"/><line x1="12" y1="21" x2="12" y2="12"/><line x1="12" y1="8" x2="12" y2="3"/><line x1="20" y1="21" x2="20" y2="16"/><line x1="20" y1="12" x2="20" y2="3"/><line x1="1" y1="14" x2="7" y2="14"/><line x1="9" y1="8" x2="15" y2="8"/><line x1="17" y1="16" x2="23" y2="16"/></svg>
      Controle de Alocação</div>
    <div class="nav-sec">Mercado</div>
    <div class="nav-item" data-tab="mkt" onclick="tab('mkt')">
      <svg viewBox="0 0 24 24"><path d="M3 3v18h18"/><path d="m19 9-5 5-4-4-3 3"/></svg>
      Mercado</div>
    <div class="nav-sec">Ferramentas</div>
    <div class="nav-item" data-tab="pdf" onclick="tab('pdf')">
      <svg viewBox="0 0 24 24"><path d="M14 2H6a2 2 0 0 0-2 2v16a2 2 0 0 0 2 2h12a2 2 0 0 0 2-2V8z"/><path d="M14 2v6h6"/></svg>
      Ler PDFs</div>
    <div class="nav-item" data-tab="mail" onclick="tab('mail')">
      <svg viewBox="0 0 24 24"><rect x="2" y="4" width="20" height="16" rx="2"/><path d="m22 7-10 6L2 7"/></svg>
      E-mail</div>
  </nav>
  <div class="side-foot">
    <div class="side-av">GE</div>
    <div class="side-nm">Guilherme Enrico<small>Assessor de Investimentos</small></div>
    <span id="excelName" style="display:none"></span>
  </div>
</aside>

<main>
  <header>
    <div><h1 id="pageTitle">Painel</h1><div class="sub" id="pageSub">Visão geral e notificações</div></div>
    <div class="grow"></div>
    <label class="hsearch"><svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round"><circle cx="11" cy="11" r="7"/><path d="m21 21-4.3-4.3"/></svg>
      <input id="hbusca" placeholder="Buscar cliente, conta…" spellcheck="false" oninput="headerBusca(this.value)"><kbd>⌘K</kbd></label>
    <button class="iconbtn" id="btnNotif" title="Notificações" onclick="toggleNotif(event)" style="position:relative">
      <svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9"/><path d="M10.3 21a1.94 1.94 0 0 0 3.4 0"/></svg>
      <span id="notifBadge" class="notif-badge" style="display:none"></span>
    </button>
    <button class="btn btn-ghost" onclick="uiRefresh()" title="Reler a planilha (pega edições feitas no Excel)">⟳&nbsp; Atualizar</button>
    <div id="notifDrop" class="notif-drop" style="display:none"></div>
  </header>

  <section class="page sel" id="page-painel">
    <div id="painelWrap"></div>
  </section>

  <section class="page" id="page-dash">
    <div id="dashWrap"></div>
  </section>

  <section class="page" id="page-pdf">
    <div class="card">
      <h3>Pasta de relatórios</h3>
      <div class="row">
        <input type="text" id="pasta" class="grow" spellcheck="false">
        <button class="btn btn-ghost" onclick="escolherPasta()">Procurar…</button>
      </div>
      <div class="progress" id="prog"><div></div></div>
    </div>
    <button class="btn btn-primary btn-big" id="btnProc" onclick="processar()" style="align-self:flex-start">
      ▶&nbsp; Ler PDFs e atualizar planilha</button>
    <div class="console" id="logPdf"></div>
  </section>

  <section class="page" id="page-mail">
    <div class="row" style="gap:14px">
      <div style="width:170px">
        <label class="fld">Mês de referência</label>
        <select id="selMes" onchange="aoMes()"></select>
      </div>
      <div style="align-self:flex-end" id="mesInfo"></div>
      <div class="grow"></div>
      <div class="save-state" id="saveState" style="align-self:flex-end"></div>
    </div>

    <div class="mail-grid">
      <div class="mail-col">
        <div class="card">
          <h3>Template</h3>
          <label class="fld">Assunto</label>
          <input type="text" id="assunto" oninput="aoEditar()" spellcheck="false">
          <label class="fld" style="margin-top:12px">Mensagem</label>
          <textarea id="corpo" rows="14" oninput="aoEditar()" spellcheck="false"></textarea>
          <div class="hint" style="margin-top:10px;margin-bottom:2px">Dados do cliente</div>
          <div class="chips" id="chipsCli"></div>
          <div class="hint" style="margin-top:10px;margin-bottom:2px">Indicadores do relatório (aba Clientes)</div>
          <div class="chips" id="chipsMkt"></div>
          <div class="hint" style="margin-top:10px;margin-bottom:2px">Blocos visuais prontos (cada um em uma linha própria da mensagem)</div>
          <div class="chips" id="chipsBlocos"></div>
          <div class="hint" style="margin-top:10px;margin-bottom:2px">Análises prontas (frases explicativas para usar no meio do texto)</div>
          <div class="chips" id="chipsFrases"></div>
          <div class="hint">Clique numa variável para inserir na posição do cursor. O template salva sozinho.</div>
        </div>
        <button class="btn btn-primary btn-big" id="btnSend" onclick="confirmarEnvio()">
          📧&nbsp; Enviar via Outlook</button>
        <div class="console" id="logMail" style="min-height:110px;flex:none;max-height:170px"></div>
      </div>

      <div class="card preview-card">
        <h3>Preview do e-mail</h3>
        <div class="preview-head">
          <button class="iconbtn" onclick="navCli(-1)">‹</button>
          <input id="selCli" data-src="mes" data-pick="mail" autocomplete="off" spellcheck="false" placeholder="Pesquisar cliente…"
                 onfocus="cbxOpen(this)" onclick="cbxOpen(this)" oninput="cbxFilter()" onkeydown="cbxKey(event)" onblur="cbxBlur()">
          <button class="iconbtn" title="Copiar conta do cliente" onclick="copiarConta(event,(clientes[pvCliIdx]||{}).conta||'')">⧉</button>
          <button class="iconbtn" onclick="navCli(1)">›</button>
        </div>
        <div class="preview-meta" id="pvMeta"></div>
        <iframe id="pvFrame"></iframe>
      </div>
    </div>
  </section>

  <section class="page" id="page-cli">
    <div id="cliResumo" class="bn-grid"></div>
    <div class="row">
      <div class="pills">
        <span class="pill sel" id="pillRel" onclick="cliView('rel')">Relatórios</span>
        <span class="pill" id="pillCad" onclick="cliView('cad')">Cadastro</span>
      </div>
      <div class="grow"></div>
      <span id="cliBadge"></span>
    </div>

    <div id="cliRel" style="display:flex;flex-direction:column;gap:16px;flex:1;min-height:0">
      <div id="cliLista" class="card" style="flex:1;display:flex;flex-direction:column;min-height:0">
        <div class="row" style="margin-bottom:12px">
          <h3 style="margin:0">Clientes</h3>
          <div class="grow"></div>
          <input id="cliBusca" placeholder="Buscar nome ou conta…" oninput="renderCliList()"
                 style="width:250px" spellcheck="false">
        </div>
        <div style="flex:1;overflow:auto;min-height:0">
          <div id="cliListWrap" style="min-width:100%">
            <div class="clihdr" id="cliHdr"></div>
            <div id="cliListBody"></div>
          </div>
        </div>
      </div>

      <div id="cliDetail" style="display:none;flex-direction:column;gap:20px;flex:1;min-height:0">
        <div class="row" style="gap:14px;align-items:center">
          <button class="btn btn-ghost" onclick="voltarClientes()">‹ Clientes</button>
          <div class="cli-av" id="cliDetAv" style="width:46px;height:46px;border-radius:12px;font-size:15px"></div>
          <div>
            <div style="display:flex;align-items:center;gap:12px">
              <div id="cliDetNome" style="font-size:19px;font-weight:700;letter-spacing:-.01em"></div>
              <span id="cliDetBday"></span>
            </div>
            <div class="small mut" id="cliDetSub" style="margin-top:3px"></div>
          </div>
        </div>
        <div class="row" style="align-items:center;gap:10px">
          <span class="small mut">Mês do relatório</span>
          <div class="pills seg" id="cliMesPills" style="flex-wrap:wrap"></div>
        </div>
        <div class="kpis" id="cliKpis"></div>
        <div class="det-grid">
          <div class="card" id="cliCompCard">
            <h3>Composição da carteira</h3>
            <div id="cliComp"></div>
          </div>
          <div class="card" id="cliExtraCard">
            <div class="row" style="margin-bottom:6px">
              <h3 style="margin:0" class="grow">Informações do cliente</h3>
              <span id="cliExtraStatus" class="small mut">salvo automaticamente</span>
            </div>
            <div class="small mut" style="margin-bottom:6px"><span style="color:#10B981">●</span> Tem · <span style="color:#F59E0B">●</span> A revisar · <span style="color:#EF4444">●</span> Não tem</div>
            <div id="cliExtra"></div>
          </div>
        </div>
        <div class="card" id="cliApCard">
          <div class="row" style="margin-bottom:6px">
            <h3 style="margin:0" class="grow">Aportes</h3>
            <span id="cliApStatus" class="small mut"></span>
          </div>
          <div id="cliApResumo" class="small" style="margin-bottom:12px;line-height:1.6"></div>
          <div class="row" style="flex-wrap:wrap">
            <input id="apData" placeholder="dd/mm/aaaa" style="width:130px" maxlength="10"
                   oninput="maskData(this)" spellcheck="false">
            <input id="apValor" placeholder="Valor (R$)" style="width:150px" spellcheck="false"
                   oninput="maskMoeda(this)" onkeydown="if(event.key==='Enter')addAporte()">
            <input id="apObs" placeholder="Observação (opcional)" class="grow" style="min-width:160px" spellcheck="false"
                   onkeydown="if(event.key==='Enter')addAporte()">
            <button class="btn btn-primary" onclick="addAporte()">＋ Registrar aporte</button>
          </div>
          <div id="cliApChart" style="margin:16px 0 4px"></div>
          <div id="cliApLista" style="margin-top:6px"></div>
          <div class="row" style="margin-top:12px">
            <span class="hint" style="margin:0">Alertar após</span>
            <input id="apAlertaDias" style="width:76px;text-align:center" onchange="salvarApAlerta(this.value)">
            <span class="hint" style="margin:0">dias sem aporte</span>
          </div>
        </div>
        <div class="card" id="cliCtCard">
          <div class="row" style="margin-bottom:6px">
            <h3 style="margin:0" class="grow">Contatos</h3>
            <span id="cliCtStatus" class="small mut"></span>
          </div>
          <div id="cliCtResumo" class="small" style="margin-bottom:12px;line-height:1.6"></div>
          <div class="row" style="flex-wrap:wrap">
            <input id="ctData" placeholder="dd/mm/aaaa" style="width:130px" maxlength="10"
                   oninput="maskData(this)" spellcheck="false">
            <select id="ctTipo" style="width:190px"></select>
            <input id="ctObs" placeholder="Observação (opcional)" class="grow" style="min-width:160px" spellcheck="false"
                   onkeydown="if(event.key==='Enter')addContato()">
            <button class="btn btn-primary" onclick="addContato()">＋ Registrar contato</button>
          </div>
          <div id="cliCtLista" style="margin-top:12px"></div>
        </div>
        <div class="card" id="cliFamCard">
          <div class="row" style="margin-bottom:6px">
            <h3 style="margin:0" class="grow">Grupo familiar</h3>
            <span id="cliFamStatus" class="small mut">salvo automaticamente</span>
          </div>
          <div id="cliFamChips" class="chips"></div>
          <div class="row" style="margin-top:14px">
            <input id="cliFamInput" list="cliFamOpts" placeholder="Nome do familiar ou cliente da base…"
                   style="max-width:340px" spellcheck="false" onfocus="famFoco(this)" onclick="famFoco(this)"
                   onkeydown="if(event.key==='Enter'){event.preventDefault();addFamiliar();}">
            <datalist id="cliFamOpts"></datalist>
            <select id="cliFamParent" style="width:160px"></select>
            <button class="btn btn-ghost" onclick="addFamiliar()">＋ Adicionar</button>
          </div>
        </div>
        <div class="hint" id="cliArq"></div>
      </div>
    </div>

    <div id="cliCad" class="card" style="display:none;flex-direction:column;flex:1;min-height:0">
      <div class="row" style="margin-bottom:10px">
        <h3 style="margin:0">Cadastro de clientes</h3>
        <span id="cadStatus" class="chip-save">salvo automaticamente</span>
        <span id="cadPendInfo"></span>
        <div class="grow"></div>
        <input id="cadBusca" placeholder="Buscar nome ou conta…" oninput="renderCadastro()"
               style="width:240px" spellcheck="false">
        <button class="btn btn-primary" onclick="addCli()">＋ Adicionar cliente</button>
      </div>
      <div style="flex:1;overflow:auto;min-height:0">
        <table class="cli" style="min-width:1980px">
          <thead><tr>
            <th>Nome completo</th><th>Conta</th><th>E-mail</th>
            <th>Termômetro</th><th>Perfil</th>
            <th>Estado</th><th>Cidade</th><th>Nascimento</th>
            <th style="cursor:pointer;user-select:none" onclick="cadSort()">Patrimônio <span id="cadSortArrow"></span></th>
            <th>Previdência</th><th>Internacional</th><th>Seguro</th><th>Plugado na mesa</th>
            <th style="width:34px"></th>
          </tr></thead>
          <tbody id="cliBody"></tbody>
        </table>
      </div>
    </div>
  </section>

  <section class="page" id="page-mkt">
    <div class="row" style="margin-bottom:2px">
      <div class="pills">
        <span class="pill sel" id="pillAlocSub" onclick="mktSub('aloc')">Carteiras Recomendadas</span>
        <span class="pill" id="pillOndeSub" onclick="mktSub('onde')">Onde Investir</span>
      </div>
      <div class="grow"></div>
      <button class="btn btn-primary" id="btnMerc" onclick="mercadoLoad()">📄&nbsp; Ler relatórios de mercado</button>
    </div>

    <div id="mktAloc" style="display:flex;flex-direction:column;gap:16px">
    <div class="row">
      <div>
        <div id="alocTitulo" style="font-size:16px;font-weight:700">Carteiras Recomendadas</div>
        <div class="small mut" id="alocSub">Use o botão “Ler relatórios de mercado” acima</div>
      </div>
    </div>

    <div id="alocVazio" class="card mut">Nenhum relatório de alocação carregado. Clique em
      “Ler relatório de alocação” e selecione o PDF de Carteiras de Alocação da XP.</div>

    <div id="alocConteudo" style="display:none;flex-direction:column;gap:16px">
      <div class="card" id="alocNotaCard">
        <h3>Destaque do mês</h3>
        <div id="alocNota" class="small" style="line-height:1.6;color:var(--txt)"></div>
      </div>
      <div class="card">
        <h3>Carteiras recomendadas por perfil</h3>
        <table class="dat"><thead><tr>
          <th>Classe de ativo</th><th>Conservadora</th><th>Moderada</th><th>Sofisticada</th>
        </tr></thead><tbody id="tbAloc"></tbody><tfoot id="tfAloc"></tfoot></table>
      </div>
      <div class="card">
        <h3>Insights do relatório</h3>
        <div id="alocInsights"></div>
      </div>
      <div class="card">
        <h3>Visão gerencial — composição por perfil</h3>
        <div id="alocChart"></div>
      </div>
      <div class="card">
        <h3>Perspectivas por classe de ativo (resumo)</h3>
        <div id="alocPersp"></div>
      </div>
    </div>
    <div class="hint"></div>
    </div><!-- /mktAloc -->

    <div id="mktOnde" style="display:none;flex-direction:column;gap:16px">
      <div class="row">
        <div>
          <div id="oiTitulo" style="font-size:16px;font-weight:700">Onde Investir</div>
          <div class="small mut" id="oiSub">Use o botão “Ler relatórios de mercado” acima</div>
        </div>
      </div>
      <div id="oiVazio" class="card mut">Nenhum relatório “Onde Investir” carregado. Clique no botão e
        selecione o PDF — isso NÃO altera o relatório de Carteiras Recomendadas.</div>
      <div id="oiConteudo" style="display:none;flex-direction:column;gap:16px">
        <div class="card"><h3>Resumo do mês — pontos-chave</h3><div id="oiResumo"></div></div>
        <div class="card"><h3>Destaques por tema</h3><div id="oiSecoes"></div></div>
      </div>
      <div class="hint"></div>
    </div>
  </section>

  <section class="page" id="page-cart">
    <div class="row" style="gap:14px;align-items:flex-end">
      <div style="width:300px">
        <label class="fld">Cliente</label>
        <div style="display:flex;align-items:center;gap:6px">
          <input id="cartCli" data-src="geral" data-pick="cart" autocomplete="off" spellcheck="false" placeholder="Pesquisar cliente…"
                 style="flex:1;min-width:0"
                 onfocus="cbxOpen(this)" onclick="cbxOpen(this)" oninput="cbxFilter()" onkeydown="cbxKey(event)" onblur="cbxBlur()">
          <span title="Copiar conta do cliente selecionado"
                style="flex:none;cursor:pointer;font-size:15px;color:#6C6C76;padding:4px"
                onmouseover="this.style.color='#F2C029'" onmouseout="this.style.color='#6C6C76'"
                onclick="copiarConta(event,$('cartCli').dataset.conta||'')">⧉</span>
        </div>
      </div>
      <div id="cartPerfil" style="display:flex;gap:8px;align-items:center;flex-wrap:wrap;padding-bottom:6px"></div>
      <div class="grow"></div>
      <span class="small mut" id="cartMes" style="padding-bottom:9px"></span>
    </div>
    <div class="row" style="align-items:center;gap:12px">
      <div class="pills">
        <span class="pill sel" id="pillComp" onclick="cartSub('comp')">Comparativo</span>
        <span class="pill" id="pillAnal" onclick="cartSub('anal')">Análise de Carteira</span>
      </div>
    </div>
    <div id="cartVazio" class="card mut">Selecione um cliente para comparar a carteira com as recomendações de mercado.</div>

    <div id="cartComp" style="display:none;flex-direction:column;gap:16px">
      <div class="kpis" id="cartKpis"></div>
      <div class="card">
        <h3>Carteira atual × recomendada (por perfil)</h3>
        <table class="dat"><thead id="cartHead"></thead><tbody id="tbCart"></tbody></table>
        <div class="hint" id="cartNaoMap" style="margin-top:8px"></div>
      </div>
      <div class="card">
        <h3>Por que essa recomendação (visão XP)</h3>
        <div id="cartPersp"></div>
      </div>
    </div>

    <div id="cartAnal" style="display:none;flex-direction:column;gap:24px">
      <div style="display:grid;grid-template-columns:340px 1fr;gap:24px;align-items:start">
        <div class="card">
          <div class="dc-h" style="margin-bottom:18px"><svg viewBox="0 0 24 24" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="9"/><circle cx="12" cy="12" r="5"/><circle cx="12" cy="12" r="1.4"/></svg><span>Aderência ao perfil</span></div>
          <div id="analGauge"></div>
          <div id="analResumo" class="small" style="margin-top:16px;line-height:1.6;color:var(--mut)"></div>
        </div>
        <div class="card">
          <div class="dc-h" style="margin-bottom:6px"><svg viewBox="0 0 24 24" stroke-linecap="round" stroke-linejoin="round"><path d="M12 2v4M12 18v4M4.93 4.93l2.83 2.83M16.24 16.24l2.83 2.83M2 12h4M18 12h4"/><circle cx="12" cy="12" r="3"/></svg><span>Recomendações inteligentes</span></div>
          <div id="analRecSub" style="font-size:11.5px;color:#62666F;margin-bottom:16px">Baseadas nos relatórios de mercado</div>
          <div id="analRecs" style="display:flex;flex-direction:column;gap:14px"></div>
        </div>
      </div>
      <div class="card">
        <h3>Carteira recomendada (percentuais-alvo por perfil)</h3>
        <table class="dat"><thead><tr>
          <th>Classe de ativo</th>
          <th id="thAnal_cons">Conservadora</th><th id="thAnal_mod">Moderada</th><th id="thAnal_sof">Sofisticada</th>
        </tr></thead><tbody id="analRecTbl"></tbody></table>
        <div class="hint" id="analRecFoot" style="margin-top:8px"></div>
      </div>
      <div class="card">
        <h3>Contexto de mercado</h3>
        <div id="analCtx" class="small" style="line-height:1.6;color:var(--txt)"></div>
      </div>
    </div>
    <div class="hint"></div>
  </section>

  <section class="page" id="page-controle">
    <div class="row" style="align-items:center;gap:10px;flex-wrap:wrap;margin-bottom:2px">
      <div class="pills">
        <span class="pill sel" id="pillCtrlMen" onclick="ctrlSub('men')">Alocação mensal</span>
        <span class="pill" id="pillCtrlMon" onclick="ctrlSub('mon')">Montagem de carteira</span>
        <span class="pill" id="pillCtrlAtv" onclick="ctrlSub('atv')">Carteira de ativos</span>
      </div>
      <div class="grow"></div>
      <select id="ctrlFPerfil" onchange="ctrlRender()" style="width:150px" title="Filtrar por perfil">
        <option value="">Todos os perfis</option><option>Conservadora</option><option>Moderada</option><option>Sofisticada</option></select>
      <select id="ctrlFTermo" onchange="ctrlRender()" style="width:140px" title="Filtrar por termômetro">
        <option value="">Todos termômetros</option><option>A</option><option>B</option><option>C</option><option>D</option><option>E</option></select>
      <span class="small mut" id="ctrlStatus">salvo automaticamente</span>
    </div>

    <div id="ctrlMen" class="card" style="overflow:auto;min-height:0">
      <div style="overflow:auto"><table class="cli" style="min-width:1280px"><thead><tr>
        <th style="width:110px">Mês</th><th style="min-width:230px">Cliente</th><th>Ativo</th>
        <th>Valor (R$)</th><th>% do patrim.</th><th>ROA %</th><th>Retorno esperado</th><th>Data liquidação</th><th style="width:34px"></th>
      </tr></thead><tbody id="ctrlMenBody"></tbody></table></div>
      <div class="row" style="margin-top:12px"><button class="btn btn-ghost" onclick="ctrlAdd('men')">＋ Adicionar alocação</button></div>
    </div>

    <div id="ctrlAtv" class="card" style="display:none;overflow:auto;min-height:0">
      <div class="row" style="gap:14px;align-items:flex-end;flex-wrap:wrap;margin-bottom:14px">
        <div style="width:280px"><label class="fld">Cliente</label>
          <input id="ctrlAtvCli" data-src="geral" data-pick="atv" autocomplete="off" spellcheck="false" placeholder="Pesquisar cliente…"
                 onfocus="cbxOpen(this)" onclick="cbxOpen(this)" oninput="cbxFilter()" onkeydown="cbxKey(event)" onblur="cbxBlur()"></div>
        <div><label class="fld">Carteira</label>
          <select id="ctrlAtvCartSel" style="width:170px" onchange="atvCart=this.value;renderCtrlAtv()"></select></div>
        <div><label class="fld">Nova variação</label>
          <div style="display:flex;gap:6px">
            <input id="ctrlAtvNova" placeholder="ex.: Conservadora" style="width:150px" spellcheck="false"
                   onkeydown="if(event.key==='Enter')ctrlAtvNovaCart()">
            <button class="btn btn-ghost" onclick="ctrlAtvNovaCart()">＋ Criar</button>
          </div></div>
        <button class="btn btn-ghost" onclick="ctrlAtvExcluirCart()"
                title="Exclui a carteira selecionada e os ativos dela">Excluir carteira</button>
        <div class="grow"></div>
        <button class="btn btn-ghost" id="btnAtvCmp" onclick="atvCmpAbrir()"
                title="Compara as variações de carteira deste cliente">⇄ Comparar carteiras</button>
      </div>
      <div id="ctrlAtvTot" class="row" style="gap:26px;flex-wrap:wrap;margin-bottom:12px;padding:12px 14px;
           background:rgba(242,192,41,.06);border:1px solid rgba(242,192,41,.18);border-radius:14px"></div>
      <div style="overflow:auto"><table class="cli" style="min-width:1180px"><thead><tr>
        <th style="min-width:170px">Ativo</th><th style="width:140px">Indexador</th><th style="min-width:160px">Segmento</th>
        <th style="width:150px">Financeiro aproximado</th><th style="width:90px">Yield (%)</th>
        <th style="width:130px">Dividendo Mensal</th><th style="width:120px">Yield anualizado</th>
        <th style="width:150px">Dividendo Anualizado</th><th style="width:34px"></th>
      </tr></thead><tbody id="ctrlAtvBody"></tbody></table></div>
      <div class="row" style="margin-top:12px"><button class="btn btn-ghost" onclick="ctrlAtvAdd()">＋ Adicionar ativo</button></div>
      <div id="ctrlAtvCmp" style="margin-top:14px"></div>
    </div>

    <div id="ctrlMon" style="display:none;flex-direction:column;gap:14px">
      <div class="row" style="gap:12px;align-items:flex-end">
        <div style="width:340px"><label class="fld">Cliente</label>
          <input id="ctrlMonCli" data-src="geral" data-pick="mon" autocomplete="off" spellcheck="false" placeholder="Pesquisar cliente…"
                 onfocus="cbxOpen(this)" onclick="cbxOpen(this)" oninput="cbxFilter()" onkeydown="cbxKey(event)" onblur="cbxBlur()"></div>
        <span title="Copiar conta do cliente selecionado"
              style="flex:none;align-self:flex-end;cursor:pointer;font-size:15px;color:#6C6C76;padding:8px 4px"
              onmouseover="this.style.color='#F2C029'" onmouseout="this.style.color='#6C6C76'"
              onclick="copiarConta(event,$('ctrlMonCli').dataset.conta||'')">⧉</span>
        <div id="ctrlMonInfo" class="small mut"></div>
      </div>
      <div class="row" style="gap:20px;align-items:flex-start">
        <div class="card" style="flex:1;min-width:0;border-color:rgba(239,68,68,.16)">
          <div style="display:flex;align-items:center;gap:9px;margin-bottom:14px">
            <span style="width:9px;height:9px;border-radius:50%;background:#EF4444;box-shadow:0 0 0 3px rgba(239,68,68,.15);flex:none"></span>
            <span style="font:700 12px 'Manrope';letter-spacing:.14em;text-transform:uppercase;color:#EF4444">Sai da Carteira</span></div>
          <table class="cli"><thead><tr><th>Ativo</th><th style="width:150px">Valor (R$)</th><th style="width:34px"></th></tr></thead>
            <tbody id="ctrlMonSai"></tbody></table>
          <div class="row" style="margin-top:14px"><button class="btn btn-ghost" onclick="ctrlMonAdd('Sai')">＋ Adicionar</button>
            <div class="grow"></div><span id="ctrlMonSaiTot"></span></div>
        </div>
        <div class="card" style="flex:1;min-width:0;border-color:rgba(16,185,129,.16)">
          <div style="display:flex;align-items:center;gap:9px;margin-bottom:14px">
            <span style="width:9px;height:9px;border-radius:50%;background:#10B981;box-shadow:0 0 0 3px rgba(16,185,129,.15);flex:none"></span>
            <span style="font:700 12px 'Manrope';letter-spacing:.14em;text-transform:uppercase;color:#10B981">Entra na Carteira</span></div>
          <table class="cli"><thead><tr><th>Ativo</th><th style="width:150px">Valor (R$)</th><th style="width:34px"></th></tr></thead>
            <tbody id="ctrlMonEntra"></tbody></table>
          <div class="row" style="margin-top:14px"><button class="btn btn-ghost" onclick="ctrlMonAdd('Entra')">＋ Adicionar</button>
            <div class="grow"></div><span id="ctrlMonEntraTot"></span></div>
        </div>
      </div>
      <div class="card" id="ctrlMonNetCard" style="display:none;padding:18px 26px">
        <span id="ctrlMonNet" style="font-size:14px;color:#C9CBD1"></span></div>
    </div>
  </section>

</main>

<div class="overlay" id="ovl">
  <div class="modal">
    <h4 id="mTitle"></h4><p id="mBody"></p>
    <div class="row">
      <button class="btn btn-ghost" onclick="modalNo()">Cancelar</button>
      <button class="btn btn-primary" id="mYes" onclick="modalYes()">Confirmar</button>
    </div>
  </div>
</div>
<div class="toast" id="toast"></div>
<div id="cbxList" class="cbx-list"></div>
<div id="cfPop" class="cbx-list" style="min-width:210px;padding:10px" onclick="event.stopPropagation()"></div>
<div id="liePop" class="cbx-list" style="min-width:160px;padding:6px" onclick="event.stopPropagation()"></div>
<div id="balaoInativos" style="display:none;position:fixed;bottom:22px;left:50%;transform:translateX(-50%);
     background:var(--card2);border:1px solid var(--border-hi);border-left:3px solid var(--err);
     border-radius:11px;padding:14px 18px;max-width:580px;z-index:70;box-shadow:0 8px 34px rgba(0,0,0,.45)">
  <div id="balaoTxt" class="small" style="line-height:1.55;color:var(--txt)"></div>
  <div class="row" style="justify-content:flex-end;margin-top:10px">
    <button class="btn btn-ghost" onclick="fecharBalao()">Fechar</button>
  </div>
</div>

<script>
const $=id=>document.getElementById(id);
const api=()=>window.pywebview.api;
let clientes=[], cadastro=[], _mCb=null, _busy=false;

const TITLES={painel:["Painel","Visão geral e notificações"],
              dash:["Dashboard","Visão consolidada da base de clientes"],
              pdf:["Ler PDFs","Extraia os relatórios e atualize a planilha"],
              mail:["E-mail","Template, preview e envio personalizado"],
              cli:["Clientes","Evolução patrimonial por cliente e cadastro"],
              cart:["Carteiras","Carteira atual × recomendada, por perfil"],
              mkt:["Mercado","Perspectivas e carteiras recomendadas"],
              controle:["Controle de Alocação","Ofertas, alocação mensal e montagem de carteira"]};
function tab(t){
  document.querySelectorAll(".nav-item").forEach(n=>n.classList.toggle("sel",n.dataset.tab===t));
  document.querySelectorAll(".page").forEach(p=>p.classList.toggle("sel",p.id==="page-"+t));
  const _pt=$("pageTitle"), _ps=$("pageSub");
  if(_pt)_pt.textContent=TITLES[t][0]; if(_ps)_ps.textContent=TITLES[t][1];
  if(t==="painel") painelLoad();
  if(t==="mail") aoMes();
  if(t==="mkt"){mktSub("aloc");mktLoad();}
  if(t==="cli"){
    cliCur=null;cliMesSel=null;
    const d=$("cliDetail"),l=$("cliLista");
    if(d)d.style.display="none"; if(l)l.style.display="flex";
    cliView("rel");
  }
  if(t==="cart"){cartSub("comp");cartReload();}
  if(t==="dash") dashLoad();
  if(t==="controle"){ctrlSubAtivo="men";controleLoad();}
}
function headerBusca(v){
  if($("cliBusca"))$("cliBusca").value=v||"";
  const navCli=document.querySelector('.nav-item[data-tab="cli"]');
  const onCli=navCli&&navCli.classList.contains("sel");
  if(!onCli) tab("cli");
  else if(typeof renderCliList==="function" && !cliCur) renderCliList();
}
let ctrlData={mensal:[],montagem:[],ativos:[]}, ctrlClientes=[], ctrlSubAtivo="men";
let ctrlMenView=[], ctrlMonSaiView=[], ctrlMonEntraView=[], _ctrlT={};
function numBR(s){s=String(s==null?"":s).trim().replace(/[^\d.,-]/g,""); if(!s)return null;
  if(s.indexOf(",")>=0)s=s.replace(/\./g,"").replace(",","."); const n=parseFloat(s); return isFinite(n)?n:null;}
function _patDe(conta){const c=ctrlClientes.find(x=>String(x.conta)===String(conta)); return c&&c.patrimonio!=null?Number(c.patrimonio):null;}
function _cliDe(conta){return ctrlClientes.find(x=>String(x.conta)===String(conta))||null;}
function _perfilNorm(p){p=(p||"").toLowerCase();
  return p.startsWith("conserv")?"Conservadora":p.startsWith("moder")?"Moderada":
    (p.startsWith("sofist")||p.startsWith("arroj")||p.startsWith("agress"))?"Sofisticada":"";}
let cliDLLista=[];
const _normTxt=s=>String(s||"").toLowerCase().normalize("NFD").replace(/[̀-ͯ]/g,"");
function cliDLFill(lista){
  if(Array.isArray(lista)&&lista.length)
    cliDLLista=lista.filter(c=>String(c.conta||"").trim())
      .map(c=>({conta:String(c.conta).trim(),nome:(c.nome||"").trim()||("Conta "+c.conta)}));
}
function cliTxtDe(conta){
  conta=String(conta||"").trim(); if(!conta)return "";
  const c=cliDLLista.find(x=>x.conta===conta);
  return c?`${c.nome} (${c.conta})`:conta;
}
let cbxEl=null, cbxItems=[], cbxIdx=-1;
function cbxItens(src){
  return src==="mes"
    ? clientes.map((c,i)=>({conta:String(c.conta||""),nome:c.nome||"(sem nome)",idx:i}))
    : cliDLLista.map(c=>({conta:c.conta,nome:c.nome}));
}
const CBX_TXT={
  mail:el=>pvCliTexto(pvCliIdx),
  cart:el=>cliTxtDe(el.dataset.conta||""),
  mon: el=>cliTxtDe(el.dataset.conta||""),
  atv: el=>cliTxtDe(el.dataset.conta||""),
  men: el=>cliTxtDe((ctrlMenView[+el.dataset.i]||{}).conta||"")
};
const CBX_PICK={
  mail:(el,it)=>{pvCliIdx=it.idx;pvSync();atualizarPreview();},
  cart:(el,it)=>{el.dataset.conta=it.conta;cartLoad();},
  mon: (el,it)=>{el.dataset.conta=it.conta;ctrlMonRender();},
  atv: (el,it)=>{el.dataset.conta=it.conta;atvCart="";atvCmpSel=null;renderCtrlAtv();},
  men: (el,it)=>{const r=ctrlMenView[+el.dataset.i];if(r){r.conta=it.conta;renderCtrlMen();ctrlSaveSoon('mensal');}}
};
function cbxPos(){
  const L=$("cbxList"); if(!cbxEl||!L)return;
  const r=cbxEl.getBoundingClientRect();
  L.style.left=r.left+"px"; L.style.top=(r.bottom+5)+"px";
  L.style.width=Math.max(r.width,280)+"px";
}
function cbxOpen(el){
  const L=$("cbxList");
  if(cbxEl===el&&L.style.display!=="none"){
    if(/\(\d+\)\s*$/.test(el.value)){el.value="";cbxFilter();}
    return;
  }
  cbxEl=el;
  el.dataset.canon=(CBX_TXT[el.dataset.pick]||(()=>""))(el)||"";
  if(/\(\d+\)\s*$/.test(el.value))el.value="";
  cbxFilter();
}
function cbxFilter(){
  const L=$("cbxList"); if(!cbxEl||!L)return;
  const q=_normTxt(cbxEl.value.trim());
  cbxItems=cbxItens(cbxEl.dataset.src).filter(it=>!q||_normTxt(it.nome).includes(q)||String(it.conta).includes(q));
  cbxIdx=cbxItems.length?0:-1;
  L.innerHTML=cbxItems.length
    ? cbxItems.slice(0,400).map((it,i)=>
      `<div class="cbx-it ${i===cbxIdx?'on':''}" onmouseenter="cbxHover(${i})" onmousedown="event.preventDefault();cbxPick(${i})">
         <span class="cbx-nm">${esc(it.nome)}</span><span class="cbx-ct">${esc(it.conta)}</span></div>`).join("")
    : `<div class="cbx-vazio">Nenhum cliente encontrado</div>`;
  cbxPos(); L.style.display="block";
}
function cbxHover(i){
  if(i===cbxIdx)return; cbxIdx=i;
  const L=$("cbxList");
  [...L.children].forEach((d,j)=>d.classList.toggle("on",j===cbxIdx));
}
function cbxFechar(){
  const L=$("cbxList"); if(L)L.style.display="none";
  cbxEl=null; cbxItems=[]; cbxIdx=-1;
}
function cbxPick(i){
  const it=cbxItems[i], el=cbxEl; if(!it||!el)return;
  el.value=`${it.nome} (${it.conta})`;
  delete el.dataset.canon;
  cbxFechar();
  (CBX_PICK[el.dataset.pick]||(()=>{}))(el,it);
}
function cbxBlur(){
  const el=cbxEl; if(!el)return;
  const v=el.value.trim();
  if(v&&v!==el.dataset.canon){
    const its=cbxItens(el.dataset.src), q=_normTxt(v);
    const m=v.match(/\((\d+)\)\s*$/);
    let hit=m?its.filter(it=>it.conta===m[1]):[];
    if(hit.length!==1)hit=its.filter(it=>it.conta===v);
    if(hit.length!==1)hit=its.filter(it=>_normTxt(it.nome)===q);
    if(hit.length!==1){const h2=its.filter(it=>_normTxt(it.nome).includes(q));if(h2.length===1)hit=h2;}
    if(hit.length===1){cbxItems=hit;cbxIdx=0;cbxPick(0);return;}
  }
  el.value=el.dataset.canon||""; delete el.dataset.canon;
  cbxFechar();
}
function cbxKey(e){
  const L=$("cbxList"); if(!cbxEl||!L||L.style.display==="none")return;
  if(e.key==="ArrowDown"||e.key==="ArrowUp"){
    e.preventDefault(); if(!cbxItems.length)return;
    cbxIdx=(cbxIdx+(e.key==="ArrowDown"?1:-1)+cbxItems.length)%cbxItems.length;
    [...L.children].forEach((d,j)=>d.classList.toggle("on",j===cbxIdx));
    const d=L.children[cbxIdx]; if(d&&d.scrollIntoView)d.scrollIntoView({block:"nearest"});
  }else if(e.key==="Enter"){e.preventDefault();if(cbxIdx>=0)cbxPick(cbxIdx);}
  else if(e.key==="Escape"){
    e.preventDefault();
    const el=cbxEl; el.value=el.dataset.canon||""; delete el.dataset.canon; cbxFechar();
  }
}
window.addEventListener("resize",()=>{if(cbxEl)cbxPos()});
window.addEventListener("scroll",()=>{if(cbxEl)cbxPos()},true);
const CBX_ATTRS=`autocomplete="off" spellcheck="false" placeholder="Pesquisar cliente…"
  onfocus="cbxOpen(this)" onclick="cbxOpen(this)" oninput="cbxFilter()" onkeydown="cbxKey(event)" onblur="cbxBlur()"`;
function famFoco(el){if(el.showPicker){try{el.showPicker()}catch(e){}}}
function fmtRS(v){const n=numBR(v); return n==null?"":rs(n);}
function maskData(el){let d=String(el.value).replace(/\D/g,"").slice(0,8), o=d;
  if(d.length>4)o=d.slice(0,2)+"/"+d.slice(2,4)+"/"+d.slice(4);
  else if(d.length>2)o=d.slice(0,2)+"/"+d.slice(2);
  el.value=o; return o;}
function maskMes(el){let d=String(el.value).replace(/\D/g,"").slice(0,6), o=d;
  if(d.length>2)o=d.slice(0,2)+"/"+d.slice(2); el.value=o; return o;}
function maskMoeda(el){
  let d=String(el.value).replace(/\D/g,"").replace(/^0+(?=\d)/,"").slice(0,15);
  if(!d){el.value="";return "";}
  el.value=(parseInt(d,10)/100).toLocaleString("pt-BR",{minimumFractionDigits:2,maximumFractionDigits:2});
  return el.value;
}
const _TCD={A:"#10B981",B:"#3B82F6",C:"#F59E0B",D:"#F97316",E:"#EF4444"};
function cliInfoLine(conta){const c=_cliDe(conta); if(!c)return "";
  const t=(c.termometro||"").toUpperCase(), pat=c.patrimonio!=null?rs(c.patrimonio):"—";
  return `Termô. <b style="color:${_TCD[t]||'var(--mut)'}">${t||"—"}</b> · Patrim. <b style="color:var(--txt)">${pat}</b>`;}
function menRetHTML(r){const val=numBR(r.valor), roa=numBR(r.roa);
  if(val==null||roa==null)return '<span class="mut">—</span>';
  const g=val*roa/100; return `<b style="color:var(--ok)">${rs(g)}</b><div class="small mut">${fBR(roa,2)}% de ${rs(val)}</div>`;}
function ctrlFiltra(rows){
  const fp=$("ctrlFPerfil").value, ftm=$("ctrlFTermo").value;
  return rows.filter(r=>{
    if(!r.conta) return true;
    const c=_cliDe(r.conta); if(!c) return true;
    if(fp && _perfilNorm(c.perfil)!==fp) return false;
    if(ftm && (c.termometro||"")!==ftm) return false;
    return true;
  });
}
async function controleLoad(){
  const r=await api().controle_get();
  if(!r||!r.ok){toast("Erro: "+((r&&r.erro)||""));return;}
  ctrlClientes=r.clientes||[];
  ctrlData={mensal:r.mensal||[],montagem:r.montagem||[],ativos:r.ativos||[]};
  cliDLFill(ctrlClientes);
  const sel=$("ctrlMonCli"), cur=sel.dataset.conta||"";
  if(!ctrlClientes.some(c=>String(c.conta)===cur)){sel.dataset.conta="";sel.value="";}
  else sel.value=cliTxtDe(cur);
  const selA=$("ctrlAtvCli"), curA=(selA&&selA.dataset.conta)||"";
  if(selA){
    if(!ctrlClientes.some(c=>String(c.conta)===curA)){selA.dataset.conta="";selA.value="";}
    else selA.value=cliTxtDe(curA);
  }
  ctrlSub(ctrlSubAtivo);
}
function ctrlSub(v){
  ctrlSubAtivo=v;
  $("pillCtrlMen").classList.toggle("sel",v==="men");
  $("pillCtrlMon").classList.toggle("sel",v==="mon");
  $("pillCtrlAtv").classList.toggle("sel",v==="atv");
  $("ctrlMen").style.display=v==="men"?"block":"none";
  $("ctrlMon").style.display=v==="mon"?"flex":"none";
  $("ctrlAtv").style.display=v==="atv"?"block":"none";
  ctrlRender();
}
function ctrlRender(){
  if(ctrlSubAtivo==="men") renderCtrlMen();
  else if(ctrlSubAtivo==="atv") renderCtrlAtv();
  else ctrlMonRender();
}
const ATV_IDX=["","Pós fixado","Prefixado","Inflação","Alternativos","Outros"];
let atvCart="", atvCmpSel=null, ctrlAtvView=[];
function _atvCalc(r){
  const fin=numBR(r.financeiro), y=numBR(r["yield"]);
  const dm=(fin!=null&&y!=null)?fin*y/100:null;
  return {fin:fin, y:y, dm:dm, ya:(y!=null)?y*12:null, da:(dm!=null)?dm*12:null};
}
function _atvTotais(rows){
  let fin=0,dm=0,da=0;
  rows.forEach(r=>{const c=_atvCalc(r); if(c.fin)fin+=c.fin; if(c.dm)dm+=c.dm; if(c.da)da+=c.da;});
  const y=fin>0?dm/fin*100:null;
  return {n:rows.length,fin,dm,da,y,ya:y!=null?y*12:null};
}
function atvContaAtual(){return ($("ctrlAtvCli")&&$("ctrlAtvCli").dataset.conta)||"";}
function atvVariantes(conta){
  return [...new Set((ctrlData.ativos||[])
    .filter(r=>String(r.conta||"").trim()===conta)
    .map(r=>(r.carteira||"Principal")))];
}
function atvRowsDe(conta,cart){
  return (ctrlData.ativos||[]).filter(r=>String(r.conta||"").trim()===conta
    &&(r.carteira||"Principal")===cart);
}
function renderCtrlAtv(){
  const conta=atvContaAtual();
  const selEl=$("ctrlAtvCartSel");
  if(!conta){
    if(selEl)selEl.innerHTML="";
    $("ctrlAtvBody").innerHTML=`<tr><td colspan=9 class="mut">Selecione um cliente para montar a carteira de ativos.</td></tr>`;
    $("ctrlAtvTot").innerHTML=""; $("ctrlAtvCmp").innerHTML=""; ctrlAtvView=[];
    return;
  }
  const vars=atvVariantes(conta);
  if(!atvCart)atvCart=vars[0]||"Principal";
  const opcoes=vars.includes(atvCart)?vars:vars.concat(atvCart);
  if(selEl)selEl.innerHTML=opcoes.map(v=>`<option value="${esc(v)}" ${v===atvCart?"selected":""}>${esc(v)}</option>`).join("");
  ctrlAtvView=atvRowsDe(conta,atvCart);
  $("ctrlAtvBody").innerHTML=ctrlAtvView.map((r,i)=>{
    const c=_atvCalc(r);
    const idxSel=ATV_IDX.map(o=>`<option value="${o}" ${o===(r.indexador||"")?"selected":""}>${o||"—"}</option>`).join("");
    return `<tr>
      <td><input value="${esc(r.ativo)}" oninput="ctrlAtvView[${i}].ativo=this.value;ctrlSaveSoon('ativos')" placeholder="Ativo"></td>
      <td><select onchange="ctrlAtvView[${i}].indexador=this.value;ctrlSaveSoon('ativos')">${idxSel}</select></td>
      <td><input value="${esc(r.segmento)}" oninput="ctrlAtvView[${i}].segmento=this.value;ctrlSaveSoon('ativos')" placeholder="Segmento"></td>
      <td><input value="${esc(fmtRS(r.financeiro))}" oninput="maskMoeda(this);ctrlAtvView[${i}].financeiro=this.value;ctrlAtvCalc(${i});ctrlSaveSoon('ativos')" placeholder="R$"></td>
      <td><input value="${esc(r["yield"])}" oninput="ctrlAtvView[${i}]['yield']=this.value;ctrlAtvCalc(${i});ctrlSaveSoon('ativos')" placeholder="%" style="text-align:center"></td>
      <td id="atvdm_${i}" class="dc-mono" style="color:var(--ok);font-weight:600">${c.dm!=null?rs(c.dm):'<span class="mut">—</span>'}</td>
      <td id="atvya_${i}" class="dc-mono" style="text-align:center">${c.ya!=null?fBR(c.ya,2)+"%":'<span class="mut">—</span>'}</td>
      <td id="atvda_${i}" class="dc-mono" style="color:var(--ok);font-weight:700">${c.da!=null?rs(c.da):'<span class="mut">—</span>'}</td>
      <td><span class="del" onclick="ctrlAtvDel(${i})">✕</span></td></tr>`;
  }).join("")||`<tr><td colspan=9 class="mut">Carteira “${esc(atvCart)}” vazia — clique em “Adicionar ativo”.</td></tr>`;
  ctrlAtvRefreshTot();
  $("ctrlAtvCmp").innerHTML=atvCmpHTML(conta,vars);
}
function ctrlAtvCalc(i){
  const c=_atvCalc(ctrlAtvView[i]);
  const dm=$("atvdm_"+i), ya=$("atvya_"+i), da=$("atvda_"+i);
  if(dm)dm.innerHTML=c.dm!=null?rs(c.dm):'<span class="mut">—</span>';
  if(ya)ya.innerHTML=c.ya!=null?fBR(c.ya,2)+"%":'<span class="mut">—</span>';
  if(da)da.innerHTML=c.da!=null?rs(c.da):'<span class="mut">—</span>';
  ctrlAtvRefreshTot();
  const conta=atvContaAtual();
  $("ctrlAtvCmp").innerHTML=atvCmpHTML(conta,atvVariantes(conta));
}
function ctrlAtvRefreshTot(){
  const t=_atvTotais(ctrlAtvView);
  const tot=(lab,val)=>`<div><div style="font:700 10px 'Manrope';letter-spacing:.12em;text-transform:uppercase;color:#9c948a;margin-bottom:2px">${lab}</div>
    <div class="dc-mono" style="font-size:15px;font-weight:700;color:#F2C029">${val}</div></div>`;
  $("ctrlAtvTot").innerHTML=t.n?
    tot("Financeiro total",rs(t.fin))
    +tot("Yield médio",t.y!=null?fBR(t.y,2)+"%":"—")
    +tot("Dividendo mensal",rs(t.dm))
    +tot("Yield anualizado",t.ya!=null?fBR(t.ya,2)+"%":"—")
    +tot("Dividendo anualizado",rs(t.da))
    +`<div class="grow"></div><span style="align-self:flex-end;font:500 10.5px 'Manrope';color:#5A5E67">${t.n} ativo(s) · carteira “${esc(atvCart)}”</span>`
    :'<span class="mut small">Adicione ativos para ver os totais.</span>';
}
let atvCmpOpen=false;
function atvCmpAbrir(){
  atvCmpOpen=!atvCmpOpen;
  if(atvCmpOpen)atvCmpSel=new Set(atvVariantes(atvContaAtual()));
  const b=$("btnAtvCmp");
  if(b){b.style.borderColor=atvCmpOpen?"rgba(242,192,41,.5)":"";b.style.color=atvCmpOpen?"#F2C029":"";}
  renderCtrlAtv();
}
function atvCmpToggle(v,on){
  if(!atvCmpSel)atvCmpSel=new Set(atvVariantes(atvContaAtual()));
  on?atvCmpSel.add(v):atvCmpSel.delete(v);
  renderCtrlAtv();
}
function atvCmpHTML(conta,vars){
  if(!atvCmpOpen)return "";
  if(vars.length<2)return `<div class="card" style="border-color:rgba(255,255,255,.08)">
    <span class="mut small">Crie uma segunda variação de carteira para comparar.</span></div>`;
  if(!atvCmpSel)atvCmpSel=new Set(vars);
  const marcadas=vars.filter(v=>atvCmpSel.has(v));
  const dados=marcadas.map(v=>({v:v,t:_atvTotais(atvRowsDe(conta,v))}));
  const melhor={};
  ["y","dm","da"].forEach(k=>{
    const vals=dados.map(d=>d.t[k]).filter(x=>x!=null);
    melhor[k]=vals.length?Math.max(...vals):null;
  });
  const cel=(val,txt,k,t)=>`<td class="dc-mono" style="font-weight:700;color:${(melhor[k]!=null&&t[k]===melhor[k]&&t[k]!=null)?'#F2C029':'#C9CBD1'}">${txt}${(melhor[k]!=null&&t[k]===melhor[k]&&t[k]!=null&&marcadas.length>1)?" ★":""}</td>`;
  const chk=vars.map(v=>`<label class="mesa-tgl-wrap" style="flex:none"><input type="checkbox" ${atvCmpSel.has(v)?"checked":""}
      onchange="atvCmpToggle('${esc(v)}',this.checked)">${esc(v)}</label>`).join(" ");
  const linhas=dados.map(d=>`<tr>
      <td style="font-weight:600;color:${d.v===atvCart?'#F2C029':'#E4E6EA'}">${esc(d.v)}${d.v===atvCart?" (atual)":""}</td>
      <td class="dc-mono">${d.t.n}</td>
      <td class="dc-mono">${rs(d.t.fin)}</td>
      ${cel(d.t.y,d.t.y!=null?fBR(d.t.y,2)+"%":"—","y",d.t)}
      ${cel(d.t.dm,rs(d.t.dm),"dm",d.t)}
      ${cel(d.t.da,rs(d.t.da),"da",d.t)}
    </tr>`).join("");
  return `<div class="card" style="margin-bottom:12px;border-color:rgba(255,255,255,.08)">
    <div class="row" style="gap:12px;flex-wrap:wrap;margin-bottom:10px">
      <span style="font:700 11px 'Manrope';letter-spacing:.1em;text-transform:uppercase;color:#9c948a">Comparar carteiras</span>${chk}</div>
    ${marcadas.length?`<table class="cli"><thead><tr><th>Carteira</th><th>Ativos</th><th>Financeiro</th><th>Yield médio</th><th>Div. mensal</th><th>Div. anualizado</th></tr></thead>
    <tbody>${linhas}</tbody></table>
    <div class="small mut" style="margin-top:8px">★ = melhor valor entre as selecionadas (maior yield/dividendo).</div>`
    :'<span class="mut small">Marque ao menos uma carteira para comparar.</span>'}</div>`;
}
function ctrlAtvNovaCart(){
  if(!atvContaAtual()){toast("Selecione um cliente primeiro.");return;}
  const el=$("ctrlAtvNova"), nome=(el&&el.value||"").trim();
  if(!nome){toast("Dê um nome para a variação (ex.: Conservadora).");return;}
  const jaExiste=atvVariantes(atvContaAtual()).includes(nome);
  atvCart=nome; if(el)el.value=""; atvCmpSel=null;
  if(jaExiste){renderCtrlAtv(); toast(`Carteira “${nome}” já existia — selecionada.`); return;}
  ctrlAtvAdd();
  toast(`Carteira “${nome}” criada e salva — preencha os ativos.`);
}
function ctrlAtvExcluirCart(){
  const conta=atvContaAtual();
  if(!conta||!atvCart)return;
  modal("Excluir carteira",`Excluir a carteira “${atvCart}” e todos os ativos dela?`,()=>{
    ctrlData.ativos=ctrlData.ativos.filter(r=>!(String(r.conta||"").trim()===conta
      &&(r.carteira||"Principal")===atvCart));
    atvCart=""; atvCmpSel=null;
    renderCtrlAtv(); ctrlSaveSoon('ativos');
  });
}
function ctrlAtvAdd(){
  const conta=atvContaAtual();
  if(!conta){toast("Selecione um cliente primeiro.");return;}
  ctrlData.ativos.push({conta:conta,carteira:atvCart||"Principal",
    ativo:"",indexador:"",segmento:"",financeiro:"","yield":""});
  renderCtrlAtv(); ctrlSaveSoon('ativos');
}
function ctrlAtvDel(i){
  const obj=ctrlAtvView[i];
  const j=ctrlData.ativos.indexOf(obj);
  if(j>=0)ctrlData.ativos.splice(j,1);
  renderCtrlAtv(); ctrlSaveSoon('ativos');
}
function ctrlSaveSoon(key){
  clearTimeout(_ctrlT[key]);
  if($("ctrlStatus"))$("ctrlStatus").textContent="salvando…";
  _ctrlT[key]=setTimeout(async()=>{
    const r=await api().controle_save(key,ctrlData[key]);
    if($("ctrlStatus"))$("ctrlStatus").textContent=(r&&r.ok)?"salvo ✓":"erro ao salvar";
    if(r&&!r.ok)toast("Erro ao salvar: "+(r.erro||""));
  },600);
}
function renderCtrlMen(){
  ctrlMenView=ctrlFiltra(ctrlData.mensal);
  $("ctrlMenBody").innerHTML=ctrlMenView.map((r,i)=>{
    return `<tr>
      <td><input value="${esc(r.mes)}" oninput="ctrlMenView[${i}].mes=maskMes(this);ctrlSaveSoon('mensal')" placeholder="mm/aaaa" maxlength="7"></td>
      <td><input value="${esc(cliTxtDe(r.conta))}" data-src="geral" data-pick="men" data-i="${i}" ${CBX_ATTRS}>
        <div class="small mut" style="margin-top:3px">${cliInfoLine(r.conta)}</div></td>
      <td><input value="${esc(r.ativo)}" oninput="ctrlMenView[${i}].ativo=this.value;ctrlSaveSoon('mensal')" placeholder="Ativo"></td>
      <td><input id="menval_${i}" value="${esc(fmtRS(r.valor))}" oninput="maskMoeda(this);ctrlMenVal(${i},this.value)" onblur="this.value=fmtRS(this.value);ctrlMenView[${i}].valor=this.value" placeholder="R$"></td>
      <td><input id="menpct_${i}" value="${esc(r.pct)}" oninput="ctrlMenPct(${i},this.value)" placeholder="%"></td>
      <td><input value="${esc(r.roa)}" oninput="ctrlMenRoa(${i},this.value)" placeholder="%"></td>
      <td id="menret_${i}">${menRetHTML(r)}</td>
      <td><input value="${esc(r.data_liq)}" oninput="ctrlMenView[${i}].data_liq=maskData(this);ctrlSaveSoon('mensal')" placeholder="dd/mm/aaaa" maxlength="10"></td>
      <td><span class="del" onclick="ctrlDel('mensal',${i})">✕</span></td></tr>`;
  }).join("")||`<tr><td colspan=9 class="mut">Sem alocações — clique em “Adicionar alocação”.</td></tr>`;
}
function _menRet(i){const r=ctrlMenView[i]; const val=numBR(r.valor), roa=numBR(r.roa);
  r.retorno=(val!=null&&roa!=null)?fBR(val*roa/100,2):"";
  const e=$("menret_"+i); if(e)e.innerHTML=menRetHTML(r);}
function ctrlMenVal(i,v){const r=ctrlMenView[i]; r.valor=v; const p=_patDe(r.conta), n=numBR(v);
  if(p&&n!=null){r.pct=fBR(n/p*100,2); const e=$("menpct_"+i); if(e)e.value=r.pct;} _menRet(i); ctrlSaveSoon('mensal');}
function ctrlMenPct(i,v){const r=ctrlMenView[i]; r.pct=v; const p=_patDe(r.conta), n=numBR(v);
  if(p&&n!=null){r.valor="R$ "+fBR(p*n/100,2); const e=$("menval_"+i); if(e)e.value=r.valor;} _menRet(i); ctrlSaveSoon('mensal');}
function ctrlMenRoa(i,v){const r=ctrlMenView[i]; r.roa=v; _menRet(i); ctrlSaveSoon('mensal');}
function ctrlAdd(t){
  ctrlData.mensal.push({mes:"",conta:"",ativo:"",valor:"",pct:"",roa:"",retorno:"",data_liq:""}); renderCtrlMen();
}
function ctrlDel(key,i){const obj=ctrlMenView[i];
  const j=ctrlData[key].indexOf(obj); if(j>=0)ctrlData[key].splice(j,1); ctrlRender(); ctrlSaveSoon(key);}
function ctrlMonRender(){
  const conta=$("ctrlMonCli").dataset.conta||"", c=_cliDe(conta), pat=_patDe(conta);
  $("ctrlMonInfo").innerHTML=c?`Perfil <b style="color:var(--yellow)">${esc(c.perfil||"—")}</b> · Termô. <b style="color:${_TCD[(c.termometro||"").toUpperCase()]||'var(--mut)'}">${esc((c.termometro||"—").toUpperCase())}</b> · Patrimônio <b style="color:var(--txt)">${pat!=null?rs(pat):"—"}</b>`:"Selecione um cliente.";
  ctrlMonSaiView=ctrlData.montagem.filter(r=>String(r.conta)===String(conta)&&r.direcao==="Sai");
  ctrlMonEntraView=ctrlData.montagem.filter(r=>String(r.conta)===String(conta)&&r.direcao==="Entra");
  $("ctrlMonSai").innerHTML=montRows(ctrlMonSaiView,"ctrlMonSaiView","Sai");
  $("ctrlMonEntra").innerHTML=montRows(ctrlMonEntraView,"ctrlMonEntraView","Entra");
  ctrlMonRefreshTot();
}
function montRows(view,vn,dir){
  return view.map((r,i)=>`<tr>
    <td><input value="${esc(r.ativo)}" oninput="${vn}[${i}].ativo=this.value;ctrlSaveSoon('montagem')" placeholder="Ativo"></td>
    <td><input value="${esc(fmtRS(r.valor))}" oninput="maskMoeda(this);${vn}[${i}].valor=this.value;ctrlMonRefreshTot();ctrlSaveSoon('montagem')" onblur="this.value=fmtRS(this.value);${vn}[${i}].valor=this.value" placeholder="R$"></td>
    <td><span class="del" onclick="ctrlMonDel('${dir}',${i})">✕</span></td></tr>`).join("")
    ||`<tr><td colspan=3 class="mut">—</td></tr>`;
}
function ctrlMonRefreshTot(){
  const sSai=ctrlMonSaiView.reduce((a,r)=>a+(numBR(r.valor)||0),0);
  const sEntra=ctrlMonEntraView.reduce((a,r)=>a+(numBR(r.valor)||0),0);
  const conta=$("ctrlMonCli").dataset.conta||"";
  const pat=_patDe(conta);
  $("ctrlMonSaiTot").innerHTML=`<span class="dc-mono" style="font-size:13.5px;font-weight:700;color:#EF4444">Total: ${rs(sSai)}</span>`;
  $("ctrlMonEntraTot").innerHTML=`<span class="dc-mono" style="font-size:13.5px;font-weight:700;color:#10B981">Total: ${rs(sEntra)}</span>`;
  const net=sEntra-sSai;
  const cor=net>0?"#10B981":(net<0?"#EF4444":"var(--mut)");
  const sinal=net>0?"+ ":(net<0?"− ":"");
  const card=$("ctrlMonNetCard");
  if(card)card.style.display=conta?"":"none";
  $("ctrlMonNet").innerHTML=conta
    ? `Saldo (entra − sai): <b class="dc-mono" style="color:${cor}">${sinal}${rs(Math.abs(net))}</b>`
      +(pat?` · <b style="color:${cor}">${net>0?"+":(net<0?"−":"")}${fBR(Math.abs(net)/pat*100,2)}%</b> <span class="mut">do patrimônio</span>`:"")
    : "";
}
function ctrlMonAdd(dir){const conta=$("ctrlMonCli").dataset.conta||""; if(!conta){toast("Selecione um cliente primeiro.");return;}
  ctrlData.montagem.push({conta:conta,direcao:dir,ativo:"",valor:""}); ctrlMonRender(); ctrlSaveSoon('montagem');}
function ctrlMonDel(dir,i){const view=dir==="Sai"?ctrlMonSaiView:ctrlMonEntraView; const obj=view[i];
  const j=ctrlData.montagem.indexOf(obj); if(j>=0)ctrlData.montagem.splice(j,1); ctrlMonRender(); ctrlSaveSoon('montagem');}
const DTC={A:"#10B981",B:"#3B82F6",C:"#F59E0B",D:"#F97316",E:"#EF4444"};
const ENQ_COR={"Bom":"#10B981","Médio":"#F59E0B","Ruim":"#F97316","Crítico":"#EF4444","s/ dados":"var(--mut2)"};
const milM=v=>v==null?"—":(v>=1e6?"R$ "+fBR(v/1e6,1)+"M":(v>=1e3?"R$ "+fBR(v/1e3,0)+"k":"R$ "+fBR(v,0)));
function nChip(n,tot,txt){
  n=n||0;
  const dif=tot!=null&&tot!==n;
  return `<div style="margin-top:12px;padding-top:8px;border-top:1px solid rgba(255,255,255,.04);
    font:500 10.5px 'Manrope';color:#5A5E67;text-align:right" title="${dif?'Clientes fora deste gráfico não têm o dado necessário':'Todos os clientes considerados'}">
    ${n}${dif?` de ${tot}`:""} ${txt||"cliente(s) analisado(s)"}</div>`;
}
function dashCard(titulo,corpo){return `<div class="bn-card"><div class="bn-lbl">${titulo}</div>${corpo}</div>`;}
function dashBarsV(items){
  const mx=Math.max(1,...items.map(x=>x[1]));
  return `<div class="bn-bars">`+items.map(([l,v,cor])=>`<div class="bn-bar">
    <div style="font-size:12px;font-weight:700;color:${cor}">${v}</div>
    <div class="col"><i style="height:${Math.round(v/mx*100)||3}%;background:linear-gradient(180deg,${cor},${cor}44)"></i></div>
    <div class="bn-sub">${l}</div></div>`).join("")+`</div>`;
}
function dashLista(items){
  return items.map(([l,v,pct,cor])=>`<div style="margin-bottom:9px">
    <div class="row" style="justify-content:space-between;font-size:12.5px;margin-bottom:3px">
      <span>${esc(l)}</span><span class="bn-sub"><b style="color:var(--txt)">${v}</b></span></div>
    ${pct!=null?`<div class="bn-track"><div class="bn-fill" style="width:${Math.max(0,Math.min(pct,100))}%;background:${cor||'var(--yellow)'}"></div></div>`:""}</div>`).join("");
}
let incluirPlugados=true;
function mesaToggleHTML(){
  return `<label class="mesa-tgl-wrap" title="Quando desligado, os clientes plugados na mesa saem apenas deste gráfico de enquadramento">
    <input type="checkbox" class="mesa-tgl" ${incluirPlugados?"checked":""} onchange="setIncluirPlugados(this.checked)">
    <span>Incluir plugados na mesa</span></label>`;
}
function setIncluirPlugados(v){
  incluirPlugados=!!v;
  try{api().set_config("incluir_plugados",incluirPlugados);}catch(e){}
  dashLoad();
}
async function dashLoad(){
  const r=await api().dashboard_data(incluirPlugados);
  if(!r||!r.ok){toast("Erro: "+((r&&r.erro)||""));return;}
  renderDash(r);
}
function kpro(icon,cor,valor,label,sub){
  return `<div class="kpro">
    <div class="ic" style="background:linear-gradient(135deg,${cor},${cor}aa)">${icon}</div>
    <div class="v">${valor}</div>
    <div class="l">${label}</div>${sub?`<div class="s">${sub}</div>`:""}</div>`;
}
const ICO={users:'<svg viewBox="0 0 24 24"><path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/></svg>',
  money:'<svg viewBox="0 0 24 24"><line x1="12" y1="1" x2="12" y2="23"/><path d="M17 5H9.5a3.5 3.5 0 0 0 0 7h5a3.5 3.5 0 0 1 0 7H6"/></svg>',
  trend:'<svg viewBox="0 0 24 24"><path d="M3 17l6-6 4 4 8-8"/><path d="M21 7v6h-6"/></svg>',
  check:'<svg viewBox="0 0 24 24"><path d="M22 11.08V12a10 10 0 1 1-5.93-9.14"/><path d="M22 4 12 14.01l-3-3"/></svg>'};
function donutSVG(segs,centerTop,centerBot){
  const tot=segs.reduce((a,s)=>a+s.value,0)||1, R=52, C=2*Math.PI*R; let off=0;
  const rings=segs.filter(s=>s.value>0).map(s=>{const len=s.value/tot*C;
    const el=`<circle cx="70" cy="70" r="${R}" fill="none" stroke="${s.color}" stroke-width="15" stroke-dasharray="${len.toFixed(2)} ${(C-len).toFixed(2)}" stroke-dashoffset="${(-off).toFixed(2)}" transform="rotate(-90 70 70)" stroke-linecap="butt"/>`;
    off+=len; return el;}).join("");
  return `<svg viewBox="0 0 140 140" width="132" height="132" style="flex:none">
    <circle cx="70" cy="70" r="52" fill="none" stroke="var(--border)" stroke-width="15"/>${rings}
    <text x="70" y="67" text-anchor="middle" fill="var(--yellow)" font-size="27" font-weight="800">${centerTop}</text>
    <text x="70" y="87" text-anchor="middle" fill="var(--mut)" font-size="10">${centerBot}</text></svg>`;
}
function dashPanel(titulo,story,corpo,sp){
  return `<div class="bn-card ${sp||'sp2'}"><div class="bn-lbl">${titulo}</div>${story?`<div class="bn-story">${story}</div>`:""}${corpo}</div>`;
}
function hbars(items){
  return items.map(([l,v,pct,cor])=>`<div class="hbar"><span class="lab">${esc(l)}</span>
    <span class="trk"><i style="width:${Math.max(2,Math.min(pct||0,100))}%;background:${cor||'var(--yellow)'}"></i></span>
    <span class="val">${v}</span></div>`).join("");
}
const _DICO={
 users:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M16 21v-2a4 4 0 0 0-4-4H6a4 4 0 0 0-4 4v2"></path><circle cx="9" cy="7" r="4"></circle><path d="M22 21v-2a4 4 0 0 0-3-3.87"></path><path d="M16 3.13a4 4 0 0 1 0 7.75"></path></svg>',
 wallet:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M21 12V7H5a2 2 0 0 1 0-4h14v4"></path><path d="M3 5v14a2 2 0 0 0 2 2h16v-5"></path><path d="M18 12a2 2 0 0 0 0 4h4v-4Z"></path></svg>',
 trend:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><polyline points="22 7 13.5 15.5 8.5 10.5 2 17"></polyline><polyline points="16 7 22 7 22 13"></polyline></svg>',
 target:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="9"></circle><circle cx="12" cy="12" r="5"></circle><circle cx="12" cy="12" r="1.4"></circle></svg>',
 pie:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M21.21 15.89A10 10 0 1 1 8 2.83"></path><path d="M22 12A10 10 0 0 0 12 2v10z"></path></svg>',
 thermo:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M14 14.76V3.5a2.5 2.5 0 0 0-5 0v11.26a4.5 4.5 0 1 0 5 0z"></path></svg>',
 bars:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><line x1="12" y1="20" x2="12" y2="10"></line><line x1="18" y1="20" x2="18" y2="4"></line><line x1="6" y1="20" x2="6" y2="16"></line></svg>',
 alert:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="9"></circle><line x1="12" y1="8" x2="12" y2="12"></line><line x1="12" y1="16" x2="12.01" y2="16"></line></svg>',
 user:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M20 21v-2a4 4 0 0 0-4-4H8a4 4 0 0 0-4 4v2"></path><circle cx="12" cy="7" r="4"></circle></svg>',
 pin:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><path d="M21 10c0 7-9 13-9 13s-9-6-9-13a9 9 0 0 1 18 0z"></path><circle cx="12" cy="10" r="3"></circle></svg>',
 activity:'<svg viewBox="0 0 24 24" fill="none" stroke-linecap="round" stroke-linejoin="round"><polyline points="22 12 18 12 15 21 9 3 6 12 2 12"></polyline></svg>'};
function renderDash(d){
  const W=$("dashWrap");
  if(!d||!d.total){W.innerHTML=`<div class="dc-card" style="color:#8A8E98">Sem clientes ainda — leia os PDFs e preencha o cadastro.</div>`;return;}
  const TO=["A","B","C","D","E"], tot=d.total, ic=_DICO;
  const CLARO="#E9EAEC", LARANJA="#E2811F", LARANJA2="#B14A17";
  const P1="#F2C029",P2="#F2C029",P3=LARANJA,P4=LARANJA2,P4L=LARANJA;
  const TC={A:"#F2C029",B:"#F2C029",C:"#F2C029",D:"#F2C029",E:"#F2C029"};
  const TCL={A:CLARO,B:CLARO,C:CLARO,D:CLARO,E:CLARO};
  const TCS={A:"rgba(255,255,255,.07)",B:"rgba(255,255,255,.07)",C:"rgba(255,255,255,.07)",D:"rgba(255,255,255,.07)",E:"rgba(255,255,255,.07)"};
  const GREEN="#F2C029",GOLD="#F2C029",ORANGE=LARANJA,RED=LARANJA,BLUE="#F2C029";
  const SB={gold:"rgba(242,192,41,.55)",green:"rgba(242,192,41,.55)",orange:"rgba(226,129,31,.55)",blue:"rgba(242,192,41,.55)",red:"rgba(226,129,31,.6)"};
  const sgn=v=>(v>0?"+":"")+fBR(v,2)+"%";
  const H=(svg,t)=>`<div class="dc-h">${svg}<span>${t}</span></div>`;
  const ST=(cor,html)=>`<div class="dc-story" style="border-left-color:${cor}">${html}</div>`;
  const enqOk=["Bom","Médio","Ruim","Crítico"].reduce((a,k)=>a+(d.enq[k]||0),0);
  const pctBom=enqOk?Math.round((d.enq["Bom"]||0)/enqOk*100):0;
  const needRebal=(d.enq["Ruim"]||0)+(d.enq["Crítico"]||0);
  const enqKpiCor=!d.tem_aloc?"#71757F":(pctBom>=80?GREEN:(pctBom>=50?GOLD:ORANGE));
  const rmNull=d.rent_media==null, rmNeg=!rmNull&&d.rent_media<0;
  const KCOL=[
   {c:"#F2C029",glow:"rgba(242,192,41,.05)",bd:"rgba(255,255,255,.06)",lab:"#9c948a"},
   {c:"#F2C029",glow:"rgba(242,192,41,.05)",bd:"rgba(255,255,255,.06)",lab:"#9c948a"},
   {c:"#F2C029",glow:"rgba(242,192,41,.05)",bd:"rgba(255,255,255,.06)",lab:"#9c948a"},
   {c:"#F2C029",glow:"rgba(242,192,41,.05)",bd:"rgba(255,255,255,.06)",lab:"#9c948a"}
  ];
  const kbox=(t,icon,numTxt,numCor,lab,extra,sub)=>`<div class="dc-kpi" style="border-color:${t.bd}">
      <div style="position:absolute;top:-54px;right:-34px;width:180px;height:180px;background:radial-gradient(circle,${t.glow},transparent 66%);pointer-events:none"></div>
      <div style="display:flex;justify-content:space-between;align-items:flex-start;position:relative"><div class="dc-ic2" style="color:${t.c}">${icon}</div>${extra||""}</div>
      <div class="num2" style="color:${numCor};position:relative">${numTxt}</div>
      <div class="lab" style="color:${t.lab};position:relative">${lab}</div>${sub?`<div class="sub" style="position:relative">${sub}</div>`:""}
    </div>`;
  const spk=(stroke,fill)=>`<svg width="70" height="30" viewBox="0 0 70 30" fill="none" style="opacity:.95"><polyline points="2,24 13,20 24,22 35,13 46,16 57,8 68,4" stroke="${stroke}" stroke-width="2.2" stroke-linecap="round" stroke-linejoin="round"></polyline><path d="M2,24 13,20 24,22 35,13 46,16 57,8 68,4 68,30 2,30 Z" fill="${fill}"></path></svg>`;
  const rmSpark=rmNull?"":(rmNeg?spk(LARANJA,"rgba(226,129,31,.12)"):spk("#F2C029","rgba(242,192,41,.12)"));
  const enqNum=!d.tem_aloc?"—":pctBom+"%";
  const enqNumCor=!d.tem_aloc?"#71757F":"#F2C029";
  const kpiRow=`<div class="dc-row4">
    ${kbox(KCOL[0], ic.users,  tot,                          CLARO, "Clientes", "", (d.pend&&d.pend.inativos)?((tot-d.pend.inativos)+" com relatório"):"")}
    ${kbox(KCOL[1], ic.wallet, milM(d.aum),                  CLARO, "Patrimônio sob gestão")}
    ${kbox(KCOL[2], ic.trend,  rmNull?"—":sgn(d.rent_media), rmNull?"#F7F8FA":(rmNeg?LARANJA:CLARO), "Rentabilidade média (mês)", rmSpark)}
    ${kbox(KCOL[3], ic.target, enqNum, enqNumCor, "Carteiras bem enquadradas", "", d.tem_aloc?"aderência ≥ 80%":"carregue a alocação")}
  </div>`;
  const enqSegs=[["Bom","≥80%","#F2C029"],["Médio","60–79%","#B9BDC6"],["Ruim","40–59%",LARANJA],["Crítico","<40%",LARANJA2]].map(([l,r,c])=>({l,r,c,v:d.enq[l]||0}));
  const enqStoryCor=needRebal?SB.orange:(pctBom>=80?SB.green:SB.gold);
  const enqStoryHtml=!d.tem_aloc?"Carregue o relatório de alocação (aba Mercado) para avaliar o enquadramento."
    :needRebal?`<b style="color:${ORANGE}">${needRebal}</b> de ${enqOk} carteira(s) abaixo de 60% — precisam de rebalanceamento.`
    :`<b style="color:${GREEN}">${pctBom}%</b> das carteiras bem enquadradas.`;
  const enqStack=`<div style="display:flex;gap:3px;height:10px;border-radius:6px;overflow:hidden;margin-bottom:18px;background:rgba(255,255,255,.05)">`+enqSegs.filter(s=>s.v>0).map(s=>`<div style="flex:${s.v};background:${s.c};border-radius:5px"></div>`).join("")+`</div>`;
  const enqList=`<div style="display:flex;flex-direction:column">`+enqSegs.map(s=>`<div class="dc-lr"><span style="display:flex;align-items:center;gap:9px"><span style="width:8px;height:8px;border-radius:50%;background:${s.c}"></span><span style="font-size:13px;font-weight:600;color:#D2D5DB">${s.l}</span><span style="font-size:12px;color:#62666F;font-weight:500">${s.r}</span></span><span class="dc-mono" style="font-size:14px;font-weight:600;color:${s.v?'#F2F3F5':'#5A5E67'}">${s.v}</span></div>`).join("")+`</div>`;
  const enqBody=d.tem_aloc?(`<div style="display:flex;align-items:baseline;gap:9px;margin-bottom:14px"><span class="dc-mono" style="font-size:34px;font-weight:700;letter-spacing:-.02em;line-height:1;color:${enqKpiCor}">${pctBom}%</span><span style="font-size:13px;color:#8A8E98;font-weight:500">bem enquadradas</span></div>`+enqStack+enqList):"";
  const cardEnq=`<div class="dc-card">${H(ic.pie,"Enquadramento das carteiras")}
    <div style="margin:-4px 0 14px">${mesaToggleHTML()}</div>
    ${ST(enqStoryCor,enqStoryHtml)}${enqBody}
    ${d.tem_aloc?nChip(enqOk+(d.enq["s/ dados"]||0),tot,"cliente(s) avaliados"):""}</div>`;
  const tmax=Math.max(1,...TO.map(k=>d.termo_qtd[k]||0));
  const tTop=TO.map(k=>[k,d.termo_qtd[k]||0]).sort((a,b)=>b[1]-a[1])[0];
  const tStory=tTop&&tTop[1]?`Maior concentração no termômetro <b style="color:${TCL[tTop[0]]}">${tTop[0]}</b>: ${tTop[1]} de ${tot} cliente(s).`:"Nenhum termômetro definido ainda.";
  const tCounts=`<div style="display:grid;grid-template-columns:repeat(5,1fr);text-align:center;margin-bottom:10px">`+TO.map(k=>{const v=d.termo_qtd[k]||0;return `<span class="dc-mono" style="font-size:15px;font-weight:700;color:${v?TCL[k]:'#5A5E67'}">${v}</span>`;}).join("")+`</div>`;
  const tBars=`<div style="display:grid;grid-template-columns:repeat(5,1fr);align-items:end;height:96px;border-bottom:1px solid rgba(255,255,255,.05)">`+TO.map(k=>{const v=d.termo_qtd[k]||0;const h=v?Math.max(12,Math.round(v/tmax*96)):7;const bg=v?`linear-gradient(180deg,${TC[k]}d9,${TC[k]}4d)`:TCS[k];const r=v?"8px 8px 0 0":"4px 4px 0 0";return `<div style="display:flex;justify-content:center"><div style="width:60%;max-width:46px;height:${h}px;border-radius:${r};background:${bg}"></div></div>`;}).join("")+`</div>`;
  const tLabs=`<div style="display:grid;grid-template-columns:repeat(5,1fr);text-align:center;margin-top:10px">`+TO.map(k=>`<span style="font-size:12px;font-weight:600;color:${(d.termo_qtd[k]||0)?'#9A9EA8':'#62666F'}">${k}</span>`).join("")+`</div>`;
  const cardTermo=`<div class="dc-card">${H(ic.thermo,"Clientes por termômetro")}${ST(SB.green,tStory)}${tCounts}${tBars}${tLabs}
    ${nChip(TO.reduce((a,k)=>a+(d.termo_qtd[k]||0),0),tot)}</div>`;
  const rentArr=TO.map(k=>[k,d.termo_rent[k]]).filter(x=>x[1]!=null).sort((a,b)=>b[1]-a[1]);
  const rentMax=Math.max(0.5,...rentArr.map(x=>Math.abs(x[1])));
  const rentStory=d.rent_media_tot==null?"Sem relatórios para calcular rentabilidade.":`Retorno médio (desde início / 24M): <b style="color:${GREEN}">${sgn(d.rent_media_tot)}</b>`+(rentArr.length?`; melhor termômetro: <b style="color:${TCL[rentArr[0][0]]}">${rentArr[0][0]}</b> (${sgn(rentArr[0][1])}).`:".");
  const rentBars=rentArr.length?rentArr.map(([k,v])=>`<div style="display:grid;grid-template-columns:22px 1fr auto;align-items:center;gap:14px;padding:10px 0"><span style="font-size:13px;font-weight:700;color:#D2D5DB">${k}</span><div style="height:8px;border-radius:999px;background:rgba(255,255,255,.05);overflow:hidden"><div style="height:100%;width:${Math.max(3,Math.abs(v)/rentMax*100).toFixed(0)}%;border-radius:999px;background:${v<0?'linear-gradient(90deg,rgba(226,129,31,.5),#E2811F)':`linear-gradient(90deg,${TC[k]}80,${TC[k]})`}"></div></div><span class="dc-mono" style="font-size:14px;font-weight:600;color:${v<0?P4L:TCL[k]}">${sgn(v)}</span></div>`).join(""):`<div style="color:#62666F;font-size:13px;padding:10px 0">Sem dados de rentabilidade por termômetro.</div>`;
  const tComCli=TO.filter(k=>d.termo_qtd[k]).length;
  const rentFoot=rentArr.length?`<div style="margin-top:18px;padding-top:14px;border-top:1px solid rgba(255,255,255,.05);font-size:12px;color:#62666F;line-height:1.5">${rentArr.length} de ${tComCli||rentArr.length} termômetro(s) com histórico suficiente (≥24M) para o cálculo.</div>`:"";
  const cardRent=`<div class="dc-card">${H(ic.trend,"Rentabilidade (desde início / 24M) · por termômetro")}${ST(SB.green,rentStory)}${rentBars}${rentFoot}
    ${nChip(d.n_rent_termo||0,tot)}</div>`;
  const pendRows=[["Sem previdência",d.pend.prev||0,"prev"],["Sem seguro",d.pend.seg||0,"seg"],
    ["Sem internacional",d.pend.intl||0,"intl"],["Não plugados na mesa",d.pend.mesa||0,"mesa"],
    ["Sem termômetro",d.pend.termo||0,"termo"],["Inativos (sem relatório)",d.pend.inativos||0,"inativos"]];
  const pendTot=pendRows.reduce((a,r)=>a+r[1],0);
  const pendStory=pendTot?`<b style="color:${RED}">${pendTot}</b> pendência(s) de cadastro a resolver.`:"Cadastro completo — nada pendente.";
  const pendBody=pendRows.map(([l,v,k])=>`<div class="dc-lr"${v?` style="cursor:pointer" title="Clique para ver esses clientes no Cadastro" onclick="dashPendGo('${k}')" onmouseover="this.style.background='rgba(255,255,255,.04)'" onmouseout="this.style.background=''"`:""}>${v?`<span style="display:flex;align-items:center;gap:9px;font-size:13.5px;font-weight:600;color:#E4E6EA"><span style="width:7px;height:7px;border-radius:50%;background:${LARANJA}"></span>${l}</span>`:`<span style="font-size:13.5px;font-weight:500;color:#9A9EA8;padding-left:16px">${l}</span>`}<span class="dc-mono" style="font-size:14px;font-weight:${v?700:600};color:${v?LARANJA:'#5A5E67'}">${v}</span></div>`).join("");
  const cardPend=`<div class="dc-card">${H(ic.alert,"Pendências de cadastro")}${ST(pendTot?SB.red:SB.green,pendStory)}${pendBody}
    ${nChip(tot)}</div>`;
  const perfArr=["Conservadora","Moderada","Sofisticada","—"].map(k=>[k,d.perfil[k]||0]).filter(x=>x[1]).sort((a,b)=>b[1]-a[1]);
  const perfMax=Math.max(1,...perfArr.map(x=>x[1]));
  const perfTop=perfArr.find(x=>x[0]!=="—")||perfArr[0];
  const perfStory=perfTop?`<b style="color:${GOLD}">${perfTop[0]==="—"?"Sem perfil":perfTop[0]}</b> é o perfil predominante (${perfTop[1]} de ${tot}).`+((d.perfil["—"]||0)?` ${d.perfil["—"]} sem perfil.`:""):"Defina o perfil dos clientes no cadastro.";
  const perfBody=perfArr.length?perfArr.map(([k,n])=>`<div style="display:grid;grid-template-columns:90px 1fr auto;align-items:center;gap:14px;padding:9px 0"><span style="font-size:13px;font-weight:600;color:#D2D5DB">${k==="—"?"Sem perfil":k}</span><div style="height:7px;border-radius:999px;background:rgba(255,255,255,.05);overflow:hidden"><div style="height:100%;width:${(n/perfMax*100).toFixed(0)}%;border-radius:999px;background:linear-gradient(90deg,#F2C02988,#F2C029)"></div></div><span class="dc-mono" style="font-size:14px;font-weight:600;color:#F2F3F5">${n}</span></div>`).join(""):`<div style="color:#62666F;font-size:13px;padding:10px 0">—</div>`;
  const cardPerf=`<div class="dc-card">${H(ic.user,"Por perfil")}${ST(SB.gold,perfStory)}${perfBody}
    ${nChip(tot)}</div>`;
  const est=d.estados||[], estMax=Math.max(1,...est.map(x=>x[1]));
  const estStory=est.length?`<b style="color:${BLUE}">${est[0][0]}</b> concentra ${est[0][1]} cliente(s)`+(est.length>1?`, seguido de ${est[1][0]} (${est[1][1]}).`:"."):"Preencha o Estado no cadastro para ver a distribuição geográfica.";
  const estBody=est.length?est.slice(0,8).map(([uf,n])=>`<div style="display:grid;grid-template-columns:34px 1fr auto;align-items:center;gap:14px;padding:9px 0"><span style="font-size:13px;font-weight:600;color:#D2D5DB">${esc(uf)}</span><div style="height:7px;border-radius:999px;background:rgba(255,255,255,.05);overflow:hidden"><div style="height:100%;width:${(n/estMax*100).toFixed(0)}%;border-radius:999px;background:linear-gradient(90deg,#F2C02988,#F2C029)"></div></div><span class="dc-mono" style="font-size:14px;font-weight:600;color:#F2F3F5">${n}</span></div>`).join(""):`<div style="color:#62666F;font-size:13px;padding:10px 0">Preencha o Estado no cadastro.</div>`;
  const cardEst=`<div class="dc-card">${H(ic.pin,"Por estado")}${ST(SB.blue,estStory)}${estBody}
    ${nChip(est.reduce((a,x)=>a+x[1],0),tot)}</div>`;
  const FXC=["#F2C029","#F2C029","#F2C029","#F2C029"],FXL=[CLARO,CLARO,CLARO,CLARO],FXS=["rgba(255,255,255,.07)","rgba(255,255,255,.07)","rgba(255,255,255,.07)","rgba(255,255,255,.07)"];
  const fx=[["<30",d.faixa["<30"]||0],["30–44",d.faixa["30–44"]||0],["45–59",d.faixa["45–59"]||0],["60+",d.faixa["60+"]||0]];
  const fxMax=Math.max(1,...fx.map(x=>x[1])), temIdade=d.faixa["—"]!==tot;
  const fxTop=fx.slice().sort((a,b)=>b[1]-a[1])[0];
  const fxStory=temIdade&&fxTop&&fxTop[1]?`Maioria na faixa <b style="color:${GREEN}">${fxTop[0]}</b> (${fxTop[1]} cliente(s)).`:"Preencha a data de nascimento no cadastro.";
  let fxBody="";
  if(temIdade){
    const fxCounts=`<div style="display:grid;grid-template-columns:repeat(4,1fr);text-align:center;margin-bottom:9px">`+fx.map(([l,n],i)=>`<span class="dc-mono" style="font-size:14px;font-weight:700;color:${n?FXL[i]:'#5A5E67'}">${n}</span>`).join("")+`</div>`;
    const fxBars=`<div style="display:grid;grid-template-columns:repeat(4,1fr);align-items:end;height:74px;border-bottom:1px solid rgba(255,255,255,.05)">`+fx.map(([l,n],i)=>{const h=n?Math.max(10,Math.round(n/fxMax*74)):7;const bg=n?`linear-gradient(180deg,${FXC[i]}d9,${FXC[i]}4d)`:FXS[i];const r=n?"8px 8px 0 0":"4px 4px 0 0";return `<div style="display:flex;justify-content:center"><div style="width:55%;max-width:42px;height:${h}px;border-radius:${r};background:${bg}"></div></div>`;}).join("")+`</div>`;
    const fxLabs=`<div style="display:grid;grid-template-columns:repeat(4,1fr);text-align:center;margin-top:10px">`+fx.map(([l,n])=>`<span style="font-size:12px;font-weight:600;color:${n?'#9A9EA8':'#62666F'}">${l}</span>`).join("")+`</div>`;
    fxBody=fxCounts+fxBars+fxLabs+(d.faixa["—"]?`<div style="margin-top:12px;font-size:12px;color:#62666F">${d.faixa["—"]} sem data de nascimento</div>`:"");
  }
  const cardFx=`<div class="dc-card">${H(ic.activity,"Faixa etária")}${ST(temIdade?SB.green:SB.gold,fxStory)}${fxBody}
    ${nChip(tot)}</div>`;
  const apm=d.ap_mes||[];
  const apmMax=Math.max(1,...apm.map(x=>x.n||0));
  const apmTem=apm.some(x=>x.n>0);
  const apmStory=apmTem
    ?`<b style="color:${GOLD}">${d.ap_cli||0}</b> cliente(s) aportaram nos últimos 12 meses.`
    :"Registre os aportes nos clientes para acompanhar aqui.";
  const apmCounts=`<div style="display:grid;grid-template-columns:repeat(12,1fr);text-align:center;margin-bottom:8px">`+apm.map(x=>`<span class="dc-mono" style="font-size:12px;font-weight:700;color:${x.n?CLARO:'#5A5E67'}">${x.n}</span>`).join("")+`</div>`;
  const apmBars=`<div style="display:grid;grid-template-columns:repeat(12,1fr);gap:8px;align-items:end;height:90px;border-bottom:1px solid rgba(255,255,255,.05)">`+apm.map(x=>{const h=x.n?Math.max(10,Math.round(x.n/apmMax*90)):5;const bg=x.n?"linear-gradient(180deg,#F2C029d9,#F2C0294d)":"rgba(255,255,255,.07)";return `<div style="display:flex;justify-content:center"><div style="width:70%;max-width:34px;height:${h}px;border-radius:7px 7px 0 0;background:${bg}"></div></div>`;}).join("")+`</div>`;
  const apmLabs=`<div style="display:grid;grid-template-columns:repeat(12,1fr);text-align:center;margin-top:8px">`+apm.map(x=>`<span style="font-size:9.5px;font-weight:600;color:${x.n?'#9A9EA8':'#62666F'}">${esc((x.label||"").slice(0,2))}</span>`).join("")+`</div>`;
  const cardApMes=`<div class="dc-card">${H(ic.bars,"Clientes aportando por mês")}${ST(SB.gold,apmStory)}${apmTem?apmCounts+apmBars+apmLabs:""}
    ${nChip(d.ap_cli||0,tot,"cliente(s) com aporte em 12M")}</div>`;
  const apt=d.ap_termo||{};
  const aptMax=Math.max(1,...TO.map(k=>apt[k]||0));
  const aptTot=TO.reduce((s,k)=>s+(apt[k]||0),0);
  const aptTop=TO.map(k=>[k,apt[k]||0]).sort((a,b)=>b[1]-a[1])[0];
  const aptStory=aptTot
    ?`Termômetro <b style="color:${GOLD}">${aptTop[0]}</b> lidera: ${aptTop[1]} de ${aptTot} cliente(s) aportando (12M).`
    :"Sem aportes registrados no período.";
  const aptCounts=`<div style="display:grid;grid-template-columns:repeat(5,1fr);text-align:center;margin-bottom:10px">`+TO.map(k=>{const v=apt[k]||0;return `<span class="dc-mono" style="font-size:15px;font-weight:700;color:${v?CLARO:'#5A5E67'}">${v}</span>`;}).join("")+`</div>`;
  const aptBars=`<div style="display:grid;grid-template-columns:repeat(5,1fr);align-items:end;height:82px;border-bottom:1px solid rgba(255,255,255,.05)">`+TO.map(k=>{const v=apt[k]||0;const h=v?Math.max(12,Math.round(v/aptMax*82)):7;const bg=v?"linear-gradient(180deg,#F2C029d9,#F2C0294d)":"rgba(255,255,255,.07)";const r=v?"8px 8px 0 0":"4px 4px 0 0";return `<div style="display:flex;justify-content:center"><div style="width:60%;max-width:46px;height:${h}px;border-radius:${r};background:${bg}"></div></div>`;}).join("")+`</div>`;
  const aptLabs=`<div style="display:grid;grid-template-columns:repeat(5,1fr);text-align:center;margin-top:10px">`+TO.map(k=>`<span style="font-size:12px;font-weight:600;color:${(apt[k]||0)?'#9A9EA8':'#62666F'}">${k}</span>`).join("")+`</div>`;
  const cardApTermo=`<div class="dc-card">${H(ic.thermo,"Clientes aportando por termômetro")}${ST(SB.gold,aptStory)}${aptTot?aptCounts+aptBars+aptLabs:""}
    ${nChip(aptTot,d.ap_cli||0,"cliente(s) analisado(s)")}</div>`;
  W.innerHTML=kpiRow
    +`<div class="dc-row3">${cardEnq}${cardTermo}${cardRent}</div>`
    +`<div class="dc-row21">${cardApMes}${cardApTermo}</div>`
    +`<div class="dc-row4">${cardPend}${cardPerf}${cardEst}${cardFx}</div>`;
}

let painelDados=null, painelRange="all", painelMetaPer="tri", painelUteisPer="mes", painelPacePer="mes", painelChartMetric="valor", painelRecPer="mes", _painelPerOk=false, _notifs=[];
function setRecPer(p){painelRecPer=p; renderPainel();}
function setChartMetric(m){painelChartMetric=m; renderPainel();}
async function salvarReceita(){
  const v=numBR(($("recValor")&&$("recValor").value)||"");
  if(v==null||v<0){toast("Informe um valor de receita válido.");return;}
  const r=await api().receita_set(($("recMes")&&$("recMes").value.trim())||"",v);
  if(r&&r.ok){toast("Receita salva ✓");painelLoad();}
  else toast("Erro ao salvar receita: "+((r&&r.erro)||""));
}
function setUteisPer(p){painelUteisPer=p; renderPainel();}
function setPacePer(p){painelPacePer=p; renderPainel();}
const PIC={
  target:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#9A8A5E" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><circle cx="12" cy="12" r="9"/><circle cx="12" cy="12" r="5"/><circle cx="12" cy="12" r="1.4"/></svg>',
  trend:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><polyline points="22 7 13.5 15.5 8.5 10.5 2 17"/><polyline points="16 7 22 7 22 13"/></svg>',
  shield:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M12 22s8-4 8-10V5l-8-3-8 3v7c0 6 8 10 8 10z"/></svg>',
  chart:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M3 3v18h18"/><path d="m19 9-5 5-4-4-3 3"/></svg>',
  pie:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M21.21 15.89A10 10 0 1 1 8 2.83"/><path d="M22 12A10 10 0 0 0 12 2v10z"/></svg>',
  bell:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="M6 8a6 6 0 0 1 12 0c0 7 3 9 3 9H3s3-2 3-9"/><path d="M10.3 21a1.94 1.94 0 0 0 3.4 0"/></svg>',
  cal:'<svg width="14" height="14" viewBox="0 0 24 24" fill="none" stroke="#6A6A74" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="3" y="4" width="18" height="17" rx="2"/><line x1="16" y1="2" x2="16" y2="6"/><line x1="8" y1="2" x2="8" y2="6"/><line x1="3" y1="10" x2="21" y2="10"/></svg>'};
const NOTIF_ICOS={
  rebal:'<svg viewBox="0 0 24 24"><polyline points="22 17 13.5 8.5 8.5 13.5 2 7"/><polyline points="16 17 22 17 22 11"/></svg>',
  aniv:'<svg viewBox="0 0 24 24"><path d="M20 21v-8a2 2 0 0 0-2-2H6a2 2 0 0 0-2 2v8"/><path d="M4 16s.5-1 2-1 2.5 2 4 2 2.5-2 4-2 2.5 2 4 2 2-1 2-1"/><path d="M2 21h20"/><path d="M7 8v3"/><path d="M12 8v3"/><path d="M17 8v3"/></svg>',
  liq:'<svg viewBox="0 0 24 24"><rect x="3" y="4" width="18" height="17" rx="2"/><line x1="16" y1="2" x2="16" y2="6"/><line x1="8" y1="2" x2="8" y2="6"/><line x1="3" y1="10" x2="21" y2="10"/></svg>',
  aporte:'<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="9"/><path d="M12 6.5v11"/><path d="M14.8 9.2c0-1.1-1.25-1.8-2.8-1.8s-2.8.7-2.8 1.8c0 2.9 5.6 1.7 5.6 4.6 0 1.1-1.25 1.9-2.8 1.9s-2.8-.7-2.8-1.8"/></svg>',
  pend:'<svg viewBox="0 0 24 24"><circle cx="12" cy="12" r="9"/><line x1="12" y1="8" x2="12" y2="12"/><line x1="12" y1="16" x2="12.01" y2="16"/></svg>',
  warn:'<svg viewBox="0 0 24 24"><path d="M10.29 3.86 1.82 18a2 2 0 0 0 1.71 3h16.94a2 2 0 0 0 1.71-3L13.71 3.86a2 2 0 0 0-3.42 0z"/><line x1="12" y1="9" x2="12" y2="13"/><line x1="12" y1="17" x2="12.01" y2="17"/></svg>'};
async function painelLoad(){
  const r=await api().painel_data();
  painelDados=(r&&r.ok)?r:null;
  if(painelDados&&painelDados.meta_per&&!_painelPerOk){painelMetaPer=painelDados.meta_per;_painelPerOk=true;}
  _notifs=(painelDados&&painelDados.notificacoes)||[];
  renderNotifBadge();
  renderPainel();
}
function setMetaPer(p){
  painelMetaPer=p; renderPainel();
  try{api().set_config("painel_meta_per",p);}catch(e){}
}
function notifAtivas(){return _notifs.filter(n=>n.ativa!==false&&!n.ignorada);}
function renderNotifBadge(){
  const b=$("notifBadge"); if(!b)return;
  const n=notifAtivas().length;
  b.style.display=n?"":"none";
  b.textContent=n>9?"9+":n;
}
function toggleNotif(ev){
  if(ev)ev.stopPropagation();
  const d=$("notifDrop"); if(!d)return;
  if(d.style.display==="none"){
    const ativas=notifAtivas();
    d.innerHTML=ativas.length
      ?ativas.map(n=>{const i=_notifs.indexOf(n);
        return `<div class="notif-item" onclick="notifGo(${i})">
          <div class="notif-ic">${NOTIF_ICOS[n.tipo]||NOTIF_ICOS.pend}</div>
          <div style="min-width:0;flex:1">
            <div style="display:flex;align-items:baseline;gap:10px"><div class="notif-tt" style="flex:1;min-width:0">${esc(n.titulo)}</div>
              <span style="font-size:10.5px;color:#B4B0A6;flex:none">${esc(n.chegou||"")}</span></div>
            <div class="notif-tx">${esc(n.texto)}</div></div></div>`;}).join("")
      :'<div class="mut" style="padding:18px;text-align:center;font-size:13px">Nenhuma notificação nova — o histórico fica no Painel ✓</div>';
    d.style.display="block";
  } else d.style.display="none";
}
function fecharNotif(){const d=$("notifDrop"); if(d)d.style.display="none";}
const COPY_ICO='<svg width="11" height="11" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><rect x="9" y="9" width="13" height="13" rx="2"/><path d="M5 15H4a2 2 0 0 1-2-2V4a2 2 0 0 1 2-2h9a2 2 0 0 1 2 2v1"/></svg>';
function _copyFallback(txt){
  const t=document.createElement("textarea");t.value=String(txt);
  t.style.cssText="position:fixed;opacity:0";document.body.appendChild(t);
  t.select();try{document.execCommand("copy");}catch(e){}document.body.removeChild(t);
}
function copiarConta(ev,conta){
  if(ev)ev.stopPropagation();
  conta=String(conta||"").trim();
  if(!conta){toast("Nenhum cliente selecionado.");return;}
  const done=()=>toast("Conta "+conta+" copiada ✓");
  if(navigator.clipboard&&navigator.clipboard.writeText)
    navigator.clipboard.writeText(String(conta)).then(done).catch(()=>{_copyFallback(conta);done();});
  else{_copyFallback(conta);done();}
}
function copyBtn(conta){
  conta=String(conta||"").trim(); if(!conta)return "";
  return `<span onclick="copiarConta(event,'${esc(conta)}')" title="Copiar conta"
    style="cursor:pointer;display:inline-flex;color:#6C6C76;vertical-align:middle;margin-left:6px"
    onmouseover="this.style.color='#F2C029'" onmouseout="this.style.color='#6C6C76'">${COPY_ICO}</span>`;
}
document.addEventListener("click",e=>{
  const d=$("notifDrop"), b=$("btnNotif");
  if(d&&d.style.display!=="none"&&!d.contains(e.target)&&!(b&&b.contains(e.target)))fecharNotif();
});
function notifGo(i){
  const n=_notifs[i]; if(!n)return;
  fecharNotif();
  if(n.acao==="cart"){tab("cart"); if(n.conta)setTimeout(()=>{const s=$("cartCli"); if(s){s.dataset.conta=String(n.conta); s.value=cliTxtDe(String(n.conta))||String(n.conta); cartLoad();}},400);}
  else if(n.acao==="cli"){tab("cli"); if(n.conta)setTimeout(()=>abrirCliente(String(n.conta)),400);}
  else if(n.acao)tab(n.acao);
}
function notifDismiss(i){
  const n=_notifs[i]; if(!n)return;
  n.ignorada=true;
  try{api().notif_ignorar(n.id,true);}catch(e){}
  renderNotifBadge();
  const el=$("painelNotifs"); if(el)el.innerHTML=painelNotifCards();
}
function notifReativar(i){
  const n=_notifs[i]; if(!n)return;
  n.ignorada=false;
  try{api().notif_ignorar(n.id,false);}catch(e){}
  renderNotifBadge();
  const el=$("painelNotifs"); if(el)el.innerHTML=painelNotifCards();
}
function salvarMeta(v){
  const n=numBR(v);
  api().set_config("painel_meta_"+painelMetaPer, n==null?"":n).then(()=>painelLoad());
}
function painelSetRange(r){painelRange=r; const el=$("painelChart");
  if(el)el.innerHTML=painelChartSVG(painelEvo());
  document.querySelectorAll("#painelRangePills .pn-pill").forEach(p=>p.classList.toggle("sel",p.dataset.r===r));}
function painelEvo(){
  const evo=(painelDados&&painelDados.evolucao)||[];
  return painelRange==="all"?evo:evo.slice(-parseInt(painelRange,10));
}
function painelChartSVG(evo){
  if(!evo.length)return '<div class="mut" style="padding:44px 0;text-align:center">Sem aportes registrados ainda — registre nos clientes para ver a evolução.</div>';
  const met=painelChartMetric;
  const fmt=met==="clientes"?(v=>String(Math.round(v||0))):milM;
  const vals=evo.map(e=>met==="clientes"?(e.cli||0):met==="rec"?(e.rec||0):(e.total||0));
  const W=800,H=250,P=18;
  const mn=Math.min(...vals), mx=Math.max(...vals), sp=(mx-mn)||1;
  const xs=i=>P+(W-2*P)*(evo.length===1?0.5:i/(evo.length-1));
  const ys=v=>H-36-(H-76)*((v-mn)/sp);
  const pts=vals.map((v,i)=>xs(i).toFixed(1)+","+ys(v).toFixed(1));
  const area="M"+pts.join(" L")+` L${xs(evo.length-1).toFixed(1)},${H-32} L${xs(0).toFixed(1)},${H-32} Z`;
  const lastV=vals[vals.length-1];
  const lx=xs(evo.length-1), ly=ys(lastV);
  const tx=Math.max(52,Math.min(lx,W-52)), ty=Math.max(38,ly-16);
  const grid=[0.18,0.5,0.82].map(f=>`<line x1="0" x2="${W}" y1="${(H-36)*f+12}" y2="${(H-36)*f+12}" stroke="#343537" stroke-width="0.6" stroke-dasharray="8 4"/>`).join("");
  const ylabs=`<text x="4" y="${((H-36)*0.18+12)-6}" fill="#6f6a5e" font-size="10" font-family="Manrope">${fmt(mx)}</text>
    <text x="4" y="${((H-36)*0.82+12)-6}" fill="#6f6a5e" font-size="10" font-family="Manrope">${fmt(mn)}</text>`;
  const labs=evo.map((e,i)=>(evo.length<=8||i%Math.ceil(evo.length/8)===0||i===evo.length-1)
    ?`<text x="${xs(i).toFixed(1)}" y="${H-10}" fill="#6f6a5e" font-size="10" text-anchor="middle" font-family="Manrope">${esc(e.label.split(" ")[0].slice(0,3))}</text>`:"").join("");
  const tip=met==="clientes"?fmt(lastV)+" cliente(s)":fmt(lastV);
  const tw=Math.max(88,tip.length*8+24);
  const hov=evo.map((e,i)=>{
    if(i===evo.length-1)return "";
    const v=vals[i], x=xs(i), y=ys(v);
    const mLab=esc((e.label||"").split(" ")[0].slice(0,3));
    const t=mLab+" · "+(met==="clientes"?fmt(v)+" cliente(s)":fmt(v));
    const w2=Math.max(84,t.length*7.5+22);
    const txc=Math.max(w2/2+4,Math.min(x,W-w2/2-4));
    const acima=y>54;
    const ry=acima?y-40:y+14, tyy=acima?y-21:y+33;
    return `<g class="pcht-pt">
      <circle cx="${x.toFixed(1)}" cy="${y.toFixed(1)}" r="3.5" fill="#F2C029" opacity=".55"/>
      <circle cx="${x.toFixed(1)}" cy="${y.toFixed(1)}" r="15" fill="transparent"/>
      <g class="tipg"><rect x="${(txc-w2/2).toFixed(1)}" y="${ry.toFixed(1)}" width="${w2}" height="26" rx="13" fill="#1f2022" stroke="rgba(255,255,255,.1)"/>
      <text x="${txc.toFixed(1)}" y="${tyy.toFixed(1)}" fill="#F2C029" font-size="12" font-weight="600" text-anchor="middle" font-family="Manrope">${t}</text></g>
    </g>`;
  }).join("");
  return `<svg viewBox="0 0 ${W} ${H}" preserveAspectRatio="none" style="width:100%;height:100%;display:block">
    <defs><linearGradient id="pgrad" x1="0" y1="0" x2="0" y2="1"><stop offset="0%" stop-color="#F2C029" stop-opacity="0.15"/><stop offset="100%" stop-color="#F2C029" stop-opacity="0"/></linearGradient></defs>
    ${grid}${ylabs}
    <path d="${area}" fill="url(#pgrad)"/>
    <polyline points="${pts.join(" ")}" fill="none" stroke="#F2C029" stroke-width="3" stroke-linecap="round" stroke-linejoin="round"/>
    ${hov}
    <rect x="${(tx-tw/2).toFixed(1)}" y="${(ty-22).toFixed(1)}" width="${tw}" height="28" rx="14" fill="#1f2022"/>
    <text x="${tx.toFixed(1)}" y="${(ty-3).toFixed(1)}" fill="#F2C029" font-size="12" font-weight="600" text-anchor="middle" font-family="Manrope">${tip}</text>
    <circle cx="${lx.toFixed(1)}" cy="${ly.toFixed(1)}" r="6" fill="#F2C029"><animate attributeName="opacity" values="1;.35;1" dur="1.8s" repeatCount="indefinite"/></circle>
    ${labs}</svg>`;
}
function painelDonut(aloc){
  if(!aloc||!aloc.length)return '<div class="mut" style="padding:24px 0">Sem composição consolidada ainda.</div>';
  const CORES=["#F2C029","#3B82F6","#10B981","#F97316","#8E8E98"];
  const R=42,CIRC=2*Math.PI*R;
  let off=0, segs=`<circle cx="50" cy="50" r="${R}" fill="transparent" stroke="#343537" stroke-width="8"/>`;
  aloc.forEach((a,i)=>{const len=Math.max(0,CIRC*a.pct/100);
    segs+=`<circle cx="50" cy="50" r="${R}" fill="transparent" stroke="${CORES[i%CORES.length]}" stroke-width="10"
      stroke-dasharray="${len.toFixed(1)} ${CIRC.toFixed(1)}" stroke-dashoffset="${(-off).toFixed(1)}"/>`;
    off+=len;});
  const top=aloc[0];
  const leg=aloc.map((a,i)=>`<div style="display:flex;justify-content:space-between;align-items:center;padding:9px 10px;border-radius:14px;transition:background .2s"
      onmouseover="this.style.background='rgba(255,255,255,.05)'" onmouseout="this.style.background=''">
    <span style="display:flex;align-items:center;gap:11px;font-size:13px;color:${i===0?'#e3e2e4':'#9c948a'};font-weight:${i===0?600:400}">
      <span style="width:11px;height:11px;border-radius:50%;background:${CORES[i%CORES.length]};flex:none;${i===0?'box-shadow:0 0 8px rgba(242,192,41,.4)':''}"></span>${esc(a.grupo)}</span>
    <span style="font-size:13px;font-weight:700;color:${i===0?'#F2C029':'#e3e2e4'}">${fBR(a.pct,0)}%</span></div>`).join("");
  return `<div style="flex:1;display:flex;align-items:center;justify-content:center;position:relative;margin:6px 0 10px">
    <div style="position:relative;width:190px;height:190px">
      <svg style="width:100%;height:100%;transform:rotate(-90deg)" viewBox="0 0 100 100">${segs}</svg>
      <div style="position:absolute;inset:0;display:flex;flex-direction:column;align-items:center;justify-content:center;pointer-events:none">
        <span style="font:700 38px/1 'Manrope';letter-spacing:-.02em;color:#e3e2e4">${fBR(top.pct,0)}%</span>
        <span style="font:600 10px 'Manrope';letter-spacing:.18em;text-transform:uppercase;color:#F2C029;margin-top:5px">${esc(top.grupo)}</span>
      </div></div></div>${leg}`;
}
let painelNotifPag=1, painelNotifFiltro="todos";
const NOTIF_TIPOS={aniv:"Aniversários",rebal:"Rebalanceamento",aporte:"Aportes",pend:"Cadastro",warn:"Avisos"};
function setNotifPag(p){painelNotifPag=p; const el=$("painelNotifs"); if(el)el.innerHTML=painelNotifCards();}
function setNotifFiltro(t){painelNotifFiltro=t; painelNotifPag=1; const el=$("painelNotifs"); if(el)el.innerHTML=painelNotifCards();}
function _notifsFiltradas(){
  return painelNotifFiltro==="todos"?_notifs:_notifs.filter(n=>n.tipo===painelNotifFiltro);
}
function notifIgnorarTodas(){
  const alvo=_notifsFiltradas().filter(n=>n.ativa!==false&&!n.ignorada);
  if(!alvo.length)return;
  alvo.forEach(n=>{n.ignorada=true; try{api().notif_ignorar(n.id,true);}catch(e){}});
  renderNotifBadge();
  const el=$("painelNotifs"); if(el)el.innerHTML=painelNotifCards();
  toast(alvo.length+" notificação(ões) ignorada(s).");
}
function painelNotifBar(){
  const tipos=[...new Set(_notifs.map(n=>n.tipo))];
  const pills=[["todos","Todas"]].concat(tipos.map(t=>[t,NOTIF_TIPOS[t]||t]));
  const temAtivas=_notifsFiltradas().some(n=>n.ativa!==false&&!n.ignorada);
  return `<div style="display:flex;align-items:center;gap:12px;flex-wrap:wrap;margin-bottom:16px">
    <div class="pn-pillbar">${pills.map(([v,l])=>
      `<span class="pn-pill ${painelNotifFiltro===v?'sel':''}" onclick="setNotifFiltro('${v}')">${esc(l)}</span>`).join("")}</div>
    <div style="flex:1"></div>
    ${temAtivas?`<button class="pn-btn pn-btn-ghost" onclick="notifIgnorarTodas()" title="Ignora todas as notificações ativas do filtro atual">Ignorar todas</button>`:""}
  </div>`;
}
function painelNotifCards(){
  if(!_notifs.length)return '<div class="pn-card mut" style="padding:24px;border-radius:24px">Nenhuma notificação — tudo em dia ✓</div>';
  const lst=_notifsFiltradas();
  if(!lst.length)return painelNotifBar()
    +'<div class="pn-card mut" style="padding:24px;border-radius:24px">Nenhuma notificação deste tema.</div>';
  const POR_PAG=5;
  const totPag=Math.max(1,Math.ceil(lst.length/POR_PAG));
  if(painelNotifPag>totPag)painelNotifPag=totPag;
  if(painelNotifPag<1)painelNotifPag=1;
  const ini=(painelNotifPag-1)*POR_PAG;
  const cards=lst.slice(ini,ini+POR_PAG).map(n=>{
    const i=_notifs.indexOf(n);
    const inativa=n.ativa===false, ign=!!n.ignorada;
    const tag=inativa
      ?'<span style="font:700 9.5px \'Manrope\';letter-spacing:.08em;text-transform:uppercase;color:#8E8E98;border:1px solid rgba(255,255,255,.14);border-radius:999px;padding:2px 9px;flex:none">resolvida</span>'
      :(ign?'<span style="font:700 9.5px \'Manrope\';letter-spacing:.08em;text-transform:uppercase;color:#9c948a;border:1px solid rgba(255,255,255,.14);border-radius:999px;padding:2px 9px;flex:none">ignorada</span>':"");
    return `<div class="pn-card" style="border-radius:24px;display:flex;gap:20px;align-items:center;padding:22px 24px;margin-bottom:16px;${(ign||inativa)?'opacity:.55':''}">
    <div class="pn-nic">${NOTIF_ICOS[n.tipo]||NOTIF_ICOS.pend}</div>
    <div style="flex:1;min-width:0">
      <div style="display:flex;align-items:baseline;gap:12px;margin-bottom:3px;flex-wrap:wrap">
        <span style="font:600 18px 'Manrope';color:#e3e2e4">${esc(n.titulo)}</span>${tag}
        <span style="font-size:11px;color:#B4B0A6;flex:none;margin-left:auto">${esc(n.chegou||"")}</span></div>
      <div style="font-size:14px;color:#9c948a;line-height:1.55">${esc(n.texto)}</div></div>
    ${inativa?"":(ign
      ?`<button class="pn-btn pn-btn-ghost" onclick="notifReativar(${i})">Reativar</button>`
      :`<button class="pn-btn pn-btn-ghost" onclick="notifDismiss(${i})">Ignorar</button>`)}
    ${n.botao?`<button class="pn-btn pn-btn-primary" onclick="notifGo(${i})">${esc(n.botao)}</button>`:""}
  </div>`;}).join("");
  let pager="";
  if(totPag>1){
    const nums=[];
    for(let p=1;p<=totPag;p++){
      if(p===1||p===totPag||Math.abs(p-painelNotifPag)<=1)nums.push(p);
      else if(nums[nums.length-1]!=="…")nums.push("…");
    }
    pager=`<div style="display:flex;justify-content:center;margin-top:4px"><div class="pn-pillbar">`
      +nums.map(p=>p==="…"
        ?`<span class="pn-pill" style="cursor:default;opacity:.5">…</span>`
        :`<span class="pn-pill ${p===painelNotifPag?'sel':''}" onclick="setNotifPag(${p})">${p}</span>`).join("")
      +`</div></div>`;
  }
  return painelNotifBar()+cards+pager;
}
function renderPainel(){
  const W=$("painelWrap"); if(!W)return;
  const d=painelDados;
  if(!d){W.innerHTML='<div class="card mut">Não consegui carregar o painel — verifique a planilha.</div>';return;}
  const sgnP=v=>(v>0?"+":"")+fBR(v,2)+"%";
  const PNOM={mes:"Mês",tri:"Trimestre",sem:"Semestre"};
  const per=painelMetaPer;
  const aporte=(d.aportes&&d.aportes[per])||0;
  const meta=(d.metas&&d.metas[per])||null;
  const perLabel=(d.periodos&&d.periodos[per])||"";
  const pct=meta?Math.min(100,Math.round(aporte/meta*100)):null;
  const falta=meta?Math.max(0,meta-aporte):null;
  const perPills=`<div class="pn-pillbar">${["mes","tri","sem"].map(p=>
    `<span class="pn-pill ${per===p?'sel':''}" onclick="setMetaPer('${p}')">${PNOM[p]}</span>`).join("")}</div>`;
  const hero=`<div class="pn-card pn-hero" style="grid-row:span 2;min-height:320px;padding:36px 40px;display:flex;flex-direction:column">
    <div style="display:flex;align-items:center;justify-content:space-between;gap:14px;row-gap:14px;flex-wrap:wrap">
      <div class="pn-htitle">${PIC.target}<span>Objetivo de Aporte</span></div>
      ${perPills}
    </div>
    <div style="margin:auto 0;padding:26px 0 4px">
      ${meta?`
        <div class="pn-glow" style="font:700 56px/1.05 'Manrope';letter-spacing:-.02em;color:#e3e2e4">
          ${falta>0?milM(falta):"Meta batida!"} <span style="font:500 26px 'Manrope';color:#9c948a">${falta>0?"para a meta":"("+milM(meta)+")"}</span></div>
        <div style="font-size:15px;color:#9c948a;margin-top:12px;max-width:520px;line-height:1.6">
          Acompanhamento da meta de captação para o ${esc(perLabel)}. Aportes somados no período:
          <b style="color:#e3e2e4">${aporte>0?milM(aporte):"R$ 0"}</b>${aporte>0?"":" — nenhum aporte registrado ainda"}.</div>
        <div style="margin-top:28px">
          <div style="display:flex;justify-content:space-between;margin-bottom:12px">
            <span class="pn-label">Progresso Atual</span>
            <span style="font:700 12px 'Manrope';letter-spacing:.05em;color:#F2C029">${pct}%</span></div>
          <div class="pn-track"><div class="pn-fill" style="width:${pct}%"></div></div></div>`
      :`
        <div class="pn-glow" style="font:700 56px/1.05 'Manrope';letter-spacing:-.02em;color:#e3e2e4">
          ${aporte>0?milM(aporte):"R$ 0"} <span style="font:500 26px 'Manrope';color:#9c948a">em aportes</span></div>
        <div style="font-size:15px;color:#9c948a;margin-top:12px;max-width:520px;line-height:1.6">
          Soma dos aportes no ${esc(perLabel)}${aporte>0?"":" — nenhum aporte registrado ainda"}.
          Defina a meta de captação do período abaixo para acompanhar o progresso.</div>`}
      <div style="display:flex;align-items:center;gap:12px;margin-top:24px">
        <span class="pn-label" style="flex:none">Meta do ${PNOM[per].toLowerCase()} (R$)</span>
        <input style="width:190px;border-radius:999px;background:#1f2022;border:1px solid rgba(255,255,255,.08);padding:9px 16px"
          placeholder="ex.: 5.000.000,00" value="${meta?fBR(meta):""}" oninput="maskMoeda(this)" onchange="salvarMeta(this.value)" spellcheck="false">
      </div>
    </div>
  </div>`;
  const pper=painelPacePer;
  const utp=(d.uteis&&d.uteis[pper])||{rest:1,total:1};
  const metaP=(d.metas&&d.metas[pper])||null;
  const aporteP=(d.aportes&&d.aportes[pper])||0;
  const faltaP=metaP?Math.max(0,metaP-aporteP):0;
  const paceP=faltaP>0?faltaP/Math.max(1,utp.rest):0;
  const pctP=metaP?Math.min(100,Math.round(aporteP/metaP*100)):0;
  const pacePills=`<div class="pn-pillbar">${["mes","tri","sem"].map(p=>
    `<span class="pn-pill ${pper===p?'sel':''}" onclick="setPacePer('${p}')">${PNOM[p]}</span>`).join("")}</div>`;
  const cardRent=`<div class="pn-card" style="padding:32px;display:flex;flex-direction:column;min-height:190px">
    <div style="display:flex;align-items:center;justify-content:space-between;gap:12px;row-gap:14px;flex-wrap:wrap;margin-bottom:18px">
      <div class="pn-htitle" style="font-size:13px;color:rgba(242,192,41,.8);text-transform:uppercase;letter-spacing:.08em;font-weight:600">${PIC.trend}<span>Pace de Captação</span></div>
      ${pacePills}</div>
    <div style="margin-top:auto;padding-top:6px">
    ${metaP?`
    <div style="font:700 34px/1.1 'Manrope';letter-spacing:-.01em;color:#e3e2e4">${faltaP>0?milM(paceP):"Meta batida!"}
      ${faltaP>0?`<span style="font:500 14px 'Manrope';color:#9c948a">/ dia útil (${utp.rest} restantes)</span>`:""}</div>
    <div style="font-size:12.5px;color:#9c948a;margin-top:14px;line-height:1.6">
      Meta: <b style="color:#e3e2e4">${milM(metaP)}</b>
      · Aportado: <b style="color:#e3e2e4">${milM(aporteP)}</b>
      · Falta: <b style="color:${faltaP>0?'#F2C029':'#e3e2e4'}">${milM(faltaP)}</b></div>
    <div class="pn-track fina" style="margin:18px 0 12px"><div class="pn-fill" style="width:${pctP}%"></div></div>
    <div class="pn-label" style="text-align:right">${pctP}%</div>`
    :'<div style="font-size:14px;color:#9c948a;line-height:1.5">Defina a meta do período no card Objetivo de Aporte para calcular o pace por dia útil.</div>'}</div></div>`;
  const ut=(d.uteis&&d.uteis[painelUteisPer])||null;
  const utLabel=(d.periodos&&d.periodos[painelUteisPer])||"";
  const utPills=`<div class="pn-pillbar">${["mes","tri","sem"].map(p=>
    `<span class="pn-pill ${painelUteisPer===p?'sel':''}" onclick="setUteisPer('${p}')">${PNOM[p]}</span>`).join("")}</div>`;
  const cardEnq=`<div class="pn-card" style="padding:32px;display:flex;flex-direction:column;min-height:190px">
    <div style="display:flex;align-items:center;justify-content:space-between;gap:12px;row-gap:14px;flex-wrap:wrap;margin-bottom:18px">
      <div class="pn-htitle" style="font-size:13px;color:rgba(242,192,41,.8);text-transform:uppercase;letter-spacing:.08em;font-weight:600">${PIC.cal}<span>Dias Úteis</span></div>
      ${utPills}</div>
    <div style="margin-top:auto;padding-top:6px">
    ${ut?`<div style="font:700 38px/1.1 'Manrope';letter-spacing:-.01em;color:#e3e2e4">${ut.rest} <span style="font:500 15px 'Manrope';color:#9c948a">restantes</span></div>
      <div class="pn-track fina" style="margin:18px 0 12px"><div class="pn-fill" style="width:${Math.round(ut.rest/Math.max(1,ut.total)*100)}%"></div></div>
      <div class="pn-label" style="display:flex;justify-content:space-between;gap:10px"><span>de ${ut.total} dias úteis</span><span>${esc(utLabel)}</span></div>`
    :'<div style="font-size:14px;color:#9c948a">—</div>'}</div></div>`;
  const rper=painelRecPer;
  const recTot=(d.receitas&&d.receitas[rper])||0;
  const recLabel=(d.periodos&&d.periodos[rper])||"";
  const recPills=`<div class="pn-pillbar">${["mes","tri","sem"].map(p=>
    `<span class="pn-pill ${rper===p?'sel':''}" onclick="setRecPer('${p}')">${PNOM[p]}</span>`).join("")}</div>`;
  const MABR=["Jan","Fev","Mar","Abr","Mai","Jun","Jul","Ago","Set","Out","Nov","Dez"];
  const hojeM=new Date().getMonth()+1;
  const perMs={mes:[hojeM],tri:[Math.floor((hojeM-1)/3)*3+1,Math.floor((hojeM-1)/3)*3+2,Math.floor((hojeM-1)/3)*3+3],
               sem:(hojeM<=6?[1,2,3,4,5,6]:[7,8,9,10,11,12])}[rper];
  const recLinhas=perMs.map(m=>{const v=(d.rec_meses||{})[String(m)];
    return `<div style="display:flex;justify-content:space-between;align-items:center;padding:8px 2px;border-bottom:1px solid rgba(255,255,255,.05)">
      <span style="font-size:13px;color:#9c948a">${MABR[m-1]}</span>
      <span class="dc-mono" style="font-size:13px;font-weight:700;color:${v?'#F2C029':'#5A5E67'}">${v?milM(v):"—"}</span></div>`;}).join("");
  const cardReceita=`<div class="pn-card" style="padding:32px;display:flex;flex-direction:column;height:400px;overflow-y:auto">
    <div style="display:flex;align-items:center;justify-content:space-between;gap:12px;row-gap:12px;flex-wrap:wrap;margin-bottom:16px">
      <span class="pn-title" style="font:500 19px 'Manrope';color:#e3e2e4">Receita</span>
      ${recPills}</div>
    <div style="font:700 36px/1.1 'Manrope';letter-spacing:-.01em;color:${recTot?'#e3e2e4':'#5A5E67'}">${recTot?milM(recTot):"R$ 0"}</div>
    <div class="pn-label" style="margin:8px 0 14px">recebido no ${esc(recLabel)}</div>
    <div style="flex:1;min-height:0;overflow-y:auto">${recLinhas}</div>
    <div style="display:flex;align-items:center;gap:8px;margin-top:14px;flex-wrap:wrap">
      <input id="recMes" placeholder="mm/aaaa" maxlength="7" oninput="maskMes(this)" spellcheck="false"
        style="width:92px;text-align:center;border-radius:999px;background:#1f2022;border:1px solid rgba(255,255,255,.08);padding:7px 8px">
      <input id="recValor" placeholder="Valor (R$)" oninput="maskMoeda(this)" spellcheck="false"
        onkeydown="if(event.key==='Enter')salvarReceita()"
        style="flex:1;min-width:110px;border-radius:999px;background:#1f2022;border:1px solid rgba(255,255,255,.08);padding:7px 14px">
      <button class="pn-btn pn-btn-primary" onclick="salvarReceita()">Salvar</button>
    </div>
  </div>`;
  W.innerHTML=`
   <div style="margin:6px 0 34px;display:flex;align-items:flex-end;justify-content:space-between;gap:16px;flex-wrap:wrap">
     <div>
       <div style="font:700 46px/1.1 'Manrope';letter-spacing:-.02em;color:#e3e2e4;margin-bottom:8px">Visão Geral</div>
       <div style="font-size:15px;color:#9c948a">Bem vindo de volta, Guilherme</div>
     </div>
     <span class="pn-chip"><span style="width:8px;height:8px;border-radius:50%;background:rgba(242,192,41,.8);flex:none"></span> Atualizado agora</span>
   </div>
   <div style="display:grid;grid-template-columns:2fr 1fr;grid-auto-rows:min-content;gap:24px;margin-bottom:40px">
     ${hero}${cardRent}${cardEnq}
   </div>
   <div style="display:grid;grid-template-columns:2fr 1fr;gap:24px;margin-bottom:44px">
     <div class="pn-card" style="padding:32px;display:flex;flex-direction:column;height:400px">
       <div style="display:flex;align-items:center;justify-content:space-between;gap:12px;row-gap:12px;flex-wrap:wrap;margin-bottom:24px">
         <span class="pn-title" style="font:500 19px 'Manrope';color:#e3e2e4">Evolução Histórica</span>
         <div style="display:flex;gap:10px;flex-wrap:wrap">
           <div class="pn-pillbar">${[["valor","Captação","Soma do valor aportado pelos clientes em cada mês"],["clientes","Clientes","Número de clientes que aportaram em cada mês"]].map(([k,l,tt])=>`<span class="pn-pill ${painelChartMetric===k?'sel':''}" title="${tt}" onclick="setChartMetric('${k}')">${l}</span>`).join("")}</div>
           <div class="pn-pillbar" id="painelRangePills">${["3","6","12","all"].map(r=>`<span class="pn-pill ${painelRange===r?'sel':''}" data-r="${r}" onclick="painelSetRange('${r}')">${r==="all"?"Tudo":r+"M"}</span>`).join("")}</div>
         </div>
       </div>
       <div id="painelChart" style="flex:1;min-height:0">${painelChartSVG(painelEvo())}</div>
       ${nChip(d.evo_cli||0,d.total,"cliente(s) com aporte registrado")}
     </div>
     ${cardReceita}
   </div>
   <div style="font:600 26px 'Manrope';letter-spacing:-.01em;color:#e3e2e4;margin-bottom:20px">Notificações</div>
   <div id="painelNotifs">${painelNotifCards()}</div>`;
}

function uiLog(which,msg,cls){
  const el=which==="pdf"?$("logPdf"):$("logMail");
  const s=document.createElement("span"); s.className=cls||""; s.textContent=msg+"\n";
  el.appendChild(s); el.scrollTop=el.scrollHeight;
}
function uiProgress(p){const pr=$("prog");pr.classList.add("on");pr.firstElementChild.style.width=p+"%";
  if(p>=100)setTimeout(()=>pr.classList.remove("on"),900);}
function uiBusy(b){_busy=b;$("btnProc").disabled=b;$("btnSend").disabled=b;}
function toast(msg,ms){const t=$("toast");t.textContent=msg;t.classList.add("on");
  clearTimeout(t._h);t._h=setTimeout(()=>t.classList.remove("on"),ms||2600);}
function avisarAniversarios(lista){
  if(!lista||!lista.length)return;
  const msg=lista.map(a=>a.status==="hoje"
    ?`HOJE é aniversário de ${a.nome}${a.idade_nova!=null?` (${a.idade_nova} anos)`:""}!`
    :`${a.nome} faz aniversário no ${a.dsem} (${a.data}) — cai no fim de semana, parabenize hoje!`);
  toast("Aniversários: "+msg.join("   ·   "),10000);
}
function modal(title,body,cb){$("mTitle").textContent=title;$("mBody").textContent=body;
  _mCb=cb;$("ovl").classList.add("on");}
function modalYes(){$("ovl").classList.remove("on");if(_mCb)_mCb();}
function modalNo(){$("ovl").classList.remove("on");_mCb=null;}

window.addEventListener("pywebviewready",async()=>{
  const st=await api().boot();
  if(st.aviso){uiLog("pdf","⚠  "+st.aviso,"warn");toast(st.aviso);}
  $("pasta").value=st.pasta; $("assunto").value=st.assunto; $("corpo").value=st.corpo;
  $("excelName").textContent=st.excel;
  cadastro=st.cadastro||[]; renderCadastro();
  setMeses(st.meses);
  avisarAniversarios(st.aniversariantes);
  const chipsCli=["{nome}","{primeiro_nome}","{conta}","{mes_ref}"];
  const chipsMkt=["{ganho}","{rentabilidade}","{cdi}","{patrimonio}","{rentabilidade_total}"];
  const chipsBlocos=["{resumo_pontos}","{composicao_carteira}","{contexto_mercado}"];
  const chipsFrases=["{analise_cdi}","{analise_ganho}","{analise_acumulado}"];
  $("chipsCli").innerHTML=chipsCli.map(c=>`<span class="chip" onclick="insChip('${c}')">${c}</span>`).join("");
  $("chipsMkt").innerHTML=chipsMkt.map(c=>`<span class="chip" onclick="insChip('${c}')">${c}</span>`).join("");
  if($("chipsBlocos"))$("chipsBlocos").innerHTML=chipsBlocos.map(c=>`<span class="chip" title="Bloco visual: use numa linha própria da mensagem" onclick="insChip('${c}')">${c}</span>`).join("");
  if($("chipsFrases"))$("chipsFrases").innerHTML=chipsFrases.map(c=>`<span class="chip" onclick="insChip('${c}')">${c}</span>`).join("");
  try{const cfg=(await api().get_config()).config||{}; incluirPlugados=cfg.incluir_plugados!==false;}catch(e){}
  painelLoad();
});
function setMeses(ms){
  const el=$("selMes"), cur=el.value;
  el.innerHTML=(ms||[]).map(m=>`<option>${m}</option>`).join("")||"<option value=''>—</option>";
  if(ms&&ms.includes(cur))el.value=cur;
}
async function uiRefresh(){
  const st=await api().boot();
  setMeses(st.meses);cadastro=st.cadastro||[];renderCadastro();
  cliCur=null; if($("cliDetail")){$("cliDetail").style.display="none";$("cliLista").style.display="flex";}
  cliLoad();
  if(typeof dashLoad==="function") dashLoad();
  if(typeof painelLoad==="function") painelLoad();
  if(st.aniversariantes&&st.aniversariantes.length)avisarAniversarios(st.aniversariantes);
  else toast(st.aviso||"Planilha atualizada");
}

async function escolherPasta(){const p=await api().choose_folder();if(p)$("pasta").value=p;}
function processar(){
  if(_busy)return;
  $("logPdf").innerHTML="";uiBusy(true);uiProgress(0);
  api().process_pdfs($("pasta").value);
}

let _deb=null;
function aoEditar(){
  $("saveState").textContent="editando…";$("saveState").classList.remove("saved");
  clearTimeout(_deb);_deb=setTimeout(salvarTemplate,1400);
  clearTimeout(window._pvDeb);window._pvDeb=setTimeout(atualizarPreview,500);
}
async function salvarTemplate(){
  const r=await api().save_template($("assunto").value,$("corpo").value);
  if(r.ok){$("saveState").textContent="✓ salvo";$("saveState").classList.add("saved");}
  else{$("saveState").textContent=r.erro?("erro: "+r.erro):"";}
}
function insChip(v){
  const t=$("corpo"),s=t.selectionStart,e=t.selectionEnd;
  t.value=t.value.slice(0,s)+v+t.value.slice(e);
  t.selectionStart=t.selectionEnd=s+v.length;t.focus();aoEditar();
}

let pvCliIdx=0;
function pvCliTexto(i){const c=clientes[i];return c?`${c.nome||"(sem nome)"} (${c.conta})`:"";}
function pvSync(){const s=$("selCli");if(s)s.value=pvCliTexto(pvCliIdx);}
async function aoMes(){
  const mes=$("selMes").value;
  if(!mes){clientes=[];pvCliIdx=0;$("selCli").value="";$("mesInfo").innerHTML="";atualizarPreview();return;}
  const r=await api().month_info(mes);
  if(!r.ok) toast("Erro ao ler clientes do mês: "+(r.erro||""));
  clientes=r.clientes||[];
  const sem=clientes.filter(c=>!c.email).length;
  $("mesInfo").innerHTML=clientes.length?
    (sem?`<span class="badge warn">👥 ${clientes.length} cliente(s) · ${sem} sem e-mail</span>`
        :`<span class="badge ok">👥 ${clientes.length} cliente(s) · todos com e-mail</span>`):"";
  pvCliIdx=0;pvSync();
  atualizarPreview();
}
function navCli(d){
  if(!clientes.length)return;
  pvCliIdx=((pvCliIdx+d)%clientes.length+clientes.length)%clientes.length;
  pvSync();atualizarPreview();
}
async function atualizarPreview(){
  const i=pvCliIdx;
  if(!clientes.length||isNaN(i)||!clientes[i]){
    $("pvFrame").srcdoc="<body style='font-family:Segoe UI;color:#999;display:flex;align-items:center;justify-content:center;height:100vh'>Selecione um mês com clientes para visualizar</body>";
    $("pvMeta").innerHTML="";return;}
  const r=await api().preview_email(clientes[i],$("assunto").value,$("corpo").value);
  if(!r.ok){$("pvMeta").textContent=r.erro;return;}
  $("pvMeta").innerHTML=`<b>Para:</b> ${r.email||"<span style='color:var(--warn)'>sem e-mail cadastrado</span>"}<br><b>Assunto:</b> ${r.assunto}`;
  $("pvFrame").srcdoc=r.html;
}

function confirmarEnvio(){
  const mes=$("selMes").value;
  if(!mes||!clientes.length){toast("Selecione um mês com clientes.");return;}
  if(!$("assunto").value.trim()||!$("corpo").value.trim()){toast("Assunto ou mensagem vazios.");return;}
  const com=clientes.filter(c=>c.email).length, sem=clientes.length-com;
  if(!com){toast("Nenhum cliente com e-mail cadastrado.");return;}
  const porEmail={};
  clientes.filter(c=>c.email).forEach(c=>{
    const e=c.email.trim().toLowerCase();
    (porEmail[e]=porEmail[e]||[]).push(c.nome||c.conta);
  });
  const dups=Object.entries(porEmail).filter(([,l])=>l.length>1);
  const aviso=dups.length
    ?`\n\n⚠ ATENÇÃO — e-mail repetido em clientes diferentes:\n`
      +dups.map(([e,l])=>`• ${e}: ${l.join(", ")}`).join("\n")
      +`\nConfira o cadastro antes de enviar.`
    :"";
  modal("Enviar e-mails",
    `Enviar para ${com} cliente(s) de ${mes}.`+(sem?`\n${sem} sem e-mail serão ignorados.`:"")+aviso,
    ()=>{$("logMail").innerHTML="";uiBusy(true);
         api().send_emails(mes,$("assunto").value,$("corpo").value,$("pasta").value);});
}

const TERMO_OPC=["","A","B","C","D","E"];
const TERMO_COR={A:"#10B981",B:"#3B82F6",C:"#F59E0B",D:"#F97316",E:"#EF4444"};
const PERFIL_OPC=["","Conservadora","Moderada","Sofisticada"];
let cadView=[], cadSortDir=0;
function cadSort(){ cadSortDir = cadSortDir===-1?1:-1; renderCadastro(); }
function cadExtraCell(c,i,d){
  const tem=flagNormInit(c[d.k+"_tem"]), est=flagEstadoJS(tem), mostra=(est==="ok"||est==="rev");
  const sel=FLAG_OPC.map(([val,lab])=>`<option value="${val}" ${val===tem?"selected":""}>${lab}</option>`).join("");
  let extra="";
  if(mostra&&!d.soTem){
    extra=`<input value="${esc(c[d.k+"_valor"]||"")}" placeholder="${d.labVal}" oninput="maskMoeda(this);cadView[${i}]['${d.k}_valor']=this.value;autoSaveCadastro()" style="width:100%;margin-top:3px">`;
    if(d.onde)extra+=`<input value="${esc(c[d.k+"_onde"]||"")}" placeholder="${d.labOnde}" oninput="cadView[${i}]['${d.k}_onde']=this.value;autoSaveCadastro()" style="width:100%;margin-top:3px">`;
  }
  return `<td style="vertical-align:top;min-width:150px">
    <div style="display:flex;align-items:center;gap:6px">
      <span style="display:inline-block;width:9px;height:9px;border-radius:50%;background:${FLAGCOR[est]||'#5C5C66'};flex:none"></span>
      <select onchange="cadView[${i}]['${d.k}_tem']=this.value;autoSaveCadastro();renderCadastro()" style="flex:1;min-width:0">${sel}</select>
    </div>${extra}</td>`;
}
let cadPendFiltro=null, _cadPendKeep=false;
const CAD_PEND={
  prev:    {lab:"Sem previdência",         f:c=>["nao",""].includes(flagEstadoJS(c.prev_tem))},
  seg:     {lab:"Sem seguro",              f:c=>["nao",""].includes(flagEstadoJS(c.seg_tem))},
  intl:    {lab:"Sem internacional",       f:c=>["nao",""].includes(flagEstadoJS(c.intl_tem))},
  mesa:    {lab:"Não plugados na mesa",    f:c=>!["ok","rev"].includes(flagEstadoJS(c.mesa_tem))},
  termo:   {lab:"Sem termômetro",          f:c=>!(c.termometro||"").trim()},
  inativos:{lab:"Inativos (sem relatório)",f:c=>c.ativo===false},
};
function dashPendGo(tipo){
  if(!CAD_PEND[tipo])return;
  cadPendFiltro=tipo; _cadPendKeep=true;
  tab("cli");
  cliView("cad");
  _cadPendKeep=false;
}
function cadPendLimpar(){cadPendFiltro=null; renderCadastro();}
async function deletarInativos(){
  const n=(cadastro||[]).filter(c=>c.ativo===false).length;
  if(!n){toast("Nenhum cliente inativo.");return;}
  modal("Deletar Inativos",
    `Excluir DEFINITIVAMENTE ${n} cliente(s) inativo(s)?\n`+
    `Serão removidos da planilha E do banco de dados, incluindo aportes, contatos e tabelas do Controle. Não há como desfazer.`,
    async()=>{
      const r=await api().deletar_inativos();
      if(r&&r.ok){
        cadastro=r.cadastro||[];
        cadPendLimpar();
        toast(`${r.n} cliente(s) inativo(s) excluído(s) ✓`);
        cliLoad(); dashLoad();
      }else toast("Erro ao excluir: "+((r&&r.erro)||""));
    });
}
function renderCadastro(){
  const q=(($("cadBusca")&&$("cadBusca").value)||"").toLowerCase().trim();
  let v=cadastro.filter(c=>!q||(c.nome||"").toLowerCase().includes(q)||String(c.conta||"").includes(q));
  const pend=cadPendFiltro&&CAD_PEND[cadPendFiltro];
  if(pend)v=v.filter(pend.f);
  const info=$("cadPendInfo");
  if(info)info.innerHTML=pend
    ?`<span class="badge warn" style="cursor:pointer" onclick="cadPendLimpar()"
        title="Clique para voltar ao cadastro completo">${esc(pend.lab)} · ${v.length} cliente(s) ✕</span>`
      +(cadPendFiltro==="inativos"&&v.length
        ?`<button class="btn" style="margin-left:10px;background:rgba(239,68,68,.14);color:#EF4444;border:1px solid rgba(239,68,68,.35)"
             onclick="deletarInativos()" title="Exclui da planilha E do banco de dados (sem volta)">🗑 Deletar Inativos</button>`:"")
    :"";
  if(cadSortDir) v=v.slice().sort((a,b)=>((a.patrimonio||0)-(b.patrimonio||0))*cadSortDir);
  cadView=v;
  if($("cadSortArrow")) $("cadSortArrow").textContent=cadSortDir===1?"▲":(cadSortDir===-1?"▼":"");
  $("cliBody").innerHTML=v.map((c,i)=>{
    const inativo=c.ativo===false;
    const tv=(c.termometro||"").toUpperCase();
    const tcor=TERMO_COR[tv]||"var(--mut)";
    const tsel=TERMO_OPC.map(t=>`<option ${t===tv?"selected":""}>${t}</option>`).join("");
    const psel=PERFIL_OPC.map(p=>`<option ${p===(c.perfil||"")?"selected":""}>${p}</option>`).join("");
    return `<tr style="${inativo?'background:rgba(239,68,68,.13)':''}">
      <td><input value="${esc(c.nome)}" oninput="cadView[${i}].nome=this.value;autoSaveCadastro()"${inativo?' style="color:var(--err)"':''}></td>
      <td><div style="display:flex;align-items:center;gap:2px"><input value="${esc(c.conta)}" oninput="cadView[${i}].conta=this.value;autoSaveCadastro()" style="min-width:0">${copyBtn(c.conta)}</div></td>
      <td><input value="${esc(c.email)}" oninput="cadView[${i}].email=this.value;autoSaveCadastro()" placeholder="email@exemplo.com"></td>
      <td><select onchange="cadView[${i}].termometro=this.value;autoSaveCadastro();renderCadastro()" style="font-weight:700;color:${tcor}">${tsel}</select></td>
      <td><select onchange="cadView[${i}].perfil=this.value;autoSaveCadastro()">${psel}</select></td>
      <td><input value="${esc(c.estado)}" oninput="cadView[${i}].estado=this.value;autoSaveCadastro()" placeholder="UF" maxlength="20"></td>
      <td><input value="${esc(c.cidade)}" oninput="cadView[${i}].cidade=this.value;autoSaveCadastro()" placeholder="Cidade"></td>
      <td><input value="${esc(c.nascimento)}" oninput="cadView[${i}].nascimento=maskData(this);autoSaveCadastro()" placeholder="dd/mm/aaaa" maxlength="10"></td>
      <td class="mut" style="white-space:nowrap">${c.patrimonio!=null?"R$ "+fBR(c.patrimonio,2):"—"}${inativo?' <span class="small" style="color:var(--err)">· sem conta</span>':''}</td>
      ${EXTRA_DEF.map(d=>cadExtraCell(c,i,d)).join("")}
      <td><span class="del" onclick="delCli(${i})">✕</span></td></tr>`;
  }).join("");
  const sem=cadastro.filter(c=>!c.email).length, inat=cadastro.filter(c=>c.ativo===false).length;
  $("cliBadge").innerHTML=cadastro.length?
    ((inat?`<span class="badge warn">${inat} sem conta</span> `:"")+
     (sem?`<span class="badge warn">${sem} sem e-mail</span>`:`<span class="badge ok">${cadastro.length} clientes</span>`)):"";
}
function esc(s){return String(s||"").replace(/&/g,"&amp;").replace(/</g,"&lt;").replace(/>/g,"&gt;").replace(/"/g,"&quot;").replace(/'/g,"&#39;");}
function addCli(){cadastro.push({nome:"",conta:"",email:"",termometro:"",perfil:"",estado:"",cidade:"",nascimento:"",patrimonio:null,ativo:true,
  prev_tem:"",prev_valor:"",prev_onde:"",intl_tem:"",intl_valor:"",intl_onde:"",seg_tem:"",seg_valor:"",mesa_tem:"",mesa_valor:"",mesa_onde:"",grupo_familiar:"",objetivo:"",apos_idade:""});
  cadSortDir=0; if($("cadBusca"))$("cadBusca").value=""; renderCadastro();
  const rows=$("cliBody").rows;if(rows.length)rows[rows.length-1].cells[0].firstElementChild.focus();}
async function delCli(i){const obj=cadView[i]; const j=cadastro.indexOf(obj); if(j>=0)cadastro.splice(j,1); renderCadastro(); const conta=String((obj&&obj.conta)||"").trim(); if(conta){try{await api().delete_cliente(conta);}catch(e){}}}
let _cadSaveT=null;
function autoSaveCadastro(){
  clearTimeout(_cadSaveT);
  if($("cadStatus"))$("cadStatus").textContent="salvando…";
  _cadSaveT=setTimeout(async()=>{
    const r=await api().save_cadastro(cadastro.filter(c=>String(c.conta||"").trim()));
    if($("cadStatus"))$("cadStatus").textContent=(r&&r.ok)?"salvo ✓":"erro ao salvar";
    if(r&&!r.ok)toast("Erro ao salvar cadastro: "+(r.erro||""));
  },700);
}
function uiInativos(lista){
  if(!lista||!lista.length)return;
  $("balaoTxt").innerHTML="⚠ Cliente(s) que não têm mais conta na base (marcados em vermelho): <b>"
    +lista.map(esc).join(", ")+"</b>";
  $("balaoInativos").style.display="block";
}
function fecharBalao(){$("balaoInativos").style.display="none";}

let cliOverview=[], cliCur=null, cliMesSel=null, cliExtraData={};
const EXTRA_DEF=[
  {k:"prev",nome:"Previdência privada",onde:true,labVal:"Valor",labOnde:"Onde está"},
  {k:"intl",nome:"Posição internacional",onde:true,labVal:"Valor",labOnde:"Onde está"},
  {k:"seg",nome:"Seguro",onde:false,labVal:"Valor da cobertura"},
  {k:"mesa",nome:"Plugado na mesa",onde:false,soTem:true}
];
const FLAG_OPC=[["","—"],["Tem","Tem"],["A revisar","A revisar"],["Não tem","Não tem"]];
function flagEstadoJS(v){
  const s=String(v||"").trim().toLowerCase();
  if(["tem","sim","s","ok","1","true","x","yes"].includes(s))return "ok";
  if(["a revisar","revisar","parcial","rev","análise","analise","em análise","em analise"].includes(s))return "rev";
  if(["não tem","nao tem","não","nao","n","0","false","no"].includes(s))return "nao";
  return "";
}
function flagNormInit(v){const e=flagEstadoJS(v);return e==="ok"?"Tem":e==="rev"?"A revisar":e==="nao"?"Não tem":"";}
function renderCliExtra(){
  const fl=(cliCur&&cliCur.flags)||{};
  cliExtraData={};
  EXTRA_DEF.forEach(d=>{const f=fl[d.k]||{};
    cliExtraData[d.k+"_tem"]=flagNormInit(f.tem!==undefined?f.tem:f.status);
    cliExtraData[d.k+"_valor"]=f.valor||"";
    if(d.onde)cliExtraData[d.k+"_onde"]=f.onde||"";});
  cliExtraData["objetivo"]=(cliCur&&cliCur.objetivo)||"";
  cliExtraData["apos_idade"]=(cliCur&&cliCur.apos_idade)||"";
  drawCliExtra();
}
const OBJ_OPC=["","Aposentadoria","Acúmulo de patrimônio","Renda passiva","Educação dos filhos","Outro"];
function objRowHTML(){
  const obj=cliExtraData["objetivo"]||"", alvoTxt=cliExtraData["apos_idade"]||"";
  const sel=OBJ_OPC.map(o=>`<option value="${o}" ${o===obj?"selected":""}>${o||"—"}</option>`).join("");
  let extra="";
  if(obj==="Aposentadoria"){
    const idade=(cliCur&&cliCur.idade!=null)?cliCur.idade:null;
    const alvo=parseInt(alvoTxt,10);
    let info;
    if(!alvo)info='<span class="small mut" style="flex:none;white-space:nowrap">defina a idade-alvo</span>';
    else if(idade==null)info='<span class="small" style="color:var(--warn);flex:none;white-space:nowrap">cadastre o nascimento (dd/mm/aaaa) para calcular</span>';
    else{
      const falta=alvo-idade;
      info=falta>0
        ?`<span class="small" style="color:var(--yellow);font-weight:700;flex:none;white-space:nowrap">faltam ${falta} ano(s) — hoje com ${idade}</span>`
        :'<span class="small" style="color:var(--ok);font-weight:700;flex:none;white-space:nowrap">idade-alvo atingida 🎉</span>';
    }
    extra=`<span class="small mut" style="flex:none;white-space:nowrap">aposentar aos</span>
      <input value="${esc(alvoTxt)}" placeholder="65" maxlength="3" style="width:64px;text-align:center;flex:none"
        oninput="cliExtraData['apos_idade']=this.value.replace(/[^0-9]/g,'');autoSaveCliExtra()"
        onchange="drawCliExtra()"> ${info}`;
  }
  const cor=obj?"#F2C029":"#5C5C66";
  return `<div class="row" style="align-items:center;gap:12px;padding:13px 0;flex-wrap:wrap">
    <span style="display:inline-block;width:11px;height:11px;border-radius:50%;background:${cor};box-shadow:0 0 0 3px ${cor}22;flex:none"></span>
    <span style="min-width:165px;font-weight:600">Objetivo</span>
    <select onchange="cliExtraData['objetivo']=this.value;drawCliExtra();autoSaveCliExtra()" style="width:200px">${sel}</select>
    ${extra}</div>`;
}
function drawCliExtra(){
  $("cliExtra").innerHTML=EXTRA_DEF.map(d=>{
    const tem=cliExtraData[d.k+"_tem"]||"", est=flagEstadoJS(tem), mostra=(est==="ok"||est==="rev");
    const sel=FLAG_OPC.map(([val,lab])=>`<option value="${val}" ${val===tem?"selected":""}>${lab}</option>`).join("");
    let campos=d.soTem?"":`<span class="small mut">${est==="nao"?"sem registro":"não informado"}</span>`;
    if(mostra&&!d.soTem){
      campos=`<input value="${esc(cliExtraData[d.k+"_valor"]||"")}" placeholder="${d.labVal}" oninput="maskMoeda(this);cliExtraData['${d.k}_valor']=this.value;autoSaveCliExtra()" style="width:150px">`;
      if(d.onde)campos+=`<input value="${esc(cliExtraData[d.k+"_onde"]||"")}" placeholder="${d.labOnde}" oninput="cliExtraData['${d.k}_onde']=this.value;autoSaveCliExtra()" style="width:190px">`;
    }
    const _c=FLAGCOR[est]||'#5C5C66';
    return `<div class="row" style="align-items:center;gap:12px;padding:13px 0;border-bottom:1px solid rgba(255,255,255,.05)">
      <span style="display:inline-block;width:11px;height:11px;border-radius:50%;background:${_c};box-shadow:0 0 0 3px ${_c}22;flex:none"></span>
      <span style="min-width:165px;font-weight:600">${d.nome}</span>
      <select onchange="cliExtraData['${d.k}_tem']=this.value;drawCliExtra();autoSaveCliExtra()" style="width:120px">${sel}</select>
      ${campos}</div>`;
  }).join("")+objRowHTML();
}
let _exSaveT=null;
function autoSaveCliExtra(){
  clearTimeout(_exSaveT);
  if($("cliExtraStatus"))$("cliExtraStatus").textContent="salvando…";
  _exSaveT=setTimeout(async()=>{
    if(!cliCur)return;
    const r=await api().save_cliente_extra(cliCur.conta,cliExtraData);
    if(r&&r.ok){
      if($("cliExtraStatus"))$("cliExtraStatus").textContent="salvo ✓";
      const fl={}; EXTRA_DEF.forEach(d=>{fl[d.k]={tem:cliExtraData[d.k+"_tem"],valor:cliExtraData[d.k+"_valor"]||"",
        onde:d.onde?(cliExtraData[d.k+"_onde"]||""):"",status:flagEstadoJS(cliExtraData[d.k+"_tem"])};});
      cliCur.flags=fl;
      cliCur.objetivo=cliExtraData["objetivo"]||"";
      cliCur.apos_idade=cliExtraData["apos_idade"]||"";
      const ov=(cliOverview||[]).find(c=>String(c.conta)===String(cliCur.conta));
      if(ov){ov.flags=fl; ov.plugado_mesa=["ok","rev"].includes(flagEstadoJS(cliExtraData["mesa_tem"]));}
      const cc=(cadastro||[]).find(c=>String(c.conta)===String(cliCur.conta));
      if(cc){EXTRA_DEF.forEach(d=>{cc[d.k+"_tem"]=cliExtraData[d.k+"_tem"]||"";
        cc[d.k+"_valor"]=cliExtraData[d.k+"_valor"]||"";
        if(d.onde)cc[d.k+"_onde"]=cliExtraData[d.k+"_onde"]||"";});
        cc.objetivo=cliExtraData["objetivo"]||""; cc.apos_idade=cliExtraData["apos_idade"]||"";}
    } else { if($("cliExtraStatus"))$("cliExtraStatus").textContent="erro ao salvar"; toast("Erro ao salvar: "+((r&&r.erro)||"")); }
  },500);
}

let cliFam=[];   // [{tipo:'cliente',conta,nome,parentesco} | {tipo:'texto',nome,parentesco}]
const FAM_ICO='<svg width="13" height="13" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" style="flex:none"><path d="M17 21v-2a4 4 0 0 0-4-4H5a4 4 0 0 0-4 4v2"/><circle cx="9" cy="7" r="4"/><path d="M23 21v-2a4 4 0 0 0-3-3.87"/><path d="M16 3.13a4 4 0 0 1 0 7.75"/></svg>';
const BDAY_ICO='<svg width="12" height="12" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.1" stroke-linecap="round" stroke-linejoin="round" style="flex:none"><path d="M20 21v-8a2 2 0 0 0-2-2H6a2 2 0 0 0-2 2v8"/><path d="M4 16s.5-1 2-1 2.5 2 4 2 2.5-2 4-2 2.5 2 4 2 2-1 2-1"/><path d="M2 21h20"/><path d="M7 8v3"/><path d="M12 8v3"/><path d="M17 8v3"/><path d="M7 4h.01"/><path d="M12 4h.01"/><path d="M17 4h.01"/></svg>';
const PARENTESCOS=["Esposa","Marido","Cônjuge","Companheira","Companheiro",
  "Filha","Filho","Mãe","Pai","Irmã","Irmão","Avó","Avô","Neta","Neto",
  "Tia","Tio","Sobrinha","Sobrinho","Prima","Primo","Genro","Nora","Sogra","Sogro",
  "Cunhada","Cunhado","Enteada","Enteado","Madrasta","Padrasto","Outro"];
function famParOpts(cur){
  const lista=(cur&&!PARENTESCOS.includes(cur))?[cur].concat(PARENTESCOS):PARENTESCOS;
  return `<option value="">${cur?"—":"parentesco…"}</option>`
    +lista.map(p=>`<option ${p===cur?"selected":""}>${p}</option>`).join("");
}
function setParentesco(i,v){
  if(!cliFam[i])return;
  cliFam[i].parentesco=v||"";
  saveGrupoFamiliar();
}
function famOpcoes(){
  const ja=new Set(cliFam.filter(f=>f.tipo==="cliente").map(f=>String(f.conta)));
  const base=(cadastro&&cadastro.length?cadastro:cliOverview)||[];
  return base.filter(c=>String(c.conta||"").trim()
      &&String(c.conta)!==String(cliCur&&cliCur.conta)&&!ja.has(String(c.conta)))
    .map(c=>({conta:String(c.conta),nome:(c.nome||"").trim()||("Conta "+c.conta)}));
}
function renderCliFam(){
  if(!$("cliFamChips"))return;
  $("cliFamChips").innerHTML=cliFam.map((f,i)=>{
    const cli=f.tipo==="cliente", nome=esc(f.nome);
    const corpo=cli
      ?`<span onclick="abrirCliente('${esc(f.conta)}')" title="Abrir o detalhe de ${nome}" style="display:inline-flex;align-items:center;gap:7px;cursor:pointer">${FAM_ICO}${nome}<span style="font:500 11px 'JetBrains Mono';opacity:.65">${esc(f.conta)}</span></span>`
      :`<span style="display:inline-flex;align-items:center;gap:7px">${nome}</span>`;
    const par=`<select class="fam-par ${f.parentesco?'':'fam-par-vazio'}" title="Parentesco de ${nome} em relação a este cliente"
      onclick="event.stopPropagation()" onchange="setParentesco(${i},this.value)">${famParOpts(f.parentesco||"")}</select>`;
    return `<span class="fam-chip ${cli?'fam-cli':''}">${corpo}${par}<span class="fam-x" title="Remover vínculo" onclick="delFamiliar(${i})">✕</span></span>`;
  }).join("")||`<span class="small mut" style="padding:4px 0">Nenhum vínculo ainda.</span>`;
  const dl=$("cliFamOpts");
  if(dl)dl.innerHTML=famOpcoes().map(o=>`<option value="${esc(o.nome)} (${esc(o.conta)})">`).join("");
  const ps=$("cliFamParent");
  if(ps&&!ps.options.length)ps.innerHTML=famParOpts("");
}
function addFamiliar(){
  const inp=$("cliFamInput"); let v=((inp&&inp.value)||"").replace(/[|;]/g," ").trim();
  if(!v||!cliCur)return;
  const par=(($("cliFamParent")&&$("cliFamParent").value)||"").trim();
  const ops=famOpcoes(); let alvo=null;
  const m=v.match(/\((\d+)\)\s*$/);
  if(m)alvo=ops.find(o=>o.conta===m[1]);
  if(!alvo)alvo=ops.find(o=>o.conta===v);
  if(!alvo){const nm=ops.filter(o=>o.nome.toLowerCase()===v.toLowerCase()); if(nm.length===1)alvo=nm[0];}
  if(alvo)cliFam.push({tipo:"cliente",conta:alvo.conta,nome:alvo.nome,parentesco:par});
  else{
    if(cliFam.some(f=>f.tipo==="texto"&&f.nome.toLowerCase()===v.toLowerCase())){inp.value="";return;}
    cliFam.push({tipo:"texto",nome:v,parentesco:par});
  }
  inp.value=""; if($("cliFamParent"))$("cliFamParent").value="";
  renderCliFam(); saveGrupoFamiliar();
}
function delFamiliar(i){cliFam.splice(i,1); renderCliFam(); saveGrupoFamiliar();}
let _famSaveT=null;
function saveGrupoFamiliar(){
  clearTimeout(_famSaveT);
  if($("cliFamStatus"))$("cliFamStatus").textContent="salvando…";
  _famSaveT=setTimeout(async()=>{
    if(!cliCur)return;
    const ent=cliFam.map(f=>(f.tipo==="cliente"?("#"+f.conta):f.nome)
      +(f.parentesco?("|"+f.parentesco):""));
    const r=await api().save_grupo_familiar(cliCur.conta,ent);
    if(r&&r.ok){
      cliFam=(r.grupo||[]).slice(); renderCliFam();
      cliCur.grupo=cliFam.slice();
      if($("cliFamStatus"))$("cliFamStatus").textContent="salvo ✓";
      if(r.avisos&&r.avisos.length)toast("⚠ "+r.avisos.join("  ·  "),6000);
      try{
        const cg=await api().cadastro_get();
        if(cg&&cg.ok&&cg.cadastro&&cg.cadastro.length)cadastro=cg.cadastro;
        const ov=await api().clients_overview();
        if(ov&&ov.ok){const map={}; (ov.clientes||[]).forEach(c=>map[String(c.conta)]=c);
          (cliOverview||[]).forEach(c=>{const n=map[String(c.conta)];
            if(n){c.grupo_n=n.grupo_n; c.grupo_txt=n.grupo_txt;}});}
      }catch(e){}
    } else {
      if($("cliFamStatus"))$("cliFamStatus").textContent="erro ao salvar";
      toast("Erro ao salvar grupo familiar: "+((r&&r.erro)||""));
    }
  },400);
}

let cliAportes=null;
async function loadAportes(){
  if(!cliCur)return;
  const r=await api().aportes_get(cliCur.conta);
  cliAportes=(r&&r.ok)?r:null;
  renderAportes();
}
function apChartHTML(mensal){
  mensal=mensal||[];
  if(!mensal.some(m=>m.total>0))return '<div class="hint">Sem aportes nos últimos 12 meses.</div>';
  const mx=Math.max(1,...mensal.map(m=>m.total||0));
  return `<div class="hint" style="margin-bottom:8px">Evolução dos aportes (últimos 12 meses)</div>
    <div style="display:grid;grid-template-columns:repeat(12,1fr);gap:6px;align-items:end;height:64px;border-bottom:1px solid rgba(255,255,255,.06)">`+
    mensal.map(m=>`<div title="${esc(m.label)}: R$ ${fBR(m.total)}" style="height:${m.total?Math.max(8,Math.round(m.total/mx*64)):3}px;border-radius:5px 5px 0 0;background:${m.total?'linear-gradient(180deg,var(--yellow),rgba(242,192,41,.35))':'rgba(255,255,255,.07)'}"></div>`).join("")+
    `</div><div style="display:grid;grid-template-columns:repeat(12,1fr);gap:6px;text-align:center;margin-top:6px">`+
    mensal.map(m=>`<span style="font-size:9px;color:#62666F">${esc((m.label||"").slice(0,2))}</span>`).join("")+`</div>`;
}
function renderAportes(){
  const R=$("cliApResumo"),L=$("cliApLista"),C=$("cliApChart");
  if(!R)return;
  const d=cliAportes;
  if(!d){R.innerHTML='<span class="mut">Não consegui carregar os aportes.</span>';
    if(L)L.innerHTML=""; if(C)C.innerHTML=""; return;}
  const itens=d.itens||[], mensal=d.mensal||[];
  const diasTxt=d.n===0
    ?'<span class="badge warn" style="background:rgba(239,68,68,.12);color:var(--err)">nunca aportou</span>'
    :(d.dias==null?"—":(d.alerta
      ?`<span class="badge warn" style="background:rgba(239,68,68,.12);color:var(--err)">há ${d.dias} dias sem aportar</span>`
      :`<span class="badge ok">último há ${d.dias} dia(s)</span>`));
  R.innerHTML=`Total aportado <b style="color:var(--yellow)">R$ ${fBR(d.total)}</b>
    &nbsp;·&nbsp; <b style="color:var(--txt)">${d.n}</b> movimentação(ões)
    &nbsp;·&nbsp; Último aporte: <b style="color:var(--txt)">${esc(d.ultimo||"—")}</b> &nbsp;${diasTxt}`;
  if(C)C.innerHTML=apChartHTML(mensal);
  if(L)L.innerHTML=itens.length?itens.map(it=>`
    <div class="row" style="padding:8px 2px;border-bottom:1px solid rgba(255,255,255,.05)">
      <span class="dc-mono" style="width:96px;flex:none;font-size:12.5px;color:#C9CBD1">${esc(it.data)}</span>
      <span class="dc-mono" style="width:140px;flex:none;font-size:13px;font-weight:700;color:var(--yellow)">R$ ${fBR(it.valor)}</span>
      <span class="small mut grow" style="min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap">${esc(it.obs||"")}</span>
      <span class="del" title="Excluir aporte" onclick="delAporte(${it.id})">✕</span>
    </div>`).join("")
    :'<div class="hint">Nenhum aporte registrado ainda — registre o primeiro acima.</div>';
  const inp=$("apAlertaDias"); if(inp&&document.activeElement!==inp)inp.value=d.alerta_dias||90;
}
async function addAporte(){
  if(!cliCur)return;
  const v=numBR(($("apValor")&&$("apValor").value)||"");
  if(!v||v<=0){toast("Informe um valor de aporte válido.");return;}
  if($("cliApStatus"))$("cliApStatus").textContent="salvando…";
  const r=await api().aporte_add(cliCur.conta,($("apData")&&$("apData").value.trim())||"",v,($("apObs")&&$("apObs").value.trim())||"");
  if(r&&r.ok){
    cliAportes=r; renderAportes();
    if($("apValor"))$("apValor").value=""; if($("apObs"))$("apObs").value=""; if($("apData"))$("apData").value="";
    if($("cliApStatus"))$("cliApStatus").textContent="salvo ✓";
    try{const ov=await api().clients_overview();
      if(ov&&ov.ok){const map={}; (ov.clientes||[]).forEach(c=>map[String(c.conta)]=c);
        (cliOverview||[]).forEach(c=>{const n=map[String(c.conta)]; if(n)Object.assign(c,{ap_n:n.ap_n,ap_total:n.ap_total,ap_ultimo:n.ap_ultimo,ap_dias:n.ap_dias,ap_alerta:n.ap_alerta,ap_lim:n.ap_lim});});}
    }catch(e){}
  } else {
    if($("cliApStatus"))$("cliApStatus").textContent="erro";
    toast("Erro ao registrar aporte: "+((r&&r.erro)||""));
  }
}
async function delAporte(id){
  if(!cliCur)return;
  const r=await api().aporte_del(cliCur.conta,id);
  if(r&&r.ok){cliAportes=r; renderAportes(); if($("cliApStatus"))$("cliApStatus").textContent="salvo ✓";}
  else toast("Erro ao excluir: "+((r&&r.erro)||""));
}
function salvarApAlerta(v){
  const n=parseInt(String(v).replace(/\D/g,""),10);
  if(!n||n<1){toast("Informe um número de dias válido.");return;}
  api().set_config("aporte_alerta_dias",n).then(()=>{loadAportes(); cliLoad();});
}

let cliContatos=null;
const CT_TIPOS=["Ligação","Mensagem","E-mail","Reunião presencial","Evento","Outro"];
async function loadContatos(){
  if(!cliCur)return;
  const r=await api().contatos_get(cliCur.conta);
  cliContatos=(r&&r.ok)?r:null;
  renderContatos();
}
function renderContatos(){
  const R=$("cliCtResumo"),L=$("cliCtLista"),S=$("ctTipo");
  if(!R)return;
  if(S&&!S.options.length)S.innerHTML=CT_TIPOS.map(t=>`<option>${t}</option>`).join("");
  const d=cliContatos;
  if(!d){R.innerHTML='<span class="mut">Não consegui carregar os contatos.</span>'; if(L)L.innerHTML=""; return;}
  const itens=d.itens||[];
  R.innerHTML=d.ultimo
    ?`Último contato: <b style="color:var(--txt)">${esc(d.ultimo)}</b> · <b style="color:var(--yellow)">${esc(d.tipo)}</b>${d.dias!=null?` · há <b style="color:var(--txt)">${d.dias}</b> dia(s)`:""} &nbsp;·&nbsp; ${d.n} contato(s) registrados`
    :'<span class="mut">Nenhum contato registrado ainda.</span>';
  if(L)L.innerHTML=itens.length?itens.map(it=>`
    <div class="row" style="padding:8px 2px;border-bottom:1px solid rgba(255,255,255,.05)">
      <span class="dc-mono" style="width:96px;flex:none;font-size:12.5px;color:#C9CBD1">${esc(it.data)}</span>
      <span style="width:160px;flex:none;font-size:12.5px;font-weight:700;color:var(--yellow)">${esc(it.tipo)}</span>
      <span class="small mut grow" style="min-width:0;overflow:hidden;text-overflow:ellipsis;white-space:nowrap">${esc(it.obs||"")}</span>
      <span class="del" title="Excluir contato" onclick="delContato(${it.id})">✕</span>
    </div>`).join(""):"";
}
async function addContato(){
  if(!cliCur)return;
  if($("cliCtStatus"))$("cliCtStatus").textContent="salvando…";
  const r=await api().contato_add(cliCur.conta,
    ($("ctData")&&$("ctData").value.trim())||"",
    ($("ctTipo")&&$("ctTipo").value)||"Outro",
    ($("ctObs")&&$("ctObs").value.trim())||"");
  if(r&&r.ok){
    cliContatos=r; renderContatos();
    if($("ctData"))$("ctData").value=""; if($("ctObs"))$("ctObs").value="";
    if($("cliCtStatus"))$("cliCtStatus").textContent="salvo ✓";
    try{const ov=await api().clients_overview();
      if(ov&&ov.ok){const map={}; (ov.clientes||[]).forEach(c=>map[String(c.conta)]=c);
        (cliOverview||[]).forEach(c=>{const n=map[String(c.conta)];
          if(n)Object.assign(c,{ct_n:n.ct_n,ct_ultimo:n.ct_ultimo,ct_tipo:n.ct_tipo,ct_dias:n.ct_dias});});}
    }catch(e){}
  } else {
    if($("cliCtStatus"))$("cliCtStatus").textContent="erro";
    toast("Erro ao registrar contato: "+((r&&r.erro)||""));
  }
}
async function delContato(id){
  if(!cliCur)return;
  const r=await api().contato_del(cliCur.conta,id);
  if(r&&r.ok){cliContatos=r; renderContatos(); if($("cliCtStatus"))$("cliCtStatus").textContent="salvo ✓";}
  else toast("Erro ao excluir: "+((r&&r.erro)||""));
}
function cliView(v){
  $("pillRel").classList.toggle("sel",v==="rel");
  $("pillCad").classList.toggle("sel",v==="cad");
  $("cliRel").style.display=v==="rel"?"flex":"none";
  $("cliCad").style.display=v==="cad"?"flex":"none";
  $("cliResumo").style.display=(v==="rel" && !cliCur) ? "" : "none";
  if(v==="rel") cliLoad();
  else if(v==="cad"){
    if(!_cadPendKeep)cadPendFiltro=null;
    const alvo=cliCur?String(cliCur.conta||"").trim():"";
    if($("cadBusca"))$("cadBusca").value=alvo;
    cadSortDir=0;
    renderCadastro();
    if(alvo)cadGarantirCliente(alvo);
  }
}
async function cadGarantirCliente(alvo){
  const acha=()=>(cadastro||[]).find(c=>String(c.conta||"").trim()===alvo);
  if(!acha()){
    try{
      const cg=await api().cadastro_get();
      if(cg&&cg.ok&&Array.isArray(cg.cadastro)&&cg.cadastro.length)cadastro=cg.cadastro;
    }catch(e){}
  }
  if(!cliCur||String(cliCur.conta||"").trim()!==alvo){renderCadastro();return;}
  let row=acha(), mudou=false;
  if(!row){
    row={nome:"",conta:alvo,email:"",termometro:"",perfil:"",estado:"",cidade:"",nascimento:"",
      patrimonio:null,ativo:true,
      prev_tem:"",prev_valor:"",prev_onde:"",intl_tem:"",intl_valor:"",intl_onde:"",
      seg_tem:"",seg_valor:"",mesa_tem:"",mesa_valor:"",mesa_onde:"",grupo_familiar:"",
      objetivo:"",apos_idade:""};
    cadastro.push(row); mudou=true;
  }
  const preenche=(campo,valor)=>{
    if(valor&&!String(row[campo]||"").trim()){row[campo]=valor;mudou=true;}
  };
  preenche("nome",cliCur.nome||"");
  preenche("email",cliCur.email||"");
  preenche("termometro",(cliCur.termometro||"").toUpperCase());
  preenche("perfil",cliCur.perfil||"");
  preenche("nascimento",cliCur.nascimento||"");
  if(row.patrimonio==null&&cliCur.patrimonio!=null){row.patrimonio=cliCur.patrimonio;mudou=true;}
  if(mudou)autoSaveCadastro();
  renderCadastro();
}
function flagsFromCad(c){
  const one=(k,onde=true)=>({tem:c[k+"_tem"]||"",valor:c[k+"_valor"]||"",
    onde:onde?(c[k+"_onde"]||""):"",status:flagEstadoJS(c[k+"_tem"])});
  return {prev:one("prev"),intl:one("intl"),seg:one("seg",false),mesa:one("mesa")};
}
async function cliLoad(){
  await cliColsLoad();
  const r=await api().clients_overview();
  cliOverview=(r&&r.ok&&r.clientes)||[];
  cliOverview.forEach(ov=>{
    const c=(cadastro||[]).find(x=>String(x.conta)===String(ov.conta));
    if(c){ov.flags=flagsFromCad(c); ov.plugado_mesa=["ok","rev"].includes(flagEstadoJS(c.mesa_tem));}
  });
  renderCliResumo();
  if(!cliCur) renderCliList();
}
const TC2={A:"#10B981",B:"#3B82F6",C:"#F59E0B",D:"#F97316",E:"#EF4444"};
const FLAGCOR={ok:"#10B981",rev:"#F59E0B",nao:"#EF4444"};
const FLAGTXT={ok:"Tem (completo)",rev:"A revisar — faltam dados",nao:"Não tem"};
function flagDot(f){
  f=f||{status:"nao"};
  const cor=FLAGCOR[f.status]||"#5C5C66";
  let tip=FLAGTXT[f.status]||"";
  if(f.status&&f.status!=="nao"){const ex=[];
    if(f.valor)ex.push("valor: "+f.valor); if(f.onde)ex.push("onde: "+f.onde);
    if(ex.length)tip+=" — "+ex.join(" · ");}
  return `<span title="${esc(tip)}" style="display:inline-block;width:11px;height:11px;border-radius:50%;background:${cor};box-shadow:0 0 0 3px ${cor}22"></span>`;
}
function renderCliResumo(){
  const cs=(cliOverview||[]).slice(), total=cs.length;
  const aum=cs.reduce((a,c)=>a+(c.patrimonio||0),0);
  const TO=["A","B","C","D","E"], tc={A:0,B:0,C:0,D:0,E:0};
  cs.forEach(c=>{const t=(c.termometro||"").toUpperCase(); if(tc[t]!==undefined)tc[t]++;});
  const termoNums=`<div style="display:grid;grid-template-columns:repeat(5,1fr);text-align:center;gap:4px;margin-top:8px">`+
    TO.map(k=>`<div><div style="font:800 22px/1 'JetBrains Mono',monospace;letter-spacing:-.02em;color:${tc[k]?TC2[k]:'#5A5E67'}">${tc[k]}</div>
      <div style="font-size:11px;font-weight:600;color:${tc[k]?TC2[k]:'#62666F'};margin-top:5px">${k}</div></div>`).join("")+`</div>`;
  const fx=[["<1M",v=>v<1e6],["1–3M",v=>v>=1e6&&v<3e6],["3–5M",v=>v>=3e6&&v<5e6],[">5M",v=>v>=5e6]];
  const fc=fx.map(([l,f])=>[l,cs.filter(c=>c.patrimonio!=null&&f(c.patrimonio)).length]);
  const fmax=Math.max(1,...fc.map(x=>x[1]));
  const fbars=fc.map(([l,v])=>`<div class="bn-bar">
     <div style="font-size:12px;font-weight:700;color:#F2C029">${v}</div>
     <div class="col"><i style="height:${Math.round(v/fmax*100)||3}%;background:linear-gradient(180deg,rgba(242,192,41,.85),rgba(242,192,41,.22))"></i></div>
     <div class="bn-sub">${l}</div></div>`).join("");
  const nPat=cs.filter(c=>c.patrimonio!=null).length;
  const nTermo=cs.filter(c=>tc[(c.termometro||"").toUpperCase()]!==undefined).length;
  $("cliResumo").innerHTML=
    `<div class="bn-card"><div class="bn-lbl">Total de clientes</div><div class="bn-big" style="color:#F3F3F5">${total}</div>${nChip(total)}</div>`+
    `<div class="bn-card"><div class="bn-lbl">Patrimônio total</div><div class="bn-big">${milM(aum)}</div>${nChip(nPat,total)}</div>`+
    `<div class="bn-card"><div class="bn-lbl">Por termômetro</div>${termoNums}${nChip(nTermo,total)}</div>`+
    `<div class="bn-card"><div class="bn-lbl">Patrimônio por faixa</div><div class="bn-bars">${fbars}</div>${nChip(nPat,total)}</div>`;
}
function _ini(nome,conta){const n=(nome||"").trim();
  if(!n)return String(conta||"?").slice(0,2);
  const p=n.split(/\s+/); return (((p[0]||"")[0]||"")+(p.length>1?(p[p.length-1][0]||""):((p[0]||"")[1]||""))).toUpperCase();}
function cobCell(f,conta,campo){f=f||{status:""}; const cor=FLAGCOR[f.status]||"#5C5C66";
  let tip=FLAGTXT[f.status]||"não informado";
  if(f.status&&f.status!=="nao"){const ex=[]; if(f.valor)ex.push("valor: "+f.valor); if(f.onde)ex.push("onde: "+f.onde); if(ex.length)tip+=" — "+ex.join(" · ");}
  const clique=conta?` onclick="lieAbrir(event,'${esc(String(conta))}','${campo}')"`:"";
  return `<div style="text-align:center${conta?';cursor:pointer':''}"${clique} title="${esc(tip)}${conta?' — clique para alterar':''}"><span style="display:inline-block;width:10px;height:10px;border-radius:50%;background:${cor};box-shadow:0 0 0 3px ${cor}1f"></span></div>`;}

const CLI_COLS={
  cliente:{lab:"Cliente",        w:"minmax(240px,2fr)",min:240},
  termo:  {lab:"Termômetro",     w:"106px",min:106,c:1},
  perfil: {lab:"Perfil",         w:"122px",min:122,c:1},
  pat:    {lab:"Patrimônio",     w:"168px",min:168,c:1},
  apos:   {lab:"Tempo para Aposentadoria",w:"186px",min:186,c:1},
  prev:   {lab:"Previdência",    w:"108px",min:108,c:1},
  intl:   {lab:"Internacional",  w:"114px",min:114,c:1},
  seg:    {lab:"Seguro",         w:"92px", min:92, c:1},
  mesa:   {lab:"Plugado na mesa",w:"132px",min:132,c:1},
  contato:{lab:"Último contato", w:"132px",min:132,c:1},
  aporte: {lab:"Último aporte",  w:"128px",min:128,c:1}
};
const CLI_COLS_DEF=["cliente","termo","perfil","pat","apos","prev","intl","seg","mesa","contato","aporte"];
let cliColOrder=CLI_COLS_DEF.slice(), _colsCfgOk=false, _dragCol=null;
async function cliColsLoad(){
  if(_colsCfgOk)return;
  try{
    const r=await api().get_config();
    const o=r&&r.ok&&r.config&&r.config.cli_cols;
    if(Array.isArray(o)&&o.length){
      const v=o.filter(k=>CLI_COLS[k]);
      CLI_COLS_DEF.forEach((k,di)=>{
        if(v.includes(k))return;
        let ins=v.length;
        for(let j=di-1;j>=0;j--){const p=v.indexOf(CLI_COLS_DEF[j]); if(p>=0){ins=p+1;break;}}
        v.splice(ins,0,k);
      });
      cliColOrder=v;
    }
  }catch(e){}
  _colsCfgOk=true;
}
function cliTpl(){return cliColOrder.map(k=>CLI_COLS[k].w).join(" ")+" 30px";}
function cliMinW(){return cliColOrder.reduce((a,k)=>a+CLI_COLS[k].min,0)+30+14*cliColOrder.length+32;}
function dropCol(k){
  if(!_dragCol||_dragCol===k||!CLI_COLS[_dragCol])return;
  const o=cliColOrder.filter(x=>x!==_dragCol);
  o.splice(o.indexOf(k),0,_dragCol);
  cliColOrder=o; _dragCol=null;
  renderCliList();
  try{api().set_config("cli_cols",cliColOrder);}catch(e){}
}
let cliSort={col:null,dir:0};
let cliFilt={};
const CF_ICON='<svg width="11" height="11" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2.6" stroke-linecap="round" stroke-linejoin="round"><polygon points="22 3 2 3 10 12.5 10 19 14 21 14 12.5 22 3"/></svg>';
const CF_ST=[["ok","Tem"],["rev","A revisar"],["nao","Não tem"],["","Não informado"]];
function _stFlag(c,k){return ((c.flags||{})[k]||{}).status||"";}
function cliFiltAtivo(k){
  const f=cliFilt[k]; if(!f)return false;
  return k==="pat"?(f.min!=null||f.max!=null):f.size>0;
}
function cliPassaFiltros(c){
  for(const k in cliFilt){
    if(!cliFiltAtivo(k))continue;
    const f=cliFilt[k];
    if(k==="pat"){
      const v=c.patrimonio;
      if(v==null)return false;
      if(f.min!=null&&v<f.min)return false;
      if(f.max!=null&&v>f.max)return false;
    }else if(k==="contato"){
      if(!f.has(c.ct_ultimo?"com":"sem"))return false;
    }else if(k==="aporte"){
      if(!f.has(!c.ap_ultimo?"sem":(c.ap_alerta?"alerta":"ok")))return false;
    }else if(k==="apos"){
      const ap=(c.apos_obj||"")==="Aposentadoria";
      const key=!ap?"sem":(c.apos_falta!=null&&c.apos_falta<=0?"atingida":"com");
      if(!f.has(key))return false;
    }else if(["prev","intl","seg","mesa"].includes(k)){
      if(!f.has(_stFlag(c,k)))return false;
    }else if(k==="termo"){
      if(!f.has((c.termometro||"").toUpperCase()||"—"))return false;
    }else if(k==="perfil"){
      if(!f.has(c.perfil||"—"))return false;
    }
  }
  return true;
}
function cliSortCol(k){
  if(cliSort.col!==k)cliSort={col:k,dir:1};
  else if(cliSort.dir===1)cliSort.dir=-1;
  else cliSort={col:null,dir:0};
  renderCliList();
}
function cliCmp(a,b){
  const k=cliSort.col, d=cliSort.dir;
  const NUM={pat:c=>c.patrimonio,contato:c=>c.ct_dias,aporte:c=>c.ap_dias,
             apos:c=>(c.apos_obj==="Aposentadoria"&&c.apos_falta!=null)?c.apos_falta:null};
  if(NUM[k]){
    const va=NUM[k](a), vb=NUM[k](b);
    if(va==null&&vb==null)return 0;
    if(va==null)return 1; if(vb==null)return -1;
    return (va-vb)*d;
  }
  let va,vb;
  if(k==="termo"){va=(a.termometro||"").toUpperCase()||"~";vb=(b.termometro||"").toUpperCase()||"~";}
  else if(k==="perfil"){va=a.perfil||"~";vb=b.perfil||"~";}
  else if(["prev","intl","seg","mesa"].includes(k)){
    const ord={ok:0,rev:1,nao:2,"":3};
    return (ord[_stFlag(a,k)]-ord[_stFlag(b,k)])*d;
  }
  else {va=(a.nome||"").toLowerCase();vb=(b.nome||"").toLowerCase();}
  return String(va).localeCompare(String(vb),"pt-BR")*d;
}
function cfOpcoes(k){
  if(k==="termo")return ["A","B","C","D","E","—"].map(v=>[v,v]);
  if(k==="perfil")return [...new Set(cliOverview.map(c=>c.perfil||"—"))].sort((a,b)=>a.localeCompare(b,"pt-BR")).map(v=>[v,v]);
  if(["prev","intl","seg","mesa"].includes(k))return CF_ST;
  if(k==="contato")return [["com","Com contato"],["sem","Sem contato"]];
  if(k==="aporte")return [["ok","Em dia"],["alerta","Em alerta"],["sem","Sem aporte"]];
  if(k==="apos")return [["com","Rumo à aposentadoria"],["atingida","Idade-alvo atingida"],["sem","Sem esse objetivo"]];
  return [];
}
function cfAbrir(ev,k){
  ev.stopPropagation();
  const P=$("cfPop");
  if(P.style.display==="block"&&P.dataset.col===k){cfFechar();return;}
  P.dataset.col=k;
  const tit=`<div style="font:700 11px 'Manrope';letter-spacing:.08em;text-transform:uppercase;color:#9c948a;padding:2px 4px 8px">${CLI_COLS[k].lab}</div>`;
  let corpo;
  if(k==="pat"){
    const f=cliFilt.pat||{};
    corpo=`<div style="display:flex;flex-direction:column;gap:8px;padding:0 4px">
      <label class="small mut">De (R$)<input value="${f.min!=null?fBR(f.min):""}" oninput="maskMoeda(this);cfPatMinMax('min',this.value)" placeholder="mínimo" style="width:100%;margin-top:3px"></label>
      <label class="small mut">Até (R$)<input value="${f.max!=null?fBR(f.max):""}" oninput="maskMoeda(this);cfPatMinMax('max',this.value)" placeholder="máximo" style="width:100%;margin-top:3px"></label></div>`;
  }else{
    const f=cliFilt[k];
    corpo=cfOpcoes(k).map(([v,lab])=>`<label class="cf-it">
      <input type="checkbox" ${f&&f.has(v)?"checked":""} onchange="cfToggle('${k}','${esc(v)}',this.checked)">
      <span>${esc(lab)}</span></label>`).join("");
  }
  P.innerHTML=tit+corpo+`<div style="display:flex;justify-content:flex-end;padding:8px 4px 2px">
    <span onclick="cfLimpar('${k}')" style="cursor:pointer;font-size:12px;font-weight:600;color:#F2C029">Limpar</span></div>`;
  const r=ev.currentTarget.getBoundingClientRect();
  P.style.left=Math.max(8,Math.min(r.left,window.innerWidth-240))+"px";
  P.style.top=(r.bottom+6)+"px";
  P.style.display="block";
}
function cfFechar(){const P=$("cfPop");if(P){P.style.display="none";P.dataset.col="";}}
function cfToggle(k,v,on){
  const f=cliFilt[k]||(cliFilt[k]=new Set());
  on?f.add(v):f.delete(v);
  renderCliList();
}
function _numFiltro(v){
  v=String(v||"").trim().replace(/[^\d.,]/g,""); if(!v)return null;
  v=v.includes(",")?v.replace(/\./g,"").replace(",","."):v.replace(/\./g,"");
  const n=parseFloat(v); return isFinite(n)?n:null;
}
function cfPatMinMax(q,v){
  const f=cliFilt.pat||(cliFilt.pat={});
  f[q]=_numFiltro(v);
  renderCliList();
}
function cfLimpar(k){delete cliFilt[k];cfFechar();renderCliList();}
let lieConta=null;
const LIE_OPS={
  termo:["","A","B","C","D","E"].map(v=>[v,v||"—",null]),
  perfil:["","Conservadora","Moderada","Sofisticada"].map(v=>[v,v||"—",null]),
  flag:[["","não informado","#5C5C66"],["Tem","Tem","#10B981"],
        ["A revisar","A revisar","#F59E0B"],["Não tem","Não tem","#EF4444"]]
};
function lieValorAtual(c,campo){
  if(campo==="termo")return (c.termometro||"").toUpperCase();
  if(campo==="perfil")return c.perfil||"";
  return flagNormInit(c[campo+"_tem"]);
}
function lieAbrir(ev,conta,campo){
  ev.stopPropagation();
  cfFechar();
  const P=$("liePop"); lieConta=String(conta);
  const c=(cadastro||[]).find(x=>String(x.conta||"").trim()===lieConta)||{};
  const cur=lieValorAtual(c,campo);
  const ops=LIE_OPS[campo]||LIE_OPS.flag;
  P.innerHTML=ops.map(([v,l,cor])=>`<div class="cbx-it ${v===cur?'on':''}"
      onmousedown="event.preventDefault();lieSet('${campo}','${v}')">
      ${cor?`<span style="width:9px;height:9px;border-radius:50%;background:${cor};flex:none;box-shadow:0 0 0 3px ${cor}22"></span>`:""}
      <span class="cbx-nm">${l}</span></div>`).join("");
  const r=ev.currentTarget.getBoundingClientRect();
  P.style.left=Math.max(8,Math.min(r.left,window.innerWidth-180))+"px";
  P.style.top=(r.bottom+5)+"px";
  P.style.display="block";
}
function lieFechar(){const P=$("liePop");if(P)P.style.display="none";lieConta=null;}
function lieSet(campo,v){
  const conta=lieConta; lieFechar();
  if(!conta)return;
  let c=(cadastro||[]).find(x=>String(x.conta||"").trim()===conta);
  if(!c){
    const ov=(cliOverview||[]).find(x=>String(x.conta)===conta)||{};
    c={nome:ov.nome||"",conta:conta,email:ov.email||"",termometro:(ov.termometro||"").toUpperCase(),
       perfil:ov.perfil||"",estado:"",cidade:"",nascimento:"",patrimonio:null,ativo:true,
       prev_tem:"",prev_valor:"",prev_onde:"",intl_tem:"",intl_valor:"",intl_onde:"",
       seg_tem:"",seg_valor:"",mesa_tem:"",mesa_valor:"",mesa_onde:"",grupo_familiar:"",
       objetivo:"",apos_idade:""};
    cadastro.push(c);
  }
  if(campo==="termo")c.termometro=v;
  else if(campo==="perfil")c.perfil=v;
  else c[campo+"_tem"]=v;
  autoSaveCadastro();
  const ov=(cliOverview||[]).find(x=>String(x.conta)===conta);
  if(ov){
    if(campo==="termo")ov.termometro=v;
    else if(campo==="perfil")ov.perfil=v;
    else{ov.flags=flagsFromCad(c);ov.plugado_mesa=["ok","rev"].includes(flagEstadoJS(c.mesa_tem));}
  }
  if(!cliCur)renderCliList();
  renderCliResumo&&cliOverview&&renderCliResumo();
}
document.addEventListener("click",e=>{
  const P=$("liePop");
  if(P&&P.style.display==="block"&&!P.contains(e.target))lieFechar();
});
document.addEventListener("keydown",e=>{if(e.key==="Escape")lieFechar();});
document.addEventListener("click",e=>{
  const P=$("cfPop");
  if(P&&P.style.display==="block"&&!P.contains(e.target))cfFechar();
});
document.addEventListener("keydown",e=>{if(e.key==="Escape")cfFechar();});
function renderCliHdr(){
  const h=$("cliHdr"); if(!h)return;
  h.style.display="grid"; h.style.gridTemplateColumns=cliTpl();
  h.style.gap="14px"; h.style.alignItems="center";
  h.innerHTML=cliColOrder.map(k=>{
    const d=CLI_COLS[k];
    const s=cliSort.col===k?cliSort.dir:0;
    const seta=s===1?"▲":(s===-1?"▼":"⇅");
    const temFiltro=cfOpcoes(k).length||k==="pat";
    return `<span draggable="true" data-col="${k}" title="Arraste para mudar a posição desta coluna"
      style="cursor:grab;white-space:nowrap;display:flex;align-items:center;gap:6px;${d.c?'justify-content:center;':''}"
      ondragstart="_dragCol='${k}'" ondragover="event.preventDefault();this.style.opacity=.45"
      ondragleave="this.style.opacity=1" ondrop="this.style.opacity=1;dropCol('${k}')"
      ondragend="_dragCol=null">
      <span onclick="event.stopPropagation();cliSortCol('${k}')" title="Clique para ordenar (crescente → decrescente → padrão)" style="cursor:pointer;display:inline-flex;align-items:center;gap:4px">${d.lab}<span style="font-size:9px;color:${s?'#F2C029':'#55565E'}">${seta}</span></span>
      ${temFiltro?`<span onclick="cfAbrir(event,'${k}')" title="Filtrar esta coluna" style="cursor:pointer;display:inline-flex;color:${cliFiltAtivo(k)?'#F2C029':'#55565E'}">${CF_ICON}</span>`:""}
    </span>`;
  }).join("")+"<span></span>";
}
function cliCell(k,c){
  const fl=c.flags||{};
  switch(k){
    case "cliente":{
      const nome=esc(c.nome)||"<span style='color:var(--warn)'>(sem nome)</span>";
      const fam=c.grupo_n?`<span class="fam-badge" title="Grupo familiar: ${esc(c.grupo_txt||"")}">${FAM_ICO}${c.grupo_n}</span>`:"";
      const bd=!c.aniv_status?"":(c.aniv_status==="hoje"
        ?`<span class="bday-badge" title="Aniversário HOJE${c.idade_nova!=null?` — faz ${c.idade_nova} anos`:""}!">${BDAY_ICO} HOJE</span>`
        :`<span class="bday-badge" title="Aniversário no ${esc(c.aniv_dsem)} (${esc(c.aniv_data)}) — cai no fim de semana, dê os parabéns hoje!">${BDAY_ICO} ${esc((c.aniv_dsem||"").slice(0,3).toUpperCase())}</span>`);
      return `<div style="display:flex;align-items:center;gap:12px;min-width:0">
        <div class="cli-av">${esc(_ini(c.nome,c.conta))}</div>
        <div style="min-width:0">
          <div style="display:flex;align-items:center;gap:8px;min-width:0"><span style="font-size:14px;font-weight:600;color:#F3F3F5;white-space:nowrap;overflow:hidden;text-overflow:ellipsis">${nome}</span>${bd}${fam}</div>
          <div style="font:500 12px 'JetBrains Mono';color:#7A7A84;display:flex;align-items:center">${esc(c.conta)}${copyBtn(c.conta)}</div></div></div>`;
    }
    case "termo":{
      const t=(c.termometro||"").toUpperCase(), tc=TC2[t]||"#5C5C66";
      return `<div style="text-align:center;cursor:pointer" title="Clique para alterar o termômetro" onclick="lieAbrir(event,'${c.conta}','termo')">${t?`<span style="display:inline-flex;width:30px;height:30px;border-radius:8px;align-items:center;justify-content:center;font:700 14px 'Manrope';color:${tc};background:${tc}1f;border:1px solid ${tc}47">${t}</span>`:'<span class="mut">—</span>'}</div>`;
    }
    case "perfil": return `<div style="font-size:13px;color:#B4B8C0;text-align:center;cursor:pointer" title="Clique para alterar o perfil" onclick="lieAbrir(event,'${c.conta}','perfil')">${esc(c.perfil)||"—"}</div>`;
    case "pat": return `<div style="font:600 13.5px 'JetBrains Mono';color:#F3F3F5;text-align:center;white-space:nowrap">${c.patrimonio!=null?"R$ "+fBR(c.patrimonio,2):'<span class="mut">—</span>'}</div>`;
    case "apos":{
      if((c.apos_obj||"")!=="Aposentadoria")
        return '<div style="text-align:center"><span class="mut">—</span></div>';
      if(c.apos_falta==null)
        return '<div style="text-align:center"><span style="font-size:11.5px;color:var(--warn);font-weight:600" title="Defina a idade-alvo e o nascimento no cliente">dados incompletos</span></div>';
      if(c.apos_falta<=0)
        return '<div style="text-align:center"><span style="font-size:12px;font-weight:700;color:var(--ok)">atingida 🎉</span></div>';
      return `<div style="text-align:center;line-height:1.3" title="Aposentadoria aos ${esc(c.apos_alvo)}">
        <div class="dc-mono" style="font-size:13.5px;font-weight:700;color:#F2C029">${c.apos_falta} ano(s)</div>
        <div style="font-size:11px;color:#9A9AA4">aos ${esc(c.apos_alvo)}</div></div>`;
    }
    case "prev": return cobCell(fl.prev,c.conta,"prev");
    case "intl": return cobCell(fl.intl,c.conta,"intl");
    case "seg":  return cobCell(fl.seg,c.conta,"seg");
    case "mesa": return cobCell(fl.mesa,c.conta,"mesa");
    case "contato":{
      if(!c.ct_ultimo)
        return `<div style="text-align:center"><span style="font-size:12px;color:#8E8E98">—</span></div>`;
      return `<div style="text-align:center;line-height:1.3" title="${esc(c.ct_tipo)}${c.ct_dias!=null?" — há "+c.ct_dias+" dia(s)":""}">
        <div class="dc-mono" style="font-size:12.5px;color:#C9CBD1">${esc(c.ct_ultimo)}</div>
        <div style="font-size:11px;font-weight:600;color:#9A9AA4">${esc(c.ct_tipo)||""}</div></div>`;
    }
    case "aporte":{
      if(!c.ap_ultimo)
        return `<div style="text-align:center"><span style="font-size:11.5px;font-weight:700;color:${c.ap_alerta?'var(--err)':'#8E8E98'}">${c.ap_alerta?'sem aporte':'—'}</span></div>`;
      const lim=c.ap_lim||90;
      const cor=c.ap_alerta?'var(--err)':(c.ap_dias>lim*0.66?'var(--warn)':'var(--ok)');
      return `<div style="text-align:center;line-height:1.3" title="Total aportado: R$ ${fBR(c.ap_total)} em ${c.ap_n} movimentação(ões)">
        <div class="dc-mono" style="font-size:12.5px;color:#C9CBD1">${esc(c.ap_ultimo)}</div>
        <div style="font-size:11px;font-weight:700;color:${cor}">há ${c.ap_dias}d</div></div>`;
    }
  }
  return "<div></div>";
}
function renderCliList(){
  const q=($("cliBusca").value||"").toLowerCase().trim();
  const lst=cliOverview.filter(c=>
    (!q||(c.nome||"").toLowerCase().includes(q)||String(c.conta).includes(q)) && cliPassaFiltros(c));
  const chev='<svg width="16" height="16" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round"><path d="m9 18 6-6-6-6"/></svg>';
  if(cliSort.col){
    lst.sort(cliCmp);
  }else{
    const _aw=c=>c.aniv_status==="hoje"?0:(c.aniv_status==="fds"?1:2);
    lst.sort((x,y)=>_aw(x)-_aw(y));
  }
  const nf=Object.keys(cliFilt).filter(cliFiltAtivo).length;
  if($("cliListWrap"))$("cliListWrap").style.minWidth=cliMinW()+"px";
  renderCliHdr();
  const tpl=cliTpl();
  $("cliListBody").innerHTML=lst.map(c=>
    `<div class="cli-row" onclick="abrirCliente('${c.conta}')" style="display:grid;grid-template-columns:${tpl};gap:14px;align-items:center">
      ${cliColOrder.map(k=>cliCell(k,c)).join("")}
      <div style="color:#4C4C55;display:flex;justify-content:center">${chev}</div>
    </div>`).join("")
    ||`<div class="mut" style="padding:22px 8px">${(nf||q)?"Nenhum cliente com os filtros atuais — use “Limpar” no funil da coluna.":"Nenhum cliente ainda — leia os PDFs na aba “Ler PDFs”."}</div>`;
}
async function abrirCliente(conta){
  const r=await api().client_detail(conta);
  if(!r.ok){toast("Erro: "+(r.erro||""));return;}
  if(!r.relatorios.length){toast("Cliente sem relatórios.");return;}
  cliCur=r; cliMesSel=r.relatorios[0].mes;
  const _cc=(cadastro||[]).find(x=>String(x.conta)===String(conta));
  if(_cc) r.flags=flagsFromCad(_cc);
  $("cliResumo").style.display="none";
  $("cliLista").style.display="none";
  $("cliDetail").style.display="flex";
  $("cliDetNome").textContent=r.nome||("Conta "+r.conta);
  if($("cliDetAv"))$("cliDetAv").textContent=_ini(r.nome,r.conta);
  const _t=(r.termometro||"").toUpperCase();
  const _p=[["Conta",`<span class="dc-mono">${esc(r.conta)}</span>${copyBtn(r.conta)}`]];
  if(_t)_p.push(["Termômetro",`<b style="color:${TC2[_t]||'var(--mut)'}">${_t}</b>`]);
  if(r.patrimonio!=null)_p.push(["Patrimônio",`<b class="dc-mono" style="color:var(--txt)">R$ ${fBR(r.patrimonio,2)}</b>`]);
  if(r.perfil)_p.push(["Perfil",`<b style="color:#F2C029">${esc(r.perfil)}</b>`]);
  if(r.idade!=null)_p.push(["Idade",`<b style="color:var(--txt)">${r.idade} anos</b>`]);
  if((r.objetivo||"")==="Aposentadoria"&&parseInt(r.apos_idade,10)&&r.idade!=null){
    const _f=parseInt(r.apos_idade,10)-r.idade;
    _p.push(["Aposentadoria",`<b style="color:#F2C029">${_f>0?`em ${_f} ano(s) (aos ${parseInt(r.apos_idade,10)})`:"idade-alvo atingida"}</b>`]);
  }
  if(!r.aniv_data&&r.nascimento)_p.push(["Nascimento",`<b style="color:var(--warn)">${esc(r.nascimento)}</b> <span class="mut">(use dd/mm/aaaa)</span>`]);
  $("cliDetSub").innerHTML=`<div style="display:flex;gap:34px;row-gap:14px;flex-wrap:wrap;margin-top:12px">`
    +_p.map(([l,v])=>`<div>
      <div style="font:700 10px 'Manrope';letter-spacing:.14em;text-transform:uppercase;color:#7A7A84;margin-bottom:4px">${l}</div>
      <div style="font-size:14.5px;display:flex;align-items:center">${v}</div></div>`).join("")+`</div>`;
  const _bd=$("cliDetBday");
  if(_bd)_bd.innerHTML=!r.aniv_status?"":(r.aniv_status==="hoje"
    ?`<span class="bday-badge big">${BDAY_ICO} Aniversário HOJE${r.idade_nova!=null?` — faz ${r.idade_nova} anos`:""}!</span>`
    :`<span class="bday-badge big" title="Cai no fim de semana — dê os parabéns hoje, no dia útil">${BDAY_ICO} Aniversário no ${esc(r.aniv_dsem)} (${esc(r.aniv_data)}) — parabenize hoje!</span>`);
  $("cliMesPills").innerHTML=r.relatorios.map(x=>
    `<span class="pill" data-mes="${x.mes}" onclick="cliMesView('${x.mes}')">${esc(x.mes_bonito)}</span>`).join("");
  renderCliCards();
  renderCliExtra();
  cliFam=(r.grupo||[]).slice(); renderCliFam();
  if($("cliFamInput"))$("cliFamInput").value="";
  if($("cliFamStatus"))$("cliFamStatus").textContent="salvo automaticamente";
  cliAportes=null; if($("cliApResumo"))$("cliApResumo").innerHTML='<span class="mut">carregando…</span>';
  if($("cliApStatus"))$("cliApStatus").textContent="salvo automaticamente";
  loadAportes();
  cliContatos=null; if($("cliCtResumo"))$("cliCtResumo").innerHTML='<span class="mut">carregando…</span>';
  if($("cliCtStatus"))$("cliCtStatus").textContent="";
  loadContatos();
}
function voltarClientes(){
  cliCur=null; cliMesSel=null;
  $("cliResumo").style.display="";
  $("cliDetail").style.display="none";
  $("cliLista").style.display="flex";
  cliLoad();
}
function cliMesView(mes){ cliMesSel=mes; renderCliCards(); }
function renderCliCards(){
  if(!cliCur)return;
  $("cliMesPills").querySelectorAll(".pill").forEach(p=>p.classList.toggle("sel",p.dataset.mes===cliMesSel));
  const m=cliCur.relatorios.find(x=>x.mes===cliMesSel)||cliCur.relatorios[0]||{};
  const per=(m.periodo_total||"").toString();
  const perLbl=per?(/in[íi]cio/i.test(per)?"(desde o início)":"("+per+")"):"(período total)";
  const cdiVal=m.cdi_mes==null?"—":
    `<span class="${m.cdi_mes>=100?'pos':(m.cdi_mes>=80?'neu':'neg')}">${fBR(m.cdi_mes,2)}%</span>`;
  const rentVal=m.rent_mes==null?"—":
    `<span class="${m.rent_mes<0?'neg':''}">${fBR(m.rent_mes,2)}%</span>`;
  const movVal=m.movimentacoes==null?"—":
    `<span class="${m.movimentacoes<0?'neg':''}">${rs(m.movimentacoes)}</span>`;
  const cards=[
    ["Ganho financeiro (mês)",rs(m.ganho_mes),"hl"],
    ["Rentabilidade (mês)",rentVal,"hl"],
    ["% do CDI (mês)",cdiVal,""],
    ["Movimentações (mês)",movVal,""],
    ["Patrimônio total bruto",rs(m.patrimonio),""],
    ["Rentabilidade "+perLbl,pct(m.rent_total),""],
    ["Ganho "+perLbl,rs(m.ganho_total),""]
  ];
  $("cliKpis").innerHTML=cards.map(([l,v,h])=>
    `<div class="kpi ${h}"><div class="lbl">${l}</div><div class="val">${v}</div></div>`).join("");
  const aloc=m.aloc||{};
  const ents=Object.keys(aloc).map(k=>[k,aloc[k]]).sort((a,b)=>b[1]-a[1]);
  $("cliCompCard").style.display=ents.length?"":"none";
  $("cliComp").innerHTML=ents.map(([k,v])=>`
    <div style="margin-bottom:9px">
      <div class="row" style="justify-content:space-between;margin-bottom:4px">
        <span>${esc(k)}</span><span class="mut">${fBR(v,2)}%</span></div>
      <div style="height:6px;background:rgba(255,255,255,.05);border-radius:4px;overflow:hidden">
        <div style="height:100%;width:${Math.max(0,Math.min(v,100))}%;background:linear-gradient(90deg,rgba(242,192,41,.55),#F2C029);border-radius:4px"></div>
      </div></div>`).join("");
  $("cliArq").textContent=m.arquivo?("Arquivo: "+m.arquivo):"";
}

const fBR=(v,nd=2)=>v==null?"—":Number(v).toLocaleString("pt-BR",{minimumFractionDigits:nd,maximumFractionDigits:nd});
const pct=(v,nd=2)=>v==null?"—":fBR(v,nd)+"%";
const rs=(v,nd=2)=>v==null?"—":"R$ "+fBR(v,nd);

let alocDados=null;
async function mktLoad(){
  const r=await api().aloc_get();
  alocDados=(r&&r.ok&&r.dados&&r.dados.classes&&r.dados.classes.length)?r.dados:null;
  renderAloc();
  oiRefresh();
}
let oiDados=null;
function mktSub(v){
  $("pillAlocSub").classList.toggle("sel",v==="aloc");
  $("pillOndeSub").classList.toggle("sel",v==="onde");
  $("mktAloc").style.display=v==="aloc"?"flex":"none";
  $("mktOnde").style.display=v==="onde"?"flex":"none";
}
async function mercadoLoad(){
  const b=$("btnMerc"); const old=b.innerHTML; b.disabled=true; b.textContent="Lendo PDFs…";
  try{
    const r=await api().mercado_load();
    if(r&&r.cancel) return;
    if(!r||!r.ok){toast("Erro: "+((r&&r.erro)||"falha ao ler"));return;}
    if(r.aloc&&r.aloc.classes&&r.aloc.classes.length){alocDados=r.aloc;}
    if(r.onde&&((r.onde.resumo||[]).length||(r.onde.secoes||[]).length)){oiDados=r.onde;}
    renderAloc(); renderOi();
    const car=(r.carregados||[]), ign=(r.ignorados||[]);
    let msg=car.length?("✓ "+car.join("  ·  ")):"Nenhum relatório de mercado reconhecido";
    if(ign.length) msg+="  ·  ignorados: "+ign.length;
    toast(msg);
  } finally { b.disabled=false; b.innerHTML=old; }
}
async function oiRefresh(){
  const r=await api().oi_get(); const d=r&&r.ok?r.dados:null;
  oiDados=(d&&((d.resumo||[]).length||(d.secoes||[]).length))?d:null;
  renderOi();
}
function renderOi(){
  if(!oiDados){$("oiVazio").style.display="";$("oiConteudo").style.display="none";
    $("oiSub").textContent="Carregue o relatório “Onde Investir” da XP";return;}
  $("oiVazio").style.display="none";$("oiConteudo").style.display="flex";
  $("oiTitulo").textContent="Onde Investir — "+(oiDados.mes_label||"");
  $("oiSub").textContent="Fonte: Onde Investir";
  $("oiResumo").innerHTML=(oiDados.resumo||[]).map(b=>
    `<div style="display:flex;gap:8px;margin-bottom:10px">
       <span style="color:var(--yellow)">▸</span>
       <span class="small" style="line-height:1.6;color:var(--txt)">${esc(b)}</span></div>`).join("")
    ||"<div class='mut'>—</div>";
  $("oiSecoes").innerHTML=(oiDados.secoes||[]).map(s=>
    `<div style="margin-bottom:14px">
       <div style="font-weight:600;color:var(--yellow);margin-bottom:3px">${esc(s.titulo)}</div>
       <div class="small" style="line-height:1.6;color:var(--txt)">${esc(s.texto)}</div></div>`).join("")
    ||"<div class='mut'>—</div>";
}
function renderAloc(){
  if(!alocDados){
    $("alocVazio").style.display=""; $("alocConteudo").style.display="none";
    $("alocSub").textContent="Carregue o relatório de alocação da XP"; return;
  }
  $("alocVazio").style.display="none"; $("alocConteudo").style.display="flex";
  $("alocTitulo").textContent="Carteiras Recomendadas — "+(alocDados.mes_label||"");
  $("alocSub").textContent="Fonte: Carteiras de Alocação";
  $("alocNota").textContent=alocDados.nota||"";
  $("alocNotaCard").style.display=alocDados.nota?"":"none";
  const dur=alocDados.duration||{};
  $("tbAloc").innerHTML=(alocDados.classes||[]).map(c=>{
    const d=dur[c.classe]!=null?` <span class="mut small">· duration ${fBR(dur[c.classe],1)} anos</span>`:"";
    return `<tr><td>${esc(c.classe)}${d}</td><td>${fBR(c.cons,1)}%</td><td>${fBR(c.mod,1)}%</td><td>${fBR(c.sof,1)}%</td></tr>`;
  }).join("");
  const rr=alocDados.retorno||{}, vv=alocDados.vol||{};
  $("tfAloc").innerHTML=
    `<tr><td><b>Retorno esperado</b></td><td><b>${rr.cons||"—"}</b></td><td><b>${rr.mod||"—"}</b></td><td><b>${rr.sof||"—"}</b></td></tr>`+
    `<tr><td><b>Volatilidade alvo</b></td><td>${vv.cons!=null?fBR(vv.cons,2)+"%":"—"}</td><td>${vv.mod!=null?fBR(vv.mod,2)+"%":"—"}</td><td>${vv.sof!=null?fBR(vv.sof,2)+"%":"—"}</td></tr>`;
  $("alocPersp").innerHTML=(alocDados.persp||[]).map(p=>
    `<div style="margin-bottom:14px">
       <div style="font-weight:600;color:var(--yellow);margin-bottom:3px">${esc(p.classe)}</div>
       <div class="small" style="color:var(--txt);line-height:1.6">${esc(p.texto)}</div></div>`).join("")
    ||"<div class='mut'>Sem comentários neste relatório.</div>";
  $("alocInsights").innerHTML=(alocDados.insights||[]).map(t=>
    `<div style="display:flex;gap:8px;margin-bottom:8px">
       <span style="color:var(--yellow)">▸</span>
       <span class="small" style="line-height:1.5;color:var(--txt)">${esc(t)}</span></div>`).join("")
    ||"<div class='mut'>—</div>";
  const buckets=[
    ["Renda Fixa",["Pós-fixado","Inflação","Prefixado","Global - Renda Fixa"],"#F2C029"],
    ["Multimercados",["Multimercados"],"#3B82F6"],
    ["Renda Variável",["Renda Variável Brasil","Global - Renda Variável"],"#10B981"],
    ["Fundos Listados",["Fundos Listados"],"#8B5CF6"],
    ["Alternativos",["Alternativos"],"#F97316"]];
  const cmap={}; (alocDados.classes||[]).forEach(c=>cmap[c.classe]=c);
  const bval=(names,k)=>names.reduce((a,n)=>a+((cmap[n]||{})[k]||0),0);
  const legenda=buckets.map(b=>`<span style="display:inline-flex;align-items:center;gap:5px;margin:0 14px 6px 0;font-size:12px">
     <span style="width:11px;height:11px;border-radius:3px;background:${b[2]};display:inline-block"></span>${b[0]}</span>`).join("");
  const barras=[["Conservadora","cons"],["Moderada","mod"],["Sofisticada","sof"]].map(([nome,k])=>{
    const segs=buckets.map(b=>{const v=bval(b[1],k); return v>0?
      `<div title="${b[0]}: ${fBR(v,1)}%" style="width:${v}%;background:${b[2]};height:100%;display:flex;align-items:center;justify-content:center;font-size:10px;color:#141414;font-weight:700;overflow:hidden">${v>=7?fBR(v,0)+"%":""}</div>`:"";}).join("");
    return `<div style="margin-bottom:12px">
      <div class="small mut" style="margin-bottom:4px">${nome}</div>
      <div style="display:flex;height:26px;border-radius:7px;overflow:hidden;border:1px solid var(--border)">${segs}</div></div>`;
  }).join("");
  $("alocChart").innerHTML=`<div style="margin-bottom:10px">${legenda}</div>${barras}`;
}

let cartComp=null;
async function cartReload(){
  const r=await api().clients_overview();
  const cli=(r&&r.ok&&r.clientes)||[];
  cliDLFill(cli);
  const sel=$("cartCli");
  let cur=(sel&&sel.dataset.conta)||"";
  if(!cli.some(c=>String(c.conta)===cur)) cur=cli.length?String(cli[0].conta):"";
  if(sel){sel.dataset.conta=cur; sel.value=cliTxtDe(cur);}
  cartLoad();
}
let cartSubAtivo="comp", cartAnalise=null;
function cartSub(v){
  cartSubAtivo=v;
  $("pillComp").classList.toggle("sel",v==="comp");
  $("pillAnal").classList.toggle("sel",v==="anal");
  const has=!!cartComp;
  $("cartComp").style.display=(has&&v==="comp")?"flex":"none";
  $("cartAnal").style.display=(has&&v==="anal")?"flex":"none";
}
function renderCartPerfil(r){
  const el=$("cartPerfil"); if(!el)return;
  if(!r){el.innerHTML="";return;}
  const t=(r.termometro||"").toUpperCase();
  const TC={A:"#10B981",B:"#3B82F6",C:"#F59E0B",D:"#F97316",E:"#EF4444"};
  const pat=r.patrimonio;
  const patTxt=pat==null?"—":(pat>=1e6?"R$ "+fBR(pat/1e6,1)+"M":(pat>=1e3?"R$ "+fBR(pat/1e3,0)+"k":"R$ "+fBR(pat,0)));
  const chip=inner=>`<span style="display:inline-flex;align-items:center;gap:6px;background:var(--card);border:1px solid var(--border);border-radius:9px;padding:5px 11px;font-size:12.5px">${inner}</span>`;
  el.innerHTML=
    chip(`<span class="mut">Perfil</span> <b style="color:var(--yellow)">${esc(r.perfil_cadastro||"—")}</b>`)+
    chip(`<span class="mut">Termô.</span> <b style="color:${TC[t]||'var(--mut)'}">${t||"—"}</b>`)+
    chip(`<span class="mut">Patrimônio</span> <b style="color:var(--txt)">${patTxt}</b>`);
}
async function cartLoad(){
  const conta=($("cartCli")&&$("cartCli").dataset.conta)||"";
  const hide=()=>{$("cartComp").style.display="none";$("cartAnal").style.display="none";};
  if(!conta){$("cartVazio").style.display="";hide();cartComp=null;renderCartPerfil(null);
    $("cartVazio").textContent="Sem clientes — leia os PDFs de cliente na aba “Ler PDFs”.";return;}
  const r=await api().carteira_comparativo(conta);
  if(!r||!r.ok){$("cartVazio").style.display="";hide();cartComp=null;renderCartPerfil(null);
    $("cartVazio").textContent=(r&&r.erro)||"Sem dados.";return;}
  cartComp=r; renderCart(); renderCartPerfil(r);
  const a=await api().analise_carteira(conta);
  cartAnalise=(a&&a.ok)?a:null; renderAnalise();
  $("cartVazio").style.display="none";
  cartSub(cartSubAtivo);
}
function renderAnalise(){
  const a=cartAnalise;
  if(!a){$("analGauge").innerHTML="";$("analResumo").textContent="";
    if($("analRecSub"))$("analRecSub").textContent="Baseadas nos relatórios de mercado";
    $("analRecs").innerHTML="<div class='mut'>Carregue os relatórios de mercado (aba Mercado) para a análise.</div>";
    $("analCtx").textContent="";return;}
  const ad=a.aderencia||0;
  const cor=ad>=85?"#10B981":(ad>=70?"#F2C029":"#EF4444");
  const cor0=ad>=85?"#059669":(ad>=70?"#D9A506":"#DC2626");
  const adp=x=>x==null?"—":fBR(x,1)+"%";
  const nmp={cons:"Conservadora",mod:"Moderada",sof:"Sofisticada"};
  const alvoK=Object.keys(nmp).find(k=>nmp[k]===(a.perfil||""))||a.closest;
  const linhas=["cons","mod","sof"].map(k=>{
    const sel=k===alvoK;
    return `<div style="display:flex;justify-content:space-between;font-size:12px">
      <span style="color:${sel?'#E9EAEE':'#9A9EA8'};font-weight:${sel?600:400}">${nmp[k]}</span>
      <span style="font-family:'JetBrains Mono';font-weight:${sel?700:400};color:${sel?cor:'#C9CBD1'}">${adp((a.aderencias||{})[k])}</span></div>`;
  }).join("");
  $("analGauge").innerHTML=
    `<div style="font:800 46px/1 'JetBrains Mono',monospace;color:${cor};letter-spacing:-.02em">${fBR(ad,1)}%</div>
     <div style="font-size:12.5px;color:#76767F;margin-top:6px">aderente ao perfil <b style="color:var(--yellow)">${esc(a.perfil||"")}</b></div>
     <div style="height:10px;background:rgba(255,255,255,.06);border-radius:6px;overflow:hidden;margin:16px 0 14px">
       <div style="height:100%;width:${Math.max(0,Math.min(ad,100))}%;background:linear-gradient(90deg,${cor0},${cor});border-radius:6px"></div></div>
     <div style="display:flex;flex-direction:column;gap:9px">${linhas}</div>`;
  $("analResumo").textContent=a.resumo||"";
  if($("analRecSub"))$("analRecSub").textContent="Baseadas nos relatórios de mercado"+(a.perfil?" · alvo perfil "+a.perfil:"");
  const patA=cartComp&&cartComp.patrimonio;
  $("analRecs").innerHTML=(a.recomendacoes||[]).map(r=>{
    const cb=r.acao==="Aumentar"?"#059669":(r.acao==="Reduzir"?"#DC2626":"#D97706");
    const ct=r.acao==="Aumentar"?"#10B981":(r.acao==="Reduzir"?"#EF4444":"#F59E0B");
    const alvo=r.acao==="Realocar"?"—":fBR(r.alvo,1)+"%";
    const movimento=(patA!=null)?(r.acao==="Reduzir"?"− ":"+ ")+rs(Math.abs(r.delta)*patA/100):fBR(r.delta,1)+" p.p.";
    return `<div style="border-left:3px solid ${cb};padding-left:13px">
      <div style="display:flex;align-items:baseline;gap:8px;flex-wrap:wrap"><b style="color:${ct};font-size:13px">${r.acao}</b><b style="font-size:13px;color:#E9EAEE">${esc(r.classe)}</b>
        <span style="font-size:11.5px;color:#76767F">· <b style="color:${ct};font-family:'JetBrains Mono'">${movimento}</b> (atual ${fBR(r.atual,1)}% → alvo ${alvo})</span></div>
      ${r.motivo?`<div style="font-size:12px;color:#8C8C96;margin-top:4px;line-height:1.55">${esc(r.motivo)}</div>`:""}</div>`;
  }).join("")||"<div class='mut'>Carteira alinhada à recomendação — sem ajustes relevantes.</div>";
  const cl=a.closest;
  ["cons","mod","sof"].forEach(k=>{const th=$("thAnal_"+k); if(th){th.style.color=k===cl?"var(--yellow)":"";th.style.fontWeight=k===cl?"700":"";}});
  const hlC=k=>k===cl?' style="background:rgba(255,255,255,.05);border-left:2px solid rgba(242,192,41,.55);color:var(--yellow-hi);font-weight:700"':'';
  $("analRecTbl").innerHTML=(a.classes_rec||[]).map(c=>
    `<tr><td>${esc(c.classe)}</td>
      <td${hlC('cons')}>${fBR(c.cons,1)}%</td>
      <td${hlC('mod')}>${fBR(c.mod,1)}%</td>
      <td${hlC('sof')}>${fBR(c.sof,1)}%</td></tr>`).join("")||"<tr><td colspan=4 class='mut'>Carregue o relatório de alocação na aba Mercado.</td></tr>";
  const rr=a.retorno||{}, vv=a.vol||{}, nm={cons:"Conservadora",mod:"Moderada",sof:"Sofisticada"};
  $("analRecFoot").textContent="Perfil-alvo desta análise: "+(a.perfil||"—")
    +(rr[cl]?("  ·  retorno esperado "+rr[cl]):"")
    +(vv[cl]!=null?("  ·  volatilidade alvo "+fBR(vv[cl],2)+"%"):"")
    +"  ·  coluna destacada = perfil mais próximo da carteira atual.";
  $("analCtx").textContent=a.contexto||"Carregue o relatório “Onde Investir” para mais contexto de mercado.";
}
function desvioMini(diff, pat){
  const a=Math.abs(diff), neutro=a<0.05;
  const cor=neutro?"var(--mut)":(diff>0?"#F97316":"#3B82F6");
  const w=Math.min(a*2.2,33);
  const seg=neutro?"":(diff>0
    ?`<span style="position:absolute;left:50%;top:1px;width:${w}px;height:6px;background:${cor};border-radius:3px"></span>`
    :`<span style="position:absolute;right:50%;top:1px;width:${w}px;height:6px;background:${cor};border-radius:3px"></span>`);
  const temPat=(pat!=null&&isFinite(pat));
  const txt=neutro?"—":(temPat?((diff>0?"− ":"+ ")+rs(a*pat/100)):((diff>0?"+":"")+fBR(diff,1)+" p.p."));
  return `<div style="display:flex;align-items:center;gap:8px;justify-content:center" title="atual − recomendado (quanto realocar)">
    <span style="color:${cor};font-weight:700;min-width:118px;text-align:right;font-size:12.5px">${txt}</span>
    <span style="position:relative;width:60px;height:8px;background:var(--border);border-radius:4px;flex:none">
      <span style="position:absolute;left:50%;top:0;width:1px;height:8px;background:var(--mut2)"></span>${seg}</span></div>`;
}
function renderCart(){
  const r=cartComp; if(!r)return;
  $("cartMes").textContent="Carteira de "+(r.mes_cliente||"")+"  ·  recomendações "+(r.mes_label||"");
  const nomes={cons:"Conservadora",mod:"Moderada",sof:"Sofisticada"};
  const dist=r.dist||{}, closest=r.closest, rr=r.retorno||{}, vv=r.vol||{};
  const patC=r.patrimonio;
  $("cartKpis").innerHTML=["cons","mod","sof"].map(k=>
    `<div class="kpi ${k===closest?'hl':''}">
       <div class="lbl">${nomes[k]}${k===closest?' · alvo':''}</div>
       <div class="val">${rr[k]||"—"}</div>
       <div class="small mut">vol alvo ${vv[k]!=null?fBR(vv[k],2)+'%':'—'} · realocar ${patC!=null?rs((dist[k]||0)*patC/100):fBR(dist[k]||0,1)+' p.p.'}</div>
     </div>`).join("");
  const cols=closest?[closest].concat(["cons","mod","sof"].filter(k=>k!==closest)):["cons","mod","sof"];
  $("cartHead").innerHTML=`<tr><th>Classe de ativo</th><th>Atual</th>`
    +`<th style="text-align:center">Ajuste p/ alvo (R$)</th>`
    +cols.map((k,i)=>`<th${i===0?' style="color:var(--yellow)"':''}>${nomes[k]}${i===0?' · perfil':''}</th>`).join("")+`</tr>`;
  const hlStyle='background:rgba(255,255,255,.05);border-left:2px solid rgba(242,192,41,.55);color:var(--yellow-hi);font-weight:700';
  $("tbCart").innerHTML=(r.linhas||[]).map(l=>{
    const rec=k=>l[k]==null?'<span class="mut">—</span>':fBR(l[k],1)+'%';
    const nome=esc(l.classe)+(l.extra?' <span class="small mut">(fora da recomendada)</span>':'');
    const alvo=(l[closest]==null?0:l[closest]);
    const diff=Math.round((l.atual-alvo)*10)/10;
    const recCells=cols.map((k,i)=>`<td${i===0?` style="${hlStyle}"`:''}>${rec(k)}</td>`).join("");
    return `<tr style="${l.extra?'background:rgba(239,68,68,.07)':''}">
      <td${l.extra?' class="mut"':''}>${nome}</td>
      <td><b>${fBR(l.atual,2)}%</b></td>
      <td>${desvioMini(diff, r.patrimonio)}</td>
      ${recCells}</tr>`;
  }).join("");
  const soma=(r.linhas||[]).reduce((a,l)=>a+(l.atual||0),0);
  $("cartNaoMap").textContent="Soma da carteira atual: "+fBR(soma,2)+"% — valores exatos da composição do cliente (aba Clientes).";
  $("cartPersp").innerHTML=(r.persp||[]).map(p=>
    `<div style="margin-bottom:12px">
       <div style="font-weight:600;color:var(--yellow);margin-bottom:3px">${esc(p.classe)}</div>
       <div class="small" style="line-height:1.6">${esc(p.texto)}</div></div>`).join("")
    ||"<div class='mut'>Carregue o relatório de alocação para ver as perspectivas.</div>";
}

(function(){
  var _c=[71,117,105,108,104,101,114,109,101,32,69,110,114,105,99,111];
  var _n=function(){return String.fromCharCode.apply(null,_c);};
  function _wm(){
    try{
      var el=document.getElementById("gh_wm");
      if(!el||!el.isConnected){
        el=el||document.createElement("div");
        el.id="gh_wm";
        var lg=document.querySelector("aside .logo");
        if(lg&&lg.parentNode)lg.parentNode.insertBefore(el,lg.nextSibling);
        else (document.body||document.documentElement).appendChild(el);
      }
      el.textContent="by "+_n();
      el.style.cssText="margin:-4px 0 14px 16px;font-family:Manrope,'Segoe UI',Arial,sans-serif;"+
        "font-size:10.5px;font-weight:200;letter-spacing:.24em;text-transform:uppercase;"+
        "color:#9A9AA4;user-select:none;pointer-events:none;white-space:nowrap;"+
        "opacity:1;visibility:visible;display:block;height:auto;overflow:visible";
    }catch(e){}
  }
  _wm();
  setInterval(_wm,1200);
  try{
    new MutationObserver(function(){
      var e=document.getElementById("gh_wm");
      if(!e||!e.isConnected)_wm();
    }).observe(document.documentElement,{childList:true,subtree:true});
  }catch(e){}
})();
</script>
</body></html>"""

if __name__ == "__main__":
    try:
        webview.create_window(
            "XPerformance Manager - Guilherme Enrico", html=HTML, js_api=Api(),
            width=1220, height=820, min_size=(980, 660),
            background_color="#0A0A0C")
        webview.start()
        _flush_agora()
    except Exception:
        with open(_LOGF, "a", encoding="utf-8") as f:
            f.write(traceback.format_exc())
